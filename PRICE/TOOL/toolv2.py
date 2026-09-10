"""
TOOL V2: CÔNG CỤ THEO DÕI, ĐỊNH GIÁ & XUẤT BÁO CÁO GIÁ TUẦN MEDIGO
===================================================================
Tính năng chuẩn hóa:
  1. Tự động nhận diện file mới nhất trong DATA/ (crawl, 28-4, banggia KV, mapping, he_so).
  2. Khớp đúng Link + ĐVT đối thủ (Long Châu / Pharmacity / An Khang) bằng Smart Matching.
  3. Bổ sung thông tin App Medigo: 'ID App', 'Tên App', 'ĐVT App' từ sheet 'import_app' & 'mapping Kiot - MedigoApp.xlsx'.
  4. Lấy đầy đủ 'Giá vốn gần nhất' và 'Giá vốn cũ' từ 22 kỳ lịch sử 'update giá nhập' + '28-4.xlsx'.
  5. Lọc đúng các sản phẩm ĐỐI THỦ THAY ĐỔI GIÁ tuần này làm Sheet 1 trình Sếp.
  6. ĐỂ TRỐNG các cột Giá Mới (Giá mới Online, Giá mới offline low, Giá mới offline high...) để người dùng/Sếp điền.
"""

import os
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path
from typing import Optional

import numpy as np
import pandas as pd

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass


# ======================================================================
# 1. CẤU HÌNH ĐƯỜNG DẪN ĐỘNG (CONFIG)
# ======================================================================
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent

DATA_DIR = PROJECT_DIR / "DATA"
OUTPUT_DIR = PROJECT_DIR / "OUTPUT"
HISTORY_DIR = SCRIPT_DIR / "history"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
HISTORY_DIR.mkdir(parents=True, exist_ok=True)

STATE_FILE = SCRIPT_DIR / "state_gia_tuan_truoc_v2.csv"


def find_latest_file(directory: Path, patterns: list[str], exclude_keywords: list[str] = None, fallback_name: Optional[str] = None) -> Path:
    candidates = []
    for pat in patterns:
        for f in directory.glob(pat):
            if exclude_keywords and any(k.lower() in f.name.lower() for k in exclude_keywords):
                continue
            candidates.append(f)
    if candidates:
        candidates.sort(key=lambda x: x.stat().st_mtime, reverse=True)
        return candidates[0]
    if fallback_name:
        fb = directory / fallback_name
        if fb.exists():
            return fb
    raise FileNotFoundError(f"Không tìm thấy file nào khớp với patterns {patterns} trong {directory}")


# ======================================================================
# 2. CHUẨN HÓA ĐƠN VỊ TÍNH (DVT NORMALIZER)
# ======================================================================
DVT_MAP = {
    "hop": "hộp", "h": "hộp", "box": "hộp", "hộp": "hộp", "hop ": "hộp",
    "vien": "viên", "v": "viên", "pill": "viên", "tab": "viên", "tablet": "viên", "viên": "viên",
    "vi": "vỉ", "blister": "vỉ", "vỉ": "vỉ",
    "tuyp": "tuýp", "t": "tuýp", "tube": "tuýp", "tuýp": "tuýp",
    "chai": "chai", "c": "chai", "bottle": "chai", "lo": "chai", "lọ": "chai", "hũ": "chai", "hu": "chai",
    "goi": "gói", "g": "gói", "sachet": "gói", "pack": "gói", "tui": "gói", "túi": "gói", "gói": "gói",
    "ong": "ống", "amp": "ống", "ampoule": "ống", "ống": "ống",
    "mieng": "miếng", "patch": "miếng", "miếng": "miếng",
    "binh": "bình", "bình": "bình",
    "cap": "cặp", "bo": "cặp", "cặp": "cặp", "bộ": "cặp",
    "cai": "cái", "cây": "cái", "chiec": "cái", "chiếc": "cái", "cái": "cái",
}


def norm_str(s: any) -> str:
    if pd.isna(s):
        return ""
    s = str(s).strip().lower()
    return re.sub(r"\s+", " ", s)


def norm_dvt(dvt: any) -> str:
    cleaned = norm_str(dvt)
    plain = (
        cleaned.replace("ộ", "o").replace("ố", "o").replace("ồ", "o").replace("ọ", "o")
        .replace("ê", "e").replace("ế", "e").replace("ề", "e").replace("ệ", "e")
        .replace("í", "i").replace("ỉ", "i").replace("ị", "i")
        .replace("ý", "y").replace("ỳ", "y").replace("ỵ", "y")
        .replace("ú", "u").replace("ù", "u").replace("ụ", "u")
        .replace("ả", "a").replace("ã", "a").replace("á", "a").replace("à", "a").replace("ạ", "a")
        .replace("đ", "d")
    )
    return DVT_MAP.get(cleaned, DVT_MAP.get(plain, cleaned))


def norm_key(*parts) -> str:
    s = "".join(str(p) for p in parts if pd.notna(p))
    return re.sub(r"\s+", "", s.strip().lower())


def read_excel_fast(path: Path, sheet_name: any = 0) -> pd.DataFrame:
    try:
        return pd.read_excel(path, sheet_name=sheet_name, engine="calamine")
    except Exception:
        return pd.read_excel(path, sheet_name=sheet_name)


# ======================================================================
# 3. TẢI DỮ LIỆU & LỊCH SỬ GIÁ VỐN + APP MAPPING
# ======================================================================
def load_lc_mapping(path: Path, app_file_path: Optional[Path] = None) -> pd.DataFrame:
    df = read_excel_fast(path, sheet_name="data_LC")
    keep = {
        "Mã Hàng": "Ma_Hang",
        "Mã sp": "Ma_sp",
        "Tên Hàng": "Ten_Hang",
        "ĐVT": "DVT",
        "Link LC": "Link_LC",
        "BG LC": "BG_LC_ghi_nhan_truoc",
        "Quy Đổi": "Quy_Doi",
    }
    available_keep = {k: v for k, v in keep.items() if k in df.columns}
    df = df[list(available_keep.keys())].rename(columns=available_keep)
    df = df.dropna(subset=["Ma_Hang", "Link_LC"]).copy()
    df["Ma_Hang"] = df["Ma_Hang"].astype(str).str.strip()
    df = df[df["Link_LC"].astype(str).str.startswith("http")].copy()
    df = df.drop_duplicates(subset="Ma_Hang", keep="first")

    if "Ma_sp" not in df.columns:
        df["Ma_sp"] = ""
    else:
        df["Ma_sp"] = df["Ma_sp"].fillna("").astype(str).str.strip()

    try:
        df_ia = read_excel_fast(path, sheet_name="import_app")
        df_ia["Mã Hàng"] = df_ia["Mã Hàng"].astype(str).str.strip()
        ia_sub = df_ia[["Mã Hàng", "Tên App", "ID App", "ĐVT App"]].dropna(subset=["Mã Hàng"]).drop_duplicates("Mã Hàng")
        ia_sub = ia_sub.rename(columns={"Mã Hàng": "Ma_Hang", "Tên App": "Ten_App", "ID App": "ID_App", "ĐVT App": "DVT_App"})
        df = df.merge(ia_sub, on="Ma_Hang", how="left")
    except Exception:
        df["Ten_App"] = np.nan
        df["ID_App"] = np.nan
        df["DVT_App"] = np.nan

    if app_file_path and app_file_path.exists():
        try:
            df_app_f = read_excel_fast(app_file_path, sheet_name=0)
            df_app_f["Mã hàng"] = df_app_f["Mã hàng"].astype(str).str.strip()
            app_dict = df_app_f.set_index("Mã hàng")["ID App"].to_dict()
            df["ID_App"] = df["ID_App"].fillna(df["Ma_Hang"].map(app_dict))
        except Exception:
            pass

    df["Ten_App"] = df["Ten_App"].fillna(df["Ten_Hang"])
    df["DVT_App"] = df["DVT_App"].fillna(df["DVT"])

    df["DVT_norm"] = df["DVT"].apply(norm_dvt)
    df["key_exact"] = [norm_key(u, d) for u, d in zip(df["Link_LC"], df["DVT_norm"])]
    df["key_url"] = [norm_key(u) for u in df["Link_LC"]]
    df["Quy_Doi"] = pd.to_numeric(df.get("Quy_Doi", 1.0), errors="coerce").fillna(1.0)
    df["BG_LC_ghi_nhan_truoc"] = pd.to_numeric(df.get("BG_LC_ghi_nhan_truoc", np.nan), errors="coerce")
    return df


def load_he_so(path: Path) -> pd.DataFrame:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    target_sheet = None
    for s in wb.sheetnames:
        if "hệ số" in s.lower() and "quy cách" in s.lower():
            target_sheet = s
            break
    wb.close()
    if not target_sheet:
        target_sheet = "hệ số  quy cách"

    df = read_excel_fast(path, sheet_name=target_sheet)
    keep = {"Mã Hàng": "Ma_Hang", "* Hệ số": "He_so"}
    for k in keep:
        if k not in df.columns:
            for c in df.columns:
                if "hệ số" in str(c).lower() and k == "* Hệ số":
                    df[k] = df[c]
                elif "mã hàng" in str(c).lower() and k == "Mã Hàng":
                    df[k] = df[c]

    df = df[["Mã Hàng", "* Hệ số"]].rename(columns=keep).dropna(subset=["Ma_Hang"])
    df["Ma_Hang"] = df["Ma_Hang"].astype(str).str.strip()
    df["He_so"] = pd.to_numeric(df["He_so"], errors="coerce").fillna(1.0)
    return df.drop_duplicates(subset="Ma_Hang", keep="first")


def load_crawl(path: Path) -> pd.DataFrame:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    sheet_names = wb.sheetnames
    wb.close()

    keep = {"name": "Ten_SP_crawl", "url": "URL_crawl", "packaging": "DVT_crawl", "price": "Gia_crawl"}
    target_sheet = sheet_names[0]
    df = None

    for s in sheet_names:
        temp_df = read_excel_fast(path, sheet_name=s)
        if all(col in temp_df.columns for col in ["url", "price"]):
            target_sheet = s
            df = temp_df
            break

    if df is None:
        df = read_excel_fast(path, sheet_name=0)

    missing = [c for c in keep if c not in df.columns]
    if missing:
        raise ValueError(f"File crawl {path.name} thiếu cột: {missing}")

    df = df[list(keep.keys())].rename(columns=keep).copy()
    df = df.dropna(subset=["URL_crawl", "Gia_crawl"]).copy()
    df["Gia_crawl"] = pd.to_numeric(df["Gia_crawl"], errors="coerce")
    df = df[df["Gia_crawl"] > 0].copy()

    df["DVT_crawl_norm"] = df["DVT_crawl"].apply(norm_dvt)
    df["key_exact"] = [norm_key(u, d) for u, d in zip(df["URL_crawl"], df["DVT_crawl_norm"])]
    df["key_url"] = [norm_key(u) for u in df["URL_crawl"]]
    return df


def load_historical_cost_prices(master_path: Path, nhap_path: Path) -> pd.DataFrame:
    df_up = read_excel_fast(master_path, sheet_name="update giá nhập")
    df_up["Mã Hàng"] = df_up["Mã Hàng"].astype(str).str.strip()
    
    vat_cols = [c for c in df_up.columns if "giá nhập vat\n" in str(c).lower() or "giá nhập vat (fix)" in str(c).lower() or "giá nhập vat cũ trước" in str(c).lower()]
    
    def extract_costs_from_history_sheet(row):
        vals = []
        for c in vat_cols:
            v = row[c]
            if pd.notna(v):
                try:
                    num = float(v)
                    if num > 0:
                        vals.append(num)
                except Exception:
                    pass
        if len(vals) >= 2:
            return vals[-2], vals[-1]
        elif len(vals) == 1:
            return vals[0], vals[0]
        return np.nan, np.nan

    res_history = [extract_costs_from_history_sheet(r) for _, r in df_up.iterrows()]
    df_up["gv_cu_sheet"] = [r[0] for r in res_history]
    df_up["gv_moi_sheet"] = [r[1] for r in res_history]
    master_costs = df_up[["Mã Hàng", "gv_cu_sheet", "gv_moi_sheet"]].drop_duplicates("Mã Hàng")

    df_nhap = read_excel_fast(nhap_path, sheet_name=0)
    df_nhap["Mã hàng"] = df_nhap["Mã hàng"].astype(str).str.strip()
    if "Trạng thái" in df_nhap.columns:
        df_nhap = df_nhap[~df_nhap["Trạng thái"].astype(str).str.contains("Hủy|huỷ", case=False, na=False)].copy()

    df_nhap["Thời gian"] = pd.to_datetime(df_nhap["Thời gian"], errors="coerce")
    df_nhap["Gia_von_VAT"] = pd.to_numeric(df_nhap["Giá nhập"], errors="coerce").fillna(0) + pd.to_numeric(df_nhap["VAT nhập hàng"], errors="coerce").fillna(0)
    df_nhap = df_nhap[df_nhap["Gia_von_VAT"] > 0].sort_values("Thời gian")

    po_latest = df_nhap.groupby("Mã hàng").tail(1).set_index("Mã hàng")["Gia_von_VAT"].to_dict()
    multi_po = df_nhap.groupby("Mã hàng").filter(lambda x: len(x) > 1)
    po_prev = multi_po.groupby("Mã hàng").nth(-2).set_index("Mã hàng")["Gia_von_VAT"].to_dict() if len(multi_po) > 0 else {}

    master_costs["gv_moi_po"] = master_costs["Mã Hàng"].map(po_latest)
    master_costs["gv_cu_po"] = master_costs["Mã Hàng"].map(po_prev)

    master_costs["Gia_von_VAT"] = master_costs["gv_moi_po"].fillna(master_costs["gv_moi_sheet"])
    master_costs["Gia_von_tuan_truoc"] = master_costs["gv_cu_po"].fillna(
        master_costs["gv_moi_sheet"].where(master_costs["gv_moi_po"].notna() & (master_costs["gv_moi_po"] != master_costs["gv_moi_sheet"]), master_costs["gv_cu_sheet"])
    )

    return master_costs[["Mã Hàng", "Gia_von_VAT", "Gia_von_tuan_truoc"]].rename(columns={"Mã Hàng": "Ma_Hang"})


def load_banggia_medigo(path: Path) -> pd.DataFrame:
    df = read_excel_fast(path, sheet_name=0)
    if "Mã hàng" not in df.columns:
        return pd.DataFrame(columns=["Ma_Hang", "Gia_ban_hien_tai_Medigo", "Nhóm hàng"])

    df["Ma_Hang"] = df["Mã hàng"].astype(str).str.strip()
    
    price_col = None
    for cand in ["Banggia online HCM", "Bảng giá chung", "Banggia online HN", "Giá bán"]:
        if cand in df.columns:
            price_col = cand
            break

    df["Gia_ban_hien_tai_Medigo"] = pd.to_numeric(df[price_col], errors="coerce") if price_col else np.nan
    df["Nhom_hang"] = df["Nhóm hàng"].astype(str) if "Nhóm hàng" in df.columns else ""

    return df[["Ma_Hang", "Gia_ban_hien_tai_Medigo", "Nhom_hang"]].drop_duplicates("Ma_Hang", keep="first")


# ======================================================================
# 4. SMART MULTI-PACKAGING MATCHING
# ======================================================================
def match_competitor_price(mapping: pd.DataFrame, crawl: pd.DataFrame, he_so: pd.DataFrame) -> pd.DataFrame:
    exact_crawl = crawl[["key_exact", "Gia_crawl", "DVT_crawl"]].drop_duplicates("key_exact", keep="first")
    exact_crawl = exact_crawl.rename(columns={"Gia_crawl": "Gia_LC_exact", "DVT_crawl": "DVT_crawl_exact"})
    out = mapping.merge(exact_crawl, on="key_exact", how="left")

    crawl_grouped = crawl.groupby("key_url").agg(
        n_pkgs=("DVT_crawl", "count"),
        dvt_list=("DVT_crawl_norm", list),
        price_list=("Gia_crawl", list),
    ).reset_index()

    out = out.merge(crawl_grouped, on="key_url", how="left")
    out = out.merge(he_so, on="Ma_Hang", how="left")
    out["He_so"] = out["He_so"].fillna(1.0)

    def resolve_fallback_price(row):
        if pd.notna(row["Gia_LC_exact"]):
            return row["Gia_LC_exact"], "khop_chinh_xac"
        
        if pd.isna(row["n_pkgs"]) or row["n_pkgs"] == 0:
            return np.nan, "khong_tim_thay"

        dvt_med = row["DVT_norm"]
        dvt_list = row["dvt_list"]
        price_list = row["price_list"]

        if dvt_med in dvt_list:
            idx = dvt_list.index(dvt_med)
            return price_list[idx], "khop_dvt_chuan_hoa"

        if len(price_list) == 1:
            return price_list[0] * row["He_so"], "fallback_url_1_dvt"

        hs = row["He_so"]
        if hs != 1.0:
            base_price = min(price_list) if hs > 1 else max(price_list)
            return base_price * hs, "fallback_he_so_multi_pkg"
        
        baseline = row["BG_LC_ghi_nhan_truoc"]
        if pd.notna(baseline) and baseline > 0:
            closest_price = min(price_list, key=lambda p: abs(p - baseline))
            return closest_price, "fallback_closest_baseline"

        return max(price_list), "fallback_default_max"

    resolved = [resolve_fallback_price(r) for _, r in out.iterrows()]
    out["BG_LC_hien_tai"] = [r[0] for r in resolved]
    out["Nguon_khop_gia"] = [r[1] for r in resolved]

    return out[[
        "Ma_Hang", "Ma_sp", "Ten_Hang", "DVT", "ID_App", "Ten_App", "DVT_App",
        "Link_LC", "BG_LC_hien_tai", "Nguon_khop_gia", "BG_LC_ghi_nhan_truoc", "Quy_Doi", "DVT_norm"
    ]]


# ======================================================================
# 5. PIPELINE CHẠY CHÍNH (ĐỂ TRỐNG CỘT GIÁ MỚI THEO YÊU CẦU)
# ======================================================================
def run_pricing_pipeline():
    print("=" * 85)
    print("      MEDIGO PRICING ENGINE V2 (MAPPING APP - KIOT - ĐỐI THỦ CHUẨN XÁC)")
    print("=" * 85)

    print("--- [1/5] Quét các file dữ liệu mới nhất trong DATA/ ---")
    crawl_file = find_latest_file(DATA_DIR, ["pharmacity_products.xlsx", "*crawl*.xlsx", "*pharmacity*.xlsx"], exclude_keywords=["mapping", "he_so", "banggia"])
    nhap_file = find_latest_file(DATA_DIR, ["28-4.xlsx", "*DanhSachChiTietNhapHang*.xlsx", "*nhaphang*.xlsx"], exclude_keywords=["mapping", "banggia_med"])
    mapping_file = find_latest_file(DATA_DIR, ["banggia_med_data_nhaphang_final.xlsx"])
    he_so_file = find_latest_file(DATA_DIR, ["crawl - mapping LC.xlsx", "crawl_-_mapping_LC.xlsx", "*he_so*.xlsx"])
    kiot_price_file = find_latest_file(DATA_DIR, ["BangGia_KV*.xlsx", "*BangGia*.xlsx"])
    
    app_mapping_file = DATA_DIR / "mapping Kiot - MedigoApp.xlsx"
    app_file_path = app_mapping_file if app_mapping_file.exists() else None

    print(f"  • File Crawl Đối thủ   : {crawl_file.name}")
    print(f"  • File Nhập hàng Tuần  : {nhap_file.name}")
    print(f"  • File Bảng giá Kiot   : {kiot_price_file.name}")
    print(f"  • File Master Mapping  : {mapping_file.name}")
    print(f"  • File Hệ số Quy cách  : {he_so_file.name}")

    print("\n--- [2/5] Nạp Master Data (App - Kiot), Lịch sử Giá Vốn & Smart Matching ---")
    mapping = load_lc_mapping(mapping_file, app_file_path)
    he_so = load_he_so(he_so_file)
    crawl = load_crawl(crawl_file)
    lc_priced = match_competitor_price(mapping, crawl, he_so)

    cost_history = load_historical_cost_prices(mapping_file, nhap_file)
    banggia_kv = load_banggia_medigo(kiot_price_file)

    snap = lc_priced.merge(cost_history, on="Ma_Hang", how="left")
    snap = snap.merge(banggia_kv, on="Ma_Hang", how="left")

    print("\n--- [3/5] So sánh biến động với Tuần trước ---")
    first_run = not STATE_FILE.exists()
    if first_run:
        print("=> Lần đầu chạy: Sử dụng 'BG LC' ghi nhận trước làm baseline đối thủ.")
        snap_prev = pd.DataFrame(columns=["Ma_Hang", "BG_LC_hien_tai"])
    else:
        backup_name = HISTORY_DIR / f"state_backup_v2_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        shutil.copy2(STATE_FILE, backup_name)
        snap_prev = pd.read_csv(STATE_FILE)
        snap_prev["Ma_Hang"] = snap_prev["Ma_Hang"].astype(str).str.strip()

    cmp_df = snap.merge(
        snap_prev[["Ma_Hang", "BG_LC_hien_tai"]],
        on="Ma_Hang", how="left", suffixes=("", "_tuan_truoc"),
    )
    cmp_df = cmp_df.rename(columns={
        "BG_LC_hien_tai_tuan_truoc": "BG_LC_tuan_truoc",
        "Gia_ban_hien_tai_Medigo": "Gia_ban_truoc_do",
    })

    dung_baseline_mapping = cmp_df["BG_LC_tuan_truoc"].isna() & cmp_df["BG_LC_ghi_nhan_truoc"].notna()
    cmp_df.loc[dung_baseline_mapping, "BG_LC_tuan_truoc"] = cmp_df.loc[dung_baseline_mapping, "BG_LC_ghi_nhan_truoc"]

    def khac_nhau(a, b, tol=0.01):
        if pd.isna(a) and pd.isna(b):
            return False
        if pd.isna(a) or pd.isna(b):
            return True
        return abs(a - b) > tol

    khac_lc = np.array([khac_nhau(a, b) for a, b in zip(cmp_df["BG_LC_hien_tai"], cmp_df["BG_LC_tuan_truoc"])])
    khac_gv = np.array([khac_nhau(a, b) for a, b in zip(cmp_df["Gia_von_VAT"], cmp_df["Gia_von_tuan_truoc"])])

    cmp_df["Doi_thu_doi_gia"] = cmp_df["BG_LC_hien_tai"].notna() & cmp_df["BG_LC_tuan_truoc"].notna() & khac_lc
    cmp_df["Mat_gia_tuan_nay"] = cmp_df["BG_LC_hien_tai"].isna() & cmp_df["BG_LC_tuan_truoc"].notna()
    cmp_df["Gia_von_doi"] = cmp_df["Gia_von_VAT"].notna() & cmp_df["Gia_von_tuan_truoc"].notna() & khac_gv

    ROUND_FACTORS = [2, 3, 4, 5, 6, 8, 10, 12, 15, 20, 24, 30, 50, 60, 100]

    def nghi_lech_dvt(gia_moi, gia_cu, quy_doi):
        if pd.isna(gia_moi) or pd.isna(gia_cu) or gia_moi == 0 or gia_cu == 0:
            return False
        ty_le = max(gia_moi, gia_cu) / min(gia_moi, gia_cu)
        candidates = ROUND_FACTORS + ([quy_doi] if pd.notna(quy_doi) and quy_doi > 1 else [])
        return any(abs(ty_le - c) / c < 0.05 for c in candidates)

    cmp_df["Nghi_lech_DVT"] = [
        nghi_lech_dvt(a, b, q)
        for a, b, q in zip(cmp_df["BG_LC_hien_tai"], cmp_df["BG_LC_tuan_truoc"], cmp_df["Quy_Doi"])
    ]
    cmp_df["Doi_thu_doi_gia_that"] = cmp_df["Doi_thu_doi_gia"] & ~cmp_df["Nghi_lech_DVT"]
    cmp_df["Nghi_lech_DVT"] = cmp_df["Doi_thu_doi_gia"] & cmp_df["Nghi_lech_DVT"]

    def gen_reason(r):
        if r["Doi_thu_doi_gia_that"]:
            bg_moi = r["BG_LC_hien_tai"]
            bg_cu = r["BG_LC_tuan_truoc"]
            pct = (bg_moi - bg_cu) / bg_cu if bg_cu > 0 else 0
            act = "tăng" if pct > 0 else "giảm"
            return f"Đối thủ {act} giá ({pct:+.1%})"
        if r["Gia_von_doi"]:
            gv_moi = r["Gia_von_VAT"]
            gv_cu = r["Gia_von_tuan_truoc"]
            pct_gv = (gv_moi - gv_cu) / gv_cu if gv_cu > 0 else 0
            act_gv = "tăng" if pct_gv > 0 else "giảm"
            return f"Giá vốn {act_gv} ({pct_gv:+.1%})"
        return ""

    cmp_df["Ly_do_dieu_chinh"] = cmp_df.apply(gen_reason, axis=1)

    # ĐỂ TRỐNG CÁC CỘT GIÁ MỚI THEO YÊU CẦU
    cmp_df["Gia_moi_Online"] = np.nan
    cmp_df["Gia_moi_offline_low"] = np.nan
    cmp_df["Gia_moi_offline_hight"] = np.nan
    cmp_df["So_voi_doi_thu"] = np.nan
    cmp_df["Margin"] = np.nan
    cmp_df["Loi_nhuan_gop_VND"] = np.nan

    # 4. Tách các bảng lọc theo điều kiện
    df_doi_thu_doi = cmp_df[cmp_df["Doi_thu_doi_gia_that"]].copy()
    df_gia_von_doi = cmp_df[cmp_df["Gia_von_doi"]].copy()
    df_tong_hop = cmp_df[cmp_df["Doi_thu_doi_gia_that"] | cmp_df["Gia_von_doi"]].copy()

    # 5. Xuất Báo cáo Excel Đa Sheet
    timestamp_str = datetime.now().strftime("%Y%m%d_%H%M")
    output_filename = OUTPUT_DIR / f"Bao_cao_import_price_V2_{timestamp_str}.xlsx"

    cols_export_template = [
        "Ma_Hang", "Ma_sp", "Ten_Hang", "DVT", "ID_App", "Ten_App", "DVT_App",
        "BG_LC_tuan_truoc", "Link_LC", "BG_LC_hien_tai",
        "Gia_von_VAT", "Gia_von_tuan_truoc", "Gia_ban_truoc_do",
        "Gia_moi_Online", "Gia_moi_offline_low", "Gia_moi_offline_hight",
        "So_voi_doi_thu", "Margin", "Loi_nhuan_gop_VND", "Ly_do_dieu_chinh"
    ]

    rename_export = {
        "Ma_Hang": "Mã Hàng (Kiot)",
        "Ma_sp": "Mã sp",
        "Ten_Hang": "Tên Hàng (Kiot)",
        "DVT": "ĐVT (Kiot)",
        "ID_App": "ID App",
        "Ten_App": "Tên trên App",
        "DVT_App": "ĐVT trên App",
        "BG_LC_tuan_truoc": "BG LC",
        "Link_LC": "Link LC",
        "BG_LC_hien_tai": "BG LC hiện tại",
        "Gia_von_VAT": "Giá vốn gần nhất",
        "Gia_von_tuan_truoc": "Giá vốn cũ",
        "Gia_ban_truoc_do": "Gía bán trước đó",
        "Gia_moi_Online": "Giá mới Online",
        "Gia_moi_offline_low": "Giá mới offline low",
        "Gia_moi_offline_hight": "Giá mới offline hight",
        "So_voi_doi_thu": "So với đối thủ",
        "Margin": "Margin",
        "Loi_nhuan_gop_VND": "Lợi nhuận gộp (VNĐ)",
        "Ly_do_dieu_chinh": "Lý do điều chỉnh",
    }

    print(f"\n--- [4/5] Xuất file Báo cáo V2: {output_filename.name} ---")
    with pd.ExcelWriter(output_filename, engine="openpyxl") as writer:
        # Sheet 1: ƯU TIÊN HIỂN THỊ CÁC SẢN PHẨM ĐỐI THỦ THAY ĐỔI GIÁ
        df_doi_thu_doi[cols_export_template].rename(columns=rename_export).to_excel(
            writer, sheet_name="Doi thu doi gia (Trinh Sep)", index=False
        )

        # Sheet 2: Các sản phẩm giá vốn thay đổi
        df_gia_von_doi[cols_export_template].rename(columns=rename_export).to_excel(
            writer, sheet_name="Gia von doi (Trinh Sep)", index=False
        )

        # Sheet 3: Tổng hợp toàn bộ mã có biến động (Đối thủ đổi HOẶC Giá vốn đổi)
        df_tong_hop[cols_export_template].rename(columns=rename_export).to_excel(
            writer, sheet_name="Tong hop bien dong", index=False
        )

        # Sheet 4: Bảng so sánh 4 giá tổng thể
        cmp_df[[
            "Ma_Hang", "Ten_Hang", "DVT", "ID_App", "Ten_App", "DVT_App",
            "BG_LC_tuan_truoc", "BG_LC_hien_tai", "Gia_von_tuan_truoc", "Gia_von_VAT", "Gia_ban_truoc_do"
        ]].rename(columns={
            "Ma_Hang": "Mã Hàng (Kiot)", "Ten_Hang": "Tên Hàng (Kiot)", "DVT": "ĐVT (Kiot)",
            "ID_App": "ID App", "Ten_App": "Tên trên App", "DVT_App": "ĐVT trên App",
            "BG_LC_tuan_truoc": "Giá đối thủ TUẦN TRƯỚC",
            "BG_LC_hien_tai": "Giá đối thủ HIỆN TẠI",
            "Gia_von_tuan_truoc": "Giá vốn có VAT TRƯỚC",
            "Gia_von_VAT": "Giá vốn có VAT SAU",
            "Gia_ban_truoc_do": "Giá bán hiện tại Medigo",
        }).to_excel(writer, sheet_name="Bang so sanh 4 gia", index=False)

        # Sheet 5: Nghi lệch ĐVT (QA)
        cmp_df[cmp_df["Nghi_lech_DVT"]][cols_export_template].rename(columns=rename_export).to_excel(
            writer, sheet_name="Nghi lech DVT (QA)", index=False
        )

        # Sheet 6: Mất giá tuần này (QA)
        cmp_df[cmp_df["Mat_gia_tuan_nay"]][cols_export_template].rename(columns=rename_export).to_excel(
            writer, sheet_name="Mat gia tuan nay (QA)", index=False
        )

        # Sheet 7: Toàn bộ danh mục
        cmp_df[cols_export_template].rename(columns=rename_export).to_excel(
            writer, sheet_name="Toan bo tuan nay", index=False
        )

    print("\n" + "=" * 125)
    print(" DANH SÁCH 10 MÃ TIÊU BIỂU ĐỐI THỦ THAY ĐỔI GIÁ (KÈM ID APP, TÊN APP, ĐVT APP)")
    print("=" * 125)
    print(f"{'Mã Kiot':<11} | {'ĐVT':<5} | {'ID App':<25} | {'ĐVT App':<8} | {'BG LC Cũ':<10} | {'BG LC Mới':<10} | {'Giá Vốn':<10} | {'Lý do điều chỉnh'}")
    print("-" * 125)
    for _, r in df_doi_thu_doi.head(10).iterrows():
        id_app_str = str(r['ID_App']) if pd.notna(r['ID_App']) else "-"
        dvt_app_str = str(r['DVT_App']) if pd.notna(r['DVT_App']) else "-"
        bg_cu = f"{r['BG_LC_tuan_truoc']:,.0f}" if pd.notna(r['BG_LC_tuan_truoc']) else "-"
        bg_moi = f"{r['BG_LC_hien_tai']:,.0f}" if pd.notna(r['BG_LC_hien_tai']) else "-"
        gv = f"{r['Gia_von_VAT']:,.0f}" if pd.notna(r['Gia_von_VAT']) else "-"
        print(f"{r['Ma_Hang']:<11} | {r['DVT']:<5} | {id_app_str:<25} | {dvt_app_str:<8} | {bg_cu:<10} | {bg_moi:<10} | {gv:<10} | {r['Ly_do_dieu_chinh']}")

    print("=" * 125)
    print(f"• File báo cáo hoàn chỉnh:                             {output_filename}")
    print(f"• Số mã Đối thủ thay đổi giá THẬT (Sheet 1 Trình Sếp): {len(df_doi_thu_doi):,} mã")
    print(f"• Số mã Giá vốn nhập hàng thay đổi (Sheet 2):          {len(df_gia_von_doi):,} mã")
    print(f"• Tổng hợp mã có biến động thực tế (Sheet 3):          {len(df_tong_hop):,} mã")
    print("=" * 125)

    return output_filename


if __name__ == "__main__":
    run_pricing_pipeline()
