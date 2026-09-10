import openpyxl
from collections import defaultdict
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import datetime
import os
import sys
import io
import json
import argparse
import calendar
import re
from kpi_styling import apply_full_kpi_styling

if sys.stdout.encoding != 'utf-8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PLANS_DIR = os.path.join(BASE_DIR, 'plans')

def to_float(val, default=0.0):
    if val is None:
        return default
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, str):
        val = val.strip().replace(',', '')
        if not val:
            return default
        try:
            return float(val)
        except Exception:
            return default
    return default

def to_int(val, default=0):
    if val is None:
        return default
    if isinstance(val, int):
        return val
    if isinstance(val, float):
        return int(val)
    if isinstance(val, str):
        val = val.strip().replace(',', '')
        if not val:
            return default
        try:
            return int(float(val))
        except Exception:
            return default
    return default

def get_col_map(header_row):
    col_map = {}
    for idx, h in enumerate(header_row):
        if h is not None:
            clean_h = str(h).strip().lower()
            col_map[clean_h] = idx
    return col_map

def find_col(col_map, candidates):
    for c in candidates:
        if c in col_map:
            return col_map[c]
    return None

def clean_branch_name(cn):
    if not cn:
        return ''
    cn = str(cn).strip()
    mapping = {
        'NT Hàng Bông 247': 'Hàng Bông',
        'NT Đường Láng 247': 'Đường Láng',
        'NT 24H Minh Châu 1': 'Minh Châu',
        'NT 24h Đỗ Quang Đẩu': 'Đỗ Quang Đẩu',
        'NT 24H Lê Bình': 'Lê Bình',
        'NT 24H Nam Hòa': 'Nam Hòa',
        'NT 24H Nguyễn Chí Thanh': 'Nguyễn Chí Thanh',
        'NT 24H Nguyễn Thị Thập': 'Nguyễn Thị Thập',
        'NT 24H Nguyễn Văn Quá': 'Nguyễn Văn Quá',
        'NT 24H Rạch Bùng Binh': 'Rạch Bùng Binh',
        'NT 24H Trường Sa': 'Trường Sa'
    }
    if cn in mapping:
        return mapping[cn]
    cleaned = re.sub(r'^NT\s+(24H\s+|24h\s+)?', '', cn, flags=re.IGNORECASE)
    cleaned = re.sub(r'\s+247$', '', cleaned)
    cleaned = re.sub(r'\s+1$', '', cleaned)
    return cleaned.strip()

# ==============================================================================
# 1. MONTHLY PLAN PACKAGE LOADER (TẦNG KẾ HOẠCH ĐẦU VÀO)
# ==============================================================================
class MonthlyPlan:
    def __init__(self, plan_file_or_period):
        if os.path.exists(plan_file_or_period):
            self.plan_path = plan_file_or_period
        else:
            # period like '2026-08' or '8'
            p_str = plan_file_or_period if '-' in plan_file_or_period else f"2026-{int(plan_file_or_period):02d}"
            self.plan_path = os.path.join(PLANS_DIR, f"KeHoachKPI_{p_str}.xlsx")
            
        if not os.path.exists(self.plan_path):
            raise FileNotFoundError(f"Không tìm thấy gói kế hoạch tháng tại: {self.plan_path}")
            
        print(f"--> [MonthlyPlan] Nạp gói kế hoạch từ: {self.plan_path}")
        self.wb = openpyxl.load_workbook(self.plan_path, data_only=True)
        self._load_staff()
        self._load_stores()
        self._load_roles()
        self._load_catalog()
        self._load_rules()
        self._load_special()

    def _load_staff(self):
        self.staff_by_key = {}  # (branch, name) -> dict
        if 'Nhân sự' in self.wb.sheetnames:
            ws = self.wb['Nhân sự']
            for r in range(2, ws.max_row + 1):
                s_id = ws.cell(r, 1).value
                s_name = ws.cell(r, 2).value
                role = ws.cell(r, 3).value or 'NV'
                st_id = ws.cell(r, 4).value
                branch = clean_branch_name(ws.cell(r, 5).value)
                region = ws.cell(r, 6).value or ('HN' if branch in ['Hàng Bông', 'Đường Láng'] else 'HCM')
                if s_name and branch:
                    self.staff_by_key[(branch, str(s_name).strip())] = {
                        'staff_id': s_id,
                        'name': str(s_name).strip(),
                        'role': str(role).strip(),
                        'store_id': st_id,
                        'branch': branch,
                        'region': str(region).strip()
                    }

    def _load_stores(self):
        self.stores = {} # branch -> dict
        if 'KPI nhà thuốc' in self.wb.sheetnames:
            ws = self.wb['KPI nhà thuốc']
            for r in range(2, ws.max_row + 1):
                st_id = ws.cell(r, 1).value
                branch = clean_branch_name(ws.cell(r, 2).value)
                region = ws.cell(r, 3).value
                kpi_rev = to_float(ws.cell(r, 4).value)
                cht = ws.cell(r, 5).value
                m1 = to_float(ws.cell(r, 6).value)
                m2 = to_float(ws.cell(r, 7).value)
                m3 = to_float(ws.cell(r, 8).value)
                if branch:
                    self.stores[branch] = {
                        'store_id': st_id,
                        'branch': branch,
                        'region': region,
                        'kpi_rev': kpi_rev,
                        'cht_name': cht,
                        'm1': m1, 'm2': m2, 'm3': m3
                    }

    def _load_roles(self):
        self.roles = {} # role -> dict
        if 'KPI nhân viên' in self.wb.sheetnames:
            ws = self.wb['KPI nhân viên']
            for r in range(2, ws.max_row + 1):
                cd = ws.cell(r, 2).value
                kpi_rev = to_float(ws.cell(r, 3).value)
                kpi_bill = to_float(ws.cell(r, 4).value)
                if cd:
                    self.roles[str(cd).strip()] = {
                        'kpi_rev': kpi_rev,
                        'kpi_tb_bill': kpi_bill,
                        'b1': to_float(ws.cell(r, 5).value),
                        'b2': to_float(ws.cell(r, 6).value),
                        'b3': to_float(ws.cell(r, 7).value)
                    }

    def _load_catalog(self):
        self.sku_catalog = {} # sku -> group
        self.name_catalog = {} # name -> group
        if 'Danh mục dự án' in self.wb.sheetnames:
            ws = self.wb['Danh mục dự án']
            for r in range(2, ws.max_row + 1):
                sku = str(ws.cell(r, 1).value or '').strip()
                name = str(ws.cell(r, 2).value or '').strip()
                grp = str(ws.cell(r, 3).value or '').strip()
                if sku and grp:
                    self.sku_catalog[sku] = grp
                if name and grp:
                    self.name_catalog[name] = grp

    def _load_rules(self):
        self.tiers_hn = []
        self.tiers_hcm = []
        self.kk_hn = {'min_daily': 2000000, 'bonus': 500000}
        self.kk_hcm = {'min_daily': 950000, 'bonus': 500000}
        
        if 'Quy tắc thưởng' in self.wb.sheetnames:
            ws = self.wb['Quy tắc thưởng']
            for r in range(2, ws.max_row + 1):
                region = str(ws.cell(r, 2).value or '').strip()
                min_ck = to_float(ws.cell(r, 4).value)
                min_cb = to_float(ws.cell(r, 5).value)
                bonus = to_float(ws.cell(r, 6).value)
                kk_min = to_float(ws.cell(r, 7).value)
                kk_bon = to_float(ws.cell(r, 8).value)
                
                tier_info = {'min_ck': min_ck, 'min_combo': min_cb, 'bonus': bonus}
                if region == 'HN':
                    self.tiers_hn.append(tier_info)
                    if kk_min > 0: self.kk_hn = {'min_daily': kk_min, 'bonus': kk_bon}
                elif region == 'HCM':
                    self.tiers_hcm.append(tier_info)
                    if kk_min > 0: self.kk_hcm = {'min_daily': kk_min, 'bonus': kk_bon}

    def _load_special(self):
        self.hot_bill_cfg = {'enabled': True, 'min_bill_val': 1000000, 'bonus_per_bill': 50000}
        if 'Chương trình đặc biệt' in self.wb.sheetnames:
            ws = self.wb['Chương trình đặc biệt']
            for r in range(2, ws.max_row + 1):
                prog = str(ws.cell(r, 1).value or '').strip()
                if 'Hot Bill' in prog:
                    self.hot_bill_cfg['enabled'] = True
                    self.hot_bill_cfg['bonus_per_bill'] = to_float(ws.cell(r, 4).value, 50000)

    def classify_item(self, sku, name):
        """
        SKU-first project item classification with intelligent prefix fallback.
        """
        if sku in self.sku_catalog:
            return self.sku_catalog[sku]
        if name in self.name_catalog:
            return self.name_catalog[name]
        
        # Fallback by prefix
        p = name.split()[0].upper() if name else ''
        if p.startswith('NY3'):
            return 'NY3'
        if p.startswith('CK') or p.startswith('CKHN'):
            return 'CK'
        if p.startswith('COMBO') or p.startswith('LIỀU') or p.startswith('LIEU'):
            return 'Combo'
        return None

# ==============================================================================
# 2. PRE-FLIGHT VALIDATOR (TẦNG ĐỐI SOÁT & KIỂM TRA CHẤT LƯỢNG ĐẦU VÀO)
# ==============================================================================
class PreFlightValidator:
    def __init__(self, plan: MonthlyPlan):
        self.plan = plan
        self.errors = []
        self.warnings = []

    def validate_schema(self, col_map, required_cols_dict, file_name):
        missing = []
        found_map = {}
        for std_name, candidates in required_cols_dict.items():
            idx = find_col(col_map, candidates)
            if idx is None:
                missing.append(std_name)
            else:
                found_map[std_name] = idx
        if missing:
            err = f"Lỗi Schema [{file_name}]: Thiếu các cột bắt buộc: {', '.join(missing)}"
            self.errors.append(err)
            raise ValueError(err)
        return found_map

    def check_staff(self, invoice_staff_set):
        for branch, seller in invoice_staff_set:
            if (branch, seller) not in self.plan.staff_by_key:
                self.warnings.append(f"Nhân sự mới chưa khai báo trong Gói Kế Hoạch: '{seller}' tại Chi nhánh '{branch}' -> Tool sẽ tự động tích hợp.")

    def check_math_reconciliation(self, raw_net_rev, total_returns, calculated_net_rev):
        expected_net = raw_net_rev - total_returns
        diff = abs(expected_net - calculated_net_rev)
        if diff > 1.0:
            err = f"Lỗi đối soát doanh thu: Doanh thu kỳ vọng (Thô {raw_net_rev:,.0f}đ - Trả hàng {total_returns:,.0f}đ = {expected_net:,.0f}đ) != Doanh thu tính toán ({calculated_net_rev:,.0f}đ), lệch {diff:,.0f}đ"
            self.errors.append(err)
            raise ValueError(err)
        print(f"✅ [PreFlight] Đối soát doanh thu toán học (Đã trừ trả hàng {total_returns:,.0f}đ): Khớp 100% ({calculated_net_rev:,.0f} VNĐ)")

    def print_report(self):
        if self.warnings:
            print("\n⚠️ [PreFlight Cảnh Báo]:")
            for w in self.warnings[:5]:
                print(f"   - {w}")
            if len(self.warnings) > 5:
                print(f"   - ... và {len(self.warnings)-5} cảnh báo nữa.")
        if self.errors:
            print("\n❌ [PreFlight Lỗi Nghiêm Trọng]:")
            for e in self.errors:
                print(f"   - {e}")

# ==============================================================================
# 3. GENERIC CALCULATION & PRESENTATION ENGINE (BỘ MÁY TÍNH TOÁN & XUẤT BÁO CÁO)
# ==============================================================================
def parse_flexible_dt(dt_val):
    if not dt_val:
        return None
    if isinstance(dt_val, datetime.datetime):
        return dt_val
    if isinstance(dt_val, datetime.date):
        return datetime.datetime(dt_val.year, dt_val.month, dt_val.day)
    s = str(dt_val).strip()
    if not s or s.lower() in ('none', 'nan', 'null', ''):
        return None
    for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%d/%m/%Y %H:%M:%S', '%d/%m/%Y', '%Y/%m/%d %H:%M:%S', '%Y/%m/%d'):
        try:
            return datetime.datetime.strptime(s[:len(fmt)], fmt)
        except Exception:
            continue
    return None

def detect_data_month_from_invoices(hoadon_file):
    """
    Tự động đọc mẫu hóa đơn đầu vào để xác định chính xác tháng giao dịch thực tế.
    Đảm bảo nếu đưa data Tháng 9 vào dù chọn chương trình Tháng 8 thì vẫn tính toán chuẩn cho Tháng 9.
    """
    try:
        wb = openpyxl.load_workbook(hoadon_file, read_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        hd_map = get_col_map(header)
        c_dt = find_col(hd_map, ['thời gian', 'thoigian', 'thời gian tạo', 'ngày bán', 'ngayban', 'thời gian xuất', 'ngày tạo'])
        if c_dt is None:
            wb.close()
            return None
        month_counts = defaultdict(int)
        for i, r in enumerate(rows):
            if i > 500:
                break
            if len(r) > c_dt and r[c_dt]:
                dt = parse_flexible_dt(r[c_dt])
                if dt:
                    month_counts[(dt.year, dt.month)] += 1
        wb.close()
        if month_counts:
            top_ym = max(month_counts.items(), key=lambda x: x[1])[0]
            return top_ym
    except Exception:
        pass
    return None

def execute_kpi_engine(hoadon_file, trahang_file, template_file, plan_file, output_file, report_date_str=None, target_month=None, target_year=None):
    # 1. Load Plan
    plan = MonthlyPlan(plan_file)
    validator = PreFlightValidator(plan)

    # 2. Determine Reporting Date Range (Auto-detect from invoices data)
    inv_ym = detect_data_month_from_invoices(hoadon_file)
    if inv_ym:
        inv_year, inv_month = inv_ym
        if target_month is not None and target_month != inv_month:
            print(f"--> [Data Month Sync] Hóa đơn nạp vào là dữ liệu Tháng {inv_month}/{inv_year} (thay vì Tháng {target_month}). Tự động đồng bộ báo cáo KPI sang Tháng {inv_month}!")
        target_year = inv_year
        target_month = inv_month
    elif target_month is not None:
        target_month = int(target_month)
        target_year = int(target_year) if target_year else 2026
    elif report_date_str:
        report_dt = datetime.datetime.strptime(report_date_str, '%Y-%m-%d')
        target_year = report_dt.year
        target_month = report_dt.month
    else:
        # Infer from plan filename e.g. KeHoachKPI_2026-08.xlsx
        m_match = re.search(r'(\d{4})-(\d{2})', str(plan.plan_path))
        if m_match:
            target_year, target_month = int(m_match.group(1)), int(m_match.group(2))
        else:
            target_month = 9
            target_year = 2026

    _, last_day = calendar.monthrange(target_year, target_month)
    report_dt = datetime.datetime(target_year, target_month, last_day, 23, 59, 59)
    month_start = datetime.datetime(target_year, target_month, 1, 0, 0, 0)
    
    if report_date_str:
        clean_date = str(report_date_str).strip().split()[0].split('T')[0]
        try:
            parsed_cutoff = datetime.datetime.strptime(clean_date, '%Y-%m-%d').replace(hour=23, minute=59, second=59)
            if parsed_cutoff.month == target_month and parsed_cutoff.year == target_year:
                report_cutoff_dt = parsed_cutoff
            else:
                report_cutoff_dt = datetime.datetime(target_year, target_month, last_day, 23, 59, 59)
        except Exception:
            report_cutoff_dt = datetime.datetime(target_year, target_month, last_day, 23, 59, 59)
    else:
        report_cutoff_dt = datetime.datetime(target_year, target_month, last_day, 23, 59, 59)

    print(f"--> [Engine] Xử lý dữ liệu Tháng {target_month}/{target_year}: {month_start} đến {report_cutoff_dt}")

    # 3. Read Invoices
    print(f"--> Đọc dữ liệu hóa đơn: {hoadon_file}")
    wb_hd = openpyxl.load_workbook(hoadon_file, read_only=True)
    ws_hd = wb_hd.active
    rows_iter = ws_hd.iter_rows(values_only=True)
    hd_header = next(rows_iter)
    hd_map = get_col_map(hd_header)

    req_hd_cols = {
        'chi nhánh': ['chi nhánh', 'chinhanh'],
        'mã hóa đơn': ['mã hóa đơn', 'mahoadon'],
        'thời gian': ['thời gian', 'thoigian', 'thời gian tạo'],
        'người bán': ['người bán', 'nguoiban'],
        'khách cần trả': ['khách cần trả', 'khachcantra'],
        'tên hàng': ['tên hàng', 'tenhang']
    }
    cols = validator.validate_schema(hd_map, req_hd_cols, os.path.basename(hoadon_file))

    c_cn = cols['chi nhánh']
    c_ma_hd = cols['mã hóa đơn']
    c_dt = cols['thời gian']
    c_seller = cols['người bán']
    c_can_tra = cols['khách cần trả']
    c_ten_hang = cols['tên hàng']

    c_channel = find_col(hd_map, ['kênh bán', 'kenhban'])
    c_tot_tien = find_col(hd_map, ['tổng tiền hàng', 'tongtienhang'])
    c_giam_gia = find_col(hd_map, ['giảm giá hóa đơn', 'giamgiahoadon', 'giảm giá'])
    c_thu_khac = find_col(hd_map, ['thu khác', 'thukhac'])
    c_ma_hang = find_col(hd_map, ['mã hàng', 'mahang'])
    c_so_luong = find_col(hd_map, ['số lượng', 'soluong'])
    c_thanh_tien = find_col(hd_map, ['thành tiền', 'thanhtien'])

    invoices = {}
    invoice_staff_set = set()
    raw_total_net = 0.0

    for r in rows_iter:
        if not r or len(r) <= max(c_cn, c_ma_hd, c_seller, c_can_tra):
            continue
        ma_hd = r[c_ma_hd]
        seller = str(r[c_seller]).strip() if r[c_seller] is not None else ''
        if not ma_hd or not seller or str(ma_hd).startswith('Mã hóa đơn') or str(r[c_cn]).startswith('Chi nhánh'):
            continue

        dt_val = r[c_dt] if c_dt is not None else None
        if dt_val:
            parsed_dt = parse_flexible_dt(dt_val)
            if parsed_dt:
                dt_val = parsed_dt
                if parsed_dt < month_start or parsed_dt > report_cutoff_dt:
                    continue
            else:
                continue
        else:
            continue

        branch = clean_branch_name(r[c_cn])
        raw_branch = str(r[c_cn]).strip()
        channel = str(r[c_channel]).strip() if c_channel is not None and r[c_channel] else 'Bán trực tiếp'
        can_tra = to_float(r[c_can_tra])
        thu_khac = to_float(r[c_thu_khac]) if c_thu_khac is not None else 0.0
        tot_tien = to_float(r[c_tot_tien]) if c_tot_tien is not None else 0.0
        gg_hd = to_float(r[c_giam_gia]) if c_giam_gia is not None else 0.0
        thanh_tien = to_float(r[c_thanh_tien]) if c_thanh_tien is not None else 0.0
        ten_hang = str(r[c_ten_hang]).strip() if c_ten_hang is not None and r[c_ten_hang] is not None else ''
        ma_hang = str(r[c_ma_hang]).strip() if c_ma_hang is not None and r[c_ma_hang] is not None else ''
        so_luong = to_float(r[c_so_luong]) if c_so_luong is not None else 1.0

        invoice_staff_set.add((branch, seller))

        if ma_hd not in invoices:
            net_inv = can_tra - thu_khac
            raw_total_net += net_inv
            invoices[ma_hd] = {
                'branch': branch,
                'raw_branch': raw_branch,
                'seller': seller,
                'channel': channel,
                'time': dt_val,
                'can_tra': can_tra,
                'thu_khac': thu_khac,
                'tong_tien': tot_tien,
                'giam_gia_hd': gg_hd,
                'items': []
            }

        if ten_hang:
            invoices[ma_hd]['items'].append({
                'ma_hang': ma_hang,
                'ten_hang': ten_hang,
                'so_luong': so_luong,
                'thanh_tien': thanh_tien
            })

    try:
        wb_hd.close()
    except Exception:
        pass

    validator.check_staff(invoice_staff_set)

    # 4. Read Returns
    returns = []
    returns_by_bill = {} # ma_tra -> {branch, seller, channel, bill_val, items: list}
    if trahang_file and os.path.exists(trahang_file):
        print(f"--> Đọc dữ liệu trả hàng: {trahang_file}")
        try:
            wb_th = openpyxl.load_workbook(trahang_file, read_only=True)
            ws_th = wb_th.active
            th_iter = ws_th.iter_rows(values_only=True)
            th_header = next(th_iter)
            th_map = get_col_map(th_header)

            th_ma_tra = find_col(th_map, ['mã trả hàng', 'matrahang', 'mã phiếu trả', 'mã đổi trả'])
            th_seller = find_col(th_map, ['người bán', 'nguoiban', 'người trả hàng'])
            th_ten_hang = find_col(th_map, ['tên hàng', 'tenhang'])
            th_ma_hang = find_col(th_map, ['mã hàng', 'mahang'])
            th_tong_sau_gg = find_col(th_map, ['tổng sau giảm giá', 'tongsaugiamgia', 'tổng sau gg', 'thành tiền', 'tổng tiền hàng trả'])
            th_dt_idx = find_col(th_map, ['thời gian', 'thoigian', 'ngày tạo'])
            th_cn_idx = find_col(th_map, ['chi nhánh', 'chinhanh'])
            th_channel_idx = find_col(th_map, ['kênh bán', 'kenhban', 'loại đơn', 'kenh_ban'])
            th_qty_idx = find_col(th_map, ['số lượng', 'soluong', 'sl'])
            th_price_idx = find_col(th_map, ['giá bán', 'giaban', 'đơn giá', 'giá nhập lại'])
            th_disc_idx = find_col(th_map, ['giảm giá', 'giamgia'])

            for r_idx, r in enumerate(th_iter, start=2):
                if not r or len(r) <= max(th_seller or 0, th_ten_hang or 0, th_tong_sau_gg or 0):
                    continue
                if th_dt_idx is not None and r[th_dt_idx]:
                    t_val = r[th_dt_idx]
                    p_dt = parse_flexible_dt(t_val)
                    if p_dt and (p_dt < month_start or p_dt > report_cutoff_dt):
                        continue
                        
                s = str(r[th_seller]).strip() if th_seller is not None and r[th_seller] is not None else ''
                th_name = str(r[th_ten_hang]).strip() if th_ten_hang is not None and r[th_ten_hang] is not None else ''
                th_sku = str(r[th_ma_hang]).strip() if th_ma_hang is not None and r[th_ma_hang] is not None else ''
                th_cn = clean_branch_name(r[th_cn_idx]) if th_cn_idx is not None and r[th_cn_idx] is not None else ''
                th_ch = str(r[th_channel_idx]).strip() if th_channel_idx is not None and r[th_channel_idx] is not None else ''
                bill_val = to_float(r[th_tong_sau_gg]) if th_tong_sau_gg is not None else 0.0
                
                ma_code = str(r[th_ma_tra]).strip() if th_ma_tra is not None and r[th_ma_tra] is not None else f"RET_{r_idx}"
                if not ma_code:
                    ma_code = f"RET_{r_idx}"

                qty = to_float(r[th_qty_idx]) if th_qty_idx is not None else 1.0
                price = to_float(r[th_price_idx]) if th_price_idx is not None else 0.0
                disc = to_float(r[th_disc_idx]) if th_disc_idx is not None else 0.0
                raw_item_val = qty * (price - disc) if price > 0 else 0.0

                if s:
                    if ma_code not in returns_by_bill:
                        returns_by_bill[ma_code] = {
                            'branch': th_cn,
                            'seller': s,
                            'channel': th_ch,
                            'bill_val': bill_val,
                            'items': []
                        }
                    returns_by_bill[ma_code]['items'].append({
                        'ma_hang': th_sku,
                        'ten_hang': th_name,
                        'raw_val': raw_item_val
                    })
            wb_th.close()
        except Exception as e:
            print(f"Warning loading returns: {e}")

    # 5. Core Aggregations with Composite Key: (branch, seller)
    data_summary = {}
    project_summary = {}
    seller_branches = {}
    hot_bills = []

    for ma_hd, inv in invoices.items():
        b = inv['branch']
        s = inv['seller']
        ch = inv['channel']
        rev = to_float(inv['can_tra']) - to_float(inv['thu_khac'])

        seller_branches.setdefault(s, set()).add(b)
        key = (b, s)

        if key not in data_summary:
            data_summary[key] = {'off_tx': 0, 'off_rev': 0.0, 'onl_tx': 0, 'onl_rev': 0.0}
        if key not in project_summary:
            project_summary[key] = {'NY3': 0.0, 'CK': 0.0, 'Combo': 0.0, 'Khác': 0.0}

        if ch == 'Medigo':
            data_summary[key]['onl_tx'] += 1
            data_summary[key]['onl_rev'] += rev
        else:
            data_summary[key]['off_tx'] += 1
            data_summary[key]['off_rev'] += rev

        tot_tien = to_float(inv['tong_tien'])
        gg_hd = to_float(inv['giam_gia_hd'])

        bill_ck_qty = 0
        bill_ck_rev = 0.0
        bill_ck_items = []

        for itm in inv['items']:
            th = itm['ten_hang']
            sku = itm['ma_hang']
            tt = to_float(itm['thanh_tien'])
            qty = itm['so_luong']
            alloc_disc = (gg_hd * tt / tot_tien) if tot_tien > 0 else 0.0
            item_net = tt - alloc_disc

            # Plan-driven classification
            grp = plan.classify_item(sku, th)
            if grp:
                if grp not in project_summary[key]:
                    project_summary[key][grp] = 0.0
                project_summary[key][grp] += item_net

            if grp == 'CK':
                bill_ck_qty += qty
                bill_ck_rev += tt
                bill_ck_items.append(f"{th} (SL: {int(qty) if qty.is_integer() else qty}, TT: {tt:,.0f}đ)")

        # Hot Bill Hanoi
        if plan.hot_bill_cfg.get("enabled", False) and b in ['Hàng Bông', 'Đường Láng']:
            min_val = plan.hot_bill_cfg.get("min_bill_val", 1000000)
            if bill_ck_rev >= min_val:
                bonus = plan.hot_bill_cfg.get("bonus_per_bill", 50000)
                hot_bills.append({
                    'ma_hd': ma_hd,
                    'time': inv['time'],
                    'raw_branch': inv['raw_branch'],
                    'branch': b,
                    'seller': s,
                    'ck_qty': bill_ck_qty,
                    'ck_rev': bill_ck_rev,
                    'bonus': bonus,
                    'details': ' | '.join(bill_ck_items)
                })

    # Apply Returns to Project Summary AND Data Summary (Unique Bill & Proportional Item)
    total_returns_val = 0.0
    for ma_tra, binfo in returns_by_bill.items():
        s = binfo['seller']
        b = binfo['branch']
        ch = binfo.get('channel', '')
        bill_val = to_float(binfo['bill_val'])
        items = binfo.get('items', [])
        total_returns_val += bill_val
        
        target_keys = []
        if b:
            target_keys = [(b, s)]
        else:
            matching_branches = seller_branches.get(s, set())
            if len(matching_branches) == 1:
                target_keys = [(list(matching_branches)[0], s)]
            else:
                target_keys = [(br, s) for br in matching_branches]

        if not target_keys:
            target_keys = [(clean_branch_name(''), s)]

        factor = 1.0 / len(target_keys) if len(target_keys) > 1 else 1.0
        
        # Calculate item-level proportional value for project summary
        sum_raw_items = sum(it.get('raw_val', 0.0) for it in items)
        
        for k in target_keys:
            # 1. Trừ vào data_summary (Tổng doanh thu chính tính % KPI - Trừ đúng 1 lần giá trị bill)
            if k in data_summary:
                val_to_deduct = bill_val * factor
                if ch == 'Medigo':
                    data_summary[k]['onl_rev'] -= val_to_deduct
                else:
                    data_summary[k]['off_rev'] -= val_to_deduct

            # 2. Trừ vào project_summary (NY3 / CK / Combo - Trừ theo giá trị từng mặt hàng)
            if k in project_summary and items:
                for it in items:
                    sku = it.get('ma_hang', '')
                    th = it.get('ten_hang', '')
                    raw_val = it.get('raw_val', 0.0)
                    
                    if sum_raw_items > 0:
                        item_net_val = (raw_val / sum_raw_items) * bill_val
                    else:
                        item_net_val = bill_val / len(items)
                        
                    grp = plan.classify_item(sku, th)
                    if grp and grp in project_summary[k]:
                        project_summary[k][grp] -= item_net_val * factor

    # Reconcile Total Math (raw_net_rev - total_returns == calc_total_net)
    calc_total_net = sum(d['off_rev'] + d['onl_rev'] for d in data_summary.values())
    validator.check_math_reconciliation(raw_total_net, total_returns_val, calc_total_net)
    validator.print_report()

    # 6. Presentation Layer: Output to Excel
    print(f"--> Nạp template mẫu: {template_file}")
    wb_out = openpyxl.load_workbook(template_file)
    proj_sheet_name = f'Dự án T{target_month}'

    # 6.1 Rename Project Sheet
    for s_name in list(wb_out.sheetnames):
        if s_name.startswith('Dự án') or s_name.startswith('Du an'):
            if s_name != proj_sheet_name:
                wb_out[s_name].title = proj_sheet_name
            break

    # 6.2 Populate Sheet 'data'
    ws_data = wb_out['data']
    ws_data.cell(1, 10, report_dt)

    for r in range(2, ws_data.max_row + 1):
        for c in range(1, 10):
            ws_data.cell(r, c).value = None
            ws_data.cell(r, c).fill = PatternFill(fill_type=None)

    rotating_sellers = {s for s, b_set in seller_branches.items() if len(b_set) > 1}
    pink_fill = PatternFill(start_color='F4CCCC', end_color='F4CCCC', fill_type='solid')

    row_idx = 2
    for (b, s), d in sorted(data_summary.items(), key=lambda x: (x[0][0], x[0][1])):
        off_tx = d['off_tx']
        off_rev = d['off_rev']
        onl_tx = d['onl_tx']
        onl_rev = d['onl_rev']
        tot_tx = off_tx + onl_tx
        tot_rev = off_rev + onl_rev

        ws_data.cell(row_idx, 1, b)
        ws_data.cell(row_idx, 2, s)
        ws_data.cell(row_idx, 3, off_tx if off_tx > 0 else None)
        ws_data.cell(row_idx, 4, off_rev if off_rev > 0 else None)
        ws_data.cell(row_idx, 5, onl_tx if onl_tx > 0 else None)
        ws_data.cell(row_idx, 6, onl_rev if onl_rev > 0 else None)
        ws_data.cell(row_idx, 7, tot_tx if tot_tx > 0 else None)
        ws_data.cell(row_idx, 8, tot_rev if tot_rev > 0 else None)
        
        if s in rotating_sellers:
            ws_data.cell(row_idx, 2).fill = pink_fill
        row_idx += 1

    # 6.3 Populate Sheet 'Hot Bill HN'
    is_hot_bill = plan.hot_bill_cfg.get("enabled", False)
    if is_hot_bill:
        if 'Hot Bill HN' in wb_out.sheetnames:
            ws_hb = wb_out['Hot Bill HN']
        else:
            ws_hb = wb_out.create_sheet('Hot Bill HN')

        hb_headers = ['Mã hóa đơn', 'Thời gian', 'Chi nhánh', 'Tên Dược Sĩ', 'Số lượng SP CK', 'Doanh thu nhóm CK+CKHN', 'Tiền Thưởng (VNĐ)', 'Chi tiết sản phẩm']
        for col_i, h in enumerate(hb_headers, 1):
            ws_hb.cell(1, col_i, h)

        for r in range(2, ws_hb.max_row + 1):
            for c in range(1, 9):
                ws_hb.cell(r, c).value = None

        hb_row = 2
        for hb in sorted(hot_bills, key=lambda x: (x['time'] if x['time'] else ''), reverse=True):
            ws_hb.cell(hb_row, 1, hb['ma_hd'])
            ws_hb.cell(hb_row, 2, hb['time'])
            ws_hb.cell(hb_row, 3, hb['branch'])
            ws_hb.cell(hb_row, 4, hb['seller'])
            ws_hb.cell(hb_row, 5, hb['ck_qty'])
            ws_hb.cell(hb_row, 6, hb['ck_rev'])
            ws_hb.cell(hb_row, 7, hb['bonus'])
            ws_hb.cell(hb_row, 8, hb['details'])
            hb_row += 1
    else:
        if 'Hot Bill HN' in wb_out.sheetnames:
            wb_out.remove(wb_out['Hot Bill HN'])

    # 6.4 Master Staff Synchronization in Sheet 'Dự án T{M}'
    ws_proj = wb_out[proj_sheet_name]
    ws_proj.cell(1, 1, report_dt)

    template_proj_map = {}
    for r in range(3, ws_proj.max_row + 1):
        b_val = ws_proj.cell(r, 1).value
        s_val = ws_proj.cell(r, 2).value
        if b_val and s_val:
            key = (clean_branch_name(b_val), str(s_val).strip())
            template_proj_map[key] = r

    # Identify new staff in data_summary not yet in template
    existing_proj_keys = set(template_proj_map.keys())
    new_staff_keys = [k for k in sorted(data_summary.keys(), key=lambda x: (x[0], x[1])) if k not in existing_proj_keys]
    
    # Auto-append new staff to ws_proj immediately after last real staff row
    ws_proj.cell(1, 1, datetime.date(target_year, target_month, last_day))
    last_real_proj = max([r for r in range(3, ws_proj.max_row + 1) if ws_proj.cell(r, 2).value and str(ws_proj.cell(r, 2).value).strip()] or [2])
    next_proj_row = last_real_proj + 1

    for (b, s) in new_staff_keys:
        template_proj_map[(b, s)] = next_proj_row
        role = plan.staff_by_key.get((b, s), {}).get('role', 'DSBC')
        ws_proj.cell(next_proj_row, 1, b)
        ws_proj.cell(next_proj_row, 2, s)
        ws_proj.cell(next_proj_row, 3, role)
        ws_proj.cell(next_proj_row, 4, f"=SUM(E{next_proj_row}:G{next_proj_row})/DAY($A$1)")
        ws_proj.cell(next_proj_row, 8, f"=F{next_proj_row}/DAY($A$1)")
        ws_proj.cell(next_proj_row, 9, f"=G{next_proj_row}/DAY($A$1)")
        ws_proj.cell(next_proj_row, 11, f'=IFERROR(IF(AND($J{next_proj_row}=0, IF(OR($A{next_proj_row}="Hàng Bông",$A{next_proj_row}="Đường Láng"),$D{next_proj_row}>=2000000,$D{next_proj_row}>=950000)), 500000, 0), 0)')
        ws_proj.cell(next_proj_row, 13, f'=IFERROR($J{next_proj_row}+$K{next_proj_row}+$L{next_proj_row}, 0)')
        next_proj_row += 1

    # Standardize Project Sheet Headers
    ws_proj.cell(2, 4, 'CK + Combo + NY3/ngày')
    ws_proj.cell(2, 8, 'Chiết khấu/ngày')
    ws_proj.cell(2, 9, 'Combo Liều/ngày')
    ws_proj.cell(2, 10, 'Thưởng dự án')
    ws_proj.cell(2, 11, 'Thưởng thêm')
    ws_proj.cell(2, 12, 'Thưởng dự án Hot Bill HN')
    ws_proj.cell(2, 13, 'Total dự án')

    # Mapping chi nhánh chính cho nhân sự xoay ca để gộp toàn bộ doanh số dự án
    STAFF_MAIN_BRANCH_MAP = {
        'Phan Công Vũ Tài': 'Minh Châu',
        'Trịnh Thị Phượng': 'Minh Châu',
        'Hồ Thị Minh Hòa': 'Lê Bình',
        'Nguyễn Trần Ngọc Phương': 'Trường Sa',
        'Phạm Nguyễn Ngọc Quý': 'Trường Sa',
        'Dương Thị Huỳnh Như': 'Nam Hòa',
        'Hoàng Lâm Gia Bảo': 'Lê Bình',
        'Nguyễn Thị Huyền Trang': 'Nam Hòa',
        'Ngô Thị Ngọc Thủy': 'Rạch Bùng Binh',
        'Cù Thị Tường Vy': 'Rạch Bùng Binh',
        'Nguyễn Trí Nghĩa': 'Nguyễn Thị Thập',
        'Võ Ngọc Giàu Sang': 'Nguyễn Văn Quá',
    }

    # Gộp 100% doanh số dự án xoay ca về chi nhánh chính
    merged_project_summary = defaultdict(lambda: defaultdict(float))
    for (b, s), p_data in project_summary.items():
        target_b = STAFF_MAIN_BRANCH_MAP.get(s, b)
        for g_k, g_v in p_data.items():
            merged_project_summary[(target_b, s)][g_k] += g_v

    # Write sales values per (branch, seller) in Project Sheet
    for (b, s), row_num in template_proj_map.items():
        ws_proj.cell(row_num, 4, f"=SUM(E{row_num}:G{row_num})/DAY($A$1)")
        
        # Nếu là dòng phụ ở chi nhánh khác của nhân sự xoay ca thì để 0 (đã gộp về chi nhánh chính)
        if s in STAFF_MAIN_BRANCH_MAP and b != STAFF_MAIN_BRANCH_MAP[s]:
            p = {'NY3': 0.0, 'CK': 0.0, 'Combo': 0.0}
        else:
            p = merged_project_summary.get((b, s), project_summary.get((b, s), {'NY3': 0.0, 'CK': 0.0, 'Combo': 0.0}))
        
        # Đổ đầy đủ doanh thu NY3 thực tế từ hóa đơn vào Cột E
        ws_proj.cell(row_num, 5, round(p.get('NY3', 0), 0))
        ws_proj.cell(row_num, 6, round(p.get('CK', 0), 0))
        ws_proj.cell(row_num, 7, round(p.get('Combo', 0), 0))
        ws_proj.cell(row_num, 8, f"=F{row_num}/DAY($A$1)")
        ws_proj.cell(row_num, 9, f"=G{row_num}/DAY($A$1)")
        
        # Apply dynamic tier formula from plan rules
        if plan.tiers_hn and plan.tiers_hcm:
            tier_formula = build_tier_formula(f'$A{row_num}', f'$H{row_num}', f'$I{row_num}', plan.tiers_hn, plan.tiers_hcm)
            ws_proj.cell(row_num, 10, tier_formula)

        ws_proj.cell(row_num, 11, f'=IFERROR(IF(AND($J{row_num}=0, IF(OR($A{row_num}="Hàng Bông",$A{row_num}="Đường Láng"),$D{row_num}>=2000000,$D{row_num}>=950000)), 500000, 0), 0)')

        if is_hot_bill:
            ws_proj.cell(row_num, 12, f"=IFERROR(SUMIFS('Hot Bill HN'!$G:$G, 'Hot Bill HN'!$C:$C, $A{row_num}, 'Hot Bill HN'!$D:$D, $B{row_num}), 0)")
        else:
            ws_proj.cell(row_num, 12, 0)

        ws_proj.cell(row_num, 13, f'=IFERROR($J{row_num}+$K{row_num}+$L{row_num}, 0)')

    # 6.5 Sheet 'kpi dược sĩ'
    if 'kpi dược sĩ' in wb_out.sheetnames:
        ws_ds = wb_out['kpi dược sĩ']
        ws_ds.cell(1, 3, datetime.date(target_year, target_month, last_day))

        template_ds_map = {}
        for r in range(3, ws_ds.max_row + 1):
            b_val = ws_ds.cell(r, 1).value
            s_val = ws_ds.cell(r, 2).value
            if b_val and s_val and str(s_val).strip():
                key = (clean_branch_name(b_val), str(s_val).strip())
                template_ds_map[key] = r

        # Auto-append new staff to ws_ds only if staff is not present anywhere in the sheet
        existing_ds_names = set(str(ws_ds.cell(r, 2).value).strip() for r in range(3, ws_ds.max_row + 1) if ws_ds.cell(r, 2).value)
        last_real_ds = max([r for r in range(3, ws_ds.max_row + 1) if ws_ds.cell(r, 2).value and str(ws_ds.cell(r, 2).value).strip()] or [2])
        next_ds_row = last_real_ds + 1
        for (b, s) in new_staff_keys:
            if s not in existing_ds_names and (b, s) not in template_ds_map:
                existing_ds_names.add(s)
                template_ds_map[(b, s)] = next_ds_row
                role = plan.staff_by_key.get((b, s), {}).get('role', 'DSBC')
                ws_ds.cell(next_ds_row, 1, b)
                ws_ds.cell(next_ds_row, 2, s)
                ws_ds.cell(next_ds_row, 3, role)
                ws_ds.cell(next_ds_row, 4, 150000)
                ws_ds.cell(next_ds_row, 5, f"=AE{next_ds_row}")
                ws_ds.cell(next_ds_row, 8, f"=Q{next_ds_row}")
                ws_ds.cell(next_ds_row, 9, f"=(AD{next_ds_row}/DAY($C$1))")
                ws_ds.cell(next_ds_row, 10, f"=I{next_ds_row}/H{next_ds_row}")
                ws_ds.cell(next_ds_row, 11, f'=_xlfn.XLOOKUP(B{next_ds_row}, \'kpi nhà thuốc\'!$A:$A, \'kpi nhà thuốc\'!$G:$G, IF(AND(OR(C{next_ds_row}="DSXC", C{next_ds_row}="DSCD", C{next_ds_row}="DSBC", C{next_ds_row}="DSTV", C{next_ds_row}="Q.CHT"), G{next_ds_row}>=F{next_ds_row}), AD{next_ds_row} * _xlfn.IFS(AND(I{next_ds_row}>=T{next_ds_row}, E{next_ds_row}>=D{next_ds_row}), 0.015, I{next_ds_row}>=T{next_ds_row}, 0.013, AND(I{next_ds_row}>=S{next_ds_row}, E{next_ds_row}>=D{next_ds_row}), 0.013, I{next_ds_row}>=S{next_ds_row}, 0.011, AND(I{next_ds_row}>=R{next_ds_row}, E{next_ds_row}>=D{next_ds_row}), 0.012, I{next_ds_row}>=R{next_ds_row}, 0.010, TRUE, 0), 0))')
                ws_ds.cell(next_ds_row, 12, f'=IF(C{next_ds_row}="DSTV","", IF(J{next_ds_row}<0.6, 0.8, 1))')
                ws_ds.cell(next_ds_row, 14, f'=N(K{next_ds_row})+M{next_ds_row}')
                ws_ds.cell(next_ds_row, 15, f'=CEILING(R{next_ds_row}, 100000)')
                ws_ds.cell(next_ds_row, 16, f'=CEILING(S{next_ds_row}, 100000)')
                ws_ds.cell(next_ds_row, 17, f'=CEILING(T{next_ds_row}, 100000)')
                ws_ds.cell(next_ds_row, 18, f'=U{next_ds_row}*$R$2')
                ws_ds.cell(next_ds_row, 19, f'=U{next_ds_row}*$S$2')
                ws_ds.cell(next_ds_row, 20, f'=U{next_ds_row}*$T$2')
                ws_ds.cell(next_ds_row, 21, f"=V{next_ds_row}/DAY($C$1)")
                ws_ds.cell(next_ds_row, 22, 270000000)
                ws_ds.cell(next_ds_row, 25, f'=IF(W{next_ds_row}>0, X{next_ds_row}/W{next_ds_row}, 0)')
                ws_ds.cell(next_ds_row, 28, f'=IF(Z{next_ds_row}>0, AA{next_ds_row}/Z{next_ds_row}, 0)')
                ws_ds.cell(next_ds_row, 29, f'=W{next_ds_row}+Z{next_ds_row}')
                ws_ds.cell(next_ds_row, 30, f'=X{next_ds_row}+AA{next_ds_row}')
                ws_ds.cell(next_ds_row, 31, f'=IF(AC{next_ds_row}>0, AD{next_ds_row}/AC{next_ds_row}, 0)')
                next_ds_row += 1

        # 6.5.1 Tính toán phân bổ chỉ tiêu hàng điểm Nhà thuốc (Cột F) cho từng nhân sự
        da_branch_tot = defaultdict(float)
        for (b_k, s_k), r_idx in template_proj_map.items():
            p_val = merged_project_summary.get((b_k, s_k), project_summary.get((b_k, s_k), {}))
            sum_val = (p_val.get('NY3', 0) + p_val.get('CK', 0) + p_val.get('Combo', 0)) / last_day
            da_branch_tot[b_k] += sum_val

        for (b, s), r in template_ds_map.items():
            # Tính chỉ tiêu Cột F phân bổ từ Kế hoạch hàng điểm nhà thuốc
            p_val = merged_project_summary.get((b, s), project_summary.get((b, s), {}))
            duan_nv = (p_val.get('NY3', 0) + p_val.get('CK', 0) + p_val.get('Combo', 0)) / last_day
            duan_nt = da_branch_tot.get(b, 0.0)
            tyle = (duan_nv / duan_nt) if duan_nt > 0 else 0.0

            store_info = plan.stores.get(b, {})
            kpi_rev = store_info.get('kpi_rev', 0.0)
            hd1 = store_info.get('m1', kpi_rev * 0.8 * 0.25)
            hd2 = store_info.get('m2', kpi_rev * 0.9 * 0.25)
            hd3 = store_info.get('m3', kpi_rev * 1.0 * 0.25)
            m2_d = (kpi_rev * 0.9 / last_day) if kpi_rev > 0 else 0.0
            m3_d = (kpi_rev * 1.0 / last_day) if kpi_rev > 0 else 0.0

            # Lấy doanh thu ngày thực tế
            dt_tot = data_summary.get((b, s), {}).get('off_rev', 0.0) + data_summary.get((b, s), {}).get('onl_rev', 0.0)
            dt_ngay = dt_tot / last_day

            if dt_ngay >= m3_d and m3_d > 0:
                hd_target = hd3
            elif dt_ngay >= m2_d and m2_d > 0:
                hd_target = hd2
            else:
                hd_target = hd1

            target_ck_f = round((hd_target / last_day) * tyle, -3)
            ws_ds.cell(r, 6, target_ck_f)

            ws_ds.cell(r, 21, f"=V{r}/DAY($C$1)")
            ws_ds.cell(r, 23, f"=IFERROR(SUMIFS(data!$C:$C, data!$A:$A, $A{r}, data!$B:$B, $B{r}), 0)")
            ws_ds.cell(r, 24, f"=IFERROR(SUMIF(data!$B:$B, $B{r}, data!$D:$D), 0)")
            ws_ds.cell(r, 26, f"=IFERROR(SUMIFS(data!$E:$E, data!$A:$A, $A{r}, data!$B:$B, $B{r}), 0)")
            ws_ds.cell(r, 27, f"=IFERROR(SUMIF(data!$B:$B, $B{r}, data!$F:$F), 0)")
            ws_ds.cell(r, 7, f"=IFERROR(SUMIF('{proj_sheet_name}'!$B:$B, $B{r}, '{proj_sheet_name}'!$D:$D), 0)")
            ws_ds.cell(r, 11, f'=_xlfn.XLOOKUP(B{r}, \'kpi nhà thuốc\'!$A:$A, \'kpi nhà thuốc\'!$G:$G, IF(AND(OR(C{r}="DSXC", C{r}="DSCD", C{r}="DSBC", C{r}="DSTV", C{r}="Q.CHT"), G{r}>=F{r}), AD{r} * _xlfn.IFS(AND(I{r}>=T{r}, E{r}>=D{r}), 0.015, I{r}>=T{r}, 0.013, AND(I{r}>=S{r}, E{r}>=D{r}), 0.013, I{r}>=S{r}, 0.011, AND(I{r}>=R{r}, E{r}>=D{r}), 0.012, I{r}>=R{r}, 0.010, TRUE, 0), 0))')
            ws_ds.cell(r, 13, f"=IFERROR(SUMIF('{proj_sheet_name}'!$B:$B, $B{r}, '{proj_sheet_name}'!$M:$M)*$L{r}, 0)")
            ws_ds.cell(r, 14, f'=N(K{r})+M{r}')

    # 6.6 Sheet 'kpi nhà thuốc'
    if 'kpi nhà thuốc' in wb_out.sheetnames:
        ws_nt = wb_out['kpi nhà thuốc']
        ws_nt.cell(1, 6, "='kpi dược sĩ'!$C$1")
        ws_nt.cell(1, 20, "='kpi dược sĩ'!$C$1")
        
        nt_row = 3
        while nt_row <= ws_nt.max_row:
            b_val = ws_nt.cell(nt_row, 2).value
            if b_val and str(b_val).strip():
                ws_nt.cell(nt_row, 21, f"=SUMIF('{proj_sheet_name}'!$A:$A, B{nt_row}, '{proj_sheet_name}'!$D:$D)*DAY($F$1)")
            nt_row += 1

    try:
        apply_full_kpi_styling(wb_out)
    except Exception as e:
        print(f"⚠️ [Styling Warning] {e}")

    try:
        wb_out.calculation.fullCalcOnLoad = True
        wb_out.calculation.calcMode = 'auto'
    except Exception:
        pass

    wb_out.save(output_file)
    wb_out.close()
    print(f"--> Báo cáo KPI xuất thành công ra: {output_file}")

    # Build rich statistics dictionary for Web App & downstream callers
    branch_stats = {}
    for b_name, b_info in plan.stores.items():
        b_rev = sum((d.get('off_rev', 0) + d.get('onl_rev', 0)) for k, d in data_summary.items() if k[0] == b_name)
        b_tx = sum((d.get('off_tx', 0) + d.get('onl_tx', 0)) for k, d in data_summary.items() if k[0] == b_name)
        b_staff = len(set(k[1] for k in data_summary.keys() if k[0] == b_name))
        mgr = b_info.get('manager', '') if isinstance(b_info, dict) else getattr(b_info, 'manager', '')
        tgt = b_info.get('target', 0.0) if isinstance(b_info, dict) else getattr(b_info, 'target', 0.0)
        branch_stats[b_name] = {
            'manager': mgr or '',
            'target': tgt,
            'rev': b_rev,
            'tx': b_tx,
            'staff_count': b_staff
        }

    # Staff list aggregation
    staff_agg = {}
    for (b_name, s_name), d in data_summary.items():
        if s_name not in staff_agg:
            staff_agg[s_name] = {
                'name': s_name,
                'branches': set(),
                'rev': 0.0,
                'tx': 0,
                'proj_rev': 0.0
            }
        staff_agg[s_name]['branches'].add(b_name)
        staff_agg[s_name]['rev'] += (d.get('off_rev', 0) + d.get('onl_rev', 0))
        staff_agg[s_name]['tx'] += (d.get('off_tx', 0) + d.get('onl_tx', 0))

    for (b_name, s_name), p in project_summary.items():
        if s_name in staff_agg:
            staff_agg[s_name]['proj_rev'] += (p.get('NY3', 0) + p.get('CK', 0) + p.get('Combo', 0))

    staff_list_sorted = []
    rotating_count = 0
    for s_name, s_data in staff_agg.items():
        is_rot = len(s_data['branches']) > 1
        if is_rot:
            rotating_count += 1
        staff_list_sorted.append({
            'name': s_name,
            'branches_str': ', '.join(sorted(list(s_data['branches']))),
            'rev': s_data['rev'],
            'tx': s_data['tx'],
            'proj_rev': s_data['proj_rev'],
            'is_rotating': is_rot
        })
    staff_list_sorted.sort(key=lambda x: x['rev'], reverse=True)

    tot_off_rev = sum(d.get('off_rev', 0) for d in data_summary.values())
    tot_onl_rev = sum(d.get('onl_rev', 0) for d in data_summary.values())
    tot_off_tx = sum(d.get('off_tx', 0) for d in data_summary.values())
    tot_onl_tx = sum(d.get('onl_tx', 0) for d in data_summary.values())
    
    tot_ny3 = sum(p.get('NY3', 0) for p in project_summary.values())
    tot_ck = sum(p.get('CK', 0) for p in project_summary.values())
    tot_combo = sum(p.get('Combo', 0) for p in project_summary.values())

    stats = {
        'total_rev': tot_off_rev + tot_onl_rev,
        'off_rev': tot_off_rev,
        'onl_rev': tot_onl_rev,
        'off_tx': tot_off_tx,
        'onl_tx': tot_onl_tx,
        'tot_ny3': tot_ny3,
        'tot_ck': tot_ck,
        'tot_combo': tot_combo,
        'invoice_count': len(invoices),
        'return_count': len(returns),
        'hot_bills_count': len(hot_bills) if is_hot_bill else 0,
        'hot_bills_total_bonus': sum(hb['bonus'] for hb in hot_bills) if is_hot_bill else 0,
        'hot_bills': hot_bills if is_hot_bill else [],
        'branches': branch_stats,
        'staff_list': staff_list_sorted,
        'rotating_staff_count': rotating_count,
        'month': target_month,
        'report_date': report_date_str or ''
    }
    return stats

def build_tier_formula(branch_cell, ck_daily_cell, combo_daily_cell, tiers_hn, tiers_hcm):
    hn_branches_cond = f'OR({branch_cell}="Hàng Bông",{branch_cell}="Đường Láng")'
    hn_parts = []
    for t in tiers_hn:
        hn_parts.append(f'AND({ck_daily_cell}>={t["min_ck"]},{combo_daily_cell}>={t["min_combo"]}),{t["bonus"]}')
    hn_parts.append('TRUE,0')
    hn_ifs = f'_xlfn.IFS({",".join(hn_parts)})'
    
    hcm_parts = []
    for t in tiers_hcm:
        hcm_parts.append(f'AND({ck_daily_cell}>={t["min_ck"]},{combo_daily_cell}>={t["min_combo"]}),{t["bonus"]}')
    hcm_parts.append('TRUE,0')
    hcm_ifs = f'_xlfn.IFS({",".join(hcm_parts)})'
    
    return f'=IFERROR(IF({hn_branches_cond},{hn_ifs},{hcm_ifs}),0)'

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Medigo Data-Driven KPI Engine")
    parser.add_argument("--plan", type=str, default="2026-08", help="Đường dẫn file kế hoạch hoặc mã kỳ (ví dụ: 2026-08)")
    parser.add_argument("--hoadon", type=str, default=None, help="Đường dẫn file hóa đơn")
    parser.add_argument("--trahang", type=str, default=None, help="Đường dẫn file trả hàng")
    parser.add_argument("--template", type=str, default=None, help="Đường dẫn file template mẫu")
    parser.add_argument("--output", type=str, default=None, help="Đường dẫn file output kết quả")
    parser.add_argument("--date", type=str, default=None, help="Ngày báo cáo (YYYY-MM-DD)")

    args = parser.parse_args()

    # Plan resolution
    plan_arg = args.plan
    if os.path.exists(plan_arg):
        plan_file = plan_arg
    else:
        p_str = plan_arg if '-' in plan_arg else f"2026-{int(plan_arg):02d}"
        plan_file = os.path.join(PLANS_DIR, f"KeHoachKPI_{p_str}.xlsx")

    # Month from plan
    m_match = re.search(r'(\d{4})-(\d{2})', plan_file)
    m = int(m_match.group(2)) if m_match else 8

    m_folder = os.path.join(BASE_DIR, f'thang{m}')
    
    hd_file = args.hoadon
    if not hd_file:
        if m == 8:
            for cf in [os.path.join(m_folder, 'DATA', 'DanhSachChiTietHoaDon_3182026.xlsx'),
                       os.path.join(m_folder, 'DanhSachChiTietHoaDon_3182026.xlsx'),
                       os.path.join(m_folder, '288', 'DanhSachChiTietHoaDon_KV02092026-113450-301.xlsx')]:
                if os.path.exists(cf): hd_file = cf; break
        elif m == 7:
            hd_file = os.path.join(m_folder, 'hoadon.xlsx')

    th_file = args.trahang
    if not th_file:
        if m == 8:
            for cth in [os.path.join(m_folder, 'DATA', 'DanhSachChiTietTraHang_3182026.xlsx'),
                        os.path.join(m_folder, 'DanhSachChiTietTraHang_3182026.xlsx'),
                        os.path.join(m_folder, '288', 'trahang288.xlsx')]:
                if os.path.exists(cth): th_file = cth; break
        elif m == 7:
            th_file = os.path.join(m_folder, 'DanhSachChiTietTraHang.xlsx')

    tmpl_file = args.template
    if not tmpl_file:
        tmpl_file = os.path.join(BASE_DIR, 'goc', f'NHÀ THUỐC THÁNG {m} 2026.xlsx')
        if not os.path.exists(tmpl_file):
            tmpl_file = os.path.join(BASE_DIR, 'goc', 'NHÀ THUỐC THÁNG 8 2026.xlsx')

    out_file = args.output
    if not out_file:
        out_file = os.path.join(BASE_DIR, f'baocaokpi_thang{m}_hoanthien.xlsx')

    print(f"=== BẮT ĐẦU CHẠY DATA-DRIVEN KPI ENGINE (KỲ {args.plan}) ===")
    execute_kpi_engine(hd_file, th_file, tmpl_file, plan_file, out_file, args.date)
    print("=== HOÀN THÀNH BÁO CÁO! ===")
