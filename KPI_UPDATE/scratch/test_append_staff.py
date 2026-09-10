import openpyxl
import subprocess
import os

wb_base = openpyxl.load_workbook('test_with_orig_ds.xlsx')
ws_ds = wb_base['kpi dược sĩ']
start_row = ws_ds.max_row + 1

# Let's get the new staff keys from baocaokpi
wb_kpi = openpyxl.load_workbook('baocaokpi_thang7_hoanthien.xlsx')
ws_kpi_ds = wb_kpi['kpi dược sĩ']

for r in range(start_row, ws_kpi_ds.max_row + 1):
    b = ws_kpi_ds.cell(r, 1).value
    s = ws_kpi_ds.cell(r, 2).value
    role = ws_kpi_ds.cell(r, 3).value
    
    # Copy all cells from row r
    for c in range(1, 35):
        ws_ds.cell(r, c, ws_kpi_ds.cell(r, c).value)
        
    tname = f'test_staff_up_to_{r}.xlsx'
    wb_base.save(tname)
    
    p = os.path.abspath(tname)
    res = subprocess.run(['powershell', '-ExecutionPolicy', 'Bypass', '-File', 'recalc_workbook.ps1', '-FilePath', p], capture_output=True, text=True)
    is_ok = 'Recalculation complete' in res.stdout
    print(f'Staff row {r} ({b} - {s}): {"SUCCESS" if is_ok else "FAILED"}')
    if not is_ok:
        print('First failed row formulas:')
        for c in range(1, 35):
            print(f'  Col {c} ({openpyxl.utils.get_column_letter(c)}): {ws_ds.cell(r, c).value}')
        break

wb_base.close()
wb_kpi.close()
