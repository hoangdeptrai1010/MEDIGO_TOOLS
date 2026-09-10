import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb_inv = openpyxl.load_workbook('thang8/DATA/DanhSachChiTietHoaDon_3182026.xlsx', data_only=True)
ws = wb_inv.active

SKUS_KAT = {'SP2725289', 'SP2725285', 'SP2725291', 'SP2725287', 'SP2725353', 'SP2725351'}
SKUS_PS = {'SP017162', 'SP2723878', 'SP2723879'}
SKUS_LC = {'SP2723124', 'SP2723125', 'SP2723127'}
SKUS_AVC = {'SP2723883', 'SP2722826', 'SP2722817', 'SP2725666'}

data_hn = {}

for r in range(2, ws.max_row + 1):
    branch = str(ws.cell(r, 1).value or '').strip()
    if 'Đường Láng' in branch or 'Hàng Bông' in branch:
        sku = str(ws.cell(r, 4).value or '').strip()
        qty = ws.cell(r, 9).value or 0
        tt = ws.cell(r, 15).value or 0
        
        b_clean = 'Đường Láng' if 'Đường Láng' in branch else 'Hàng Bông'
        if b_clean not in data_hn:
            data_hn[b_clean] = {'kat': {}, 'ps': {}, 'lc': {}, 'avc': {}}
            
        if sku in SKUS_KAT:
            data_hn[b_clean]['kat'][sku] = data_hn[b_clean]['kat'].get(sku, 0) + qty
        if sku in SKUS_PS:
            data_hn[b_clean]['ps'][sku] = data_hn[b_clean]['ps'].get(sku, 0) + qty
        if sku in SKUS_LC:
            data_hn[b_clean]['lc'][sku] = data_hn[b_clean]['lc'].get(sku, 0) + tt
        if sku in SKUS_AVC:
            data_hn[b_clean]['avc'][sku] = data_hn[b_clean]['avc'].get(sku, 0) + qty

print('=== Hanoi Sales Summary by Branch ===')
for b, d in data_hn.items():
    kat_tot = sum(d['kat'].values())
    ps_tot = sum(d['ps'].values())
    lc_tot = sum(d['lc'].values())
    avc_tot = sum(d['avc'].values())
    print(f'Branch: {b}')
    print(f"  KAT: total qty = {kat_tot}, details = {d['kat']}")
    print(f"  PS: total qty = {ps_tot}, details = {d['ps']}")
    print(f"  LadyCare: total rev = {lc_tot:,.0f}, details = {d['lc']}")
    print(f"  AVC: total qty = {avc_tot}, details = {d['avc']}")
