import openpyxl
import os
import sys
sys.stdout.reconfigure(encoding='utf-8')

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_bangluong = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final.xlsx")
f_ngoc = os.path.join(latvat_dir, "Gởi Ms. Ngoc - 07.09. Bang Luong T8.2026.xlsx")

wb_bl = openpyxl.load_workbook(f_bangluong, data_only=True)
wb_ng = openpyxl.load_workbook(f_ngoc, data_only=True)

print("=== CHECKING FORMULAS & REFERENCES IN BANGLUONG Sheet1 ===")
wb_bl_formulas = openpyxl.load_workbook(f_bangluong, data_only=False)
ws_s1_f = wb_bl_formulas["Sheet1"]
for r in range(1, ws_s1_f.max_row + 1):
    c_name = ws_s1_f.cell(r, 3).value
    c_stk = ws_s1_f.cell(r, 5).value
    c_bank = ws_s1_f.cell(r, 6).value
    if c_name is not None:
        print(f"Row {r:2d} | Name: {c_name!r:<30} | STK: {str(c_stk):<25} | Bank: {str(c_bank):<25}")

print("\n=== CHECKING FORMULAS IN BANGLUONG Med ===")
ws_med_f = wb_bl_formulas["Med"]
for r in range(9, 30):
    c_name = ws_med_f.cell(r, 4).value
    c_stk = ws_med_f.cell(r, 9).value
    c_bank = ws_med_f.cell(r, 10).value
    if c_name is not None:
        print(f"Row {r:2d} | Name: {c_name!r:<30} | STK: {str(c_stk):<25} | Bank: {str(c_bank):<25}")
