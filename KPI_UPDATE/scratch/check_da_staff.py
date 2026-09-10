import openpyxl

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws_da = wb['Dự án']

for r in range(2, ws_da.max_row + 1):
    n = ws_da.cell(r, 2).value
    if 'Thủy' in str(n) or 'Thao' in str(n) or 'Tài' in str(n):
        print(f"Row {r}: {[ws_da.cell(r, c).value for c in range(1, 13)]}")
