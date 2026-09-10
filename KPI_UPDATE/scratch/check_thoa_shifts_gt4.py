import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws_nc = wb['Ngày công']

print("=== KIỂM TRA SỐ CA ĐÊM CỦA HỨA THỊ KIM THOA VỚI ĐIỀU KIỆN > 4 TIẾNG ===")
thoa_night_gt4 = 0
thoa_night_le4 = 0
for r in range(2, ws_nc.max_row+1):
    name = str(ws_nc.cell(r, 2).value or '')
    if 'thoa' in name.lower():
        ca = str(ws_nc.cell(r, 4).value or '')
        hrs = ws_nc.cell(r, 6).value
        if 'đêm' in ca.lower():
            if isinstance(hrs, (int, float)):
                if hrs > 4:
                    thoa_night_gt4 += 1
                    print(f"Row {r:4d} | Ca đêm > 4h: {ca:<40} | Giờ: {hrs}")
                else:
                    thoa_night_le4 += 1
                    print(f"Row {r:4d} | Ca đêm <= 4h: {ca:<40} | Giờ: {hrs}")

print(f"\nTổng ca đêm > 4h của Thoa: {thoa_night_gt4}")
print(f"Tổng ca đêm <= 4h của Thoa: {thoa_night_le4}")
