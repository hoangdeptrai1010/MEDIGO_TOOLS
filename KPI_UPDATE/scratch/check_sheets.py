import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook('thang8/tinhcongnhungthuongchia.xlsx', data_only=False)
print("Sheets in tinhcongnhungthuongchia.xlsx:", wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"Sheet '{sname}': max_row={ws.max_row}, max_column={ws.max_column}")
