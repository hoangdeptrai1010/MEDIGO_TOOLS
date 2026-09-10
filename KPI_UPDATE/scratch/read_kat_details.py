import openpyxl

for p in [
    'thang8/Pharmacy_retail_Store_KPIs_August 2026/Hà Nội/Dự án Mini/Chương trình/HN_KAT Project August_2026.xlsx',
    'thang8/Pharmacy_retail_Store_KPIs_August 2026/Hồ Chí Minh/Dự án mini/Chương trình/HCM_KAT Project August_2026.xlsx',
    'thang8/Pharmacy_retail_Store_KPIs_August 2026/Hà Nội/Dự án Mini/Chương trình/HN_AVC project_August_2026.xlsx'
]:
    wb = openpyxl.load_workbook(p, data_only=True)
    print(f"=== FILE: {p} ===")
    ws = wb.active
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(x is not None for x in vals):
            print(f"  r{r}: {vals}")
