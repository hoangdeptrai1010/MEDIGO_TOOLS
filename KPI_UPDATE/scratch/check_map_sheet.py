import openpyxl

for fn in ['thang7/BẢNG LƯƠNG THÁNG 7 2026.xlsx', 'thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx']:
    wb = openpyxl.load_workbook(fn, read_only=True, data_only=True)
    if 'Map' in wb.sheetnames:
        ws = wb['Map']
        rows = [r for r in ws.iter_rows(values_only=True) if any(x is not None for x in r)]
        print(f"{fn} -> Map rows: {len(rows)}")
        for r in rows[:4]:
            print("  ", r[:6])
