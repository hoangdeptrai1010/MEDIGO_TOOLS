import openpyxl

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws = wb['BẢNG LƯƠNG']

print("=== thang8/bangluong_thang8_hoanthien.xlsx: First 15 staff ===")
print(f"{'CN':<15} | {'Họ tên':<24} | {'Chức vụ':<8} | {'Giờ ngày (E)':<12} | {'Giờ đêm (F)':<12} | {'Ngày làm (J)':<12} | {'Ngày đêm (K)':<12} | {'Ngày công (R)':<12}")
print("-" * 115)

for r in range(3, 18):
    stt = ws.cell(r, 1).value
    cn = str(ws.cell(r, 2).value or '')
    name = str(ws.cell(r, 3).value or '')
    cv = str(ws.cell(r, 4).value or '')
    gn = ws.cell(r, 5).value
    gd = ws.cell(r, 6).value
    sn = ws.cell(r, 10).value
    sd = ws.cell(r, 11).value
    nc = ws.cell(r, 18).value
    print(f"{cn:<15} | {name:<24} | {cv:<8} | {str(gn):<12} | {str(gd):<12} | {str(sn):<12} | {str(sd):<12} | {str(nc):<12}")
