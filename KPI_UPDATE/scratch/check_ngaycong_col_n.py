import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=False)
ws_nc = wb['Ngày công']

print("--- SHEET 'Ngày công' SUMMARY TABLE (Cols K - O) ---")
for r in range(1, 20):
    k_val = ws_nc.cell(r, 11).value
    l_val = ws_nc.cell(r, 12).value
    m_val = ws_nc.cell(r, 13).value
    n_val = ws_nc.cell(r, 14).value
    o_val = ws_nc.cell(r, 15).value
    print(f"Row {r:2d}: K={str(k_val):<15} | L={str(l_val):<25} | M={str(m_val):<35} | N={str(n_val):<25} | O={str(o_val)}")

print("\n--- Any formulas in Col N? ---")
for r in range(2, ws_nc.max_row + 1):
    n_val = ws_nc.cell(r, 14).value
    if n_val and str(n_val).startswith('='):
        print(f"Row {r}: N formula = {n_val}")
        break
else:
    print("Col N has NO formulas, they are hardcoded values!")
