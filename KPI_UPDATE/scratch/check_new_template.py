import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('d:/MEDIGO/KPI_UPDATE/thang9/output/NHÀ THUỐC THÁNG 9 2026_new.xlsx', data_only=False)
ws = wb['kpi dược sĩ']
print("Row 3 (Trường Sa - Phương):")
print("  Col J (% KPI):", ws.cell(3, 10).value)
print("  Col K (Thưởng KPI):", ws.cell(3, 11).value[:80] + "...")
print("  Col V (Monthly target):", ws.cell(3, 22).value)

print("\nHanoi Staff Rows:")
for r in range(3, ws.max_row + 1):
    st = ws.cell(r, 1).value
    name = ws.cell(r, 2).value
    role = ws.cell(r, 3).value
    col_j = ws.cell(r, 10).value
    col_v = ws.cell(r, 22).value
    if st in ['Hàng Bông', 'Đường Láng']:
        print(f"  Row {r:2d} | {st:10s} | {name:25s} | Role: {role:10s} | Monthly: {col_v:>11,d} đ | % KPI: {col_j}")
wb.close()
