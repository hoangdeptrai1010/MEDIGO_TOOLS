import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb8 = openpyxl.load_workbook(r'thang8\bangluong_thang8_hoanthien.xlsx', data_only=True)
ws8 = wb8['BẢNG LƯƠNG']

print("--- [BẢNG LƯƠNG THÁNG 8] Header Columns ---")
for c in range(1, ws8.max_column + 1):
    h = ws8.cell(2, c).value
    if h:
        print(f"Col {openpyxl.utils.get_column_letter(c)} ({c}): {repr(str(h).strip())}")

print("\n--- Rows 3 to 10 in BẢNG LƯƠNG T8 ---")
for r in range(3, 11):
    name = ws8.cell(r, 3).value
    vals = {openpyxl.utils.get_column_letter(c): ws8.cell(r, c).value for c in range(25, 45) if ws8.cell(r, c).value is not None}
    print(f"Row {r} ({name}): {vals}")
