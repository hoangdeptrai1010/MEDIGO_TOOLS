import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

f_gen = r"d:\MEDIGO\KPI_UPDATE\TOOL_KPISHEET\output\NHÀ THUỐC THÁNG 9 2026_new.xlsx"
wb = openpyxl.load_workbook(f_gen, data_only=False)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n=======================================================")
    print(f"VERIFYING COLORS IN: '{sname}'")
    print(f"=======================================================")
    for r in range(1, min(4, ws.max_row + 1)):
        row_fills = []
        for c in range(1, min(ws.max_column + 1, 15)):
            cell = ws.cell(r, c)
            fill_rgb = cell.fill.start_color.rgb if cell.fill and cell.fill.start_color else None
            font_bold = cell.font.bold if cell.font else False
            val = str(cell.value or '')[:15]
            row_fills.append(f"{val} ({fill_rgb})")
        print(f"Row {r:2d}: {row_fills}")
