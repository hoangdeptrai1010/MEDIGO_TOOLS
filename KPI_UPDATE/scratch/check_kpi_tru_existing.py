import openpyxl

for p in ['thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', 'thang8/bangluong_thang8_hoanthien.xlsx']:
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    if 'KPI trừ' in wb.sheetnames:
        ws = wb['KPI trừ']
        rows = [r for r in ws.iter_rows(values_only=True) if any(x is not None for x in r)]
        print(f"{p} -> KPI trừ rows: {len(rows)}")
        for r in rows[:5]:
            print("  ", r[:8])
