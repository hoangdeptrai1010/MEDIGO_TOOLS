import openpyxl
import os
import sys
import unicodedata
import re
import glob

sys.stdout.reconfigure(encoding='utf-8')

def remove_accents(input_str):
    if not input_str:
        return ""
    nfkd_form = unicodedata.normalize('NFKD', str(input_str))
    return "".join([c for c in nfkd_form if not unicodedata.combining(c)])

def normalize_name(s):
    if not s:
        return ""
    s = unicodedata.normalize('NFC', str(s))
    return " ".join(s.lower().strip().split())

target_names = [
    "Trần Thị Kim Khánh",
    "Phạm Thị Diên",
    "Nguyễn Ngọc Đài Trang",
    "Phan Công Vũ Tài",
    "Hạ Ngân Khánh",
    "Đỗ Thị Thu Hương",
    "Nguyễn Đăng An",
    "Nguyễn Minh Trí",
    "Nguyễn Trần Vĩnh Khang",
    "Nguyễn Mẫn Tiệp",
    "Trần Gia Tiến",
    "Nguyễn Thị Mai Duyên",
    "Đỗ Thu Huệ"
]

target_norms = {normalize_name(n): n for n in target_names}
target_norms_noacc = {normalize_name(remove_accents(n)): n for n in target_names}

workspace = r"d:\MEDIGO\KPI_UPDATE"
excel_files = glob.glob(os.path.join(workspace, "**/*.xlsx"), recursive=True)

results = {}

for f in excel_files:
    if "~$" in f or "scratch" in f:
        continue
    rel = os.path.relpath(f, workspace)
    try:
        wb = openpyxl.load_workbook(f, data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            for r in range(1, ws.max_row + 1):
                for c in range(1, min(20, ws.max_column + 1)):
                    val = ws.cell(r, c).value
                    if val:
                        n_norm = normalize_name(val)
                        n_noacc = normalize_name(remove_accents(val))
                        matched_target = target_norms.get(n_norm) or target_norms_noacc.get(n_noacc)
                        if matched_target:
                            row_vals = [ws.cell(r, col).value for col in range(1, min(15, ws.max_column + 1))]
                            if matched_target not in results:
                                results[matched_target] = []
                            results[matched_target].append((rel, sname, r, row_vals))
    except Exception as e:
        pass

for name, occurrences in results.items():
    print(f"\n================ Target: {name} ================")
    for rel, sname, r, row_vals in occurrences[:10]:
        print(f"[{rel} | {sname} | r{r}]: {row_vals}")
