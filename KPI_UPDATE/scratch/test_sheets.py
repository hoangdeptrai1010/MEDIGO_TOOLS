import openpyxl
import subprocess
import os

wb = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx')
for s in wb.sheetnames:
    wb_single = openpyxl.Workbook()
    ws_new = wb_single.active
    ws_new.title = s
    ws_src = wb[s]
    
    for r in range(1, ws_src.max_row+1):
        for c in range(1, ws_src.max_column+1):
            cell = ws_src.cell(r, c)
            if cell.value is not None:
                ws_new.cell(r, c, cell.value)
                
    fn = f'test_sheet_{s.replace(" ", "_")}.xlsx'
    wb_single.save(fn)
    wb_single.close()
    
    res = subprocess.run(['powershell', '-ExecutionPolicy', 'Bypass', '-File', 'recalc_workbook.ps1', '-FilePath', fn], capture_output=True, text=True)
    success = 'successfully' in res.stdout
    print(f'Sheet [{s}]: {"SUCCESS" if success else "FAILED"}')
    if os.path.exists(fn): os.remove(fn)
