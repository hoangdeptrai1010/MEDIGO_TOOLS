import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws_bl = wb['BẢNG LƯƠNG']
ws_gc = wb['Giờ công']
ws_nc = wb['Ngày công']

headers_bl = [ws_bl.cell(2, c).value for c in range(1, ws_bl.max_column+1)]

print("=== 1. KIỂM TRA TRÊN SHEET BẢNG LƯƠNG ===")
for r in range(3, ws_bl.max_row+1):
    name = str(ws_bl.cell(r, 3).value or '')
    if any(k in name.lower() for k in ['thoa', 'tâm', 'tam']):
        cn = ws_bl.cell(r, 2).value
        role = ws_bl.cell(r, 4).value
        row_vals = {f"Col {openpyxl.utils.get_column_letter(c)} ({headers_bl[c-1]})": ws_bl.cell(r, c).value for c in range(1, 43) if c-1 < len(headers_bl) and headers_bl[c-1]}
        print(f"\nRow {r}: Chi nhánh={cn}, Tên={name}, Vị trí={role}")
        for k, v in row_vals.items():
            if v is not None and v != '' and v != 0:
                print(f"  {k}: {v}")

print("\n=== 2. KIỂM TRA TRÊN SHEET GIỜ CÔNG (Bảng tổng hợp G:K) ===")
for r in range(2, ws_gc.max_row+1):
    name = str(ws_gc.cell(r, 8).value or '')
    if any(k in name.lower() for k in ['thoa', 'tâm', 'tam']):
        cn = ws_gc.cell(r, 7).value
        print(f"Row {r} | CN: {cn} | Tên: {name} | Giờ ngày: {ws_gc.cell(r, 9).value} | Giờ đêm: {ws_gc.cell(r, 10).value} | Tổng: {ws_gc.cell(r, 11).value}")

print("\n=== 3. KIỂM TRA TỪNG CA LÀM TRONG SHEET GIỜ CÔNG (Cột A:E) ===")
for r in range(2, ws_gc.max_row+1):
    name = str(ws_gc.cell(r, 2).value or '')
    if any(k in name.lower() for k in ['thoa', 'tâm', 'tam']):
        cn = ws_gc.cell(r, 1).value
        ca = ws_gc.cell(r, 3).value
        gio_txt = ws_gc.cell(r, 4).value
        gio_num = ws_gc.cell(r, 5).value
        print(f"Giờ công chi tiết Row {r:3d} | CN: {cn} | Tên: {name:<22} | Ca: {str(ca):<35} | Giờ txt: {gio_txt} | Giờ số: {gio_num}")

