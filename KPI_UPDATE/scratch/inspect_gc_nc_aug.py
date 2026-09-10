import openpyxl

wb_aug = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=True)
ws_gc = wb_aug['Giờ công']
ws_nc = wb_aug['Ngày công']

print("=== Giờ công sample rows (2-15) ===")
for r in range(2, 16):
    row_vals = [ws_gc.cell(r, c).value for c in range(1, 6)]
    print(f"R{r}: {row_vals}")

print("\n=== Ngày công sample rows (2-15) ===")
for r in range(2, 16):
    row_vals = [ws_nc.cell(r, c).value for c in range(1, 6)]
    print(f"R{r}: {row_vals}")
