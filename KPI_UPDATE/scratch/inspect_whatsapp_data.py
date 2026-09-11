import os, sys, io, glob
import openpyxl
import docx

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

# Inspect docx files
for docx_path in glob.glob('d:/MEDIGO/KPI_UPDATE/thang8/**/KPIs Whatsapp*.docx', recursive=True):
    print("=== DOCX:", docx_path)
    doc = docx.Document(docx_path)
    for p in doc.paragraphs:
        if p.text.strip():
            print("  ", p.text)
    for t in doc.tables:
        print("  --- TABLE ---")
        for row in t.rows:
            print("   ", [c.text.strip().replace('\n', ' ') for c in row.cells])

# Inspect Bang Luong T8 sheets
for bl_path in [
    'd:/MEDIGO/KPI_UPDATE/thang8/output/bangluong_thang8_hoanthien.xlsx',
    'd:/MEDIGO/KPI_UPDATE/thang8/latvat/Gởi Ms. Ngoc - 07.09. Bang Luong T8.2026.xlsx',
    'd:/MEDIGO/KPI_UPDATE/thang8/latvat/BANGLUONGT8FIXV31 final.xlsx'
]:
    if os.path.exists(bl_path):
        print("\n=== BANGLUONG:", bl_path)
        wb = openpyxl.load_workbook(bl_path, data_only=True)
        print("Sheets:", wb.sheetnames)
        for sname in wb.sheetnames:
            if 'whatsapp' in sname.lower() or 'phiếu' in sname.lower() or 'phieu' in sname.lower() or 'thưởng' in sname.lower() or 'thuong' in sname.lower() or 'tổng' in sname.lower() or 'tong' in sname.lower():
                ws = wb[sname]
                print(f"--- Sheet {sname} (rows={ws.max_row}, cols={ws.max_column}) ---")
                for r in range(1, min(10, ws.max_row+1)):
                    vals = [ws.cell(r, c).value for c in range(1, min(20, ws.max_column+1))]
                    if any(vals):
                        print(f"  Row {r}:", [v for v in vals if v is not None])
        wb.close()
