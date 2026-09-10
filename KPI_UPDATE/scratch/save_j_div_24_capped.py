# -*- coding: utf-8 -*-
import openpyxl, os, sys
sys.stdout.reconfigure(encoding='utf-8')

for p in ['thang8/bangluong_thang8_hoanthien.xlsx', 'thang8/BANGLUONGTHANG8.xlsx']:
    if not os.path.exists(p):
        continue
    print(f"--> Updating {p}...")
    try:
        wb = openpyxl.load_workbook(p)
        ws = wb['BẢNG LƯƠNG']
        ws.cell(2, 10, "Số ngày làm ngày")
        ws.cell(2, 11, "Số ca làm đêm")
        for r in range(3, 59):
            # Cột J: Quy đổi số giờ ca ngày (G) sang ngày theo quy ước 24 giờ = 1 ngày (ROUND(G/24, 0))
            # và chặn trần để J + K <= 31 ngày
            ws.cell(r, 10, f"=MIN(ROUND(G{r}/24, 0), MAX(0, 31-K{r}))")
            # Cột K: Ca đêm >= 4h
            ws.cell(r, 11, f"=SUMIFS('Ngày công'!I:I, 'Ngày công'!B:B, C{r})")
            # Cột X: Phụ cấp ca đêm
            ws.cell(r, 24, f'=IF(ISNUMBER(SEARCH("DSCD", D{r})), ROUND(1500000/28*K{r}, 0), IF(K{r}>20, ROUND(1500000/28*K{r}, 0), 0))')
        wb.save(p)
        print(f"✅ Đã lưu thành công: {p}")
    except PermissionError:
        print(f"❌ PermissionError: File {p} đang được mở trong Excel!")
    except Exception as e:
        print(f"Error {p}: {e}")
