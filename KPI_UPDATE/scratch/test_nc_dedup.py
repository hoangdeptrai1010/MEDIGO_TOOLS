import openpyxl, time

wb = openpyxl.load_workbook('d:/MEDIGO/KPI_UPDATE/thang8/bangluong_thang8_hoanthien.xlsx')
ws_nc = wb['Ngày công']
ws_nc.cell(1, 8, 'Công thực tế')
ws_nc.cell(1, 9, 'Ca đêm chuẩn')

t0 = time.time()
for r in range(3, 1874):
    f_h = f'=IF(AND(ISNUMBER(F{r}), F{r}>0, COUNTIFS($A$3:A{r}, A{r}, $B$3:B{r}, B{r}, $C$3:C{r}, C{r})=1), 1, 0)'
    f_i = f'=IF(AND(ISNUMBER(F{r}), F{r}>0, ISNUMBER(SEARCH("đêm", D{r})), COUNTIFS($A$3:A{r}, A{r}, $B$3:B{r}, B{r}, $C$3:C{r}, C{r}, $D$3:D{r}, "*đêm*")=1), 1, 0)'
    ws_nc.cell(r, 8, f_h)
    ws_nc.cell(r, 9, f_i)

for r in range(2, 72):
    ws_nc.cell(r, 13, f'=SUMIFS(H:H, A:A, K{r}, B:B, L{r})')

ws_bl = wb['BẢNG LƯƠNG']
for r in range(3, 73):
    ws_bl.cell(r, 10, f"=SUMIFS('Ngày công'!H:H, 'Ngày công'!A:A, B{r}, 'Ngày công'!B:B, C{r})")
    ws_bl.cell(r, 11, f"=SUMIFS('Ngày công'!I:I, 'Ngày công'!A:A, B{r}, 'Ngày công'!B:B, C{r})")
    ws_bl.cell(r, 18, f"=SUMIFS('Ngày công'!M:M, 'Ngày công'!K:K, B{r}, 'Ngày công'!L:L, C{r})")

wb.save('d:/MEDIGO/KPI_UPDATE/scratch/test_nc_calc.xlsx')
print(f'Saved test file in {time.time()-t0:.2f}s')
