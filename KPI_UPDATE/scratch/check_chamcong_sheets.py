import openpyxl

for path in [
    'thang8/DATA/BangChiTietChamCong_thang8.xlsx',
    'thang7/DATA/BangChiTietChamCong_thang7.xlsx',
    'thang7/BangChiTietChamCong_thang7.xlsx',
]:
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        print(f"=== {path} ===")
        print("Sheets:", wb.sheetnames)
        for sname in wb.sheetnames:
            ws = wb[sname]
            print(f"  Sheet '{sname}': {ws.max_row} rows, {ws.max_column} cols")
    except Exception as e:
        print(f"Error reading {path}: {e}")
