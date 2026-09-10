import openpyxl

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)

for sname in ['MiniKat - HN', 'MiniKat - HCM']:
    if sname in wb.sheetnames:
        ws = wb[sname]
        print(f"=== SHEET: {sname} ({ws.max_row} rows x {ws.max_column} cols) ===")
        # Print headers
        headers = [f"{openpyxl.utils.get_column_letter(c)}: {ws.cell(1, c).value}" for c in range(1, ws.max_column+1) if ws.cell(1, c).value]
        print("Headers:\n  " + "\n  ".join(headers))
        
        # Print rows with non-zero bonus
        print("\nNon-zero rows:")
        for r in range(2, min(ws.max_row+1, 100)):
            vals = [ws.cell(r, c).value for c in range(1, ws.max_column+1)]
            # check if any bonus column > 0
            if any(isinstance(v, (int, float)) and v > 0 for v in vals[4:]):
                print(f"Row {r:2d}: " + " | ".join([f"{str(v)}" for v in vals[:15] if v is not None]))
