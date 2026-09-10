import sys
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=False)
ws_bl = wb['BẢNG LƯƠNG']

# List of secondary rows to remove (identified by (branch, name))
SECONDARY_STAFF_BRANCH = {
    ('Trường Sa', 'Hồ Thị Minh Hòa'),
    ('Lê Bình', 'Nguyễn Trần Ngọc Phương'),
    ('Trường Sa', 'Trịnh Thị Phượng'),
    ('Lê Bình', 'Phạm Nguyễn Ngọc Quý'),
    ('Đỗ Quang Đẩu', 'Dương Thị Huỳnh Như'),
    ('Đỗ Quang Đẩu', 'Hoàng Lâm Gia Bảo'),
    ('Nam Hòa', 'Hoàng Lâm Gia Bảo'),
    ('Nguyễn Thị Thập', 'Hoàng Lâm Gia Bảo'),
    ('Nam Hòa', 'Phan Công Vũ Tài'),
    ('Nguyễn Thị Thập', 'Phan Công Vũ Tài'),
    ('Rạch Bùng Binh', 'Phan Công Vũ Tài'),
    ('Lê Bình', 'Nguyễn Thị Huyền Trang'),
    ('Nguyễn Chí Thanh', 'Ngô Thị Ngọc Thủy'),
    ('Nguyễn Thị Thập', 'Cù Thị Tường Vy'),
}

rows_to_delete = []
for r in range(3, 73):
    cn = str(ws_bl.cell(r, 2).value or '').strip()
    name = str(ws_bl.cell(r, 3).value or '').strip()
    if (cn, name) in SECONDARY_STAFF_BRANCH:
        rows_to_delete.append(r)

print('Rows to delete:', rows_to_delete, 'Count:', len(rows_to_delete))

# Delete in reverse order
for r in sorted(rows_to_delete, reverse=True):
    ws_bl.delete_rows(r, 1)

print('New max row:', ws_bl.max_row)
for r in range(3, 62):
    cn = ws_bl.cell(r, 2).value
    name = ws_bl.cell(r, 3).value
    role = ws_bl.cell(r, 4).value
    if name:
        print(f'Row {r}: STT={r-2}, Branch={cn}, Name={name}, Role={role}')
    else:
        print(f'Row {r}: [Empty or Summary] Col D={role}, Col E={ws_bl.cell(r, 5).value}')
        break
