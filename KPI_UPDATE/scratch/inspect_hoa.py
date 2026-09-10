import openpyxl

wb = openpyxl.load_workbook('thang8/tinhcongnhungthuongchia.xlsx', data_only=False)
ws = wb['BẢNG LƯƠNG']

headers = [ws.cell(2, c).value for c in range(1, 45)]

with open('scratch/hoa_check.txt', 'w', encoding='utf-8') as f:
    f.write(f"Total rows in BẢNG LƯƠNG: {ws.max_row}\n")
    for r in range(3, ws.max_row+1):
        name = ws.cell(r, 3).value
        cn = ws.cell(r, 2).value
        if name and 'Hòa' in str(name):
            f.write(f"\n--- Row {r}: Chi nhánh={cn}, Tên={name} ---\n")
            for c in range(1, 45):
                h = headers[c-1] if c-1 < len(headers) else f"Col{c}"
                v = ws.cell(r, c).value
                if v is not None:
                    f.write(f"  Col {openpyxl.utils.get_column_letter(c)} ({h}): {v}\n")
print("Done writing to scratch/hoa_check.txt")
