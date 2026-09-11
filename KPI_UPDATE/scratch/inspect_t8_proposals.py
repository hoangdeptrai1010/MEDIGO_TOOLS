import openpyxl
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

hn_path = 'thang9/data/kpi_thang9/e_xuat_KPI_Quy_3.26.xlsx'
wb_hn = openpyxl.load_workbook(hn_path, data_only=True)
ws_t8 = wb_hn['T8.26']

print("=== CHI TIẾT SHEET T8.26 FILE ĐỀ XUẤT HÀ NỘI ===")
for r in range(1, ws_t8.max_row + 1):
    vals = [ws_t8.cell(r, c).value for c in range(1, ws_t8.max_column + 1)]
    if any(v is not None for v in vals):
        print(f"Row {r:02d}: {vals}")
