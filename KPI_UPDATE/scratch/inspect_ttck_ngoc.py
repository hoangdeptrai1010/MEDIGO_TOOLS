import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_ngoc = os.path.join(latvat_dir, "Gởi Ms. Ngoc - 07.09. Bang Luong T8.2026.xlsx")
wb = openpyxl.load_workbook(f_ngoc, data_only=True)

if "TT CK CNT" in wb.sheetnames:
    ws = wb["TT CK CNT"]
    print(f"=== TT CK CNT in Ms. Ngoc (rows={ws.max_row}) ===")
    for r in range(1, min(45, ws.max_row + 1)):
        vals = [ws.cell(r, c).value for c in range(1, min(12, ws.max_column + 1))]
        if any(v is not None for v in vals):
            print(f"Row {r:2d}: {vals}")

if "TTCK CNT" in wb.sheetnames:
    ws = wb["TTCK CNT"]
    print(f"\n=== TTCK CNT in Ms. Ngoc (rows={ws.max_row}) ===")
    for r in range(1, min(45, ws.max_row + 1)):
        vals = [ws.cell(r, c).value for c in range(1, min(12, ws.max_column + 1))]
        if any(v is not None for v in vals):
            print(f"Row {r:2d}: {vals}")
