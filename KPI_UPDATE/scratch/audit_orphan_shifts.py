import openpyxl
import sys
from collections import defaultdict

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws_nc = wb['Ngày công']
ws_bl = wb['BẢNG LƯƠNG']

# Get all (branch, staff) pairs in BẢNG LƯƠNG
bl_pairs = set()
for r in range(3, ws_bl.max_row + 1):
    cn = ws_bl.cell(r, 2).value
    name = ws_bl.cell(r, 3).value
    if name and str(name).strip():
        bl_pairs.add((str(cn).strip(), str(name).strip()))

# Also get (branch, staff) pairs in Ngày công summary
nc_sum_pairs = set()
for r in range(2, 75):
    cn = ws_nc.cell(r, 11).value
    name = ws_nc.cell(r, 12).value
    if name and str(name).strip():
        nc_sum_pairs.add((str(cn).strip(), str(name).strip()))

print(f"Total pairs in BẢNG LƯƠNG: {len(bl_pairs)}")
print(f"Total pairs in Ngày công summary: {len(nc_sum_pairs)}")

# Scan raw rows in Ngày công
unmatched_shifts = []
raw_worked_by_staff_branch = defaultdict(lambda: {'hours': 0.0, 'shifts': 0, 'dates': set()})

for r in range(3, ws_nc.max_row + 1):
    cn = ws_nc.cell(r, 1).value
    name = ws_nc.cell(r, 2).value
    date_val = ws_nc.cell(r, 3).value
    shift = ws_nc.cell(r, 4).value
    h = ws_nc.cell(r, 6).value or 0
    
    if not name or not str(name).strip():
        continue
    name = str(name).strip()
    cn = str(cn).strip() if cn else ''
    
    # Check if this is the dummy header row
    if shift == '-' or not date_val:
        continue
        
    try:
        hf = float(h)
    except:
        hf = 0.0
        
    raw_worked_by_staff_branch[(cn, name)]['hours'] += hf
    raw_worked_by_staff_branch[(cn, name)]['shifts'] += 1
    raw_worked_by_staff_branch[(cn, name)]['dates'].add(str(date_val))
    
    if (cn, name) not in bl_pairs:
        unmatched_shifts.append((r, cn, name, date_val, shift, hf))

print(f"\nTotal raw shifts not matching any (Branch, Staff) in BẢNG LƯƠNG: {len(unmatched_shifts)}")
orphan_by_person = defaultdict(list)
for item in unmatched_shifts:
    orphan_by_person[item[2]].append(item)

for person, items in orphan_by_person.items():
    tot_h = sum(x[5] for x in items)
    print(f"\n--> Nhân viên: {person} (Tổng giờ bị mất: {tot_h:.2f}h trong {len(items)} ca)")
    for x in items:
        print(f"    Dòng {x[0]}: Chi nhánh '{x[1]}' | Ngày: {x[3]} | Ca: {x[4]} | Giờ: {x[5]}h")
