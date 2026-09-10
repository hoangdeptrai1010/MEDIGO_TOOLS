import openpyxl

wb_ref = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx', data_only=True)
wb_tool = openpyxl.load_workbook('test_tool_output.xlsx', data_only=True)

print("="*85)
print("ĐỐI CHIẾU TRỰC DIỆN GIỮA FILE BÁO CÁO THỰC TẾ VÀ TOOL KPI_ENGINE VỪA TẠO")
print("="*85)

for sname in wb_ref.sheetnames:
    if sname not in wb_tool.sheetnames:
        print(f"\n❌ Sheet '{sname}' CÓ TRONG FILE THỰC TẾ NHƯNG THIẾU TRONG TOOL OUTPUT!")
        continue
        
    ws_r = wb_ref[sname]
    ws_t = wb_tool[sname]
    
    print(f"\n--- SHEET: {sname} ---")
    print(f"  File thực tế: {ws_r.max_row} rows x {ws_r.max_column} cols")
    print(f"  Tool output : {ws_t.max_row} rows x {ws_t.max_column} cols")
    
    diffs = []
    # Check rows
    max_r = max(ws_r.max_row, ws_t.max_row)
    max_c = max(ws_r.max_column, ws_t.max_column)
    
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            vr = ws_r.cell(r, c).value
            vt = ws_t.cell(r, c).value
            
            if vr != vt:
                # float tolerance
                if isinstance(vr, (int, float)) and isinstance(vt, (int, float)):
                    if abs(vr - vt) > 0.5:
                        diffs.append((r, c, vr, vt))
                else:
                    diffs.append((r, c, vr, vt))
                    
    print(f"  Số ô lệch: {len(diffs)}")
    if diffs:
        print("  Top 10 ô lệch tiêu biểu:")
        for r, c, vr, vt in diffs[:10]:
            col_letter = openpyxl.utils.get_column_letter(c)
            # identify person/store if possible
            entity = ws_r.cell(r, 2).value or ws_r.cell(r, 1).value or f"Row {r}"
            print(f"    {col_letter}{r} ({entity}): File_Ref={vr} | Tool_Output={vt}")
