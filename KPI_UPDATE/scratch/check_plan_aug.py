import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('plans/KeHoachKPI_2026-08.xlsx', data_only=True)
print("Sheet names in KeHoachKPI_2026-08.xlsx:", wb.sheetnames)

for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"\n--- Sheet: {sn} (max_row={ws.max_row}, max_col={ws.max_column}) ---")
    for r in range(1, 15):
        vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column+1))]
        if any(vals):
            print(f"Row {r:2d}: {vals}")
