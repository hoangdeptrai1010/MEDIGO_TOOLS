import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
ws = wb.active

print("=== CHI TIẾT TỪ FILE GỐC BangChiTietChamCong_thang8.xlsx ===")
for r in range(195, 220):
    row_meta = [ws.cell(r, c).value for c in range(1, 8)]
    # count non-empty punch days
    punch_days = 0
    for col in range(8, ws.max_column, 2):
        v1 = ws.cell(r, col).value
        v2 = ws.cell(r, col+1).value
        if v1 or v2:
            punch_days += 1
    print(f"Row {r:3d}: {row_meta} | Số ngày quẹt thẻ: {punch_days}")
