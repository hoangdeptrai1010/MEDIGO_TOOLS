# -*- coding: utf-8 -*-
import sys, os
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

for p in ['thang8/Pharmacy_retail_Store_KPIs_August 2026/Hồ Chí Minh/Dự án/Chương trình/HCM_Projects_August_2026_v1.pptx',
          'thang8/Pharmacy_retail_Store_KPIs_August 2026/Hà Nội/Dự án/Chương trình/Hanoi_Projects_August_2026_v1.pptx']:
    if os.path.exists(p):
        print(f"\n==================== PPTX: {p} ====================")
        prs = Presentation(p)
        for idx, slide in enumerate(prs.slides):
            print(f"\n--- Slide {idx+1} ---")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        t = paragraph.text.strip()
                        if t: print("  Text:", t)
                elif shape.has_table:
                    for row in shape.table.rows:
                        print("  Table row:", [cell.text.strip() for cell in row.cells])
