import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.path.insert(0, 'd:/MEDIGO/KPI_UPDATE/TOOL_KPISHEET')
from builder_engine import remove_accents

wb = openpyxl.load_workbook('d:/MEDIGO/KPI_UPDATE/thang9/data/kpi_thang9/e_xuat_KPI_Quy_3.26.xlsx', data_only=True)
ws = wb['T9.26']
for c in range(1, ws.max_column+1):
    v = ws.cell(1, c).value
    norm_v = remove_accents(str(v)).lower().strip() if v else ''
    print(f'Col {c}: orig="{v}", norm="{norm_v}"')
