import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb_raw = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
ws_raw = wb_raw.active

print("=== QUÉT TOÀN BỘ NHÂN VIÊN TỪ BANGCHITIETCHAMCONG_THANG8.XLSX ===")

staff_shifts = {} # (cn, name) -> list of shifts
curr_staff = ""
curr_branch = ""

for r in range(5, ws_raw.max_row+1):
    c_name = ws_raw.cell(r, 3).value
    c_branch = ws_raw.cell(r, 6).value
    c_shift = ws_raw.cell(r, 7).value
    
    if c_name and str(c_name).strip():
        curr_staff = str(c_name).strip()
    if c_branch and str(c_branch).strip():
        curr_branch = str(c_branch).strip()
    
    if curr_staff and c_shift:
        # Check all days
        for day in range(1, 32):
            col_in = 8 + (day-1)*2
            col_out = col_in + 1
            v_in = str(ws_raw.cell(r, col_in).value or '').strip()
            v_out = str(ws_raw.cell(r, col_out).value or '').strip()
            if v_in and v_out:
                try:
                    h_in, m_in = map(int, v_in.split(' ')[0].split(':')[:2])
                    h_out, m_out = map(int, v_out.split(' ')[0].split(':')[:2])
                    t_in = h_in * 60 + m_in
                    t_out = h_out * 60 + m_out
                    if t_out < t_in:
                        t_out += 24 * 60
                    dur_h = (t_out - t_in) / 60.0
                    key = (curr_branch, curr_staff)
                    if key not in staff_shifts:
                        staff_shifts[key] = []
                    staff_shifts[key].append({
                        'day': day, 'shift': str(c_shift).strip(), 'in': v_in, 'out': v_out, 'hrs': dur_h
                    })
                except Exception as e:
                    pass

print(f"Tổng số nhân sự có chấm công: {len(staff_shifts)}")

for (cn, name), shifts in staff_shifts.items():
    if 'thoa' in name.lower() or 'tâm' in name.lower():
        night_gt4 = sum(1 for s in shifts if 'đêm' in s['shift'].lower() and s['hrs'] > 4)
        night_all = sum(1 for s in shifts if 'đêm' in s['shift'].lower())
        day_gt4 = sum(1 for s in shifts if 'đêm' not in s['shift'].lower() and s['hrs'] > 4)
        tot_hrs = sum(s['hrs'] for s in shifts)
        unique_days = len(set(s['day'] for s in shifts if s['hrs'] > 0))
        print(f"\n{name} ({cn}):")
        print(f"  - Số ca đêm > 4h: {night_gt4}")
        print(f"  - Số ca đêm tất cả: {night_all}")
        print(f"  - Số ca ngày > 4h: {day_gt4}")
        print(f"  - Tổng ngày công thực tế: {unique_days}")
        print(f"  - Tổng giờ làm: {tot_hrs:.2f}h")
