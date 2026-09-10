import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl, pandas as pd
from collections import defaultdict

# 1. Đọc Danh mục Kế hoạch Tháng 8
wb_plan = openpyxl.load_workbook('plans/KeHoachKPI_2026-08.xlsx', data_only=True)
ws_cat = wb_plan['Danh mục dự án']
sku_cat = {}
name_cat = {}
for r in range(2, ws_cat.max_row+1):
    sku = str(ws_cat.cell(r, 1).value or '').strip()
    name = str(ws_cat.cell(r, 2).value or '').strip()
    grp = str(ws_cat.cell(r, 3).value or '').strip()
    if sku and grp: sku_cat[sku] = grp
    if name and grp: name_cat[name] = grp

print(f"=== 1. TỔNG QUAN DANH MỤC DỰ ÁN (PLANS) ===")
print(f"Tổng số SKU trong Danh mục: {len(sku_cat)}")
print(f" - NY3:   {sum(1 for g in sku_cat.values() if g=='NY3')} SKU")
print(f" - CK:    {sum(1 for g in sku_cat.values() if g=='CK')} SKU")
print(f" - Combo: {sum(1 for g in sku_cat.values() if g=='Combo')} SKU")

# 2. Đọc file Hóa đơn và Trả hàng tháng 8
df_hd = pd.read_excel('thang8/DATA/DanhSachChiTietHoaDon_3182026.xlsx')
df_th = pd.read_excel('thang8/DATA/DanhSachChiTietTraHang_3182026.xlsx')

print(f"\n=== 2. TỔNG QUAN DỮ LIỆU BÁN HÀNG THÁNG 8 ===")
print(f"Tổng số dòng Hóa đơn chi tiết: {len(df_hd):,}")
print(f"Tổng số dòng Trả hàng chi tiết: {len(df_th):,}")

# Phân loại toàn bộ sản phẩm bán ra
sales_summary = defaultdict(lambda: {'total_off': 0.0, 'total_onl': 0.0, 'NY3': 0.0, 'CK': 0.0, 'Combo': 0.0, 'Khac': 0.0})
branch_summary = defaultdict(lambda: {'total_rev': 0.0, 'NY3': 0.0, 'CK': 0.0, 'Combo': 0.0, 'Khac': 0.0})

for _, row in df_hd.iterrows():
    sku = str(row.get('Mã hàng', '')).strip()
    th = str(row.get('Tên hàng', '')).strip()
    branch = str(row.get('Chi nhánh', '')).strip()
    seller = str(row.get('Người bán', '')).strip()
    channel = str(row.get('Kênh bán', '')).strip().lower()
    tt = float(row.get('Thành tiền', 0) or 0)
    
    grp = sku_cat.get(sku) or name_cat.get(th)
    if not grp:
        p = th.split()[0].upper() if th else ''
        if p.startswith('NY3'): grp = 'NY3'
        elif p.startswith('CK') or p.startswith('CKHN'): grp = 'CK'
        elif p.startswith('COMBO') or p.startswith('LIỀU') or p.startswith('LIEU'): grp = 'Combo'
        else: grp = 'Khac'
    
    key = (branch, seller)
    if 'online' in channel or 'onl' in channel or 'app' in channel:
        sales_summary[key]['total_onl'] += tt
    else:
        sales_summary[key]['total_off'] += tt
        
    sales_summary[key][grp] += tt
    branch_summary[branch]['total_rev'] += tt
    branch_summary[branch][grp] += tt

# Trừ trả hàng
for _, row in df_th.iterrows():
    sku = str(row.get('Mã hàng', '')).strip()
    th = str(row.get('Tên hàng', '')).strip()
    branch = str(row.get('Chi nhánh', '')).strip()
    seller = str(row.get('Người bán', '')).strip()
    channel = str(row.get('Kênh bán', '')).strip().lower()
    tt = float(row.get('Thành tiền', 0) or 0)
    
    grp = sku_cat.get(sku) or name_cat.get(th)
    if not grp:
        p = th.split()[0].upper() if th else ''
        if p.startswith('NY3'): grp = 'NY3'
        elif p.startswith('CK') or p.startswith('CKHN'): grp = 'CK'
        elif p.startswith('COMBO') or p.startswith('LIỀU') or p.startswith('LIEU'): grp = 'Combo'
        else: grp = 'Khac'
    
    key = (branch, seller)
    if 'online' in channel or 'onl' in channel or 'app' in channel:
        sales_summary[key]['total_onl'] -= tt
    else:
        sales_summary[key]['total_off'] -= tt
        
    sales_summary[key][grp] -= tt
    branch_summary[branch]['total_rev'] -= tt
    branch_summary[branch][grp] -= tt

print(f"\n=== 3. TỔNG DOANH SỐ THỰC TẾ THEO NHÓM SẢN PHẨM TRÊN TOÀN HỆ THỐNG ===")
tot_all_rev = sum(b['total_rev'] for b in branch_summary.values())
tot_ny3 = sum(b['NY3'] for b in branch_summary.values())
tot_ck = sum(b['CK'] for b in branch_summary.values())
tot_combo = sum(b['Combo'] for b in branch_summary.values())

print(f"Tổng Doanh Thu Toàn Hệ Thống: {tot_all_rev:15,.0f} VNĐ (100.0%)")
print(f" - Nhóm NY3:                   {tot_ny3:15,.0f} VNĐ ({tot_ny3/tot_all_rev*100:4.1f}%)")
print(f" - Nhóm Chiết Khấu (CK):       {tot_ck:15,.0f} VNĐ ({tot_ck/tot_all_rev*100:4.1f}%)")
print(f" - Nhóm Combo Liều:            {tot_combo:15,.0f} VNĐ ({tot_combo/tot_all_rev*100:4.1f}%)")
print(f" - Nhóm Thuốc & Hàng Khác:     {tot_all_rev-tot_ny3-tot_ck-tot_combo:15,.0f} VNĐ")

print(f"\n=== 4. PHÂN BỔ DOANH SỐ NY3 THEO CHI NHÁNH ===")
for b, data in sorted(branch_summary.items(), key=lambda x: x[1]['NY3'], reverse=True):
    print(f"  {b:<25}: NY3 = {data['NY3']:10,.0f} đ | CK = {data['CK']:12,.0f} đ | Combo = {data['Combo']:10,.0f} đ")
