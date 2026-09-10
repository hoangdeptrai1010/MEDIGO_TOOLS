import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=False)
ws_gc = wb['Giờ công']

print("--- SHEET 'Giờ công' SUMMARY TABLE (Cols G - K) ---")
for r in range(1, 15):
    g = ws_gc.cell(r, 7).value
    h = ws_gc.cell(r, 8).value
    i = ws_gc.cell(r, 9).value
    j = ws_gc.cell(r, 10).value
    k = ws_gc.cell(r, 11).value
    print(f"Row {r:2d}: G={str(g):<15} | H={str(h):<25} | I={str(i):<35} | J={str(j):<35} | K={str(k)}")
