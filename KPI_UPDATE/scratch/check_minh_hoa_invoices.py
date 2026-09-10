# -*- coding: utf-8 -*-
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb_inv = openpyxl.load_workbook('thang8/DATA/DanhSachChiTietHoaDon_3182026.xlsx', data_only=True)
ws_inv = wb_inv.active

print("=== INVOICES OF HỒ THỊ MINH HÒA IN AUGUST ===")
branches = {}
for r in range(2, ws_inv.max_row+1):
    seller = str(ws_inv.cell(r, 4).value or '').strip()
    branch = str(ws_inv.cell(r, 3).value or '').strip()
    sku = str(ws_inv.cell(r, 7).value or '').strip()
    name = str(ws_inv.cell(r, 8).value or '').strip()
    qty = ws_inv.cell(r, 9).value or 0
    
    if 'Hồ Thị Minh Hòa' in seller or 'Minh Hòa' in seller:
        if branch not in branches:
            branches[branch] = []
        branches[branch].append((sku, name, qty))

for b, items in branches.items():
    print(f"Chi nhánh: {b} ({len(items)} items)")
    for sku, name, qty in items[:5]:
        print(f"  {sku} | {name} | SL={qty}")
