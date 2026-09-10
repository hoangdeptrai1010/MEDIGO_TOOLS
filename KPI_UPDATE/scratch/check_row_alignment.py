import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws_bl = wb['BẢNG LƯƠNG']
ws_gc = wb['Giờ công']
ws_nc = wb['Ngày công']

print("--- Comparison of Rows between BẢNG LƯƠNG, Giờ công summary, and Ngày công summary ---")
print(f"BẢNG LƯƠNG rows: {ws_bl.max_row}, Giờ công summary rows: {ws_gc.max_row}, Ngày công summary rows: {ws_nc.max_row}")

mismatches = []
for r in range(3, 73):
    bl_cn = ws_bl.cell(r, 2).value
    bl_name = ws_bl.cell(r, 3).value
    
    # In Giờ công summary, rows start at 2 (r-1)
    gc_r = r - 1
    gc_cn = ws_gc.cell(gc_r, 7).value
    gc_name = ws_gc.cell(gc_r, 8).value
    
    nc_r = r - 1
    nc_cn = ws_nc.cell(nc_r, 11).value
    nc_name = ws_nc.cell(nc_r, 12).value
    
    match = (str(bl_cn).strip() == str(gc_cn).strip() == str(nc_cn).strip()) and (str(bl_name).strip() == str(gc_name).strip() == str(nc_name).strip())
    if not match:
        mismatches.append((r, (bl_cn, bl_name), (gc_cn, gc_name), (nc_cn, nc_name)))

print(f"Total row alignment mismatches between BẢNG LƯƠNG and Giờ công/Ngày công: {len(mismatches)}")
for m in mismatches:
    print(f"Row {m[0]}: BL={m[1]} | GC={m[2]} | NC={m[3]}")
