import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/BANGLUONGTHANG8.xlsx', data_only=True)
ws_bl = wb['BẢNG LƯƠNG']

print("=== AUDIT ALL 70 STAFF IN BANGLUONGTHANG8.xlsx ===")
print(f"{'STT':<4} | {'Chi nhánh':<18} | {'Tên':<26} | {'Chức danh':<10} | {'Giờ Ngày(G)':<12} | {'Giờ Đêm(H)':<12} | {'Ca Ngày(J)':<10} | {'Ca Đêm(K)':<10} | {'Tổng Công(R)':<12} | {'PC Đêm(X)':<12}")
print("-" * 130)

night_staff = []

for r in range(3, 75):
    stt = ws_bl.cell(r, 1).value
    cn = ws_bl.cell(r, 2).value
    name = ws_bl.cell(r, 3).value
    role = ws_bl.cell(r, 4).value
    
    if not isinstance(stt, (int, float)):
        continue

    try:
        g = float(ws_bl.cell(r, 7).value or 0)
        h = float(ws_bl.cell(r, 8).value or 0)
        j = float(ws_bl.cell(r, 10).value or 0)
        k = float(ws_bl.cell(r, 11).value or 0)
        r_val = float(ws_bl.cell(r, 18).value or 0)
        x = float(ws_bl.cell(r, 24).value or 0)
        print(f"{int(stt):<4} | {str(cn):<18} | {str(name):<26} | {str(role):<10} | {g:<12.2f} | {h:<12.2f} | {j:<10.1f} | {k:<10.0f} | {r_val:<12.0f} | {x:<12,.0f}")
    except Exception as e:
        print(f"Row {r} error: {e}")
