import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'thang8\bangluong_thang8_hoanthien.xlsx', data_only=True)
ws = wb['BẢNG LƯƠNG']

print(f"{'STT':<4} | {'Họ và tên':<26} | {'Chức danh':<10} | {'Giờ ngày (G)':<14} | {'G/8 (chính xác)':<16} | {'ROUND(G/8,0)':<14} | {'Hiện tại (J)':<14} | {'Ca đêm (K)':<10}")
print("-" * 125)

for r in range(3, ws.max_row + 1):
    stt = ws.cell(r, 1).value
    name = ws.cell(r, 3).value
    chucdanh = ws.cell(r, 4).value
    g = ws.cell(r, 7).value or 0
    j = ws.cell(r, 10).value or 0
    k = ws.cell(r, 11).value or 0
    
    if name and str(name).strip() and "Tổng" not in str(name):
        g_val = float(g) if isinstance(g, (int, float)) else 0.0
        g_div_8 = g_val / 8.0
        g_div_8_rnd = round(g_div_8, 0)
        
        print(f"{str(stt):<4} | {str(name):<26} | {str(chucdanh):<10} | {g_val:>12.2f} h | {g_div_8:>14.2f} ng | {g_div_8_rnd:>12.0f} ng | {str(j):>12} ng | {str(k):>8}")
