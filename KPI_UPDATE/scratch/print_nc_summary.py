import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws_nc = wb['Ngày công']

print(f"{'Row':<4} | {'Chi nhánh':<18} | {'Tên nhân viên':<25} | {'Ngày công':<10} | {'Giờ công':<10} | {'TB Giờ/công':<12}")
print("-" * 90)

for r in range(2, 72):
    cn = ws_nc.cell(r, 11).value
    name = ws_nc.cell(r, 12).value
    ngay_cong = ws_nc.cell(r, 13).value
    gio_cong = ws_nc.cell(r, 14).value
    tb = ws_nc.cell(r, 15).value
    
    nc_str = f"{float(ngay_cong):.1f}" if ngay_cong is not None else "None"
    gc_str = f"{float(gio_cong):.2f}" if gio_cong is not None else "None"
    tb_str = f"{float(tb):.2f}" if isinstance(tb, (int, float)) else str(tb)
    
    print(f"{r:<4} | {str(cn):<18} | {str(name):<25} | {nc_str:<10} | {gc_str:<10} | {tb_str:<12}")
