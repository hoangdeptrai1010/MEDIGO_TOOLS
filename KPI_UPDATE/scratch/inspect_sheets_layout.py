import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/tinhcongnhungthuongchia.xlsx', data_only=False)

print("=== SHEET GIỜ CÔNG ===")
ws_gc = wb['Giờ công']
print("Max row:", ws_gc.max_row, "Max col:", ws_gc.max_column)
for r in range(1, 10):
    row_vals = [ws_gc.cell(r, c).value for c in range(1, 15)]
    print(f"Row {r}: {row_vals}")

print("\n=== SHEET NGÀY CÔNG ===")
ws_nc = wb['Ngày công']
print("Max row:", ws_nc.max_row, "Max col:", ws_nc.max_column)
for r in range(1, 10):
    row_vals = [ws_nc.cell(r, c).value for c in range(1, 18)]
    print(f"Row {r}: {row_vals}")

wb.close()
