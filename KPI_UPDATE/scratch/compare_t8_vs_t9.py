import openpyxl, sys, io
if sys.stdout.encoding != 'utf-8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb_goc8 = openpyxl.load_workbook('goc/NHÀ THUỐC THÁNG 8 2026.xlsx', data_only=True)
wb_out9 = openpyxl.load_workbook('TOOL_KPISHEET/output/NHÀ THUỐC THÁNG 9 2026.xlsx', data_only=True)

print("=== GOC THÁNG 8 STAFF TARGETS (First 15) ===")
ws8 = wb_goc8['kpi dược sĩ']
for r in range(3, 18):
    st = ws8.cell(r, 1).value
    nm = ws8.cell(r, 2).value
    tgt_thang = ws8.cell(r, 22).value # Col V
    tb_bill = ws8.cell(r, 4).value
    print(f"T8 Row {r:2d}: [{st}] {nm} | TB_Bill={tb_bill} | Target_Thang={tgt_thang}")

print("\n=== OUTPUT THÁNG 9 STAFF TARGETS (First 15) ===")
ws9 = wb_out9['kpi dược sĩ']
for r in range(3, 18):
    st = ws9.cell(r, 1).value
    nm = ws9.cell(r, 2).value
    tgt_thang = ws9.cell(r, 22).value # Col V
    tb_bill = ws9.cell(r, 4).value
    print(f"T9 Row {r:2d}: [{st}] {nm} | TB_Bill={tb_bill} | Target_Thang={tgt_thang}")
