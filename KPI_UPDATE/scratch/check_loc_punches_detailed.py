import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb_raw = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
ws_raw = wb_raw.active

wb_bl = openpyxl.load_workbook('thang8/BANGLUONGTHANG8.xlsx', data_only=True)
ws_bl = wb_bl['BẢNG LƯƠNG']
ws_gc = wb_bl['Giờ công']
ws_nc = wb_bl['Ngày công']

print("=== DETAILED CHECK: NGUYỄN THỊ PHÚC LỘC PUNCHES & SHIFTS ===")
# Raw rows for Loc
for r in range(5, ws_raw.max_row + 1):
    c_name = ws_raw.cell(r, 3).value
    c_branch = ws_raw.cell(r, 6).value
    c_shift = ws_raw.cell(r, 7).value
    c_code = ws_raw.cell(r, 2).value
    
    # Check if this row belongs to Loc
    # (Loc is row 146 in ws_raw, row 147 is her Ca 3 Dem)
    if r in [146, 147]:
        print(f"\nRow {r}: Code={c_code} | Name={c_name} | Branch={c_branch} | Shift={c_shift}")
        total_m = 0
        for day in range(1, 32):
            col_in = 8 + (day - 1) * 2
            col_out = col_in + 1
            v_in = ws_raw.cell(r, col_in).value
            v_out = ws_raw.cell(r, col_out).value
            if v_in or v_out:
                vin_s = str(v_in).strip()
                vout_s = str(v_out).strip()
                p_in = vin_s.split(' ')[0].split(':')
                p_out = vout_s.split(' ')[0].split(':')
                hin, min_val = int(p_in[0]), int(p_in[1])
                hout, mout_val = int(p_out[0]), int(p_out[1])
                tin = hin * 60 + min_val
                tout = hout * 60 + mout_val
                if 'đêm' in str(c_shift).lower() or tout < tin:
                    if tout < tin:
                        tout += 24 * 60
                dur = tout - tin
                total_m += dur
                print(f"  Day {day:02d}: In={vin_s} -> Out={vout_s} | Duration = {dur//60}h{dur%60:02d}p ({dur/60:.2f} hrs)")
        print(f"Total row duration: {total_m//60}h{total_m%60:02d}p = {total_m/60:.2f} hrs ({round(total_m/60, 2)} hrs)")
