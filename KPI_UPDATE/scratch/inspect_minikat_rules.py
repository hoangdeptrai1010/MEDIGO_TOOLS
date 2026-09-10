import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

files_to_check = [
    r'thang8\Pharmacy_retail_Store_KPIs_August 2026\Hà Nội\Dự án Mini\Chương trình\HN_KAT Project August_2026.xlsx',
    r'thang8\Pharmacy_retail_Store_KPIs_August 2026\Hà Nội\Dự án Mini\Chương trình\HN_Party Smart_Project _August _2026.xlsx',
    r'thang8\Pharmacy_retail_Store_KPIs_August 2026\Hà Nội\Dự án Mini\Chương trình\HN_Ladycare Project_August _2026.xlsx',
    r'thang8\Pharmacy_retail_Store_KPIs_August 2026\Hà Nội\Dự án Mini\Chương trình\HN_AVC project_August_2026.xlsx',
    r'thang8\Pharmacy_retail_Store_KPIs_August 2026\Hồ Chí Minh\Dự án mini\Chương trình\HCM_KAT Project August_2026.xlsx',
    r'thang8\Pharmacy_retail_Store_KPIs_August 2026\Hồ Chí Minh\Dự án mini\Chương trình\HCM_Party Smart_Project _August _2026.xlsx',
    r'thang8\Pharmacy_retail_Store_KPIs_August 2026\Hồ Chí Minh\Dự án mini\Chương trình\HCM_Ladycare Project_August_2026.xlsx',
]

for fp in files_to_check:
    print("=" * 80)
    print(f"FILE: {fp}")
    print("=" * 80)
    if os.path.exists(fp):
        wb = openpyxl.load_workbook(fp, data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            print(f"--- Sheet: {sname} ---")
            for r in range(1, ws.max_row + 1):
                row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
                if any(v is not None for v in row_vals):
                    print(f"Row {r:2d}: {row_vals}")
