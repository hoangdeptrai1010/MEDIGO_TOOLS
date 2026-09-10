import openpyxl
import sys
from collections import defaultdict

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws_nc = wb['Ngày công']
ws_gc = wb['Giờ công']
ws_bl = wb['BẢNG LƯƠNG']

# 1. Total hours in Ngày công raw (excluding dummy '-' rows)
nc_raw_tot = 0.0
nc_raw_by_person = defaultdict(float)
for r in range(3, ws_nc.max_row + 1):
    shift = ws_nc.cell(r, 4).value
    name = ws_nc.cell(r, 2).value
    h = ws_nc.cell(r, 6).value or 0
    if shift != '-' and name and str(name).strip():
        try:
            hf = float(h)
            nc_raw_tot += hf
            nc_raw_by_person[str(name).strip()] += hf
        except:
            pass

# 2. Total hours in Giờ công summary (Col K)
gc_sum_tot = 0.0
gc_sum_by_person = defaultdict(float)
for r in range(2, 72):
    name = ws_gc.cell(r, 8).value
    tot = ws_gc.cell(r, 11).value or 0
    if name and str(name).strip():
        try:
            tot_f = float(tot)
            gc_sum_tot += tot_f
            gc_sum_by_person[str(name).strip()] += tot_f
        except:
            pass

# 3. Total hours in Ngày công summary (Col N)
nc_sum_tot = 0.0
nc_sum_by_person = defaultdict(float)
for r in range(2, 72):
    name = ws_nc.cell(r, 12).value
    tot = ws_nc.cell(r, 14).value or 0
    if name and str(name).strip():
        try:
            tot_f = float(tot)
            nc_sum_tot += tot_f
            nc_sum_by_person[str(name).strip()] += tot_f
        except:
            pass

# 4. Total hours in BẢNG LƯƠNG (Col E ca ngày + Col F ca đêm)
bl_tot = 0.0
bl_by_person = defaultdict(float)
for r in range(3, 73):
    name = ws_bl.cell(r, 3).value
    e = ws_bl.cell(r, 5).value or 0
    f = ws_bl.cell(r, 6).value or 0
    if name and str(name).strip():
        try:
            tot_f = float(e) + float(f)
            bl_tot += tot_f
            bl_by_person[str(name).strip()] += tot_f
        except:
            pass

print(f"1. Tổng giờ công thực tế từ chi tiết quẹt thẻ (Ngày công raw): {nc_raw_tot:.2f} giờ")
print(f"2. Tổng giờ công bảng tóm tắt Giờ công:                      {gc_sum_tot:.2f} giờ")
print(f"3. Tổng giờ công bảng tóm tắt Ngày công:                     {nc_sum_tot:.2f} giờ")
print(f"4. Tổng giờ công trong BẢNG LƯƠNG:                            {bl_tot:.2f} giờ")
print(f"--> Tổng số giờ chênh lệch toàn công ty:                     {nc_raw_tot - bl_tot:.2f} giờ")

print("\n--- Chi tiết những nhân viên bị chênh lệch giữa Raw và BẢNG LƯƠNG ---")
for p in sorted(nc_raw_by_person.keys()):
    r_h = nc_raw_by_person[p]
    b_h = bl_by_person[p]
    if abs(r_h - b_h) > 0.05:
        print(f"Nhân viên: {p:<30} | Thực làm (Raw): {r_h:>7.2f}h | Bảng lương: {b_h:>7.2f}h | Chênh lệch: {r_h - b_h:>6.2f}h")
