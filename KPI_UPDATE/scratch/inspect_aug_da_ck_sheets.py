# -*- coding: utf-8 -*-
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)

print("=== SHEET 'Dự án' in bangluong_thang8_hoanthien.xlsx ===")
ws_da = wb['Dự án']
for r in range(1, 20):
    row_vals = [f"Col{c}:{ws_da.cell(r, c).value}" for c in range(1, 14) if ws_da.cell(r, c).value is not None]
    if row_vals:
        print(f"Row {r:2d}: {row_vals}")

print("\n=== SHEET 'Thưởng CK' in bangluong_thang8_hoanthien.xlsx ===")
ws_ck = wb['Thưởng CK']
for r in range(1, 20):
    row_vals = [f"Col{c}:{ws_ck.cell(r, c).value}" for c in range(1, 14) if ws_ck.cell(r, c).value is not None]
    if row_vals:
        print(f"Row {r:2d}: {row_vals}")
