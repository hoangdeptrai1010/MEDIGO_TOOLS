import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_bl = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final_backup.xlsx")
wb = openpyxl.load_workbook(f_bl, data_only=False)

print("=== INSPECTING Sheet1 (Formulas & Columns) ===")
ws1 = wb["Sheet1"]
print(f"Max row: {ws1.max_row}, Max col: {ws1.max_column}")
for r in range(1, ws1.max_row + 1):
    vals = [ws1.cell(r, c).value for c in range(1, ws1.max_column + 1)]
    if any(v is not None for v in vals):
        print(f"Row {r:2d}: {vals}")

print("\n=== INSPECTING Med (Headers & Sample Rows) ===")
ws_med = wb["Med"]
for r in range(9, 15):
    vals = [ws_med.cell(r, c).value for c in range(1, min(15, ws_med.max_column + 1))]
    print(f"Med Row {r:2d}: {vals}")

# Let's see what column is "Thực lãnh" / "Số tiền" in Med
for c in range(1, ws_med.max_column + 1):
    h = ws_med.cell(9, c).value
    if h and any(k in str(h).lower() for k in ['thực', 'lãnh', 'chuyển khoản', 'tổng thu', 'số tiền', 'lương']):
        print(f"Med Col {c} ({openpyxl.utils.get_column_letter(c)}): {repr(h)}")
