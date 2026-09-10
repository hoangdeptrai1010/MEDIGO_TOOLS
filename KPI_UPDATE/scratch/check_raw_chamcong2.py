import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb_raw = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
for sname in wb_raw.sheetnames:
    ws = wb_raw[sname]
    print(f"\n--- Sheet: {sname} (max_row={ws.max_row}, max_col={ws.max_column}) ---")
    for r in range(1, 6):
        print(f"Row {r}: {[ws.cell(r, c).value for c in range(1, min(12, ws.max_column+1))]}")

wb_orig = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=True)
print("\n--- Original Template 'BẢNG LƯƠNG THÁNG 8 2026.xlsx' sheets: ---", wb_orig.sheetnames)
ws_nc = wb_orig['Ngày công']
print(f"Original 'Ngày công': max_row={ws_nc.max_row}")
for r in range(1, 6):
    print(f"Row {r}: {[ws_nc.cell(r, c).value for c in range(1, 10)]}")

ws_gc = wb_orig['Giờ công']
print(f"\nOriginal 'Giờ công': max_row={ws_gc.max_row}")
for r in range(1, 6):
    print(f"Row {r}: {[ws_gc.cell(r, c).value for c in range(1, 10)]}")
