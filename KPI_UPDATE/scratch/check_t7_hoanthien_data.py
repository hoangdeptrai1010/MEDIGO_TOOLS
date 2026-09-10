import openpyxl

for fn in ['thang7/bangluong_thang7_hoanthien.xlsx', 'thang7/minikat']:
    pass

wb = openpyxl.load_workbook('thang7/bangluong_thang7_hoanthien.xlsx', data_only=True)
for sn in ['MiniKat - HN', 'MiniKat - HCM', 'whatsapp', 'Thưởng CK', 'KPI', 'Cận date']:
    ws = wb[sn]
    non_zeros = []
    for r in range(2, min(50, ws.max_row + 1)):
        v = [ws.cell(r, c).value for c in range(1, min(12, ws.max_column + 1))]
        if any(isinstance(x, (int, float)) and x > 0 for x in v):
            non_zeros.append((r, v))
    print(f"Sheet {sn}: total rows={ws.max_row}, non-zero sample rows={len(non_zeros)}")
    for r, v in non_zeros[:3]:
        print(f"  r{r}: {v}")
