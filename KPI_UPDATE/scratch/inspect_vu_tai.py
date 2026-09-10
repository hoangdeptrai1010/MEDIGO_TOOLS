import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

print("================================================================================")
print("CHI TIẾT VŨ TÀI TRONG BAOCAOKPI_THANG8_HOANTHIEN.XLSX")
print("================================================================================")

wb_kpi = openpyxl.load_workbook(r'baocaokpi_thang8_hoanthien.xlsx', data_only=True)
wb_kpi_f = openpyxl.load_workbook(r'baocaokpi_thang8_hoanthien.xlsx', data_only=False)

for sname in ['Dự án T8', 'kpi dược sĩ', 'data']:
    if sname in wb_kpi.sheetnames:
        ws = wb_kpi[sname]
        ws_f = wb_kpi_f[sname]
        print(f"\n--- Sheet: {sname} ---")
        for r in range(1, ws.max_row + 1):
            row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            row_f = [ws_f.cell(r, c).value for c in range(1, ws.max_column + 1)]
            if any('Tài' in str(v) or 'Phan Công Vũ Tài' in str(v) for v in row_vals):
                print(f"Row {r} HEADERS: {[ws.cell(2, c).value for c in range(1, len(row_vals)+1)]}")
                print(f"Row {r} FORMULAS: {row_f}")
                print(f"Row {r} VALUES:   {row_vals}")

print("\n================================================================================")
print("CHI TIẾT VŨ TÀI TRONG BANGLUONG_THANG8_HOANTHIEN.XLSX")
print("================================================================================")

wb_bl = openpyxl.load_workbook(r'thang8\bangluong_thang8_hoanthien.xlsx', data_only=True)
wb_bl_f = openpyxl.load_workbook(r'thang8\bangluong_thang8_hoanthien.xlsx', data_only=False)

for sname in ['BẢNG LƯƠNG', 'Dự án', 'Thưởng CK', 'KPI']:
    if sname in wb_bl.sheetnames:
        ws = wb_bl[sname]
        ws_f = wb_bl_f[sname]
        print(f"\n--- Sheet: {sname} ---")
        for r in range(1, ws.max_row + 1):
            row_vals = [ws.cell(r, c).value for c in range(1, min(45, ws.max_column + 1))]
            row_f = [ws_f.cell(r, c).value for c in range(1, min(45, ws.max_column + 1))]
            if any('Tài' in str(v) or 'Phan Công Vũ Tài' in str(v) for v in row_vals):
                print(f"Row {r} VALUES:   {row_vals}")
                print(f"Row {r} FORMULAS: {row_f}")
