import openpyxl

wb_july = openpyxl.load_workbook('thang7/BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=False)
ws_bl = wb_july['BẢNG LƯƠNG']

print("=== BẢNG LƯƠNG ROW 1 & 2 HEADERS ===")
for c in range(1, 30):
    print(f"Col {c} ({openpyxl.utils.get_column_letter(c)}): {repr(ws_bl.cell(1, c).value)} | {repr(ws_bl.cell(2, c).value)}")

print("\n=== BẢNG LƯƠNG ROW 3 FORMULAS ===")
for c in range(1, 30):
    val = ws_bl.cell(3, c).value
    if val is not None:
        print(f"Col {c} ({openpyxl.utils.get_column_letter(c)}): {repr(val)}")
