import openpyxl

wb = openpyxl.load_workbook('NHÀ THUỐC THÁNG 8 2026.xlsx', data_only=True)
ws_ds = wb['kpi dược sĩ']
ws_da = wb['Dự án T8']

da_branch = {}
for r in range(3, ws_da.max_row+1):
    nv = ws_da.cell(r, 2).value
    b = ws_da.cell(r, 1).value
    if nv: da_branch[nv] = b

mismatches = []
for r in range(3, ws_ds.max_row+1):
    nv = ws_ds.cell(r, 2).value
    b = ws_ds.cell(r, 1).value
    if nv and nv in da_branch:
        if da_branch[nv] != b:
            mismatches.append((nv, b, da_branch[nv]))
            print(f"Branch mismatch for {nv}: KPI Dược Sĩ='{b}', Dự án T8='{da_branch[nv]}'")

if not mismatches:
    print("No other branch mismatches found!")
