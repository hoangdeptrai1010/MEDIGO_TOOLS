import openpyxl

wb7 = openpyxl.load_workbook('thang7/BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=True)
ws7_nc = wb7['Ngày công']

print("=== July Ngày công rows 1550-1560 ===")
for r in range(1550, 1561):
    vals = [ws7_nc.cell(r, c).value for c in range(1, 9)]
    if any(v is not None for v in vals):
        print(f"R{r}: {vals}")

wb8 = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=True)
ws8_nc = wb8['Ngày công']
print("\n=== August Ngày công rows 1550-1560 ===")
for r in range(1550, 1561):
    vals = [ws8_nc.cell(r, c).value for c in range(1, 9)]
    if any(v is not None for v in vals):
        print(f"R{r}: {vals}")

print(f"\nAugust Ngày công row 2677: {[ws8_nc.cell(2677, c).value for c in range(1, 9)]}")
print(f"August Ngày công row 2678: {[ws8_nc.cell(2678, c).value for c in range(1, 9)]}")
