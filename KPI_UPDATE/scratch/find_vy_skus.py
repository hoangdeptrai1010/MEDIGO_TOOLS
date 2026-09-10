import openpyxl

wb = openpyxl.load_workbook('thang7/DATAKIOT/hoadon.xlsx', data_only=True, read_only=True)
ws = wb.active
headers = next(ws.iter_rows(max_row=1, values_only=True))
col_map = {h: idx for idx, h in enumerate(headers) if h}

c_seller = col_map.get('Người bán', 6)
c_date = col_map.get('Thời gian', 2)
c_sku = col_map.get('Mã hàng', 13)
c_name = col_map.get('Tên hàng', 15)
c_exp = col_map.get('Hạn sử dụng', 18)
c_val = col_map.get('Thành tiền', 25)

vy_candidates = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r: continue
    s = str(r[c_seller] or '').strip()
    if s == 'Cù Thị Tường Vy':
        exp = r[c_exp]
        val = float(r[c_val] or 0)
        if exp and val > 0 and str(exp)[:10] <= '2027-02-28':
            vy_candidates.append({
                'date': str(r[c_date])[:10],
                'exp': str(exp)[:10],
                'sku': r[c_sku],
                'name': r[c_name],
                'val': val
            })

target = 4911200.0
print(f"Total candidates: {len(vy_candidates)}, sum: {sum(c['val'] for c in vy_candidates)}")

# Group by SKU
by_sku = {}
for c in vy_candidates:
    sku = c['sku']
    if sku not in by_sku:
        by_sku[sku] = {'val': 0, 'items': []}
    by_sku[sku]['val'] += c['val']
    by_sku[sku]['items'].append(c)

for sku, d in by_sku.items():
    print(f"SKU {sku} (HSD {d['items'][0]['exp']}): {d['val']:>10,.0f}đ - {d['items'][0]['name']}")
