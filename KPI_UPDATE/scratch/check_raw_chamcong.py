import openpyxl
import sys
from collections import defaultdict

sys.stdout.reconfigure(encoding='utf-8')

wb_raw = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
print('Sheet names in BangChiTietChamCong_thang8.xlsx:', wb_raw.sheetnames)
ws_raw = wb_raw.active
print(f'Max row in BangChiTietChamCong: {ws_raw.max_row}')
for r in range(1, 10):
    print(r, [ws_raw.cell(r, c).value for c in range(1, 12)])

# Calculate raw total hours and shifts from BangChiTietChamCong_thang8.xlsx
raw_hours_by_person = defaultdict(float)
raw_shifts_by_person = defaultdict(int)
raw_dates_by_person = defaultdict(set)

# Find header row
header_row = 1
for r in range(1, 10):
    vals = [str(ws_raw.cell(r, c).value) for c in range(1, 15)]
    if any('Tên nhân viên' in v or 'Mã nhân viên' in v or 'Nhân viên' in v for v in vals):
        header_row = r
        break

print(f"Header row in BangChiTietChamCong: {header_row}")
headers = [ws_raw.cell(header_row, c).value for c in range(1, 20)]
print("Headers:", headers)
