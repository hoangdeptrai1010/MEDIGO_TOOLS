import openpyxl

wb_aug = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=False)
print("Sheets in thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx:")
for s in wb_aug.sheetnames:
    ws = wb_aug[s]
    print(f"  {s:<20}: {ws.max_row} rows, {ws.max_column} cols")
