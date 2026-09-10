# -*- coding: utf-8 -*-
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/BANGLUONGTHANG8.xlsx', data_only=True)
ws_bl = wb['BẢNG LƯƠNG']

print(f"{'STT':<3} | {'Chi nhánh':<15} | {'Họ và tên':<24} | {'Chức danh':<10} | {'Giờ ngày (G)':<12} | {'Giờ đêm (H)':<12} | {'Ca ngày (J)':<12} | {'Ca đêm (K)':<10} | {'PC Đêm (X)':<14}")
print("-" * 125)

for r in range(3, 59):
    stt = ws_bl.cell(r, 1).value
    cn = ws_bl.cell(r, 2).value or ''
    name = ws_bl.cell(r, 3).value or ''
    role = str(ws_bl.cell(r, 4).value or '').strip()
    g = ws_bl.cell(r, 7).value or 0
    h = ws_bl.cell(r, 8).value or 0
    j = ws_bl.cell(r, 10).value or 0
    k = ws_bl.cell(r, 11).value or 0
    x = ws_bl.cell(r, 24).value or 0
    
    g_str = f"{g:,.2f}h" if isinstance(g, (int, float)) else str(g)
    h_str = f"{h:,.2f}h" if isinstance(h, (int, float)) else str(h)
    x_str = f"{x:,.0f} đ" if isinstance(x, (int, float)) and x > 0 else "-"
    
    print(f"{stt:<3} | {cn:<15} | {name:<24} | {role:<10} | {g_str:<12} | {h_str:<12} | {j:<12} | {k:<10} | {x_str:<14}")
