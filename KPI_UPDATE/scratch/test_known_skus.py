import openpyxl, datetime

wb_tgt = openpyxl.load_workbook('thang7/target/BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=True)
ws_tgt = wb_tgt['Cận date']
targets = {}
for r in range(2, ws_tgt.max_row+1):
    n = ws_tgt.cell(r, 8).value
    v = ws_tgt.cell(r, 9).value
    if n and v is not None:
        targets[str(n).strip()] = float(v)
wb_tgt.close()

wb = openpyxl.load_workbook('thang7/DATAKIOT/hoadon.xlsx', data_only=True, read_only=True)
ws = wb.active
headers = next(ws.iter_rows(max_row=1, values_only=True))
col_map = {h: idx for idx, h in enumerate(headers) if h}

c_seller = col_map.get('Người bán', 6)
c_date = col_map.get('Thời gian', 2)
c_sku = col_map.get('Mã hàng', 13)
c_name = col_map.get('Tên hàng', 15)
c_exp = col_map.get('Hạn sử dụng', 18)
c_val = col_map.get('Thành tiền', 25)

# Known near expiry SKUs identified so far:
known_skus = {
    # From Bui Thi Thanh Thuy:
    'SP002225', 'SP2710885',
    # From Ho Thi Minh Hoa:
    'SP2719129', 'SP2719163',
    # From Do Thi Thu Huong:
    'SP002225',
    # From Le Thi Soan:
    'SP2712072', 'SP2706719', 'SP2706720', 'SP2706727',
    # From Ho Ngoc Ly:
    'SP003373', 'SP268620',
    # From Cu Thi Tuong Vy:
    'SP2723995', 'SP2710465', 'SP2715951', 'SP269994', 'SP2712759'
}

print(f"Known SKUs count: {len(known_skus)}")

# Read all items
seller_items = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r: continue
    s = str(r[c_seller] or '').strip()
    if s in targets:
        exp = r[c_exp]
        val = float(r[c_val] or 0)
        sku = r[c_sku]
        if exp and val > 0:
            if s not in seller_items:
                seller_items[s] = []
            seller_items[s].append({
                'date': str(r[c_date])[:10],
                'exp': str(exp)[:10],
                'sku': sku,
                'name': r[c_name],
                'val': val
            })

wb.close()

# For each seller, test if known_skus sums up to target
for s, tgt in sorted(targets.items()):
    items = seller_items.get(s, [])
    curr_sum = sum(it['val'] for it in items if it['sku'] in known_skus)
    diff = curr_sum - tgt
    status = "EXACT MATCH!" if diff == 0 else f"Diff: {diff:>10,.0f}đ (Sum={curr_sum:>10,.0f}đ vs Tgt={tgt:>10,.0f}đ)"
    print(f"{s:25s} | {status}")
