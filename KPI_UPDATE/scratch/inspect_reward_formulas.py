import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb7 = openpyxl.load_workbook(r'thang7\target\BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=False)
ws7 = wb7['BẢNG LƯƠNG']

print("--- THÁNG 7 TARGET BẢNG LƯƠNG FORMULAS (Row 3, 5, 8, 11) ---")
for r in [3, 5, 8, 11]:
    name = ws7.cell(r, 3).value
    for col_letter in ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI']:
        c = openpyxl.utils.column_index_from_string(col_letter)
        f = ws7.cell(r, c).value
        print(f"Row {r} ({name}) {col_letter} ({ws7.cell(2, c).value}): {f}")

wb8 = openpyxl.load_workbook(r'thang8\bangluong_thang8_hoanthien.xlsx', data_only=False)
ws8 = wb8['BẢNG LƯƠNG']

print("\n--- THÁNG 8 HOÀN THIỆN BẢNG LƯƠNG FORMULAS (Row 3, 7, 8, 11) ---")
for r in [3, 7, 8, 11]:
    name = ws8.cell(r, 3).value
    for col_letter in ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI']:
        c = openpyxl.utils.column_index_from_string(col_letter)
        f = ws8.cell(r, c).value
        print(f"Row {r} ({name}) {col_letter} ({ws8.cell(2, c).value}): {f}")
