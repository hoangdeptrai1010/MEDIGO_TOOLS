import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb7 = openpyxl.load_workbook(r'thang7\target\BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=False)
wb7_v = openpyxl.load_workbook(r'thang7\target\BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=True)

ws = wb7['BẢNG LƯƠNG']
wsv = wb7_v['BẢNG LƯƠNG']

print("--- THÁNG 7 TARGET BẢNG LƯƠNG: CỘT G, H, J, K, R ---")
print(f"{'Tên':<25} | {'Giờ ngày G':<12} | {'Giờ đêm H':<12} | {'Col J Formula':<35} | {'J Val':<8} | {'K Val':<8} | {'R Val':<8}")
print("-" * 125)

for r in range(3, 20):
    name = ws.cell(r, 3).value
    g = wsv.cell(r, 7).value
    h = wsv.cell(r, 8).value
    jf = str(ws.cell(r, 10).value)
    jv = wsv.cell(r, 10).value
    kv = wsv.cell(r, 11).value
    rv = wsv.cell(r, 18).value
    
    g_s = f"{g:,.2f}" if isinstance(g, (int, float)) else str(g)
    h_s = f"{h:,.2f}" if isinstance(h, (int, float)) else str(h)
    print(f"{str(name):<25} | {g_s:<12} | {h_s:<12} | {jf:<35} | {str(jv):<8} | {str(kv):<8} | {str(rv):<8}")
