import openpyxl

wb_aug = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=False)

for sname in ['Giờ công', 'Ngày công', 'Tăng ca lễ']:
    if sname in wb_aug.sheetnames:
        ws = wb_aug[sname]
        print(f"=== AUGUST EXISTING: {sname} ===")
        print(f"max_row = {ws.max_row}, max_column = {ws.max_column}")
        for r in range(1, 10):
            print(f"R{r}: {[ws.cell(r, c).value for c in range(1, 17)]}")
