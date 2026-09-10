# -*- coding: utf-8 -*-
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx', data_only=True)
print("=== SHEET 'kpi dược sĩ' in baocaokpi_thang8_hoanthien.xlsx ===")
ws = wb['kpi dược sĩ']
for r in range(1, 15):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None]
    if vals:
        print(f"Row {r:2d}: {vals}")
