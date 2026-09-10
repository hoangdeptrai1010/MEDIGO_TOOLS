import openpyxl, glob, os, sys

sys.stdout.reconfigure(encoding='utf-8')

print("=== FINDING ALL RELEVANT EXCEL FILES ===")
files = glob.glob('**/*.xlsx', recursive=True)
for f in files:
    if 'nhà thuốc' in f.lower() or 'nha thuoc' in f.lower() or 'kpi' in f.lower() or 'thang8' in f.lower():
        print(f"File: {f}")

goc_file = 'goc/NHÀ THUỐC THÁNG 8 2026.xlsx'
if os.path.exists(goc_file):
    print(f"\n{'='*20} INSPECTING {goc_file} {'='*20}")
    wb = openpyxl.load_workbook(goc_file, data_only=False)
    print("Sheets:", wb.sheetnames)
    
    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f"\n--- Sheet: '{sname}' (Max row: {ws.max_row}, Max col: {ws.max_column}) ---")
        for r in range(1, min(ws.max_row + 1, 6)):
            row_vals = []
            for c in range(1, min(ws.max_column + 1, 18)):
                val = ws.cell(r, c).value
                val_str = str(val)[:40] if val is not None else ""
                row_vals.append(f"{openpyxl.utils.get_column_letter(c)}{r}: {val_str}")
            print(f"Row {r}: {', '.join(row_vals[:10])}")
    wb.close()
