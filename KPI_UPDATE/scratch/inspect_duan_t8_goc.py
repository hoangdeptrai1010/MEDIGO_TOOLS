import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('goc/NHÀ THUỐC THÁNG 8 2026.xlsx', data_only=False)
ws_da = wb['Dự án T8']

print("=== SHEET DỰ ÁN T8 IN GOC/NHÀ THUỐC THÁNG 8 2026.XLSX ===")
print("Headers:")
for c in range(1, 20):
    val = ws_da.cell(2, c).value
    print(f"  Col {openpyxl.utils.get_column_letter(c)} ({c}): {val}")

print("\nFormulas in Rows 3..6:")
for r in range(3, 7):
    print(f"\n--- Row {r}: {ws_da.cell(r, 1).value} - {ws_da.cell(r, 2).value} ({ws_da.cell(r, 3).value}) ---")
    for c in range(1, 15):
        val = ws_da.cell(r, c).value
        print(f"  Col {openpyxl.utils.get_column_letter(c)}: {val}")

wb.close()
