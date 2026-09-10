import openpyxl

wb = openpyxl.load_workbook('thang7/BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=False)

sheets_to_check = ['Bảng đánh giá', 'BẢNG LƯƠNG', 'MiniKat - HN', 'MiniKat - HCM', 'whatsapp', 'Cận date', 'KPI', 'Thưởng CK', 'Dự án']

for name in sheets_to_check:
    if name in wb.sheetnames:
        ws = wb[name]
        print(f"\n==========================================")
        print(f"=== SHEET: {name} (max_row={ws.max_row}, max_column={ws.max_column}) ===")
        print(f"==========================================")
        for r in range(1, min(6, ws.max_row + 1)):
            row_vals = [ws.cell(r, c).value for c in range(1, min(20, ws.max_column + 1))]
            if any(v is not None for v in row_vals):
                print(f"  Row {r}: {row_vals}")
