import openpyxl

wb = openpyxl.load_workbook('scratch/test_baug_recalc.xlsx', data_only=True)
ws_nc = wb['Ngày công']

print("=== Ngày công summary table (Cols K-O) in test_baug_recalc.xlsx ===")
for r in range(2, 20):
    k = ws_nc.cell(r, 11).value
    l = ws_nc.cell(r, 12).value
    m = ws_nc.cell(r, 13).value
    n = ws_nc.cell(r, 14).value
    o = ws_nc.cell(r, 15).value
    print(f"R{r}: CN={k}, Ten={l}, NgayCongChuan(M)={m}, GioCong(N)={n}, TB(O)={o}")
