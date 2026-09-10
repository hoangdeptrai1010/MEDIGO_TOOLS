# -*- coding: utf-8 -*-
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx', data_only=True)
print("=== SHEETS IN baocaokpi_thang8_hoanthien.xlsx ===")
print(wb.sheetnames)

staff_to_check = ['Phan Công Vũ Tài', 'Trần Thị Kim Khánh', 'Đinh Thị Lan Anh', 'Vũ Thanh Hằng', 'Nguyễn Trần Ngọc Phương']

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n==================== SHEET: {sname} ====================")
    headers = [ws.cell(1, c).value or ws.cell(2, c).value for c in range(1, min(20, ws.max_column + 1))]
    print(f"Headers: {headers[:12]}")
    for r in range(1, ws.max_row + 1):
        row_str = " ".join([str(ws.cell(r, c).value or '') for c in range(1, min(20, ws.max_column + 1))])
        for st in staff_to_check:
            if st in row_str:
                row_vals = [f"C{c}({ws.cell(2, c).value or ws.cell(1, c).value}): {ws.cell(r, c).value}" for c in range(1, min(15, ws.max_column + 1)) if ws.cell(r, c).value is not None]
                print(f"  Row {r:3d} [{st}]: {row_vals}")
