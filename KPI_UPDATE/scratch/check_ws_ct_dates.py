import openpyxl

wb_cc = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
ws_ct = wb_cc['Bảng chi tiết chấm công']

print(f"ws_ct max_row: {ws_ct.max_row}")
has_date = 0
no_date = 0
for r in range(4, ws_ct.max_row + 1):
    d = ws_ct.cell(r, 6).value # Col F
    if d is not None and str(d).strip() != '':
        has_date += 1
    else:
        no_date += 1

print(f"Has date (Col F): {has_date}")
print(f"No date (Col F): {no_date}")
print(f"Total rows (from row 4): {ws_ct.max_row - 3}")
