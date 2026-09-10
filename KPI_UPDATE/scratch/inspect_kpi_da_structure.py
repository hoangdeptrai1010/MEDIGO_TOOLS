import openpyxl

wb = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=False)

for sname in ['KPI', 'Dự án', 'Thưởng CK', 'Cận date']:
    ws = wb[sname]
    print(f"=== {sname} headers (row 1) and row 2 ===")
    print("R1:", [ws.cell(1, c).value for c in range(1, 15)])
    print("R2:", [ws.cell(2, c).value for c in range(1, 15)])
