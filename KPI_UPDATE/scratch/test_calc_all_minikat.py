import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb_inv = openpyxl.load_workbook('thang8/DATA/DanhSachChiTietHoaDon_3182026.xlsx', read_only=True)
ws_inv = wb_inv.active

SKUS_PARTY_SMART = {'SP017162', 'SP2723878', 'SP2723879'}
SKUS_KAT = {'SP2725289', 'SP2725285', 'SP2725291', 'SP2725287', 'SP2725353', 'SP2725351'}
SKUS_LADYCARE = {'SP2723124', 'SP2723125', 'SP2723127'}
SKUS_AVC = {'SP2723883', 'SP2722826', 'SP2722817', 'SP2725666'}

seller_sales = {}
branch_sales = {}

for idx, r in enumerate(ws_inv.iter_rows(values_only=True)):
    if idx == 0 or not r or len(r) < 27: continue
    branch = str(r[0]).strip() if r[0] else ''
    seller = str(r[7]).strip() if r[7] else ''
    sku = str(r[14]).strip() if r[14] else ''
    try:
        qty = float(r[21] or 0)
        tt = float(r[26] or 0)
    except:
        continue
        
    if sku in (SKUS_PARTY_SMART | SKUS_KAT | SKUS_LADYCARE | SKUS_AVC):
        if seller not in seller_sales:
            seller_sales[seller] = {'branch': branch, 'ps_qty': 0.0, 'kat_qty': 0.0, 'lady_qty': 0.0, 'lady_rev': 0.0, 'kat_rev': 0.0, 'avc_qty': 0.0, 'avc_rev': 0.0}
        if branch not in branch_sales:
            branch_sales[branch] = {'ps_qty': 0.0, 'kat_qty': 0.0, 'lady_rev': 0.0}
            
        if sku in SKUS_PARTY_SMART:
            seller_sales[seller]['ps_qty'] += qty
            branch_sales[branch]['ps_qty'] += qty
        elif sku in SKUS_KAT:
            seller_sales[seller]['kat_qty'] += qty
            seller_sales[seller]['kat_rev'] += tt
            branch_sales[branch]['kat_qty'] += qty
        elif sku in SKUS_LADYCARE:
            seller_sales[seller]['lady_qty'] += qty
            seller_sales[seller]['lady_rev'] += tt
            branch_sales[branch]['lady_rev'] += tt
        elif sku in SKUS_AVC:
            seller_sales[seller]['avc_qty'] += qty
            seller_sales[seller]['avc_rev'] += tt

# Check returns
wb_ret = openpyxl.load_workbook('thang8/DATA/DanhSachChiTietTraHang_3182026.xlsx', read_only=True)
ws_ret = wb_ret.active
for idx, r in enumerate(ws_ret.iter_rows(values_only=True)):
    if idx == 0 or not r or len(r) < 37: continue
    branch = str(r[0]).strip() if r[0] else ''
    seller = str(r[6]).strip() if r[6] else ''
    sku = str(r[28]).strip() if r[28] else ''
    try:
        qty = float(r[35] or 0)
        price = float(r[36] or 0)
    except:
        continue
    if sku in SKUS_PARTY_SMART and seller in seller_sales: seller_sales[seller]['ps_qty'] -= qty
    if sku in SKUS_KAT and seller in seller_sales: seller_sales[seller]['kat_qty'] -= qty
    if sku in SKUS_LADYCARE and seller in seller_sales: 
        seller_sales[seller]['lady_qty'] -= qty
        seller_sales[seller]['lady_rev'] -= qty * price
    if sku in SKUS_AVC and seller in seller_sales: 
        seller_sales[seller]['avc_qty'] -= qty
        seller_sales[seller]['avc_rev'] -= qty * price
    
    if sku in SKUS_PARTY_SMART and branch in branch_sales: branch_sales[branch]['ps_qty'] -= qty
    if sku in SKUS_KAT and branch in branch_sales: branch_sales[branch]['kat_qty'] -= qty
    if sku in SKUS_LADYCARE and branch in branch_sales: branch_sales[branch]['lady_rev'] -= qty * price

print('=== QUALIFYING SELLERS ===')
for s, d in seller_sales.items():
    bonus = 0
    reasons = []
    is_hn = 'Hàng Bông' in d['branch'] or 'Đường Láng' in d['branch']
    if is_hn:
        if d['kat_qty'] >= 17: 
            bonus += d['kat_qty'] * 30000
            reasons.append('KAT')
        elif d['kat_qty'] >= 15: 
            bonus += d['kat_qty'] * 25000
            reasons.append('KAT')
        if d['ps_qty'] >= 80: 
            bonus += d['ps_qty'] * 6000
            reasons.append('PS')
        if d['lady_rev'] > 7000000: 
            bonus += d['lady_rev'] * 0.06
            reasons.append('Lady')
        elif d['lady_rev'] > 5000000: 
            bonus += d['lady_rev'] * 0.04
            reasons.append('Lady')
        if d['avc_rev'] > 9000000: 
            bonus += d['avc_qty'] * 6000
            reasons.append(f"AVC 6k ({d['avc_qty'] * 6000:,.0f})")
        elif d['avc_rev'] > 7000000: 
            bonus += d['avc_qty'] * 4000
            reasons.append(f"AVC 4k ({d['avc_qty'] * 4000:,.0f})")
    else:
        if d['kat_qty'] >= 10: 
            bonus += d['kat_qty'] * 30000
            reasons.append('KAT')
        elif d['kat_qty'] >= 7: 
            bonus += d['kat_qty'] * 25000
            reasons.append('KAT')
        if d['ps_qty'] >= 80: 
            bonus += d['ps_qty'] * 6000
            reasons.append('PS')
        if d['lady_rev'] > 7000000: 
            bonus += d['lady_rev'] * 0.06
            reasons.append('Lady')
        elif d['lady_rev'] > 5000000: 
            bonus += d['lady_rev'] * 0.04
            reasons.append('Lady')
        elif d['lady_rev'] > 3000000: 
            bonus += d['lady_rev'] * 0.03
            reasons.append('Lady')
        
    if bonus > 0:
        print(f"{s:<25} ({d['branch']}) : Bonus = {bonus:,.0f} đ | Reasons: {reasons}")

print('\n=== QUALIFYING BRANCHES ===')
for b, d in branch_sales.items():
    is_hn = 'Hàng Bông' in b or 'Đường Láng' in b
    b_bonus = 0
    b_reasons = []
    if d['kat_qty'] >= 17: 
        b_bonus += 300000
        b_reasons.append(f"KAT 17h ({d['kat_qty']})")
    if d['ps_qty'] >= 180: 
        b_bonus += 300000
        b_reasons.append(f"PS 180h ({d['ps_qty']})")
    elif d['ps_qty'] >= 150: 
        b_bonus += 150000
        b_reasons.append(f"PS 150h ({d['ps_qty']})")
    if is_hn:
        if d['lady_rev'] > 30000000: 
            b_bonus += 700000
            b_reasons.append('Lady 30M')
        elif d['lady_rev'] > 22000000: 
            b_bonus += 500000
            b_reasons.append('Lady 22M')
    else:
        if d['lady_rev'] > 22000000: 
            b_bonus += 500000
            b_reasons.append('Lady 22M')
        elif d['lady_rev'] > 15000000: 
            b_bonus += 200000
            b_reasons.append('Lady 15M')
    print(f"{b:<30} : Bonus = {b_bonus:,.0f} đ | KAT={d['kat_qty']:.0f} | PS={d['ps_qty']:.0f} | Lady={d['lady_rev']:,.0f} | Reasons: {b_reasons}")
