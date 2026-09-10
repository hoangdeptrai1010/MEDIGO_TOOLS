# -*- coding: utf-8 -*-
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/BANGLUONGTHANG8.xlsx', data_only=True)
ws_gc = wb['Giờ công']

print("=== CHECKING ALL ROWS IN GIỜ CÔNG ===")
for r in range(2, ws_gc.max_row+1):
    cn = ws_gc.cell(r, 1).value
    name = ws_gc.cell(r, 2).value
    ca = str(ws_gc.cell(r, 3).value or '').strip()
    raw = ws_gc.cell(r, 4).value
    hours = ws_gc.cell(r, 5).value
    
    if not name or ca == '-':
        continue
        
    has_qua_dem = 'qua đêm' in ca.lower()
    has_dem = 'đêm' in ca.lower()
    
    # Check if there is any shift with 'đêm' but not 'qua đêm' or vice-versa
    if has_dem != has_qua_dem:
        print(f"Row {r:3d}: CN={cn} | Name={name} | Shift={ca} | Hours={hours} | has_dem={has_dem} | has_qua_dem={has_qua_dem}")
