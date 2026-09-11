import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('d:/MEDIGO/KPI_UPDATE/thang9/data/kpi_thang9/e_xuat_KPI_Quy_3.26.xlsx', data_only=True)
ws = wb['T9.26']
for r in range(1, ws.max_row + 1):
    c1 = ws.cell(r, 1).value
    c2 = ws.cell(r, 2).value
    c7 = ws.cell(r, 7).value
    c17 = ws.cell(r, 17).value
    c18 = ws.cell(r, 18).value
    c19 = ws.cell(r, 19).value
    print(f"Row {r:2d}: {str(c1):12s} | {str(c2):25s} | C7 (daily)={c7} | C17 (T8)={c17} | C18 (T9)={c18} | C19 (NT)={c19}")
