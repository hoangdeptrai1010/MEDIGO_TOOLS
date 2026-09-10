import openpyxl
import sys
import io

if sys.stdout.encoding != 'utf-8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb_goc = openpyxl.load_workbook('goc/NHÀ THUỐC THÁNG 8 2026.xlsx', data_only=False)
wb_out = openpyxl.load_workbook('TOOL_KPISHEET/output/NHÀ THUỐC THÁNG 8 2026.xlsx', data_only=False)

print("=== COMPARING SHEET 'kpi dược sĩ' ===")
ws_goc_ds = wb_goc['kpi dược sĩ']
ws_out_ds = wb_out['kpi dược sĩ']

print("Row 2 Headers comparison:")
for c in range(1, 32):
    g_val = ws_goc_ds.cell(2, c).value
    o_val = ws_out_ds.cell(2, c).value
    match = "✅" if g_val == o_val else f"❌ [Goc: '{g_val}' vs Out: '{o_val}']"
    print(f"  Col {c:2d} ({openpyxl.utils.get_column_letter(c)}): {match}")

print("\nRow 3 (First Staff) comparison:")
for c in range(1, 32):
    g_val = ws_goc_ds.cell(3, c).value
    o_val = ws_out_ds.cell(3, c).value
    match = "✅" if str(g_val) == str(o_val) else f"❌ [Goc: {g_val} vs Out: {o_val}]"
    print(f"  Col {c:2d} ({openpyxl.utils.get_column_letter(c)}): {match}")

print("\n=== COMPARING SHEET 'kpi nhà thuốc' ===")
ws_goc_nt = wb_goc['kpi nhà thuốc']
ws_out_nt = wb_out['kpi nhà thuốc']

print("Row 2 Headers comparison:")
for c in range(1, 33):
    g_val = ws_goc_nt.cell(2, c).value
    o_val = ws_out_nt.cell(2, c).value
    match = "✅" if g_val == o_val else f"❌ [Goc: '{g_val}' vs Out: '{o_val}']"
    print(f"  Col {c:2d} ({openpyxl.utils.get_column_letter(c)}): {match}")

print("\nRow 3 (First Store) comparison:")
for c in range(1, 33):
    g_val = ws_goc_nt.cell(3, c).value
    o_val = ws_out_nt.cell(3, c).value
    match = "✅" if str(g_val) == str(o_val) else f"❌ [Goc: {g_val} vs Out: {o_val}]"
    print(f"  Col {c:2d} ({openpyxl.utils.get_column_letter(c)}): {match}")

print("\n=== COMPARING SHEET 'Dự án T8' ===")
ws_goc_p = wb_goc['Dự án T8']
ws_out_p = wb_out['Dự án T8']

print("Row 2 Headers comparison:")
for c in range(1, 14):
    g_val = ws_goc_p.cell(2, c).value
    o_val = ws_out_p.cell(2, c).value
    match = "✅" if g_val == o_val else f"❌ [Goc: '{g_val}' vs Out: '{o_val}']"
    print(f"  Col {c:2d} ({openpyxl.utils.get_column_letter(c)}): {match}")

print("\nRow 3 (First Project Row) comparison:")
for c in range(1, 14):
    g_val = ws_goc_p.cell(3, c).value
    o_val = ws_out_p.cell(3, c).value
    match = "✅" if str(g_val) == str(o_val) else f"❌ [Goc: {g_val} vs Out: {o_val}]"
    print(f"  Col {c:2d} ({openpyxl.utils.get_column_letter(c)}): {match}")
