import openpyxl

wb_aug = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=False)
ws_bl = wb_aug['BẢNG LƯƠNG']

print("Row 3 formulas in thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx:")
for c in range(1, 45):
    val = ws_bl.cell(3, c).value
    if val is not None:
        print(f"Col {c} ({openpyxl.utils.get_column_letter(c)}): {repr(val)}")

print("\nRow 4 formulas:")
for c in range(1, 45):
    val = ws_bl.cell(4, c).value
    if val is not None:
        print(f"Col {c} ({openpyxl.utils.get_column_letter(c)}): {repr(val)}")
