import openpyxl, datetime

inv_file = 'thang8/DATA/DanhSachChiTietHoaDon_3182026.xlsx'
ret_file = 'thang8/DATA/DanhSachChiTietTraHang_3182026.xlsx'

SKUS_PARTY_SMART = {'SP017162'}
SKUS_KAT = {'SP2725289', 'SP2725285', 'SP2725291', 'SP2725287', 'SP2725353', 'SP2725351'}
SKUS_LADYCARE = {'SP2723124', 'SP2723125', 'SP2723127'}

HN_BRANCHES = {'Hàng Bông', 'Đường Láng', 'NT Hàng Bông 247', 'NT Đường Láng 247'}

# 1. Parse returns first
returns_by_seller_sku = {} # (seller, sku) -> qty, net_amount
wb_ret = openpyxl.load_workbook(ret_file, read_only=True, data_only=True)
ws_ret = wb_ret.active
ret_iter = ws_ret.iter_rows(values_only=True)
ret_hdr = next(ret_iter)
for r in ret_iter:
    if not r or len(r) < 37: continue
    seller = str(r[6]).strip() if r[6] else ''
    sku = str(r[28]).strip() if r[28] else ''
    qty = float(r[35] or 0)
    price = float(r[36] or 0)
    if sku in (SKUS_PARTY_SMART | SKUS_KAT | SKUS_LADYCARE):
        k = (seller, sku)
        if k not in returns_by_seller_sku:
            returns_by_seller_sku[k] = {'qty': 0.0, 'amount': 0.0}
        returns_by_seller_sku[k]['qty'] += qty
        returns_by_seller_sku[k]['amount'] += qty * price
wb_ret.close()
print("Returns found for MiniKat SKUs:", returns_by_seller_sku)

# 2. Parse invoices
sales_by_seller = {} # (branch, seller) -> {ps_qty, kat_qty, kat_rev, lady_qty, lady_rev}
wb_inv = openpyxl.load_workbook(inv_file, read_only=True, data_only=True)
ws_inv = wb_inv.active
inv_iter = ws_inv.iter_rows(values_only=True)
inv_hdr = next(inv_iter)

for r in inv_iter:
    if not r or len(r) < 27: continue
    branch = str(r[0]).strip() if r[0] else ''
    seller = str(r[7]).strip() if r[7] else ''
    sku = str(r[14]).strip() if r[14] else ''
    qty = float(r[21] or 0)
    tt = float(r[26] or 0) # Thành tiền
    
    if sku in (SKUS_PARTY_SMART | SKUS_KAT | SKUS_LADYCARE):
        k = (branch, seller)
        if k not in sales_by_seller:
            sales_by_seller[k] = {'ps_qty': 0.0, 'kat_qty': 0.0, 'kat_rev': 0.0, 'lady_qty': 0.0, 'lady_rev': 0.0}
        
        if sku in SKUS_PARTY_SMART:
            sales_by_seller[k]['ps_qty'] += qty
        elif sku in SKUS_KAT:
            sales_by_seller[k]['kat_qty'] += qty
            sales_by_seller[k]['kat_rev'] += tt
        elif sku in SKUS_LADYCARE:
            sales_by_seller[k]['lady_qty'] += qty
            sales_by_seller[k]['lady_rev'] += tt
wb_inv.close()

# Net out returns
for (branch, seller), d in sales_by_seller.items():
    for sku in SKUS_PARTY_SMART:
        ret = returns_by_seller_sku.get((seller, sku), {'qty': 0, 'amount': 0})
        d['ps_qty'] -= ret['qty']
    for sku in SKUS_KAT:
        ret = returns_by_seller_sku.get((seller, sku), {'qty': 0, 'amount': 0})
        d['kat_qty'] -= ret['qty']
        d['kat_rev'] -= ret['amount']
    for sku in SKUS_LADYCARE:
        ret = returns_by_seller_sku.get((seller, sku), {'qty': 0, 'amount': 0})
        d['lady_qty'] -= ret['qty']
        d['lady_rev'] -= ret['amount']

print("\n--- SAMPLE SALES BY SELLER ---")
for (branch, seller), d in sorted(sales_by_seller.items())[:15]:
    print(f"{branch} | {seller}: PS={d['ps_qty']}, KAT={d['kat_qty']} (rev={d['kat_rev']}), LadyCare={d['lady_qty']} (rev={d['lady_rev']})")
