import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb7 = openpyxl.load_workbook(r'baocaokpi_thang7_hoanthien.xlsx', data_only=True)
ws7_da = wb7['Dự án T7']

print("--- baocaokpi_thang7_hoanthien.xlsx Sheet 'Dự án T7' ---")
print(f"{'Nhà thuốc':<15} | {'Nhân viên':<25} | {'CK/ngày (H)':<15} | {'Combo/ngày (I)':<15} | {'Thưởng DA (J)':<15} | {'Thưởng thêm (K)':<15} | {'Total DA (M)':<15}")
print("-" * 120)
for r in range(3, 15):
    nt = ws7_da.cell(r, 1).value
    nv = ws7_da.cell(r, 2).value
    ck_ng = ws7_da.cell(r, 8).value
    cb_ng = ws7_da.cell(r, 9).value
    t_da = ws7_da.cell(r, 10).value
    t_them = ws7_da.cell(r, 11).value
    tot_da = ws7_da.cell(r, 13).value
    
    ck_s = f"{ck_ng:,.0f}" if isinstance(ck_ng, (int, float)) else str(ck_ng)
    cb_s = f"{cb_ng:,.0f}" if isinstance(cb_ng, (int, float)) else str(cb_ng)
    tda_s = f"{t_da:,.0f}" if isinstance(t_da, (int, float)) else str(t_da)
    tthem_s = f"{t_them:,.0f}" if isinstance(t_them, (int, float)) else str(t_them)
    tot_s = f"{tot_da:,.0f}" if isinstance(tot_da, (int, float)) else str(tot_da)
    print(f"{str(nt):<15} | {str(nv):<25} | {ck_s:<15} | {cb_s:<15} | {tda_s:<15} | {tthem_s:<15} | {tot_s:<15}")

wb8 = openpyxl.load_workbook(r'baocaokpi_thang8_hoanthien.xlsx', data_only=True)
ws8_da = wb8['Dự án T8']

print("\n--- baocaokpi_thang8_hoanthien.xlsx Sheet 'Dự án T8' ---")
print(f"{'Nhà thuốc':<15} | {'Nhân viên':<25} | {'CK/ngày (H)':<15} | {'Combo/ngày (I)':<15} | {'Thưởng DA (J)':<15} | {'Thưởng thêm (K)':<15} | {'Hot Bill (L)':<15} | {'Total DA (M)':<15}")
print("-" * 135)
for r in range(3, 15):
    nt = ws8_da.cell(r, 1).value
    nv = ws8_da.cell(r, 2).value
    ck_ng = ws8_da.cell(r, 8).value
    cb_ng = ws8_da.cell(r, 9).value
    t_da = ws8_da.cell(r, 10).value
    t_them = ws8_da.cell(r, 11).value
    hb = ws8_da.cell(r, 12).value
    tot_da = ws8_da.cell(r, 13).value
    
    ck_s = f"{ck_ng:,.0f}" if isinstance(ck_ng, (int, float)) else str(ck_ng)
    cb_s = f"{cb_ng:,.0f}" if isinstance(cb_ng, (int, float)) else str(cb_ng)
    tda_s = f"{t_da:,.0f}" if isinstance(t_da, (int, float)) else str(t_da)
    tthem_s = f"{t_them:,.0f}" if isinstance(t_them, (int, float)) else str(t_them)
    hb_s = f"{hb:,.0f}" if isinstance(hb, (int, float)) else str(hb)
    tot_s = f"{tot_da:,.0f}" if isinstance(tot_da, (int, float)) else str(tot_da)
    print(f"{str(nt):<15} | {str(nv):<25} | {ck_s:<15} | {cb_s:<15} | {tda_s:<15} | {tthem_s:<15} | {hb_s:<15} | {tot_s:<15}")

