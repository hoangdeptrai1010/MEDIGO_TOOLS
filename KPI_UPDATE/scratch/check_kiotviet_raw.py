import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
ws_ct = wb['Bảng chi tiết chấm công']
ws_h = wb['Bảng chấm công theo giờ']

print("--- Inspect 'Bảng chi tiết chấm công' rows 4 to 12 ---")
for r in range(4, 13):
    print(r, [ws_ct.cell(r, c).value for c in range(1, 12)])

print("\n--- Inspect 'Bảng chấm công theo giờ' rows 4 to 10 ---")
for r in range(4, 11):
    print(r, [ws_h.cell(r, c).value for c in range(1, 15)])
