import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

# Output path
out_dir = 'd:/MEDIGO/KPI_UPDATE/thang8/output'
os.makedirs(out_dir, exist_ok=True)
out_file = os.path.join(out_dir, 'TONG_HOP_DATA_PHIEU_WHATSAPP_THANG_8_2026.xlsx')

wb = openpyxl.Workbook()

# Fonts & Colors
font_title = Font(name='Calibri', size=14, bold=True, color='1F4E79')
font_subtitle = Font(name='Calibri', size=11, italic=True, color='595959')
font_header_white = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
font_header_dark = Font(name='Calibri', size=11, bold=True, color='1E293B')
font_header_green = Font(name='Calibri', size=11, bold=True, color='166534')
font_bold = Font(name='Calibri', size=11, bold=True)
font_regular = Font(name='Calibri', size=11)

fill_navy = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
fill_soft_blue = PatternFill(start_color='DDEBF7', end_color='DDEBF7', fill_type='solid')
fill_light_blue = PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid')
fill_yellow = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
fill_emerald = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
fill_gray = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
fill_highlight = PatternFill(start_color='EBF9F1', end_color='EBF9F1', fill_type='solid')

border_thin = Border(
    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
)
border_header = Border(
    left=Side(style='thin', color='B0C4DE'), right=Side(style='thin', color='B0C4DE'),
    top=Side(style='medium', color='1F4E79'), bottom=Side(style='medium', color='1F4E79')
)
border_total = Border(
    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='1F4E79'), bottom=Side(style='double', color='1F4E79')
)

# ==============================================================================
# SHEET 1: TỔNG HỢP THƯỞNG WHATSAPP THÁNG 8 (THEO NHÂN SỰ)
# ==============================================================================
ws1 = wb.active
ws1.title = 'Tổng hợp thưởng Whatsapp'
ws1.views.sheetView[0].showGridLines = True

# Title
ws1.merge_cells('A1:G1')
c1 = ws1['A1']
c1.value = 'BẢNG TỔNG HỢP DOANH THU & THƯỞNG ĐƠN WHATSAPP THÁNG 08/2026'
c1.font = font_title
c1.alignment = Alignment(horizontal='center', vertical='center')
ws1.row_dimensions[1].height = 30

ws1.merge_cells('A2:G2')
c2 = ws1['A2']
c2.value = 'Trích xuất từ Hóa đơn bán hàng & Sheet Whatsapp trong Bảng lương Tháng 8'
c2.font = font_subtitle
c2.alignment = Alignment(horizontal='center', vertical='center')
ws1.row_dimensions[2].height = 20

# Headers
headers1 = ['STT', 'Chi nhánh', 'Họ và tên Dược sĩ', 'Chức danh', 'Doanh thu đơn Whatsapp (VNĐ)', 'Tỷ lệ thưởng (%)', 'Tiền thưởng Whatsapp (VNĐ)']
ws1.append([])
ws1.append(headers1)
ws1.row_dimensions[4].height = 28

for col_idx in range(1, 8):
    cell = ws1.cell(4, col_idx)
    cell.font = font_header_white
    cell.fill = fill_navy
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = border_header

# Data
staff_wa = [
    {'stt': 1, 'branch': 'Đỗ Quang Đẩu', 'name': 'Ngô Thị Thanh Thắm', 'role': 'CHT', 'rev': 67046300, 'rate': 0.04},
    {'stt': 2, 'branch': 'Đỗ Quang Đẩu', 'name': 'Hoàng Thanh Thủy', 'role': 'DSXC', 'rev': 92439700, 'rate': 0.04},
    {'stt': 3, 'branch': 'Đỗ Quang Đẩu', 'name': 'Lê Thị Huyền Trân', 'role': 'DSXC', 'rev': 58824400, 'rate': 0.04},
    {'stt': 4, 'branch': 'Đỗ Quang Đẩu', 'name': 'Trần Thiên Phát', 'role': 'DSCD', 'rev': 9377800, 'rate': 0.04},
    {'stt': 5, 'branch': 'Đỗ Quang Đẩu', 'name': 'Phạm Thị Nghĩa Hương', 'role': 'DSCD', 'rev': 8501100, 'rate': 0.04},
    {'stt': 6, 'branch': 'Hàng Bông', 'name': 'Đinh Thị Khánh Ly', 'role': 'DSXC', 'rev': 19489000, 'rate': 0.015},
]

start_row = 5
for idx, s in enumerate(staff_wa, start=start_row):
    ws1.append([
        s['stt'],
        s['branch'],
        s['name'],
        s['role'],
        s['rev'],
        s['rate'],
        f"=E{idx}*F{idx}"
    ])
    ws1.row_dimensions[idx].height = 22
    for c_i in range(1, 8):
        cell = ws1.cell(idx, c_i)
        cell.font = font_regular
        cell.border = border_thin
        if c_i == 1:
            cell.alignment = Alignment(horizontal='center')
        elif c_i in (2, 3):
            cell.alignment = Alignment(horizontal='left')
        elif c_i == 4:
            cell.alignment = Alignment(horizontal='center')
        elif c_i == 5:
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal='right')
        elif c_i == 6:
            cell.number_format = '0.0%'
            cell.alignment = Alignment(horizontal='center')
        elif c_i == 7:
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal='right')
            cell.font = font_bold
            cell.fill = fill_highlight

# Total row
tot_r = start_row + len(staff_wa)
ws1.append(['', 'TỔNG CỘNG', '', '', f"=SUM(E{start_row}:E{tot_r-1})", '', f"=SUM(G{start_row}:G{tot_r-1})"])
ws1.row_dimensions[tot_r].height = 26
ws1.merge_cells(f'B{tot_r}:D{tot_r}')

for c_i in range(1, 8):
    cell = ws1.cell(tot_r, c_i)
    cell.font = font_bold
    cell.fill = fill_emerald
    cell.border = border_total
    if c_i == 2:
        cell.alignment = Alignment(horizontal='center')
    elif c_i == 5:
        cell.number_format = '#,##0'
        cell.alignment = Alignment(horizontal='right')
    elif c_i == 7:
        cell.number_format = '#,##0'
        cell.alignment = Alignment(horizontal='right')
        cell.font = Font(name='Calibri', size=11, bold=True, color='15803D')

# Set column widths
ws1.column_dimensions['A'].width = 8
ws1.column_dimensions['B'].width = 18
ws1.column_dimensions['C'].width = 28
ws1.column_dimensions['D'].width = 14
ws1.column_dimensions['E'].width = 26
ws1.column_dimensions['F'].width = 18
ws1.column_dimensions['G'].width = 26

# ==============================================================================
# SHEET 2: THỐNG KÊ DOANH SỐ & TỶ LỆ THEO NHÀ THUỐC
# ==============================================================================
ws2 = wb.create_sheet('Thống kê theo Nhà thuốc')
ws2.views.sheetView[0].showGridLines = True

ws2.merge_cells('A1:F1')
ws2['A1'].value = 'THỐNG KÊ DOANH THU & XÉT MỨC THƯỞNG WHATSAPP THEO NHÀ THUỐC'
ws2['A1'].font = font_title
ws2['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws2.row_dimensions[1].height = 30

ws2.append([])
headers2 = ['STT', 'Nhà thuốc', 'Doanh thu Whatsapp (VNĐ)', 'Mốc KPI Cửa hàng đạt', 'Tỷ lệ thưởng áp dụng', 'Tổng tiền thưởng (VNĐ)']
ws2.append(headers2)
ws2.row_dimensions[3].height = 28

for col_idx in range(1, 7):
    cell = ws2.cell(3, col_idx)
    cell.font = font_header_white
    cell.fill = fill_navy
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = border_header

store_summary = [
    {
        'stt': 1,
        'store': 'Đỗ Quang Đẩu',
        'rev': 236189300,
        'tier': 'Mức 2 (>= 200 Triệu)',
        'rate': 0.04,
        'bonus': 9447572
    },
    {
        'stt': 2,
        'store': 'Hàng Bông',
        'rev': 19489000,
        'tier': 'Chưa đạt mức 1 (< 150 Triệu)',
        'rate': 0.015,
        'bonus': 292335
    }
]

for idx, s in enumerate(store_summary, start=4):
    ws2.append([
        s['stt'],
        s['store'],
        s['rev'],
        s['tier'],
        s['rate'],
        s['bonus']
    ])
    ws2.row_dimensions[idx].height = 24
    for c_i in range(1, 7):
        cell = ws2.cell(idx, c_i)
        cell.font = font_regular
        cell.border = border_thin
        if c_i == 1: cell.alignment = Alignment(horizontal='center')
        elif c_i == 2: cell.alignment = Alignment(horizontal='left')
        elif c_i == 3:
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal='right')
        elif c_i == 4: cell.alignment = Alignment(horizontal='center')
        elif c_i == 5:
            cell.number_format = '0.0%'
            cell.alignment = Alignment(horizontal='center')
        elif c_i == 6:
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal='right')
            cell.font = font_bold

# Store total
s_tot = 4 + len(store_summary)
ws2.append(['', 'TỔNG CỘNG', f"=SUM(C4:C{s_tot-1})", '', '', f"=SUM(F4:F{s_tot-1})"])
ws2.row_dimensions[s_tot].height = 26
for c_i in range(1, 7):
    cell = ws2.cell(s_tot, c_i)
    cell.font = font_bold
    cell.fill = fill_emerald
    cell.border = border_total
    if c_i == 2: cell.alignment = Alignment(horizontal='center')
    elif c_i in (3, 6):
        cell.number_format = '#,##0'
        cell.alignment = Alignment(horizontal='right')

ws2.column_dimensions['A'].width = 8
ws2.column_dimensions['B'].width = 22
ws2.column_dimensions['C'].width = 28
ws2.column_dimensions['D'].width = 30
ws2.column_dimensions['E'].width = 22
ws2.column_dimensions['F'].width = 26

# ==============================================================================
# SHEET 3: CHÍNH SÁCH THƯỞNG WHATSAPP
# ==============================================================================
ws3 = wb.create_sheet('Chính sách thưởng Whatsapp')
ws3.views.sheetView[0].showGridLines = True

ws3.merge_cells('A1:D1')
ws3['A1'].value = 'QUY ĐỊNH CHƯƠNG TRÌNH THƯỞNG ĐƠN HÀNG QUA ỨNG DỤNG WHATSAPP'
ws3['A1'].font = font_title
ws3['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws3.row_dimensions[1].height = 30

ws3.append([])
headers3 = ['Mức KPI Doanh thu Cửa hàng', 'Ngưỡng Doanh thu (VNĐ)', 'Tỷ lệ thưởng Dược sĩ (%)', 'Ghi chú điều kiện']
ws3.append(headers3)
ws3.row_dimensions[3].height = 28

for col_idx in range(1, 5):
    cell = ws3.cell(3, col_idx)
    cell.font = font_header_white
    cell.fill = fill_navy
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = border_header

policies = [
    ['Mức 1', '150,000,000 đ', '3.0%', 'Cửa hàng đạt mốc 150 triệu -> Thưởng 3% doanh thu đơn cá nhân'],
    ['Mức 2', '200,000,000 đ', '4.0%', 'Cửa hàng đạt mốc 200 triệu -> Thưởng 4% doanh thu đơn cá nhân'],
    ['Mức 3', '280,000,000 đ', '6.0%', 'Cửa hàng đạt mốc 280 triệu -> Thưởng 6% doanh thu đơn cá nhân'],
    ['Cơ bản (Không đạt)', '< 150,000,000 đ', '1.5%', 'Cửa hàng không đạt KPI -> Thưởng mức cơ bản 1.5% doanh thu đơn cá nhân']
]

for idx, p in enumerate(policies, start=4):
    ws3.append(p)
    ws3.row_dimensions[idx].height = 24
    for c_i in range(1, 5):
        cell = ws3.cell(idx, c_i)
        cell.font = font_regular
        cell.border = border_thin
        if c_i in (1, 3): cell.alignment = Alignment(horizontal='center')
        elif c_i == 2: cell.alignment = Alignment(horizontal='right')
        else: cell.alignment = Alignment(horizontal='left')

# Additional Notes
notes_row = 9
ws3.cell(notes_row, 1, 'ĐIỀU KIỆN GHI NHẬN ĐƠN HÀNG HỢP LỆ:').font = font_bold
notes_items = [
    '1. Dược sĩ lên đơn trên phần mềm KiotViet, bắt buộc phải ghi chú "Whatsapp" trong đơn hàng.',
    '2. Khách hàng có tương tác trực tiếp qua WhatsApp (hỏi giá, tư vấn dùng thuốc, xác nhận đơn hàng...).',
    '3. Đơn hàng được tạo sau cuộc trao đổi và truy vết được (SĐT, tên khách, thời gian tạo, ghi chú trên KiotViet).',
    '4. Chỉ áp dụng cho tệp khách hàng quốc tế kết nối qua QR WhatsApp của nhà thuốc.'
]
for offset, n_txt in enumerate(notes_items, start=1):
    ws3.cell(notes_row + offset, 1, n_txt).font = font_regular

ws3.column_dimensions['A'].width = 24
ws3.column_dimensions['B'].width = 26
ws3.column_dimensions['C'].width = 24
ws3.column_dimensions['D'].width = 65

wb.save(out_file)
wb.close()
print(f"--> Đã xuất thành công file tổng hợp Whatsapp ra: {out_file}")
