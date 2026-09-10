import openpyxl

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws_nc = wb['Ngày công']

hidden_rows = [r for r, dim in ws_nc.row_dimensions.items() if dim.hidden]
print(f"Total hidden rows in Ngày công: {len(hidden_rows)}")

branches = {}
for r in range(2, ws_nc.max_row + 1):
    cn = ws_nc.cell(r, 1).value
    name = ws_nc.cell(r, 2).value
    date_val = ws_nc.cell(r, 3).value
    if cn and name:
        branches[cn] = branches.get(cn, set())
        branches[cn].add(name)

print("\n=== CÁC CHI NHÁNH VÀ SỐ NHÂN VIÊN TRONG SHEET NGÀY CÔNG ===")
for cn, staff in sorted(branches.items()):
    print(f"Chi nhánh '{cn}': {len(staff)} nhân viên -> {list(staff)[:3]}...")

print(f"\nTổng số dòng dữ liệu quẹt thẻ thực tế: {ws_nc.max_row - 1}")
