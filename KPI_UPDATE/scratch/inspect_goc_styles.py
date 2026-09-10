import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

f_nt8 = r"d:\MEDIGO\KPI_UPDATE\goc\NHÀ THUỐC THÁNG 8 2026.xlsx"
wb = openpyxl.load_workbook(f_nt8, data_only=False)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n==========================================================================")
    print(f"SHEET: '{sname}' STYLES")
    print(f"==========================================================================")
    for r in range(1, min(4, ws.max_row + 1)):
        print(f"--- Row {r} ---")
        for c in range(1, min(ws.max_column + 1, 35)):
            cell = ws.cell(r, c)
            val = cell.value
            fill = cell.fill.start_color.rgb if cell.fill and cell.fill.start_color else None
            fill_type = cell.fill.fill_type if cell.fill else None
            font_color = cell.font.color.rgb if cell.font and cell.font.color else None
            font_bold = cell.font.bold if cell.font else None
            if val is not None or fill is not None:
                print(f"  Col {c:2d} ({openpyxl.utils.get_column_letter(c):>3s}): val={repr(val)[:30]:<30} | fill={fill} ({fill_type}) | font_bold={font_bold} | font_col={font_color}")
