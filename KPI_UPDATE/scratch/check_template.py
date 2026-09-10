import openpyxl

wb = openpyxl.load_workbook('thang8/tinhcongnhungthuongchia.xlsx', data_only=False)
ws = wb['BẢNG LƯƠNG']

with open('scratch/template_structure.txt', 'w', encoding='utf-8') as f:
    f.write(f"Max row: {ws.max_row}\n")
    for r in range(1, 95):
        row_vals = [ws.cell(r, c).value for c in range(1, 10)]
        f.write(f"Row {r:2d}: {row_vals}\n")

print("Done writing template_structure.txt")
