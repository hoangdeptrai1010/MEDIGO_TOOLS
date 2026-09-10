import sys
import openpyxl
from collections import Counter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=False)
ws_bl = wb['BẢNG LƯƠNG']

staff_rows = []
for r in range(3, ws_bl.max_row + 1):
    stt = ws_bl.cell(r, 1).value
    cn = ws_bl.cell(r, 2).value
    name = ws_bl.cell(r, 3).value
    role = ws_bl.cell(r, 4).value
    if name and str(name).strip():
        staff_rows.append((r, stt, str(cn).strip(), str(name).strip(), str(role).strip()))

counts = Counter(s[3] for s in staff_rows)
dups = {k: v for k, v in counts.items() if v > 1}
print(f'Total rows with name: {len(staff_rows)}, Unique names: {len(counts)}')
print('Duplicate staff in BẢNG LƯƠNG:', len(dups))
for name, cnt in dups.items():
    print(f'\n--- {name} ({cnt} rows) ---')
    for s in staff_rows:
        if s[3] == name:
            r = s[0]
            # print some columns like lương cơ bản, các cột giờ, công thức
            lcb = ws_bl.cell(r, 19).value # Cột S lương cơ bản hoặc cột khác
            c_g = ws_bl.cell(r, 7).value
            c_h = ws_bl.cell(r, 8).value
            c_j = ws_bl.cell(r, 10).value
            c_k = ws_bl.cell(r, 11).value
            c_r = ws_bl.cell(r, 18).value
            print(f'  Row {r}: STT={s[1]}, Branch={s[2]}, Role={s[4]} | G={c_g}, H={c_h}, J={c_j}, K={c_k}, R={c_r}')
