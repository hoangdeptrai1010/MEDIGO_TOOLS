import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

f_nt7 = r"d:\MEDIGO\KPI_UPDATE\goc\NHÀ THUỐC THÁNG 7 2026.xlsx"
f_nt8 = r"d:\MEDIGO\KPI_UPDATE\goc\NHÀ THUỐC THÁNG 8 2026.xlsx"

for m, fpath in [("T7", f_nt7), ("T8", f_nt8)]:
    print(f"\n==========================================================================")
    print(f"                         WORKBOOK GỐC {m}")
    print(f"==========================================================================")
    wb = openpyxl.load_workbook(fpath, data_only=True)
    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f"\n--- SHEET: '{sname}' ---")
        # Find header row
        for r in range(1, min(5, ws.max_row + 1)):
            vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            # filter non-empty
            non_empty = [(c, v) for c, v in enumerate(vals, 1) if v is not None]
            if non_empty:
                print(f"  Row {r:2d} ({len(non_empty)} cols):")
                for c, v in non_empty[:25]:
                    print(f"    Col {c:2d} ({openpyxl.utils.get_column_letter(c):>3s}): {repr(v)}")
