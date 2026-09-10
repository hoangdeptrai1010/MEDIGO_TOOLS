import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl, subprocess

wb = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx', data_only=False)
ws_ds = wb['kpi dược sĩ']

# Update Col F with fixed LET formula for all rows:
for r in range(3, ws_ds.max_row + 1):
    nv = ws_ds.cell(r, 2).value
    if not nv or not str(nv).strip():
        continue
    formula = (
        f'=IFERROR(ROUND(_xlfn.LET('
        f'nt, $A{r}, '
        f'nv, $B{r}, '
        f'songay, DAY(EOMONTH($C$1, 0)), '
        f'doanhthungay, $I{r}, '
        f'duan_nv, SUMIFS(\'Dự án T8\'!$D:$D, \'Dự án T8\'!$A:$A, nt, \'Dự án T8\'!$B:$B, nv), '
        f'duan_nt, SUMIF(\'Dự án T8\'!$A:$A, nt, \'Dự án T8\'!$D:$D), '
        f'tyle, IF(duan_nt>0, duan_nv/duan_nt, 0), '
        f'muc2_ngay, _xlfn.XLOOKUP(nt, \'kpi nhà thuốc\'!$B:$B, \'kpi nhà thuốc\'!$P:$P, 0), '
        f'muc3_ngay, _xlfn.XLOOKUP(nt, \'kpi nhà thuốc\'!$B:$B, \'kpi nhà thuốc\'!$T:$T, 0), '
        f'hangdiem1, _xlfn.XLOOKUP(nt, \'kpi nhà thuốc\'!$B:$B, \'kpi nhà thuốc\'!$I:$I, 0), '
        f'hangdiem2, _xlfn.XLOOKUP(nt, \'kpi nhà thuốc\'!$B:$B, \'kpi nhà thuốc\'!$N:$N, 0), '
        f'hangdiem3, _xlfn.XLOOKUP(nt, \'kpi nhà thuốc\'!$B:$B, \'kpi nhà thuốc\'!$R:$R, 0), '
        f'_xlfn.SWITCH(TRUE, '
        f'doanhthungay>=muc3_ngay, hangdiem3, '
        f'doanhthungay>=muc2_ngay, hangdiem2, '
        f'TRUE, hangdiem1'
        f')/songay*tyle'
        f'), -3), 0)'
    )
    ws_ds.cell(r, 6, formula)

wb.save('scratch/test_let_col_f.xlsx')
print('Saved scratch/test_let_col_f.xlsx')
