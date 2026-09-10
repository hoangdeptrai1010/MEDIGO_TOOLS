import openpyxl

wb = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=True)
print("Sheetnames:", wb.sheetnames)

for sn in ['BẢNG LƯƠNG', 'MiniKat - HN', 'MiniKat - HCM', 'whatsapp', 'KPI', 'Thưởng CK', 'Dự án', 'Giờ công', 'Ngày công', 'Cận date']:
    if sn in wb.sheetnames:
        ws = wb[sn]
        print(f"\n=== Sheet {sn} (rows={ws.max_row}, cols={ws.max_column}) ===")
        # Print non-empty sample rows
        count = 0
        for r in range(1, ws.max_row + 1):
            row_v = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
            if any(x is not None for x in row_v):
                count += 1
                if count <= 5 or r == ws.max_row:
                    print(f"  r{r}: {row_v[:8]}")
        print(f"  -> Total non-empty rows: {count}")
