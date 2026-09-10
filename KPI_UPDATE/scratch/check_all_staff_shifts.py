import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'thang8\bangluong_thang8_hoanthien.xlsx', data_only=True)
ws = wb['BẢNG LƯƠNG']

print(f"{'STT':<4} | {'Họ và tên':<26} | {'Chức danh':<10} | {'Giờ ngày (G)':<14} | {'Ngày làm ngày(J)':<16} | {'Ca làm đêm(K)':<14} | {'Tổng công (R)':<14}")
print("-" * 105)

for r in range(3, ws.max_row + 1):
    stt = ws.cell(r, 1).value
    name = ws.cell(r, 3).value
    cd = ws.cell(r, 4).value
    g = ws.cell(r, 7).value or 0
    j = ws.cell(r, 10).value or 0
    k = ws.cell(r, 11).value or 0
    r_val = ws.cell(r, 18).value or 0
    
    if name and str(name).strip() and "Tổng" not in str(name):
        g_val = float(g) if isinstance(g, (int, float)) else 0.0
        j_val = float(j) if isinstance(j, (int, float)) else 0.0
        k_val = float(k) if isinstance(k, (int, float)) else 0.0
        r_num = float(r_val) if isinstance(r_val, (int, float)) else 0.0
        print(f"{str(stt):<4} | {str(name):<26} | {str(cd):<10} | {g_val:>12.2f} h | {j_val:>14.0f} ng | {k_val:>12.0f} ca | {r_num:>12.0f} ng")
