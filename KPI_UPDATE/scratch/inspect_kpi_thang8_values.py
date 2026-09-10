import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

for fname in ['baocaokpi_thang8_hoanthien.xlsx', 'thang8/baocaokpi_thang8_hoanthien.xlsx']:
    try:
        wb = openpyxl.load_workbook(fname, data_only=True)
        print(f"\n{'='*20} {fname} (DATA ONLY) {'='*20}")
        print("Sheets:", wb.sheetnames)
        
        ws_ds = wb['kpi dược sĩ']
        print(f"\n--- Sheet 'kpi dược sĩ' (Max row: {ws_ds.max_row}) ---")
        headers = [ws_ds.cell(2, c).value for c in range(1, 15)]
        print("Headers Col 1..14:", headers)
        
        for r in range(3, min(ws_ds.max_row + 1, 15)):
            cn = ws_ds.cell(r, 1).value
            nv = ws_ds.cell(r, 2).value
            cd = ws_ds.cell(r, 3).value
            tb_bill_kpi = ws_ds.cell(r, 4).value
            tb_bill_act = ws_ds.cell(r, 5).value
            kpi_ck = ws_ds.cell(r, 6).value
            act_ck = ws_ds.cell(r, 7).value
            kpi_rev = ws_ds.cell(r, 8).value
            act_rev = ws_ds.cell(r, 9).value
            pct = ws_ds.cell(r, 10).value
            bonus = ws_ds.cell(r, 11).value
            print(f"Row {r:2d}: [{cn}] {nv} ({cd}) | F(KPI CK)={kpi_ck} | G(CK ngày)={act_ck} | H(KPI DT)={kpi_rev} | I(DT thực)={act_rev} | J(% KPI)={pct} | K(Thưởng)={bonus}")
        
        wb.close()
    except Exception as e:
        print(f"Error reading {fname}: {e}")
