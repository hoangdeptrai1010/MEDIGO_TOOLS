import sys
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=False)
ws = wb['BẢNG LƯƠNG']

wb_val = openpyxl.load_workbook('thang8/BANGLUONGTHANG8.xlsx', data_only=True)
ws_val = wb_val['BẢNG LƯƠNG']

print("=== CHECKING FORMULAS & VALUES IN BẢNG LƯƠNG ===")
for r in range(4, ws.max_row + 1):
    stt = ws.cell(r, 1).value
    name = ws.cell(r, 3).value
    branch = ws.cell(r, 2).value
    if not name and not stt:
        continue
    g_form = ws.cell(r, 7).value
    h_form = ws.cell(r, 8).value
    r_form = ws.cell(r, 18).value
    
    g_val = ws_val.cell(r, 7).value
    h_val = ws_val.cell(r, 8).value
    e_val = ws_val.cell(r, 5).value
    f_val = ws_val.cell(r, 6).value
    r_val = ws_val.cell(r, 18).value
    
    # Let's print rows for people who worked multiple branches, or sample rows
    if r in [4, 5, 6, 24, 25, 26, 60, 61, 62] or (stt and ('Hòa' in str(name) or 'Phương' in str(name) or 'Phượng' in str(name) or 'TỔNG' in str(stt) or 'TỔNG' in str(name))):
        print(f"Row {r:2d} | STT: {stt} | Branch: {branch} | Name: {name}")
        print(f"   Formulas: G={g_form} | H={h_form} | R={r_form}")
        print(f"   Values:   E(tot day)={e_val} | G(day)={g_val} | F(tot night)={f_val} | H(night)={h_val} | R(days)={r_val}")
