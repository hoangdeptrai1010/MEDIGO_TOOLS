import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_bangluong = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final.xlsx")
wb = openpyxl.load_workbook(f_bangluong, data_only=True)

print("=== ALL ROWS IN Sheet1 ===")
ws1 = wb["Sheet1"]
for r in range(1, ws1.max_row + 1):
    vals = [ws1.cell(r, c).value for c in range(1, ws1.max_column + 1)]
    if any(v is not None for v in vals):
        print(f"Row {r:2d}: {vals}")

print("\n=== ALL ROWS IN Med WITH NAMES ===")
ws_med = wb["Med"]
for r in range(9, ws_med.max_row + 1):
    stt = ws_med.cell(r, 1).value
    code = ws_med.cell(r, 3).value
    name = ws_med.cell(r, 4).value
    store = ws_med.cell(r, 6).value
    role = ws_med.cell(r, 7).value
    stk = ws_med.cell(r, 9).value
    bank = ws_med.cell(r, 10).value
    if name is not None and str(name).strip() != "":
        print(f"Row {r:3d} | Code: {str(code):<25} | Name: {str(name):<30} | Store: {str(store):<15} | STK: {str(stk):<20} | Bank: {str(bank):<25}")
