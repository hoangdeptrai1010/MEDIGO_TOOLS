import os
import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

for root, dirs, files in os.walk('thang8'):
    for f in files:
        if f.endswith('.docx') and 'Chuong' in f or 'kpi' in f.lower():
            full_p = os.path.join(root, f)
            print(f"\n==================== {full_p} ====================")
            try:
                doc = docx.Document(full_p)
                for p in doc.paragraphs:
                    if p.text.strip():
                        print(p.text.strip())
                for t in doc.tables:
                    for row in t.rows:
                        print(" | ".join(c.text.strip() for c in row.cells))
            except Exception as e:
                print(f"Error reading {full_p}: {e}")
