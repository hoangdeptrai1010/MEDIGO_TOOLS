import openpyxl

wb = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx', data_only=True)
print("baocaokpi_thang8_hoanthien.xlsx sheetnames:", wb.sheetnames)
ws_da = wb['Dự án']
print("Dự án headers:")
for r in range(1, 4):
    print([ws_da.cell(r, c).value for c in range(1, 15)])
print("\nDự án sample rows:")
for r in range(4, 12):
    print([ws_da.cell(r, c).value for c in range(1, 15)])
