import openpyxl

wb_plan = openpyxl.load_workbook('plans/KeHoachKPI_2026-08.xlsx', data_only=False)
print("Plans sheetnames:", wb_plan.sheetnames)

for sname in wb_plan.sheetnames:
    ws = wb_plan[sname]
    print(f"\n--- Sheet: {sname} (max_row={ws.max_row}, max_col={ws.max_column}) ---")
    for r in range(1, min(6, ws.max_row + 1)):
        vals = [ws.cell(r, c).value for c in range(1, min(10, ws.max_column + 1))]
        if any(v is not None for v in vals):
            print(f"  Row {r}: {vals}")
