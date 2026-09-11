import openpyxl, sys, io, glob

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

for f in glob.glob('d:/MEDIGO/KPI_UPDATE/thang8/**/*.xlsx', recursive=True):
    if '~$' in f: continue
    try:
        wb = openpyxl.load_workbook(f, data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            found = []
            for r in range(1, min(ws.max_row+1, 200)):
                row_str = " ".join([str(ws.cell(r, c).value or '') for c in range(1, min(ws.max_column+1, 30))])
                if 'whatsapp' in row_str.lower() or 'hđ' in row_str.lower() or 'hd' in row_str.lower() and '67046300' in row_str:
                    found.append((r, row_str[:120]))
            if found:
                print(f"\nFile: {f} | Sheet: {sname}")
                for r, txt in found[:5]:
                    print(f"  Row {r}: {txt}")
        wb.close()
    except Exception as e:
        pass
