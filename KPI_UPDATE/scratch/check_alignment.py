import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook('thang8/tinhcongnhungthuongchia.xlsx', data_only=False)
ws_nc = wb['Ngày công']
ws_gc = wb['Giờ công']
ws_bl = wb['BẢNG LƯƠNG']

print(f"BẢNG LƯƠNG: {ws_bl.max_row} rows")
for r in range(3, 73):
    bl_cn = ws_bl.cell(r, 2).value
    bl_name = ws_bl.cell(r, 3).value
    
    nc_r = r - 1 # row 2 to 71
    nc_cn = ws_nc.cell(nc_r, 11).value
    nc_name = ws_nc.cell(nc_r, 12).value
    
    gc_cn = ws_gc.cell(nc_r, 7).value
    gc_name = ws_gc.cell(nc_r, 8).value
    
    if (bl_cn != nc_cn or bl_name != nc_name) or (bl_cn != gc_cn or bl_name != gc_name):
        print(f"Mismatch at BL row {r}: BL=({bl_cn}, {bl_name}), NC=({nc_cn}, {nc_name}), GC=({gc_cn}, {gc_name})")

print("Verification between BL, NC, GC summary complete!")
