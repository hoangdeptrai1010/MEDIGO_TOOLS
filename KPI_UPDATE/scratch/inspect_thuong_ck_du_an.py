import sys, io
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

def inspect_file(path):
    print(f"\n==================== {path} ====================")
    wb = openpyxl.load_workbook(path, data_only=False)
    print("Sheets:", wb.sheetnames)
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        print(f"\n--- Sheet: {sheetname} (max_row={ws.max_row}, max_col={ws.max_column}) ---")
        for r in range(1, min(6, ws.max_row + 1)):
            row_vals = [ws.cell(r, c).value for c in range(1, min(30, ws.max_column + 1))]
            if any(v is not None for v in row_vals):
                print(f"Row {r}: {row_vals}")

inspect_file("thang8/bangluong_thang8_hoanthien.xlsx")
inspect_file("thang8/baocaokpi_thang8_hoanthien.xlsx")
