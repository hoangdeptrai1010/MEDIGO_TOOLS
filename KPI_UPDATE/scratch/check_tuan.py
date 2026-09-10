import openpyxl
import datetime

# Target for Nguyễn Mạnh Tuấn
TARGET = 1577380
SELLER = "Nguyễn Mạnh Tuấn"

print(f"Loading invoice data for month 8...")
wb = openpyxl.load_workbook('thang8/DanhSachChiTietHoaDon_3182026.xlsx', data_only=True, read_only=True)
ws = wb.active
headers = next(ws.iter_rows(max_row=1, values_only=True))

col_map = {h: i for i, h in enumerate(headers) if h}
c_seller = col_map.get('Người bán', 6)
c_sku = col_map.get('Mã hàng', 13)
c_name = col_map.get('Tên hàng', 15)
c_val = col_map.get('Thành tiền', 25)
c_exp = col_map.get('Hạn sử dụng', 18)

items = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r: continue
    seller = str(r[c_seller] or '').strip()
    if seller != SELLER: continue
    
    sku = str(r[c_sku] or '').strip()
    val = float(r[c_val] or 0)
    name = str(r[c_name] or '').strip()
    
    if not sku or val <= 0: continue
    items.append({'sku': sku, 'name': name, 'val': val, 'exp': r[c_exp]})
wb.close()

# Group by SKU
by_sku = {}
for it in items:
    sku = it['sku']
    if sku not in by_sku:
        by_sku[sku] = {'val': 0, 'name': it['name']}
    by_sku[sku]['val'] += it['val']

candidates = [(v['val'], k, v['name']) for k, v in by_sku.items() if v['val'] <= TARGET + 1]

# DFS
candidates.sort(key=lambda x: x[0], reverse=True)
suffix_sums = [0] * len(candidates)
curr_sum = 0
for i in range(len(candidates)-1, -1, -1):
    curr_sum += candidates[i][0]
    suffix_sums[i] = curr_sum

def dfs(idx, current_sum, path):
    if abs(current_sum - TARGET) < 1.0:
        return path
    if idx >= len(candidates):
        return None
    if current_sum + suffix_sums[idx] < TARGET - 1.0:
        return None
    
    val = candidates[idx][0]
    if current_sum + val <= TARGET + 1.0:
        res = dfs(idx + 1, current_sum + val, path + [candidates[idx]])
        if res: return res
        
    res = dfs(idx + 1, current_sum, path)
    if res: return res
    return None

match = dfs(0, 0, [])
if match:
    print(f"\n✅ Found exact match for {SELLER} (Target: {TARGET}):")
    for val, sku, name in match:
        print(f"   - {sku}: {name} ({val:,.0f}đ)")
else:
    print(f"\n❌ Could not find exact match for {SELLER} (Target: {TARGET})")
    print(f"Top candidates:")
    for val, sku, name in candidates[:10]:
        print(f"   - {sku}: {name} ({val:,.0f}đ)")
