import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

f_nt8 = r"d:\MEDIGO\KPI_UPDATE\goc\NHÀ THUỐC THÁNG 8 2026.xlsx"
wb = openpyxl.load_workbook(f_nt8, data_only=False)
ws = wb["kpi dược sĩ"]

print("=== kpi dược sĩ: Headers (Row 2) and Sample Formulas (Row 3, 4) ===")
for r in [2, 3, 4]:
    print(f"\n--- Row {r} ---")
    for c in range(1, ws.max_column + 1):
        v = ws.cell(r, c).value
        if v is not None:
            print(f"  Col {c:2d} ({openpyxl.utils.get_column_letter(c):>3s}): {repr(v)}")
