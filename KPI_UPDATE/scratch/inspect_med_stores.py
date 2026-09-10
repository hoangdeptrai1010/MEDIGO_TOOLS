import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_bl = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final_backup.xlsx")
wb = openpyxl.load_workbook(f_bl, data_only=True)
ws_med = wb["Med"]

current_store = None
stores = {}

for r in range(10, ws_med.max_row + 1):
    stt = ws_med.cell(r, 1).value
    code = ws_med.cell(r, 3).value
    name = ws_med.cell(r, 4).value
    store = ws_med.cell(r, 6).value
    role = ws_med.cell(r, 7).value
    stk = ws_med.cell(r, 9).value
    bank = ws_med.cell(r, 10).value
    total_net = ws_med.cell(r, 80).value # Col CB: Tổng thực nhận
    ck_tk = ws_med.cell(r, 85).value     # Col CG: CK TK
    
    if store and str(store).strip():
        current_store = str(store).strip()
        if current_store not in stores:
            stores[current_store] = []
            
    if name and str(name).strip() and not str(name).strip().startswith("CHUỖI") and not str(name).strip().startswith("Tổng"):
        stores[current_store].append({
            'med_row': r,
            'stt': stt,
            'code': code,
            'name': str(name).strip(),
            'role': role,
            'stk': str(stk).strip() if stk else "",
            'bank': str(bank).strip() if bank else "",
            'total_net': total_net,
            'ck_tk': ck_tk
        })

for sname, staff_list in stores.items():
    print(f"\n=== CỬA HÀNG: {sname} ({len(staff_list)} nhân sự) ===")
    for emp in staff_list:
        print(f"  Row {emp['med_row']:2d} | STT: {str(emp['stt']):<4} | {emp['name']:<26} | STK: {emp['stk']:<16} | Bank: {emp['bank']:<22} | Net: {emp['total_net']} | CK: {emp['ck_tk']}")
