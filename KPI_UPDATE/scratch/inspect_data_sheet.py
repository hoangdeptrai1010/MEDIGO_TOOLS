import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

f_nt8 = r"d:\MEDIGO\KPI_UPDATE\goc\NHÀ THUỐC THÁNG 8 2026.xlsx"
wb = openpyxl.load_workbook(f_nt8, data_only=True)
ws = wb["data"]

print("=== data Sheet Headers and Sample Rows ===")
for r in range(1, 6):
    row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    print(f"\nRow {r}:")
    for c, v in enumerate(row_vals, 1):
        if v is not None:
            print(f"  Col {c:2d} ({openpyxl.utils.get_column_letter(c):>3s}): {repr(v)}")
