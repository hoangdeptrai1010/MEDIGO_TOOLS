import openpyxl

wb_cc = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
ws_th = wb_cc['Bảng tổng hợp chấm công']
ws_ct = wb_cc['Bảng chi tiết chấm công']

wb_baug = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=True)
ws_gc = wb_baug['Giờ công']
ws_nc = wb_baug['Ngày công']

# 1. Compare total rows in Bảng chi tiết chấm công vs Ngày công
print("=== KIỂM TRA BẢNG CHI TIẾT CHẤM CÔNG vs NGÀY CÔNG ===")
ct_rows = [r for r in range(4, ws_ct.max_row + 1) if ws_ct.cell(r, 6).value is not None]
nc_rows = [r for r in range(2, ws_nc.max_row + 1) if ws_nc.cell(r, 3).value is not None]
print(f"Số dòng chấm công chi tiết có ngày trong BangChiTietChamCong: {len(ct_rows)}")
print(f"Số dòng chấm công chi tiết có ngày trong Ngày công: {len(nc_rows)}")

# 2. Check names in Bảng tổng hợp chấm công vs Giờ công
print("\n=== KIỂM TRA BẢNG TỔNG HỢP CHẤM CÔNG vs GIỜ CÔNG ===")
th_names = set()
for r in range(4, ws_th.max_row + 1):
    ten = ws_th.cell(r, 3).value
    cn = ws_th.cell(r, 6).value
    if ten and str(ten).strip():
        th_names.add((str(cn).strip(), str(ten).strip()))

gc_names = set()
for r in range(2, ws_gc.max_row + 1):
    ten = ws_gc.cell(r, 2).value
    cn = ws_gc.cell(r, 1).value
    if ten and str(ten).strip():
        gc_names.add((str(cn).strip(), str(ten).strip()))

print(f"Số (Chi nhánh, Nhân viên) trong Bảng tổng hợp: {len(th_names)}")
print(f"Số (Chi nhánh, Nhân viên) trong Giờ công: {len(gc_names)}")
diff = th_names - gc_names
if diff:
    print(f"Có trong Bảng tổng hợp nhưng thiếu trong Giờ công: {diff}")
else:
    print("Tất cả nhân sự trong Bảng tổng hợp đều có trong Giờ công!")
