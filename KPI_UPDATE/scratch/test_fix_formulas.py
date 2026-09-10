import openpyxl

wb = openpyxl.load_workbook('scratch/test_baug_recalc.xlsx', data_only=False)
ws_bl = wb['BẢNG LƯƠNG']

for r in range(3, 73):
    # Fix Col E: SUM instead of +
    ws_bl.cell(r, 5, f"=SUM(G{r}, I{r}, L{r}, O{r}, S{r}, T{r})")
    # Fix Col F: SUM
    ws_bl.cell(r, 6, f"=SUM(H{r}, P{r})")
    # Fix Col J: count in Ngày công cols A and B
    ws_bl.cell(r, 10, f"=COUNTIFS('Ngày công'!A:A, B{r}, 'Ngày công'!B:B, C{r})")
    # Fix Col K: count in Ngày công cols A, B and D
    ws_bl.cell(r, 11, f'=COUNTIFS(\'Ngày công\'!A:A, B{r}, \'Ngày công\'!B:B, C{r}, \'Ngày công\'!D:D, "*đêm*")')

wb.save('scratch/test_fixed_formulas.xlsx')
print("Saved test_fixed_formulas.xlsx successfully!")
