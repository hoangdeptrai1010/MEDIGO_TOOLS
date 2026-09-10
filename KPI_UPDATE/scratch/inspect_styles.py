import openpyxl
import sys
import io

if sys.stdout.encoding != 'utf-8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

def inspect_file(path):
    print(f"\n=================================================================")
    print(f"INSPECTING: {path}")
    print(f"=================================================================")
    wb = openpyxl.load_workbook(path, data_only=False)
    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f"\n--- Sheet: {sname} (Rows: {ws.max_row}, Cols: {ws.max_column}) ---")
        h_row = 1 if sname in ('data', 'Hot Bill HN') else 2
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(h_row, c)
            fill_type = cell.fill.fill_type if cell.fill else None
            start_c = cell.fill.start_color.index if cell.fill and cell.fill.start_color else None
            theme = getattr(cell.fill.start_color, 'theme', None) if cell.fill and cell.fill.start_color else None
            tint = getattr(cell.fill.start_color, 'tint', None) if cell.fill and cell.fill.start_color else None
            rgb = getattr(cell.fill.start_color, 'rgb', None) if cell.fill and cell.fill.start_color else None
            font_color = cell.font.color.rgb if cell.font and cell.font.color else (getattr(cell.font.color, 'theme', None) if cell.font and cell.font.color else None)
            bold = cell.font.bold if cell.font else None
            print(f"  Col {c:2d} ({cell.coordinate}): val='{str(cell.value)[:20]}' | fill_type={fill_type}, rgb={rgb}, theme={theme}, tint={tint} | font_c={font_color}, bold={bold}")

inspect_file('goc/NHÀ THUỐC THÁNG 8 2026.xlsx')
inspect_file('goc/NHÀ THUỐC THÁNG 7 2026.xlsx')
