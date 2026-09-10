import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

SKUS_PARTY_SMART = {'SP017162', 'SP2723878', 'SP2723879'}
SKUS_KAT = {'SP2725289', 'SP2725285', 'SP2725291', 'SP2725287', 'SP2725353', 'SP2725351'}
SKUS_LADYCARE = {'SP2723124', 'SP2723125', 'SP2723127'}
SKUS_AVC = {'SP2723883', 'SP2722826', 'SP2722817', 'SP2725666'}

wb_inv = openpyxl.load_workbook('thang8/DATA/DanhSachChiTietHoaDon_3182026.xlsx', data_only=True)
ws = wb_inv.active

headers = [ws.cell(1, c).value for c in range(1, ws.max_column+1)]
c_branch = headers.index('Chi nhánh') + 1
c_sku = headers.index('Mã hàng') + 1
c_qty = headers.index('Số lượng') + 1
c_tt = headers.index('Thành tiền') + 1
c_seller = headers.index('Người bán') + 1

sellers = {}

for r in range(2, ws.max_row + 1):
    seller = str(ws.cell(r, c_seller).value or '').strip()
    if not seller: continue
    sku = str(ws.cell(r, c_sku).value or '').strip()
    qty = ws.cell(r, c_qty).value or 0
    tt = ws.cell(r, c_tt).value or 0
    
    if seller not in sellers:
        sellers[seller] = {'kat_qty': 0, 'kat_rev': 0, 'ps_qty': 0, 'lc_qty': 0, 'lc_rev': 0, 'avc_qty': 0, 'avc_rev': 0}
        
    if sku in SKUS_KAT:
        sellers[seller]['kat_qty'] += qty
        sellers[seller]['kat_rev'] += tt
    elif sku in SKUS_PARTY_SMART:
        sellers[seller]['ps_qty'] += qty
    elif sku in SKUS_LADYCARE:
        sellers[seller]['lc_qty'] += qty
        sellers[seller]['lc_rev'] += tt
    elif sku in SKUS_AVC:
        sellers[seller]['avc_qty'] += qty
        sellers[seller]['avc_rev'] += tt

print('=== TOTAL SALES PER PHARMACIST (All branches combined) ===')
for s, d in sorted(sellers.items(), key=lambda x: x[1]['kat_qty'], reverse=True):
    if d['kat_qty'] >= 5 or d['ps_qty'] >= 40 or d['lc_rev'] >= 2000000 or d['avc_qty'] >= 10:
        kq = d['kat_qty']
        kr = d['kat_rev']
        pq = d['ps_qty']
        lcr = d['lc_rev']
        aq = d['avc_qty']
        ar = d['avc_rev']
        print(f'{s:<25}: KAT={kq:2.0f} ({kr:,.0f}d), PS={pq:2.0f}, LC_rev={lcr:,.0f}d, AVC={aq:2.0f} ({ar:,.0f}d)')
