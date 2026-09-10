import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_ngoc = os.path.join(latvat_dir, "Gởi Ms. Ngoc - 07.09. Bang Luong T8.2026.xlsx")
wb = openpyxl.load_workbook(f_ngoc, data_only=False)

if "TTCK CNT" in wb.sheetnames:
    ws = wb["TTCK CNT"]
    print("=== TTCK CNT FORMULAS IN MS. NGOC ===")
    for r in range(1, 15):
        vals = [ws.cell(r, c).value for c in range(1, 9)]
        print(f"Row {r:2d}: {vals}")
