import openpyxl

wb = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=False)

ws_gc = wb['Giờ công']
ws_nc = wb['Ngày công']
ws_bl = wb['BẢNG LƯƠNG']

print("--- Inspecting Giờ công summary table (Cols G-K) ---")
for r in range(2, 10):
    print(f"R{r}: G={repr(ws_gc.cell(r, 7).value)}, H={repr(ws_gc.cell(r, 8).value)}, I={repr(ws_gc.cell(r, 9).value)}, J={repr(ws_gc.cell(r, 10).value)}, K={repr(ws_gc.cell(r, 11).value)}")

print("\n--- Inspecting Ngày công summary table (Cols K-O) ---")
for r in range(2, 10):
    print(f"R{r}: K={repr(ws_nc.cell(r, 11).value)}, L={repr(ws_nc.cell(r, 12).value)}, M={repr(ws_nc.cell(r, 13).value)}, N={repr(ws_nc.cell(r, 14).value)}, O={repr(ws_nc.cell(r, 15).value)}")

print("\n--- Inspecting BẢNG LƯƠNG formulas (Rows 3-6) ---")
for r in range(3, 7):
    print(f"R{r}: B={repr(ws_bl.cell(r, 2).value)}, C={repr(ws_bl.cell(r, 3).value)}")
    print(f"      G (giờ ngày)={repr(ws_bl.cell(r, 7).value)}")
    print(f"      H (giờ đêm) ={repr(ws_bl.cell(r, 8).value)}")
    print(f"      J (ngày làm)={repr(ws_bl.cell(r, 10).value)}")
    print(f"      K (ngày đêm)={repr(ws_bl.cell(r, 11).value)}")
    print(f"      R (ngày công)={repr(ws_bl.cell(r, 18).value)}")
