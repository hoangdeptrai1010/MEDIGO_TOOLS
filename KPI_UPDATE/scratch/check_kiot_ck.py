import openpyxl

base_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\Pharmacy_retail_Store_KPIs_August 2026"

# Load CK-HCM
wb_ck_hcm = openpyxl.load_workbook(f"{base_dir}/Hồ Chí Minh/Dự án/Chương trình/CK-HCM.xlsx", data_only=True)
wb_ck_hn = openpyxl.load_workbook(f"{base_dir}/Hà Nội/Dự án/Chương trình/CK-HN.xlsx", data_only=True)

mini_skus = {
    'SP2725289': 'Calcium Organika',
    'SP2725285': 'Hair plus Organika',
    'SP2725291': 'Liver pro Organika',
    'SP2725287': 'Super iq Organika',
    'SP2725353': 'Maca #2 Organika',
    'SP2725351': 'Maca #6 Organika',
    'SP017162': 'Party smart Himalaya',
    'SP2723124': 'Ladycare extra aloe',
    'SP2723125': 'Ladycare classic 190ml',
    'SP2723127': 'Ladycare classic 90ml',
    'SP2723883': 'Foslin avc',
    'SP2722826': 'Trùng thảo avc',
    'SP2722817': 'Thanh nhiệt mát gan avc',
    'SP2725666': 'Ceregold brain avc'
}

print("=== CHIẾT KHẤU TRÊN KIOT CỦA CÁC SẢN PHẨM DỰ ÁN MINI ===")

for wb_name, wb in [('CK-HCM.xlsx', wb_ck_hcm), ('CK-HN.xlsx', wb_ck_hn)]:
    print(f"\n--- {wb_name} ---")
    for sname in wb.sheetnames:
        ws = wb[sname]
        for r in range(2, ws.max_row+1):
            sku = str(ws.cell(r, 2).value or '')
            if sku in mini_skus:
                name = ws.cell(r, 3).value
                price = ws.cell(r, 5).value
                ck = ws.cell(r, 6).value
                print(f"  [{sname}] SKU={sku:10s} | {name:40s} | Giá={str(price):>7} | CK Kiot={str(ck):>7}")
