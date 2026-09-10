import openpyxl
import sys
import io
import os

if sys.stdout.encoding != 'utf-8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

print("=== 1. AUDIT INPUT HCM PROPOSAL ===")
hcm_path = 'TOOL_KPISHEET/uploads/hcm_proposal_T8_KPI CNT HCM Tháng 09.xlsx'
wb_hcm = openpyxl.load_workbook(hcm_path, data_only=True)
print("HCM Sheets:", wb_hcm.sheetnames)
for sname in wb_hcm.sheetnames:
    ws = wb_hcm[sname]
    print(f"Sheet: {sname}, Rows: {ws.max_row}, Cols: {ws.max_column}")
    for r in range(1, min(6, ws.max_row + 1)):
        vals = [ws.cell(r, c).value for c in range(1, min(12, ws.max_column + 1))]
        if any(v is not None for v in vals):
            print(f"  Row {r}: {vals}")

print("\n=== 2. AUDIT INPUT HN PROPOSAL ===")
hn_path = 'TOOL_KPISHEET/uploads/hn_proposal_T8_e_xuat_KPI_Quy_3.26.xlsx'
wb_hn = openpyxl.load_workbook(hn_path, data_only=True)
print("HN Sheets:", wb_hn.sheetnames)
for sname in wb_hn.sheetnames:
    ws = wb_hn[sname]
    print(f"Sheet: {sname}, Rows: {ws.max_row}, Cols: {ws.max_column}")
    for r in range(1, min(6, ws.max_row + 1)):
        vals = [ws.cell(r, c).value for c in range(1, min(12, ws.max_column + 1))]
        if any(v is not None for v in vals):
            print(f"  Row {r}: {vals}")

print("\n=== 3. AUDIT OUTPUT GENERATED vs PRODUCTION GOC ===")
out_path = 'TOOL_KPISHEET/output/NHÀ THUỐC THÁNG 8 2026.xlsx'
goc_path = 'goc/NHÀ THUỐC THÁNG 8 2026.xlsx'

if os.path.exists(out_path):
    wb_out = openpyxl.load_workbook(out_path, data_only=False)
    print("Output Sheets:", wb_out.sheetnames)
    for sname in wb_out.sheetnames:
        ws = wb_out[sname]
        print(f"  Out Sheet {sname}: Rows={ws.max_row}, Cols={ws.max_column}")
        for r in range(1, min(6, ws.max_row + 1)):
            vals = [str(ws.cell(r, c).value) for c in range(1, min(10, ws.max_column + 1))]
            print(f"    Row {r}: {vals}")

if os.path.exists(goc_path):
    wb_goc = openpyxl.load_workbook(goc_path, data_only=False)
    print("Goc Sheets:", wb_goc.sheetnames)
    for sname in wb_goc.sheetnames:
        ws = wb_goc[sname]
        print(f"  Goc Sheet {sname}: Rows={ws.max_row}, Cols={ws.max_column}")
        for r in range(1, min(6, ws.max_row + 1)):
            vals = [str(ws.cell(r, c).value) for c in range(1, min(10, ws.max_column + 1))]
            print(f"    Row {r}: {vals}")
