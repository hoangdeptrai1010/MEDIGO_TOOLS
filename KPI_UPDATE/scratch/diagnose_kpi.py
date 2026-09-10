import openpyxl
import subprocess
import os

def test_open(fname):
    p = os.path.abspath(fname)
    res = subprocess.run(['powershell', '-ExecutionPolicy', 'Bypass', '-File', 'recalc_workbook.ps1', '-FilePath', p], capture_output=True, text=True)
    print(f'Test {fname}: Return {res.returncode}, Output: {res.stdout.strip()[:140]}')

wb = openpyxl.load_workbook('goc/NHÀ THUỐC THÁNG 7 2026.xlsx')

# Step 1: Add Hot Bill HN
ws_hb = wb.create_sheet('Hot Bill HN')
ws_hb.cell(1, 1, 'Mã HD')
wb.save('test_step1.xlsx')
test_open('test_step1.xlsx')

# Step 2: Modify data sheet
ws_d = wb['data']
ws_d.cell(2, 1, 'Hàng Bông')
ws_d.cell(2, 2, 'Lê Ngọc Anh')
ws_d.cell(2, 3, 10)
ws_d.cell(2, 4, 1000000)
wb.save('test_step2.xlsx')
test_open('test_step2.xlsx')

# Step 3: Populate project data in 'Dự án T7'
ws_p = wb['Dự án T7']
ws_p.cell(3, 5, 100000)
ws_p.cell(3, 6, 200000)
ws_p.cell(3, 7, 300000)
wb.save('test_step3.xlsx')
test_open('test_step3.xlsx')

# Step 4: Add new staff in ws_proj
nr_p = ws_p.max_row + 1
ws_p.cell(nr_p, 1, 'Hàng Bông')
ws_p.cell(nr_p, 2, 'New Staff')
ws_p.cell(nr_p, 3, 'BC')
ws_p.cell(nr_p, 4, f'=SUM(E{nr_p}:G{nr_p})/DAY($A$1)')
ws_p.cell(nr_p, 5, 10000)
ws_p.cell(nr_p, 6, 20000)
ws_p.cell(nr_p, 7, 30000)
ws_p.cell(nr_p, 8, f'=E{nr_p}/DAY($A$1)')
ws_p.cell(nr_p, 9, f'=F{nr_p}/DAY($A$1)')
ws_p.cell(nr_p, 10, '=0')
ws_p.cell(nr_p, 11, 0)
ws_p.cell(nr_p, 12, '=0')
ws_p.cell(nr_p, 13, f'=IFERROR(J{nr_p}+K{nr_p}+L{nr_p}, 0)')
wb.save('test_step4.xlsx')
test_open('test_step4.xlsx')

# Step 5: Add new staff in ws_ds
ws_ds = wb['kpi dược sĩ']
nr_d = ws_ds.max_row + 1
ws_ds.cell(nr_d, 1, 'Hàng Bông')
ws_ds.cell(nr_d, 2, 'New Staff')
ws_ds.cell(nr_d, 3, 'BC')
ws_ds.cell(nr_d, 23, f'=IFERROR(SUMIFS(data!$C:$C, data!$A:$A, $A{nr_d}, data!$B:$B, $B{nr_d}), 0)')
ws_ds.cell(nr_d, 24, f'=IFERROR(SUMIFS(data!$D:$D, data!$A:$A, $A{nr_d}, data!$B:$B, $B{nr_d}), 0)')
ws_ds.cell(nr_d, 7, f'=IFERROR(SUMIFS(\'Dự án T7\'!$D:$D, \'Dự án T7\'!$A:$A, $A{nr_d}, \'Dự án T7\'!$B:$B, $B{nr_d}), 0)')
wb.save('test_step5.xlsx')
test_open('test_step5.xlsx')
