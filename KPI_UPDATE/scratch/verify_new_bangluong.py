import openpyxl

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws_bl = wb['BẢNG LƯƠNG']

print("=== KIỂM TRA BẢNG LƯƠNG THÁNG 8 HOÀN THIỆN MỚI ===")
print(f"{'CN':<15} | {'Họ tên':<24} | {'Chức vụ':<8} | {'Giờ ngày':<9} | {'Giờ đêm':<8} | {'Ngày làm':<8} | {'Ngày đêm':<8} | {'Ngày công':<9} | {'Dự án':<10} | {'KPI':<10} | {'Thưởng CK':<10} | {'Cận date':<9} | {'Trừ':<8}")
print("-" * 165)

sample_check = [
    'Hồ Thị Minh Hòa', 'Nguyễn Trần Ngọc Phương', 'Nguyễn Ngọc Anh Thư', 'Trịnh Thị Phượng',
    'Ngô Thị Thanh Thắm', 'Phạm Thị Nghĩa Hương', 'Trần Thiên Phát', 'Hoàng Thanh Thủy',
    'Lê Thị Huyền Trân', 'Thái Thùy Linh', 'Lê Thị Soạn', 'Nguyễn Thị Mai Duyên',
    'Đinh Thị Lan Anh', 'Vũ Thanh Hằng', 'Nguyễn Mạnh Tuấn'
]

for r in range(3, ws_bl.max_row + 1):
    name = ws_bl.cell(r, 3).value
    if name in sample_check:
        cn = str(ws_bl.cell(r, 2).value or '')
        cv = str(ws_bl.cell(r, 4).value or '')
        gn = ws_bl.cell(r, 5).value or 0
        gd = ws_bl.cell(r, 6).value or 0
        sn = ws_bl.cell(r, 10).value or 0
        sd = ws_bl.cell(r, 11).value or 0
        nc = ws_bl.cell(r, 18).value or 0
        da = ws_bl.cell(r, 30).value or 0
        kpi = ws_bl.cell(r, 31).value or 0
        ck = ws_bl.cell(r, 32).value or 0
        cd = ws_bl.cell(r, 33).value or 0
        tru = ws_bl.cell(r, 36).value or 0
        print(f"{cn:<15} | {name:<24} | {cv:<8} | {gn:<9.1f} | {gd:<8.1f} | {sn:<8} | {sd:<8} | {nc:<9.1f} | {da:<10,.0f} | {kpi:<10,.0f} | {ck:<10,.0f} | {cd:<9,.0f} | {tru:<8,.0f}")
