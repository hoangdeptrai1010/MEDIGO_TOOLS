import openpyxl

wb = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx', data_only=True)
ws = wb['Dự án T8']

print(f"{'STT':<4} | {'Chi nhánh':<15} | {'Họ tên':<25} | {'CK/ngày':>12} | {'Combo/ngày':>12} | {'CK+CB/ng':>12} | {'Bậc cũ':>10} | {'BẬC MỚI':>10} | {'Thêm cũ':>10} | {'THÊM MỚI':>10} | {'HotBill':>8} | {'TỔNG MỚI':>10}")
print("-" * 145)

for r in range(3, 54):
    b = ws.cell(r, 1).value
    name = ws.cell(r, 2).value
    if not name: continue
    
    ck = ws.cell(r, 6).value or 0
    cb = ws.cell(r, 7).value or 0
    ny3 = ws.cell(r, 5).value or 0
    hb = ws.cell(r, 12).value or 0
    
    old_j = ws.cell(r, 10).value or 0
    old_k = ws.cell(r, 11).value or 0
    old_m = ws.cell(r, 13).value or 0
    
    ck_d = ck / 31.0
    cb_d = cb / 31.0
    ck_cb_d = (ck + cb) / 31.0
    
    # New Tier Calculation according to August Slide
    is_hn = b in ['Hàng Bông', 'Đường Láng']
    new_j = 0
    if is_hn:
        if ck_d >= 3400000 and cb_d >= 900000: new_j = 4000000
        elif ck_d >= 3200000 and cb_d >= 750000: new_j = 2800000
        elif ck_d >= 2800000 and cb_d >= 650000: new_j = 2200000
        elif ck_d >= 2500000 and cb_d >= 520000: new_j = 1600000
        elif ck_d >= 2200000 and cb_d >= 450000: new_j = 1000000
    else:
        if ck_d >= 1500000 and cb_d >= 700000: new_j = 4000000
        elif ck_d >= 1200000 and cb_d >= 650000: new_j = 2800000
        elif ck_d >= 1000000 and cb_d >= 580000: new_j = 2200000
        elif ck_d >= 850000 and cb_d >= 480000: new_j = 1600000
        elif ck_d >= 630000 and cb_d >= 350000: new_j = 1000000
        
    # New Extra 500k Calculation (CK + Combo, no NY3)
    new_k = 0
    if new_j == 0:
        min_extra = 2000000 if is_hn else 950000
        if ck_cb_d >= min_extra:
            new_k = 500000
            
    new_m = new_j + new_k + hb
    
    diff_mark = " *** THAY ĐỔI ***" if new_m != old_m else ""
    print(f"{r:<4} | {str(b):<15} | {str(name):<25} | {ck_d:>12,.0f} | {cb_d:>12,.0f} | {ck_cb_d:>12,.0f} | {old_j:>10,.0f} | {new_j:>10,.0f} | {old_k:>10,.0f} | {new_k:>10,.0f} | {hb:>8,.0f} | {new_m:>10,.0f}{diff_mark}")
