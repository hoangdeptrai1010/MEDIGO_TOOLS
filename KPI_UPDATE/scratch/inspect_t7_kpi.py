import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb7_kpi = openpyxl.load_workbook(r'baocaokpi_thang7_hoanthien.xlsx', data_only=True)
print("Sheets baocaokpi_thang7_hoanthien:", wb7_kpi.sheetnames)

wb7_goc = openpyxl.load_workbook(r'goc\NHÀ THUỐC THÁNG 7 2026.xlsx', data_only=True)
print("Sheets goc T7:", wb7_goc.sheetnames)

for sname in wb7_goc.sheetnames:
    if any(k in sname.lower() for k in ['ck', 'chiết khấu', 'dự án', 'thưởng']):
        ws = wb7_goc[sname]
        print(f"\n[GỐC T7] {sname}:")
        for r in range(1, 6):
            print(f"  R{r}: {[ws.cell(r, c).value for c in range(1, 12)]}")
