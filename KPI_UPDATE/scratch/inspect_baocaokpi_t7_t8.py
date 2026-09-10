import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

print("==================== baocaokpi_thang7_hoanthien.xlsx ====================")
wb7 = openpyxl.load_workbook(r'baocaokpi_thang7_hoanthien.xlsx', data_only=False)
wb7_v = openpyxl.load_workbook(r'baocaokpi_thang7_hoanthien.xlsx', data_only=True)
print("Sheets T7:", wb7.sheetnames)

for sname in ['Dự án T7', 'Thưởng CK', 'kpi nhà thuốc', 'Chiết khấu - combo - NY3']:
    if sname in wb7.sheetnames:
        ws = wb7[sname]
        wsv = wb7_v[sname]
        print(f"\n--- Sheet: {sname} ---")
        for r in range(1, 8):
            f_row = [ws.cell(r, c).value for c in range(1, 15)]
            v_row = [wsv.cell(r, c).value for c in range(1, 15)]
            print(f"R{r} FORMULA: {f_row}")
            print(f"R{r} VALUE:   {v_row}")

print("\n==================== baocaokpi_thang8_hoanthien.xlsx ====================")
wb8 = openpyxl.load_workbook(r'baocaokpi_thang8_hoanthien.xlsx', data_only=False)
wb8_v = openpyxl.load_workbook(r'baocaokpi_thang8_hoanthien.xlsx', data_only=True)
print("Sheets T8:", wb8.sheetnames)

for sname in ['Dự án T8', 'Thưởng CK', 'kpi nhà thuốc', 'Chiết khấu - combo - NY3']:
    if sname in wb8.sheetnames:
        ws = wb8[sname]
        wsv = wb8_v[sname]
        print(f"\n--- Sheet: {sname} ---")
        for r in range(1, 8):
            f_row = [ws.cell(r, c).value for c in range(1, 15)]
            v_row = [wsv.cell(r, c).value for c in range(1, 15)]
            print(f"R{r} FORMULA: {f_row}")
            print(f"R{r} VALUE:   {v_row}")

