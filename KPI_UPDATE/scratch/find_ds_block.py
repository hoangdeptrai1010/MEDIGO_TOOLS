import openpyxl
import subprocess
import os

wb_orig = openpyxl.load_workbook('goc/NHÀ THUỐC THÁNG 7 2026.xlsx')

def test_wb(wb, label):
    fname = f'test_{label}.xlsx'
    wb.save(fname)
    wb.close()
    p = os.path.abspath(fname)
    res = subprocess.run(['powershell', '-ExecutionPolicy', 'Bypass', '-File', 'recalc_workbook.ps1', '-FilePath', p], capture_output=True, text=True)
    is_ok = 'Recalculation complete' in res.stdout
    print(f'Test {label}: {"SUCCESS" if is_ok else "FAILED"}')

# Sub-test A: Only update template rows in ws_ds with SUMIFS
wb_a = openpyxl.load_workbook('test_with_orig_ds.xlsx')
ws_a = wb_a['kpi dược sĩ']
for r in range(3, 60):
    ws_a.cell(r, 23, f"=IFERROR(SUMIFS(data!$C:$C, data!$A:$A, $A{r}, data!$B:$B, $B{r}), 0)")
    ws_a.cell(r, 24, f"=IFERROR(SUMIFS(data!$D:$D, data!$A:$A, $A{r}, data!$B:$B, $B{r}), 0)")
    ws_a.cell(r, 26, f"=IFERROR(SUMIFS(data!$E:$E, data!$A:$A, $A{r}, data!$B:$B, $B{r}), 0)")
    ws_a.cell(r, 27, f"=IFERROR(SUMIFS(data!$F:$F, data!$A:$A, $A{r}, data!$B:$B, $B{r}), 0)")
    ws_a.cell(r, 7, f"=IFERROR(SUMIFS('Dự án T7'!$D:$D, 'Dự án T7'!$A:$A, $A{r}, 'Dự án T7'!$B:$B, $B{r}), 0)")
    ws_a.cell(r, 13, f"=IFERROR(SUMIFS('Dự án T7'!$M:$M, 'Dự án T7'!$A:$A, $A{r}, 'Dự án T7'!$B:$B, $B{r})*$L{r}, 0)")
test_wb(wb_a, 'part_a_template_sumifs')

# Sub-test B: Append new staff rows
wb_b = openpyxl.load_workbook('test_with_orig_ds.xlsx')
ws_b = wb_b['kpi dược sĩ']
next_ds_row = ws_b.max_row + 1
ws_b.cell(next_ds_row, 1, 'Hàng Bông')
ws_b.cell(next_ds_row, 2, 'New Pharmacist')
ws_b.cell(next_ds_row, 3, 'BC')
ws_b.cell(next_ds_row, 4, 150000)
ws_b.cell(next_ds_row, 5, f"=AE{next_ds_row}")
ws_b.cell(next_ds_row, 8, f"=Q{next_ds_row}")
ws_b.cell(next_ds_row, 9, f"=(AD{next_ds_row}/DAY($C$1))")
ws_b.cell(next_ds_row, 10, f"=I{next_ds_row}/H{next_ds_row}")
ws_b.cell(next_ds_row, 11, f'=IF(C{next_ds_row}="DSTV", 0, IF(J{next_ds_row}>=1, AD{next_ds_row}*0.015, IF(J{next_ds_row}>=0.9, AD{next_ds_row}*0.013, IF(J{next_ds_row}>=0.8, AD{next_ds_row}*0.01, 0))))')
ws_b.cell(next_ds_row, 12, f'=IF(C{next_ds_row}="DSTV","", IF(J{next_ds_row}<0.6, 0.8, 1))')
ws_b.cell(next_ds_row, 14, f'=K{next_ds_row}+M{next_ds_row}*L{next_ds_row}')
ws_b.cell(next_ds_row, 15, f'=CEILING(R{next_ds_row}, 100000)')
ws_b.cell(next_ds_row, 16, f'=CEILING(S{next_ds_row}, 100000)')
ws_b.cell(next_ds_row, 17, f'=CEILING(T{next_ds_row}, 100000)')
ws_b.cell(next_ds_row, 18, f'=U{next_ds_row}*$R$2')
ws_b.cell(next_ds_row, 19, f'=U{next_ds_row}*$S$2')
ws_b.cell(next_ds_row, 20, f'=U{next_ds_row}*$T$2')
ws_b.cell(next_ds_row, 21, f"=V{next_ds_row}/DAY($C$1)")
ws_b.cell(next_ds_row, 22, 270000000)
ws_b.cell(next_ds_row, 23, f"=IFERROR(SUMIFS(data!$C:$C, data!$A:$A, $A{next_ds_row}, data!$B:$B, $B{next_ds_row}), 0)")
ws_b.cell(next_ds_row, 24, f"=IFERROR(SUMIFS(data!$D:$D, data!$A:$A, $A{next_ds_row}, data!$B:$B, $B{next_ds_row}), 0)")
ws_b.cell(next_ds_row, 25, f'=IF(W{next_ds_row}>0, X{next_ds_row}/W{next_ds_row}, 0)')
ws_b.cell(next_ds_row, 26, f"=IFERROR(SUMIFS(data!$E:$E, data!$A:$A, $A{next_ds_row}, data!$B:$B, $B{next_ds_row}), 0)")
ws_b.cell(next_ds_row, 27, f"=IFERROR(SUMIFS(data!$F:$F, data!$A:$A, $A{next_ds_row}, data!$B:$B, $B{next_ds_row}), 0)")
ws_b.cell(next_ds_row, 28, f'=IF(Z{next_ds_row}>0, AA{next_ds_row}/Z{next_ds_row}, 0)')
ws_b.cell(next_ds_row, 29, f'=W{next_ds_row}+Z{next_ds_row}')
ws_b.cell(next_ds_row, 30, f'=X{next_ds_row}+AA{next_ds_row}')
ws_b.cell(next_ds_row, 31, f'=IF(AC{next_ds_row}>0, AD{next_ds_row}/AC{next_ds_row}, 0)')
test_wb(wb_b, 'part_b_new_staff')
