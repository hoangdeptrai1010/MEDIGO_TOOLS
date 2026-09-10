import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import sys
import unicodedata
import re

sys.stdout.reconfigure(encoding='utf-8')

def remove_accents(input_str):
    if not input_str:
        return ""
    nfkd_form = unicodedata.normalize('NFKD', str(input_str))
    res = "".join([c for c in nfkd_form if not unicodedata.combining(c)])
    # replace đ, Đ
    res = res.replace('đ', 'd').replace('Đ', 'D')
    return res

def normalize_text(s):
    if not s:
        return ""
    s = unicodedata.normalize('NFC', str(s))
    return " ".join(s.lower().strip().split())

def clean_stk(stk):
    if stk is None:
        return ""
    stk_str = str(stk).strip()
    if stk_str in ["#N/A", "#REF!", "None", "nan", ""]:
        return ""
    stk_clean = re.sub(r'[\s\.\-_]', '', stk_str)
    if stk_clean.endswith(".0"):
        stk_clean = stk_clean[:-2]
    return stk_clean

def normalize_bank(b):
    if not b:
        return ""
    b_str = str(b).strip()
    if b_str in ["#N/A", "#REF!", "None", "nan", ""]:
        return ""
    b_norm = remove_accents(b_str).lower()
    b_norm = re.sub(r'[^a-z0-9]', '', b_norm)
    return b_norm

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_bl = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final.xlsx")
f_backup = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final_backup.xlsx")

# Load old Sheet1 from backup
wb_old = openpyxl.load_workbook(f_backup, data_only=True)
ws1_old = wb_old["Sheet1"]

old_records = {} # key: (norm_store, norm_name) -> old data dict
for r in range(2, ws1_old.max_row + 1):
    stt = ws1_old.cell(r, 1).value
    store = ws1_old.cell(r, 2).value
    name = ws1_old.cell(r, 3).value
    acc_holder = ws1_old.cell(r, 4).value
    stk = ws1_old.cell(r, 5).value
    bank = ws1_old.cell(r, 6).value
    amount = ws1_old.cell(r, 7).value
    content = ws1_old.cell(r, 8).value
    
    if name is not None and str(name).strip():
        n_name = normalize_text(name)
        n_store = normalize_text(remove_accents(store))
        key = (n_store, n_name)
        old_records[key] = {
            'old_row': r,
            'stt': stt,
            'store': store,
            'name': name,
            'acc_holder': acc_holder,
            'stk': str(stk).strip() if stk is not None else "",
            'bank': str(bank).strip() if bank is not None else "",
            'amount': amount,
            'content': content
        }
        # Also map by name alone for fallback matching
        if n_name not in old_records:
            old_records[n_name] = old_records[key]

print(f"Loaded {len(old_records)} old records from old Sheet1.")

# Load main workbook to modify
wb = openpyxl.load_workbook(f_bl)
ws_med = wb["Med"]
ws1 = wb["Sheet1"]

# Read all real employee rows grouped by store from Med
stores_med = {}
current_store = ""

for r in range(10, ws_med.max_row + 1):
    stt = ws_med.cell(r, 1).value
    code = ws_med.cell(r, 3).value
    name = ws_med.cell(r, 4).value
    store = ws_med.cell(r, 6).value
    role = ws_med.cell(r, 7).value
    stk = ws_med.cell(r, 9).value
    bank = ws_med.cell(r, 10).value
    net_amount = ws_med.cell(r, 85).value # Col CG: CK TK
    if net_amount is None:
        net_amount = ws_med.cell(r, 80).value # Col CB
        
    if store and str(store).strip():
        current_store = str(store).strip()
        if current_store not in stores_med:
            stores_med[current_store] = []
            
    if name and str(name).strip() and not str(name).strip().startswith("CHUỖI") and not str(name).strip().startswith("Tổng"):
        stores_med[current_store].append({
            'med_row': r,
            'stt': stt,
            'code': str(code).strip() if code else "",
            'name': str(name).strip(),
            'store': current_store,
            'role': str(role).strip() if role else "",
            'stk': str(stk).strip() if stk is not None else "",
            'bank': str(bank).strip() if bank is not None else "",
            'amount': net_amount
        })

print(f"Loaded {len(stores_med)} stores with {sum(len(v) for v in stores_med.values())} employees from Med.")

# Clear old rows in Sheet1 (keep header row 1)
for r in range(2, ws1.max_row + 50):
    for c in range(1, 15):
        ws1.cell(r, c).value = None
        ws1.cell(r, c).fill = PatternFill(fill_type=None)
        ws1.cell(r, c).font = Font(name="Calibri", size=11)
        ws1.cell(r, c).border = Border()

# Setup styles
red_fill = PatternFill(start_color="FFFF9999", end_color="FFFF9999", fill_type="solid")
red_font = Font(name="Calibri", size=11, bold=True, color="FFCC0000")
normal_font = Font(name="Calibri", size=11)
bold_font = Font(name="Calibri", size=11, bold=True)
thin_border = Border(
    left=Side(style='thin', color='FFD9D9D9'),
    right=Side(style='thin', color='FFD9D9D9'),
    top=Side(style='thin', color='FFD9D9D9'),
    bottom=Side(style='thin', color='FFD9D9D9')
)

current_r = 2
running_stt = 1
comparison_report = []

for store_name, emp_list in stores_med.items():
    store_clean_noacc = remove_accents(store_name).replace(" ", "")
    start_store_row = current_r
    
    for emp in emp_list:
        emp_name = emp['name']
        emp_code = emp['code'] or remove_accents(emp_name)
        emp_stk = emp['stk']
        emp_bank = emp['bank']
        emp_amt = emp['amount']
        
        # Check against old records
        n_store = normalize_text(remove_accents(store_name))
        n_name = normalize_text(emp_name)
        old_data = old_records.get((n_store, n_name)) or old_records.get(n_name)
        
        is_existing_in_old = (old_data is not None)
        stk_mismatch = False
        bank_mismatch = False
        amt_mismatch = False
        
        if is_existing_in_old:
            old_stk_clean = clean_stk(old_data['stk'])
            new_stk_clean = clean_stk(emp_stk)
            if old_stk_clean != new_stk_clean or not old_stk_clean:
                stk_mismatch = True
                
            old_bank_norm = normalize_bank(old_data['bank'])
            new_bank_norm = normalize_bank(emp_bank)
            if old_bank_norm != new_bank_norm or not old_bank_norm:
                if old_bank_norm not in new_bank_norm and new_bank_norm not in old_bank_norm:
                    bank_mismatch = True
                    
            old_amt = old_data['amount']
            try:
                if old_amt is None or str(old_amt).startswith("#") or abs(float(old_amt) - float(emp_amt or 0)) > 1.0:
                    amt_mismatch = True
            except:
                amt_mismatch = True
                
        # Fill cells in Sheet1
        c_stt = ws1.cell(current_r, 1, value=running_stt)
        c_store = ws1.cell(current_r, 2, value=store_name)
        c_name = ws1.cell(current_r, 3, value=emp_name)
        c_acc = ws1.cell(current_r, 4, value=emp_code)
        c_stk = ws1.cell(current_r, 5, value=emp_stk)
        c_bank = ws1.cell(current_r, 6, value=emp_bank)
        c_amt = ws1.cell(current_r, 7, value=emp_amt)
        
        content_val = f"Med  TT Luong T8.2026 -{store_clean_noacc}-{emp_code}"
        c_content = ws1.cell(current_r, 8, value=content_val)
        c_tag = ws1.cell(current_r, 9, value=store_name)
        
        # Apply basic formatting
        for cell in [c_stt, c_store, c_name, c_acc, c_stk, c_bank, c_amt, c_content, c_tag]:
            cell.font = normal_font
            cell.border = thin_border
            
        c_stt.alignment = Alignment(horizontal="center")
        c_amt.number_format = '#,##0'
        
        # Apply Red highlight if mismatched
        if is_existing_in_old:
            if stk_mismatch:
                c_stk.fill = red_fill
                c_stk.font = red_font
            if bank_mismatch:
                c_bank.fill = red_fill
                c_bank.font = red_font
            if amt_mismatch:
                c_amt.fill = red_fill
                c_amt.font = red_font
                
            comparison_report.append({
                'row': current_r,
                'store': store_name,
                'name': emp_name,
                'type': 'EXISTING_IN_OLD',
                'stk_diff': stk_mismatch,
                'old_stk': old_data['stk'],
                'new_stk': emp_stk,
                'bank_diff': bank_mismatch,
                'old_bank': old_data['bank'],
                'new_bank': emp_bank,
                'amt_diff': amt_mismatch,
                'old_amt': old_data['amount'],
                'new_amt': emp_amt
            })
        else:
            comparison_report.append({
                'row': current_r,
                'store': store_name,
                'name': emp_name,
                'type': 'NEW_FROM_MED',
                'stk_diff': False,
                'old_stk': '(Mới)',
                'new_stk': emp_stk,
                'bank_diff': False,
                'old_bank': '(Mới)',
                'new_bank': emp_bank,
                'amt_diff': False,
                'old_amt': '(Mới)',
                'new_amt': emp_amt
            })
            
        running_stt += 1
        current_r += 1
        
    # Add Subtotal Row for this store
    end_store_row = current_r - 1
    c_subtotal_amt = ws1.cell(current_r, 7, value=f"=SUM(G{start_store_row}:G{end_store_row})")
    c_subtotal_amt.font = bold_font
    c_subtotal_amt.number_format = '#,##0'
    c_subtotal_amt.border = thin_border
    
    ws1.cell(current_r, 2, value=f"Tổng {store_name}").font = bold_font
    current_r += 1

# Save the updated workbook
wb.save(f_bl)
print(f"Saved updated Sheet1 with {current_r - 2} rows to {f_bl}")

print("\n==========================================================================================")
print(f"{'Row':<4} | {'Cửa hàng':<15} | {'Họ và tên':<24} | {'Trạng thái':<14} | {'STK Cũ -> Mới':<28} | {'Ngân hàng Cũ -> Mới':<32} | {'Số tiền Cũ -> Mới':<24}")
print("==========================================================================================")
for rep in comparison_report:
    stk_str = f"{str(rep['old_stk'])} -> {str(rep['new_stk'])}" if rep['stk_diff'] else str(rep['new_stk'])
    bank_str = f"{str(rep['old_bank'])} -> {str(rep['new_bank'])}" if rep['bank_diff'] else str(rep['new_bank'])
    amt_str = f"{str(rep['old_amt'])} -> {str(rep['new_amt'])}" if rep['amt_diff'] else str(rep['new_amt'])
    status = "ĐÃ ĐỐI SOÁT" if rep['type'] == 'EXISTING_IN_OLD' else "THÊM MỚI"
    if rep['stk_diff'] or rep['bank_diff'] or rep['amt_diff']:
        status = "KHÁC NHAU (ĐỎ)"
    print(f"{rep['row']:<4} | {rep['store']:<15} | {rep['name']:<24} | {status:<14} | {stk_str:<28} | {bank_str:<32} | {amt_str:<24}")
