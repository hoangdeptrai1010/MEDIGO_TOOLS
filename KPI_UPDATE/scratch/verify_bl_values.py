import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws = wb['BẢNG LƯƠNG']

headers = [ws.cell(2, c).value for c in range(1, 45)]

print(f"Total rows in BẢNG LƯƠNG: {ws.max_row}")

target_names = ['Hồ Thị Minh Hòa', 'Phan Công Vũ Tài', 'Hoàng Lâm Gia Bảo', 'Nguyễn Trần Ngọc Phương']

for r in range(3, 75):
    name = ws.cell(r, 3).value
    cn = ws.cell(r, 2).value
    role = ws.cell(r, 4).value
    
    if any(tn in str(name) for tn in target_names) or r == 74:
        stt = ws.cell(r, 1).value
        gio_ngay = ws.cell(r, 7).value
        gio_dem = ws.cell(r, 8).value
        chuyen_can = ws_cell_val = ws.cell(r, 9).value
        ngay_thuc = ws.cell(r, 10).value
        ngay_dem = ws.cell(r, 11).value
        ngay_cong = ws.cell(r, 18).value
        pc_cht = ws.cell(r, 22).value
        pc_dem = ws.cell(r, 24).value
        kat_ds = ws.cell(r, 27).value
        kat_cht = ws.cell(r, 28).value
        wa = ws.cell(r, 29).value
        kpi = ws.cell(r, 31).value
        ck = ws.cell(r, 32).value
        can_date = ws.cell(r, 33).value
        
        print(f"\nRow {r:2d} | STT={stt} | CN={cn} | Tên={name} | Vị trí={role}:")
        print(f"  - Giờ ca ngày (Col G): {gio_ngay}, Giờ ca đêm (Col H): {gio_dem}")
        print(f"  - Thưởng Chuyên cần (Col I): {chuyen_can}")
        print(f"  - Ngày làm ngày (Col J): {ngay_thuc}, Ngày làm đêm (Col K): {ngay_dem}, Ngày công thực tế (Col R): {ngay_cong}")
        print(f"  - Phụ cấp CHT (Col V): {pc_cht:,.0f} đ" if isinstance(pc_cht, (int, float)) else f"  - Phụ cấp CHT: {pc_cht}")
        print(f"  - Phụ cấp ca đêm (Col X): {pc_dem:,.0f} đ" if isinstance(pc_dem, (int, float)) else f"  - Phụ cấp ca đêm: {pc_dem}")
        print(f"  - Thưởng MiniKat DS (Col AA): {kat_ds:,.0f} đ" if isinstance(kat_ds, (int, float)) else f"  - Thưởng MiniKat DS: {kat_ds}")
        print(f"  - Thưởng MiniKat CHT (Col AB): {kat_cht:,.0f} đ" if isinstance(kat_cht, (int, float)) else f"  - Thưởng MiniKat CHT: {kat_cht}")
        print(f"  - Thưởng WhatsApp (Col AC): {wa:,.0f} đ" if isinstance(wa, (int, float)) else f"  - Thưởng WhatsApp: {wa}")
        print(f"  - Thưởng KPI (Col AE): {kpi:,.0f} đ" if isinstance(kpi, (int, float)) else f"  - Thưởng KPI: {kpi}")
        print(f"  - Thưởng CK (Col AF): {ck:,.0f} đ" if isinstance(ck, (int, float)) else f"  - Thưởng CK: {ck}")
        print(f"  - Thưởng Cận date (Col AG): {can_date:,.0f} đ" if isinstance(can_date, (int, float)) else f"  - Thưởng Cận date: {can_date}")

print("\n--- Chi nhánh Summary Table (Rows 77-88) ---")
for r in range(77, 89):
    bname = ws.cell(r, 5).value
    g_ngay = ws.cell(r, 7).value
    g_dem = ws.cell(r, 8).value
    g_tot = ws.cell(r, 9).value
    print(f"Row {r:2d} | CN: {bname:<16} | Giờ ngày: {g_ngay:8.2f} | Giờ đêm: {g_dem:8.2f} | Tổng giờ: {g_tot:8.2f}" if isinstance(g_ngay, (int, float)) else f"Row {r:2d} | CN: {bname} | {g_ngay}")
