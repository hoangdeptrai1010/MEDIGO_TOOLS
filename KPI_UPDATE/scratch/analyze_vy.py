import openpyxl

wb = openpyxl.load_workbook('thang7/DATAKIOT/hoadon.xlsx', data_only=True, read_only=True)
ws = wb.active
headers = next(ws.iter_rows(max_row=1, values_only=True))
col_map = {h: idx for idx, h in enumerate(headers) if h}

c_seller = col_map.get('Người bán', 6)
c_date = col_map.get('Thời gian', 2)
c_sku = col_map.get('Mã hàng', 13)
c_name = col_map.get('Tên hàng', 15)
c_exp = col_map.get('Hạn sử dụng', 18)
c_val = col_map.get('Thành tiền', 25)

vy_items = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r: continue
    s = str(r[c_seller] or '').strip()
    if s == 'Cù Thị Tường Vy':
        exp = r[c_exp]
        val = float(r[c_val] or 0)
        if val > 0:
            vy_items.append({
                'date': str(r[c_date])[:10],
                'exp': str(exp)[:10] if exp else 'NO_EXP',
                'sku': r[c_sku],
                'name': r[c_name],
                'val': val
            })

tot_val = sum(it['val'] for it in vy_items)
print('Total items sold by Cu Thi Tuong Vy:', len(vy_items))
print(f'Total revenue: {tot_val:,.0f}đ')

# Target is 4,911,200đ
sku_sums = {}
for it in vy_items:
    sku = it['sku']
    if sku not in sku_sums:
        sku_sums[sku] = {'val': 0, 'name': it['name'], 'exp': it['exp']}
    sku_sums[sku]['val'] += it['val']

sorted_skus = sorted(sku_sums.items(), key=lambda x: x[1]['val'], reverse=True)
print('\nTop SKUs sold by Cu Thi Tuong Vy:')
for sku, info in sorted_skus[:35]:
    exp_str = info['exp']
    val_str = f"{info['val']:>10,.0f}đ"
    name_str = str(info['name'])[:35]
    print(f"SKU: {sku:12s} | HSD: {exp_str:10s} | {val_str} | {name_str}")
