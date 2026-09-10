import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
ws = wb.active

for r in range(5, ws.max_row + 1):
    c_name = str(ws.cell(r, 3).value or '').strip()
    if 'Thoa' in c_name:
        sh = ws.cell(r, 7).value
        br = ws.cell(r, 6).value
        print(f"\nRow {r}: {c_name} - {br} - Shift: {sh}")
        for day in range(1, 32):
            col_in = 8 + (day - 1) * 2
            col_out = col_in + 1
            v_in = ws.cell(r, col_in).value
            v_out = ws.cell(r, col_out).value
            if v_in and v_out:
                print(f"  Day {day}: In {v_in} -> Out {v_out}")

wb.close()
