# -*- coding: utf-8 -*-
import openpyxl, sys, re
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/BANGLUONGTHANG8.xlsx', data_only=True)
ws_nc = wb['Ngày công']

print("=== CHECKING ALL SHIFT DEFINITIONS IN NGÀY CÔNG ===")
mismatches = 0
for r in range(3, ws_nc.max_row+1):
    c_name = ws_nc.cell(r, 2).value
    c_ca = str(ws_nc.cell(r, 4).value or '').strip()
    c_h = ws_nc.cell(r, 6).value
    c_night = ws_nc.cell(r, 9).value # Col I: Ca đêm chuẩn
    c_day = ws_nc.cell(r, 10).value # Col J: Ca ngày chuẩn
    
    if not c_ca or c_ca == '-':
        continue
        
    m = re.search(r'(\d{1,2}):(\d{2})\s*-\s*(\d{1,2}):(\d{2})', c_ca)
    if m and isinstance(c_h, (int, float)) and c_h >= 4:
        start_h = int(m.group(1))
        is_night_time = (start_h >= 22 or start_h < 6 or 'qua đêm' in c_ca.lower())
        
        if is_night_time and c_night != 1:
            print(f"Mismatch Night Row {r}: {c_name} | {c_ca} | h={c_h} | Col I={c_night}")
            mismatches += 1
        elif not is_night_time and c_day != 1:
            print(f"Mismatch Day Row {r}: {c_name} | {c_ca} | h={c_h} | Col J={c_day}")
            mismatches += 1

print(f"Total Mismatches in Ngày công: {mismatches}")
