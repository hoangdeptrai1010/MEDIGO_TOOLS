import openpyxl

wb = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)

for sname in ['Bảng tổng hợp chấm công', 'Bảng chi tiết chấm công']:
    ws = wb[sname]
    print(f"=== {sname} for Hồ Thị Minh Hòa ===")
    for r in range(1, ws.max_row + 1):
        row_str = " ".join([str(ws.cell(r, c).value or '') for c in range(1, 15)])
        if 'Hồ Thị Minh Hòa' in row_str or 'Minh Hòa' in row_str:
            vals = [ws.cell(r, c).value for c in range(1, 15) if ws.cell(r, c).value is not None]
            print(f"Row {r}: {vals}")
