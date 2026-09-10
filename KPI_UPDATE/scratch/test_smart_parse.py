import openpyxl
import re
import unicodedata
import calendar
import sys, io

if sys.stdout.encoding != 'utf-8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

def remove_accents(input_str):
    if not input_str:
        return ""
    nfkd_form = unicodedata.normalize('NFKD', str(input_str))
    res = "".join([c for c in nfkd_form if not unicodedata.combining(c)])
    return res.replace('đ', 'd').replace('Đ', 'D')

def clean_branch_name(cn):
    if not cn:
        return ''
    cn = str(cn).strip()
    mapping = {
        'NT Hàng Bông 247': 'Hàng Bông',
        'NT Đường Láng 247': 'Đường Láng',
        'NT 24H Minh Châu 1': 'Minh Châu',
        'NT 24h Đỗ Quang Đẩu': 'Đỗ Quang Đẩu',
        'NT 24H Lê Bình': 'Lê Bình',
        'NT 24H Nam Hòa': 'Nam Hòa',
        'NT 24H Nguyễn Chí Thanh': 'Nguyễn Chí Thanh',
        'NT 24H Nguyễn Thị Thập': 'Nguyễn Thị Thập',
        'NT 24H Nguyễn Văn Quá': 'Nguyễn Văn Quá',
        'NT 24H Rạch Bùng Binh': 'Rạch Bùng Binh',
        'NT 24H Trường Sa': 'Trường Sa'
    }
    if cn in mapping:
        return mapping[cn]
    cleaned = re.sub(r'^NT\s+(24H\s+|24h\s+)?', '', cn, flags=re.IGNORECASE)
    cleaned = re.sub(r'\s+247$', '', cleaned)
    cleaned = re.sub(r'\s+1$', '', cleaned)
    return cleaned.strip()

def smart_parse_hcm_proposal(wb, month_num=9):
    # Select sheet
    target_sheet = None
    for name in wb.sheetnames:
        n_low = remove_accents(name).lower()
        if 'kpi' in n_low and (f'thang {month_num:02d}' in n_low or f'thang {month_num}' in n_low or f't{month_num}' in n_low):
            target_sheet = name
            break
    if not target_sheet:
        for name in wb.sheetnames:
            n_low = remove_accents(name).lower()
            if 'kpi' in n_low:
                target_sheet = name
                break
    if not target_sheet:
        target_sheet = wb.sheetnames[-1]

    ws = wb[target_sheet]
    print(f"--> Using HCM Sheet: '{target_sheet}'")

    # Detect header row (row 1 to 4)
    h_row = 1
    for r in range(1, 5):
        row_str = " ".join([str(ws.cell(r, c).value or '') for c in range(1, 15)])
        row_str_norm = remove_accents(row_str).lower()
        if 'nhan vien' in row_str_norm or 'duoc si' in row_str_norm:
            h_row = r
            break

    headers = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(h_row, c).value
        if v:
            headers[c] = remove_accents(str(v)).lower().strip()

    print("Detected headers:", headers)

    col_store = next((c for c, h in headers.items() if 'nha thuoc' in h or 'chi nhanh' in h), 1)
    col_name = next((c for c, h in headers.items() if 'nhan vien' in h or 'duoc si' in h), 2)
    col_role = next((c for c, h in headers.items() if 'chuc danh' in h or 'vi tri' in h), 3)
    col_tb = next((c for c, h in headers.items() if 'trung binh bill' in h or 'tb bill' in h or 'tbbill' in h), 4)
    col_gd = next((c for c, h in headers.items() if 'giao dich' in h or 'gd' in h), 5)
    
    # Target month: match f"thang {month_num}" or f"thang {month_num:02d}" or fallback
    col_target = next((c for c, h in headers.items() if f'thang {month_num}' in h or f'thang {month_num:02d}' in h or f't{month_num}' in h), None)
    if not col_target:
        col_target = next((c for c, h in headers.items() if 'doanh thu/thang' in h or 'kpi thang' in h), 7)
        
    col_nt = next((c for c, h in headers.items() if 'nt thang' in h or 'nha thuoc thang' in h or 'target nt' in h), None)

    print(f"Columns mapped: store={col_store}, name={col_name}, role={col_role}, tb_bill={col_tb}, gd={col_gd}, target={col_target}, nt_target={col_nt}")

    staff_list = []
    store_targets = {}
    current_store = ''

    for r in range(h_row + 1, ws.max_row + 1):
        st_v = ws.cell(r, col_store).value
        name_v = ws.cell(r, col_name).value
        role_v = ws.cell(r, col_role).value
        tb_v = ws.cell(r, col_tb).value
        gd_v = ws.cell(r, col_gd).value if col_gd else 0
        tgt_v = ws.cell(r, col_target).value if col_target else 0
        nt_v = ws.cell(r, col_nt).value if col_nt else 0

        if st_v:
            current_store = clean_branch_name(str(st_v))

        if nt_v and current_store:
            try:
                store_targets[current_store] = float(nt_v)
            except Exception:
                pass

        if name_v and current_store:
            s_name = str(name_v).strip()
            if s_name and s_name.lower() not in ('nhan vien', 'tong', 'none', 'tong cong'):
                staff_list.append({
                    'store': current_store,
                    'name': s_name,
                    'role': str(role_v or 'NV').strip(),
                    'tb_bill': float(tb_v or 0) if isinstance(tb_v, (int, float)) else 0.0,
                    'giao_dich_ngay': float(gd_v or 0) if isinstance(gd_v, (int, float)) else 0.0,
                    'kpi_thang': float(tgt_v or 0) if isinstance(tgt_v, (int, float)) else 0.0,
                    'region': 'HCM'
                })

    return staff_list, store_targets

wb = openpyxl.load_workbook('TOOL_KPISHEET/KPI CNT HCM Tháng 09.xlsx', data_only=True)
s9, t9 = smart_parse_hcm_proposal(wb, 9)
print(f"Month 9 result: {len(s9)} staff, {len(t9)} stores")
for s in s9[:5]:
    print(" ", s)
print("Stores:", t9)
