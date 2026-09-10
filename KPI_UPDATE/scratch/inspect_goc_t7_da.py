import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'goc\NHÀ THUỐC THÁNG 7 2026.xlsx', data_only=True)
ws = wb['Dự án T7']

print(f"{'Nhà thuốc':<15} | {'Nhân viên':<25} | {'Chức danh':<10} | {'CK+Combo+NY3/ng':<15} | {'Chiết khấu':<12} | {'Combo':<10} | {'Thưởng DA':<12} | {'Thưởng thêm':<12} | {'Total DA (Col M)':<15}")
print("-" * 130)

for r in range(2, ws.max_row + 1):
    row_vals = [ws.cell(r, c).value for c in range(1, 15)]
    if any(v is not None for v in row_vals):
        nt = str(ws.cell(r, 1).value or '')
        nv = str(ws.cell(r, 2).value or '')
        cd = str(ws.cell(r, 3).value or '')
        avg_rev = ws.cell(r, 4).value
        ck = ws.cell(r, 6).value
        combo = ws.cell(r, 7).value
        thuong_da = ws.cell(r, 10).value
        thuong_them = ws.cell(r, 11).value
        total_da = ws.cell(r, 13).value
        
        avg_str = f"{avg_rev:,.0f}" if isinstance(avg_rev, (int, float)) else str(avg_rev)
        ck_str = f"{ck:,.0f}" if isinstance(ck, (int, float)) else str(ck)
        combo_str = f"{combo:,.0f}" if isinstance(combo, (int, float)) else str(combo)
        tda_str = f"{thuong_da:,.0f}" if isinstance(thuong_da, (int, float)) else str(thuong_da)
        tt_str = f"{thuong_them:,.0f}" if isinstance(thuong_them, (int, float)) else str(thuong_them)
        tot_str = f"{total_da:,.0f}" if isinstance(total_da, (int, float)) else str(total_da)
        
        print(f"{nt:<15} | {nv:<25} | {cd:<10} | {avg_str:<15} | {ck_str:<12} | {combo_str:<10} | {tda_str:<12} | {tt_str:<12} | {tot_str:<15}")
