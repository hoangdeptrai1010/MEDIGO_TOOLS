import openpyxl
from collections import Counter
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/BANGLUONGTHANG8.xlsx', data_only=True)
ws_hcm = wb['MiniKat - HCM']

print("=== MiniKat - HCM: All Rows (Cols A-J) ===")
names = []
for r in range(1, 60):
    val = [ws_hcm.cell(r, c).value for c in range(1, 11)]
    s_name = ws_hcm.cell(r, 1).value
    if s_name:
        s_name_str = str(s_name).strip()
        names.append((r, s_name_str))
        print(f"Row {r:2d}: {s_name_str:28s} | PS={val[1]:2} | KAT={val[2]:2} | Lady={val[3]:2} | LadyRev={val[4]:9} | KatRev={val[5]:8} | ThuongKAT={val[6]} | ThuongPS={val[7]} | ThuongLady={val[8]} | TongThuong={val[9]}")

print("\n=== Duplicate Names Check in MiniKat - HCM ===")
counts = Counter([n for r, n in names if n != 'Người bán'])
dup_found = False
for name, cnt in counts.items():
    if cnt > 1:
        dup_found = True
        row_indices = [r for r, n in names if n == name]
        print(f"  [ERROR DUPLICATE] ({cnt} times): {name} -> Rows {row_indices}")
if not dup_found:
    print("  ✅ Tuyệt vời: KHÔNG CÒN BẤT KỲ NHÂN VIÊN NÀO BỊ TRÙNG LẶP (0 trùng lặp)!")

print("\n=== Branch Summary Table (Cols M-U) ===")
for r in range(1, 13):
    val = [ws_hcm.cell(r, c).value for c in range(13, 22)]
    if any(x is not None for x in val):
        print(f"Row {r:2d}: {val[0]:18s} | KAT_Q={val[1]} | PS_Q={val[2]} | Th_KAT={val[3]} | Th_PS180={val[4]} | Th_PS150={val[5]} | Lady_Rev={val[6]} | Th_Lady={val[7]} | Tong={val[8]}")

print("\n=== BẢNG LƯƠNG: MiniKat DS (Cột AA) & MiniKat CHT (Cột AB) ===")
ws_bl = wb['BẢNG LƯƠNG']
for r in range(3, ws_bl.max_row + 1):
    c_name = ws_bl.cell(r, 3).value
    if c_name:
        b_name = ws_bl.cell(r, 2).value
        role = ws_bl.cell(r, 4).value
        aa_val = ws_bl.cell(r, 27).value
        ab_val = ws_bl.cell(r, 28).value
        if aa_val or ab_val:
            print(f"Row {r:2d}: {str(b_name):20s} | {str(c_name):28s} | {str(role):10s} | Col AA (DS) = {aa_val:,} | Col AB (CHT) = {ab_val:,}")





