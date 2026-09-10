import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'baocaokpi_thang7_hoanthien.xlsx', data_only=False)
ws = wb['Dự án T7']

print("--- baocaokpi_thang7_hoanthien.xlsx Sheet 'Dự án T7' Formulas ---")
for r in range(2, 10):
    print(f"Row {r}:")
    for c in range(1, ws.max_column + 1):
        f = ws.cell(r, c).value
        h = ws.cell(2, c).value
        print(f"  Col {openpyxl.utils.get_column_letter(c)} ({h}): {f}")
