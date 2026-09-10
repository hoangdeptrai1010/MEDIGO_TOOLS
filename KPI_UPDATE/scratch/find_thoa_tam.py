import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
ws = wb.active

for r in range(1, 15):
    row_vals = [ws.cell(r, c).value for c in range(1, 12)]
    print(f"Row {r:2d}: {row_vals}")

# Search for Thoa and Tam anywhere in the sheet
print("\n--- Search for Thoa and Tam ---")
for r in range(1, ws.max_row+1):
    for c in range(1, ws.max_column+1):
        v = str(ws.cell(r, c).value or '')
        if 'thoa' in v.lower() or 'tâm' in v.lower() or 'tam' in v.lower():
            print(f"Cell ({r}, {c}) = '{v}'")
            row_full = [ws.cell(r, col).value for col in range(1, 12)]
            print(f"  Row {r}: {row_full}")
            break
