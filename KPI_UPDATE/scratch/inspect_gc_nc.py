import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook('thang8/tinhcongnhungthuongchia.xlsx', data_only=False)

for sname in ['Giờ công', 'Ngày công']:
    ws = wb[sname]
    print(f"\n--- Sheet '{sname}' (max_row={ws.max_row}, max_col={ws.max_column}) ---")
    for r in range(1, 15):
        row_vals = [ws.cell(r, c).value for c in range(1, 16)]
        print(f"Row {r:2d}: {row_vals}")

print("Done inspection.")
