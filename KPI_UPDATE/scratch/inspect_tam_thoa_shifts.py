import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws_nc = wb['Ngày công']

print("=== CHI TIẾT TỪNG NGÀY / CA CỦA NGUYỄN THỊ TÂM TRONG SHEET NGÀY CÔNG ===")
tam_rows = []
for r in range(2, ws_nc.max_row+1):
    name = str(ws_nc.cell(r, 2).value or '')
    if 'tâm' in name.lower() and 'đoàn' not in name.lower():
        cn = ws_nc.cell(r, 1).value
        dt = ws_nc.cell(r, 3).value
        ca = ws_nc.cell(r, 4).value
        txt = ws_nc.cell(r, 5).value
        hrs = ws_nc.cell(r, 6).value
        col_g = ws_nc.cell(r, 7).value
        col_h = ws_nc.cell(r, 8).value
        col_i = ws_nc.cell(r, 9).value
        col_j = ws_nc.cell(r, 10).value
        tam_rows.append((r, cn, dt, ca, txt, hrs, col_g, col_h, col_i, col_j))
        print(f"Row {r:4d} | Ngày: {str(dt)[:10]} | Ca: {str(ca):<35} | Giờ: {hrs} | G(<=4h)={col_g} | H(Thực)={col_h} | I(Đêm)={col_i} | J(Ngày)={col_j}")

print(f"\nTổng số dòng của Tâm: {len(tam_rows)}")
print(f"Tổng ca đêm (I): {sum(x[8] for x in tam_rows if isinstance(x[8], (int, float)))}")
print(f"Tổng ca ngày (J): {sum(x[9] for x in tam_rows if isinstance(x[9], (int, float)))}")
print(f"Tổng công thực tế (H): {sum(x[7] for x in tam_rows if isinstance(x[7], (int, float)))}")

print("\n=== CHI TIẾT TỪNG NGÀY / CA CỦA HỨA THỊ KIM THOA TRONG SHEET NGÀY CÔNG ===")
thoa_rows = []
for r in range(2, ws_nc.max_row+1):
    name = str(ws_nc.cell(r, 2).value or '')
    if 'thoa' in name.lower():
        cn = ws_nc.cell(r, 1).value
        dt = ws_nc.cell(r, 3).value
        ca = ws_nc.cell(r, 4).value
        txt = ws_nc.cell(r, 5).value
        hrs = ws_nc.cell(r, 6).value
        col_g = ws_nc.cell(r, 7).value
        col_h = ws_nc.cell(r, 8).value
        col_i = ws_nc.cell(r, 9).value
        col_j = ws_nc.cell(r, 10).value
        thoa_rows.append((r, cn, dt, ca, txt, hrs, col_g, col_h, col_i, col_j))
        print(f"Row {r:4d} | Ngày: {str(dt)[:10]} | Ca: {str(ca):<35} | Giờ: {hrs} | G(<=4h)={col_g} | H(Thực)={col_h} | I(Đêm)={col_i} | J(Ngày)={col_j}")

print(f"\nTổng số dòng của Thoa: {len(thoa_rows)}")
print(f"Tổng ca đêm (I): {sum(x[8] for x in thoa_rows if isinstance(x[8], (int, float)))}")
print(f"Tổng ca ngày (J): {sum(x[9] for x in thoa_rows if isinstance(x[9], (int, float)))}")
print(f"Tổng công thực tế (H): {sum(x[7] for x in thoa_rows if isinstance(x[7], (int, float)))}")
