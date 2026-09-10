import openpyxl

wb = openpyxl.load_workbook('thang7/BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=False)

for sname in ['MiniKat - HN', 'MiniKat - HCM']:
    ws = wb[sname]
    print(f"=== {sname} FORMULAS ===")
    for c in range(1, ws.max_column + 1):
        v = ws.cell(2, c).value
        hdr = ws.cell(1, c).value
        if v is not None:
            print(f"  Col {c} ({hdr}): {v}")
    # Also check rows 2-5 for column 13-19
    for r in range(2, 6):
        vals_extra = [f"Col {c} ({ws.cell(1, c).value}): {ws.cell(r, c).value}" for c in range(13, 20) if ws.cell(r, c).value is not None]
        if vals_extra:
            print(f"  Row {r} store table: {vals_extra}")
