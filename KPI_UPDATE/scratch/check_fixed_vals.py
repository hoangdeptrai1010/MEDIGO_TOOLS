import openpyxl

wb = openpyxl.load_workbook('scratch/test_fixed_formulas.xlsx', data_only=True)
ws_bl = wb['BẢNG LƯƠNG']

print(f"{'CN':<15} | {'Họ tên':<24} | {'Chức vụ':<8} | {'Giờ ngày (E)':<12} | {'Giờ đêm (F)':<12} | {'Ngày làm (J)':<12} | {'Ngày đêm (K)':<12} | {'Ngày công (R)':<12}")
print("-" * 115)

for r in range(3, 25):
    stt = ws_bl.cell(r, 1).value
    cn = str(ws_bl.cell(r, 2).value or '')
    name = str(ws_bl.cell(r, 3).value or '')
    cv = str(ws_bl.cell(r, 4).value or '')
    gn = ws_bl.cell(r, 5).value
    gd = ws_bl.cell(r, 6).value
    sn = ws_bl.cell(r, 10).value
    sd = ws_bl.cell(r, 11).value
    nc = ws_bl.cell(r, 18).value
    print(f"{cn:<15} | {name:<24} | {cv:<8} | {str(gn):<12} | {str(gd):<12} | {str(sn):<12} | {str(sd):<12} | {str(nc):<12}")
