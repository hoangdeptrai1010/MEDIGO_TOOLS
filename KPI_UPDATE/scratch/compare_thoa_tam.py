import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

files = {
    'bangluong_thang8_hoanthien.xlsx': 'thang8/bangluong_thang8_hoanthien.xlsx',
    'tinhcongnhungthuongchia.xlsx': 'thang8/tinhcongnhungthuongchia.xlsx',
    'BANGLUONGTHANG8_HOANG_.xlsx': 'thang8/BANGLUONGTHANG8_HOANG_.xlsx',
}

for fname, fpath in files.items():
    print(f"\n==================== {fname} ====================")
    wb = openpyxl.load_workbook(fpath, data_only=True)
    ws = wb['BẢNG LƯƠNG']
    for r in range(3, ws.max_row+1):
        name = str(ws.cell(r, 3).value or '')
        if 'thoa' in name.lower() or ('tâm' in name.lower() and 'đoàn' not in name.lower()):
            cn = ws.cell(r, 2).value
            g_ngay = ws.cell(r, 7).value
            g_dem = ws.cell(r, 8).value
            chuyen_can = ws.cell(r, 9).value
            col_j = ws.cell(r, 10).value
            col_k = ws.cell(r, 11).value
            col_r = ws.cell(r, 18).value
            tot_g = ws.cell(r, 5).value
            tot_h = ws.cell(r, 6).value
            pc_dem = ws.cell(r, 24).value
            print(f"Row {r:2d} | {name:<20} | CN: {cn:<12} | Giờ ngày(G): {g_ngay} | Giờ đêm(H): {g_dem} | Chuyên cần(I): {chuyen_can} | Cột J: {col_j} | Cột K: {col_k} | Cột R: {col_r} | Phụ cấp đêm(X): {pc_dem}")
