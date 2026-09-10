import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

print("=" * 110)
print("THÁNG 7: TỔNG HỢP SO SÁNH THƯỞNG CK VÀ DỰ ÁN")
print("=" * 110)

wb7_target = openpyxl.load_workbook(r'thang7\target\BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=True)
ws7_bl = wb7_target['BẢNG LƯƠNG']

print(f"{'STT':<4} | {'Chi nhánh':<15} | {'Tên':<25} | {'MiniKat AA':<12} | {'MiniKat AB':<12} | {'Dự án AD':<12} | {'Thưởng CK AF':<15} | {'Thưởng khác AI':<15}")
print("-" * 110)
for r in range(3, ws7_bl.max_row + 1):
    stt = ws7_bl.cell(r, 1).value
    cn = ws7_bl.cell(r, 2).value
    name = ws7_bl.cell(r, 3).value
    aa = ws7_bl.cell(r, 27).value
    ab = ws7_bl.cell(r, 28).value
    ad = ws7_bl.cell(r, 30).value
    af = ws7_bl.cell(r, 32).value
    ai = ws7_bl.cell(r, 35).value
    if name and str(name).strip() and "Tổng" not in str(name):
        aa_s = f"{aa:,.0f}" if isinstance(aa, (int, float)) else str(aa or 0)
        ab_s = f"{ab:,.0f}" if isinstance(ab, (int, float)) else str(ab or 0)
        ad_s = f"{ad:,.0f}" if isinstance(ad, (int, float)) else str(ad or 0)
        af_s = f"{af:,.0f}" if isinstance(af, (int, float)) else str(af or 0)
        ai_s = f"{ai:,.0f}" if isinstance(ai, (int, float)) else str(ai or 0)
        print(f"{str(stt):<4} | {str(cn):<15} | {str(name):<25} | {aa_s:<12} | {ab_s:<12} | {ad_s:<12} | {af_s:<15} | {ai_s:<15}")

print("\n" + "=" * 110)
print("THÁNG 8: TỔNG HỢP SO SÁNH THƯỞNG CK VÀ DỰ ÁN")
print("=" * 110)

wb8_bl = openpyxl.load_workbook(r'thang8\bangluong_thang8_hoanthien.xlsx', data_only=True)
ws8_bl = wb8_bl['BẢNG LƯƠNG']

print(f"{'STT':<4} | {'Chi nhánh':<15} | {'Tên':<25} | {'MiniKat AA':<12} | {'MiniKat AB':<12} | {'Dự án AD':<12} | {'Thưởng CK AF':<15} | {'Thưởng khác AI':<15}")
print("-" * 110)
for r in range(3, ws8_bl.max_row + 1):
    stt = ws8_bl.cell(r, 1).value
    cn = ws8_bl.cell(r, 2).value
    name = ws8_bl.cell(r, 3).value
    aa = ws8_bl.cell(r, 27).value
    ab = ws8_bl.cell(r, 28).value
    ad = ws8_bl.cell(r, 30).value
    af = ws8_bl.cell(r, 32).value
    ai = ws8_bl.cell(r, 35).value
    if name and str(name).strip() and "Tổng" not in str(name):
        aa_s = f"{aa:,.0f}" if isinstance(aa, (int, float)) else str(aa or 0)
        ab_s = f"{ab:,.0f}" if isinstance(ab, (int, float)) else str(ab or 0)
        ad_s = f"{ad:,.0f}" if isinstance(ad, (int, float)) else str(ad or 0)
        af_s = f"{af:,.0f}" if isinstance(af, (int, float)) else str(af or 0)
        ai_s = f"{ai:,.0f}" if isinstance(ai, (int, float)) else str(ai or 0)
        print(f"{str(stt):<4} | {str(cn):<15} | {str(name):<25} | {aa_s:<12} | {ab_s:<12} | {ad_s:<12} | {af_s:<15} | {ai_s:<15}")
