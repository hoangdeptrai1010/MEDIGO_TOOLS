import openpyxl

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws_bl = wb['BẢNG LƯƠNG']

print("=== BẢNG LƯƠNG Row 74 (Tổng) ===")
cols = [(5, "Tổng giờ ca ngày (E)"), (6, "Tổng giờ ca đêm (F)"), (7, "Giờ ca ngày (G)"), (8, "Giờ ca đêm (H)"), (9, "Tăng ca (I)")]
for c, name in cols:
    val = ws_bl.cell(74, c).value
    print(f"{name}: {val:,.2f}" if isinstance(val, (int, float)) else f"{name}: {val}")

print("\n=== Chi nhánh summary table (Rows 76-87) ===")
for r in range(76, 88):
    cn = ws_bl.cell(r, 5).value # Col E
    g_ngay = ws_bl.cell(r, 7).value # Col G
    g_dem = ws_bl.cell(r, 8).value # Col H
    g_tot = ws_bl.cell(r, 9).value # Col I
    print(f"R{r}: CN={cn} | G_ngay={g_ngay} | G_dem={g_dem} | G_tot={g_tot}")
