import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws_bl = wb['BẢNG LƯƠNG']

print("=== KIỂM TRA BẢNG LƯƠNG HOÀN THIỆN THÁNG 8 ===")
print("=" * 110)
print(f"{'STT':<4} | {'Chi nhánh':<15} | {'Tên NV':<24} | {'Giờ ngày(G)':<11} | {'Giờ đêm(H)':<10} | {'Ngày(J)':<8} | {'Ca đêm(K)':<10} | {'Công TT(R)':<10}")
print("=" * 110)

for r in range(3, 73):
    stt = ws_bl.cell(r, 1).value
    cn = ws_bl.cell(r, 2).value
    name = ws_bl.cell(r, 3).value
    g = ws_bl.cell(r, 7).value
    h = ws_bl.cell(r, 8).value
    j = ws_bl.cell(r, 10).value
    k = ws_bl.cell(r, 11).value
    res_r = ws_bl.cell(r, 18).value
    
    g_str = f"{g:.2f}" if isinstance(g, (int, float)) else str(g)
    h_str = f"{h:.2f}" if isinstance(h, (int, float)) else str(h)
    j_str = f"{j:.1f}" if isinstance(j, (int, float)) else str(j)
    k_str = str(k)
    r_str = str(res_r)
    
    print(f"{str(stt):<4} | {str(cn):<15} | {str(name):<24} | {g_str:<11} | {h_str:<10} | {j_str:<8} | {k_str:<10} | {r_str:<10}")

print("\n" + "=" * 110)
print("DÒNG TỔNG CỘNG (Row 74):")
tot_g = ws_bl.cell(74, 7).value
tot_h = ws_bl.cell(74, 8).value
tot_j = ws_bl.cell(74, 10).value
tot_k = ws_bl.cell(74, 11).value
tot_r = ws_bl.cell(74, 18).value
print(f"Tổng Giờ ngày (G): {tot_g} | Tổng Giờ đêm (H): {tot_h} | Tổng Ngày quy đổi (J): {tot_j} | Tổng Ca đêm (K): {tot_k} | Tổng Công TT (R): {tot_r}")

wb.close()
