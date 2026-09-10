import openpyxl

for p in ['thang7/BẢNG LƯƠNG THÁNG 7 2026.xlsx', 'thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx']:
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb['BẢNG LƯƠNG']
    print(f"=== {p} ===")
    for r in range(3, ws.max_row + 1):
        name = str(ws.cell(r, 3).value or '')
        if 'Hồ Thị Minh Hòa' in name or 'Nguyễn Trần Ngọc Phương' in name:
            stt = ws.cell(r, 1).value
            cn = ws.cell(r, 2).value
            cv = ws.cell(r, 4).value
            print(f"Row {r}: STT={stt}, CN={cn}, Name={name}, Role={cv}")
