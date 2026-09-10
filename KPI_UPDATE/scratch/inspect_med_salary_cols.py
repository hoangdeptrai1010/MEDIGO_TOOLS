import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_bl = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final_backup.xlsx")
wb = openpyxl.load_workbook(f_bl, data_only=True)
ws_med = wb["Med"]

for r in [11, 12, 13, 14, 15, 16, 17]:
    name = ws_med.cell(r, 4).value
    stk = ws_med.cell(r, 9).value
    bank = ws_med.cell(r, 10).value
    cb = ws_med.cell(r, 80).value # Col CB (80)
    cc = ws_med.cell(r, 81).value # Col CC (81)
    cd = ws_med.cell(r, 82).value # Col CD (82)
    ce = ws_med.cell(r, 83).value # Col CE (83): Tổng thu nhập còn lại
    cf = ws_med.cell(r, 84).value # Col CF (84): CK TM
    cg = ws_med.cell(r, 85).value # Col CG (85): CK TK
    print(f"Row {r:2d} | {str(name):<25} | STK: {str(stk):<15} | Bank: {str(bank):<20} | CB(80): {cb} | CE(83): {ce} | CF(84): {cf} | CG(85): {cg}")
