import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

inv_file = 'thang8/DATA/DanhSachChiTietHoaDon_3182026.xlsx'
ret_file = 'thang8/DATA/DanhSachChiTietTraHang_3182026.xlsx'

SKUS_PARTY_SMART = {'SP017162', 'SP2723878', 'SP2723879'}
SKUS_KAT = {'SP2725289', 'SP2725285', 'SP2725291', 'SP2725287', 'SP2725353', 'SP2725351'}
SKUS_LADYCARE = {'SP2723124', 'SP2723125', 'SP2723127'}

# Read returns
returns_by_seller = {}
returns_by_branch = {}
wb_ret = openpyxl.load_workbook(ret_file, read_only=True, data_only=True)
ws_ret = wb_ret.active
for idx, r in enumerate(ws_ret.iter_rows(values_only=True)):
    if idx == 0 or not r or len(r) < 37: continue
    branch = str(r[0]).strip() if r[0] else ''
    seller = str(r[6]).strip() if r[6] else ''
    sku = str(r[28]).strip() if r[28] else ''
    try:
        qty = float(r[35] or 0)
        price = float(r[36] or 0)
    except:
        continue
    amt = qty * price
    if sku in (SKUS_PARTY_SMART | SKUS_KAT | SKUS_LADYCARE):
        k_s = (seller, sku)
        k_b = (branch, sku)
        returns_by_seller[k_s] = returns_by_seller.get(k_s, 0.0) + qty
        returns_by_branch[k_b] = returns_by_branch.get(k_b, 0.0) + qty

# Read sales
sales_by_seller = {}
sales_by_branch = {}
wb_inv = openpyxl.load_workbook(inv_file, read_only=True, data_only=True)
ws_inv = wb_inv.active
for idx, r in enumerate(ws_inv.iter_rows(values_only=True)):
    if idx == 0 or not r or len(r) < 27: continue
    branch = str(r[0]).strip() if r[0] else ''
    seller = str(r[7]).strip() if r[7] else ''
    sku = str(r[14]).strip() if r[14] else ''
    try:
        qty = float(r[21] or 0)
        tt = float(r[26] or 0)
    except:
        continue
    if sku in (SKUS_PARTY_SMART | SKUS_KAT | SKUS_LADYCARE):
        if seller not in sales_by_seller:
            sales_by_seller[seller] = {'ps_qty': 0.0, 'kat_qty': 0.0, 'kat_rev': 0.0, 'lady_qty': 0.0, 'lady_rev': 0.0}
        if branch not in sales_by_branch:
            sales_by_branch[branch] = {'ps_qty': 0.0, 'kat_qty': 0.0, 'lady_rev': 0.0}
        
        if sku in SKUS_PARTY_SMART:
            sales_by_seller[seller]['ps_qty'] += qty
            sales_by_branch[branch]['ps_qty'] += qty
        elif sku in SKUS_KAT:
            sales_by_seller[seller]['kat_qty'] += qty
            sales_by_seller[seller]['kat_rev'] += tt
            sales_by_branch[branch]['kat_qty'] += qty
        elif sku in SKUS_LADYCARE:
            sales_by_seller[seller]['lady_qty'] += qty
            sales_by_seller[seller]['lady_rev'] += tt
            sales_by_branch[branch]['lady_rev'] += tt

# Net out returns
for s, d in sales_by_seller.items():
    for sku in SKUS_PARTY_SMART:
        d['ps_qty'] -= returns_by_seller.get((s, sku), 0.0)
    for sku in SKUS_KAT:
        d['kat_qty'] -= returns_by_seller.get((s, sku), 0.0)
    for sku in SKUS_LADYCARE:
        d['lady_qty'] -= returns_by_seller.get((s, sku), 0.0)

for b, d in sales_by_branch.items():
    for sku in SKUS_PARTY_SMART:
        d['ps_qty'] -= returns_by_branch.get((b, sku), 0.0)
    for sku in SKUS_KAT:
        d['kat_qty'] -= returns_by_branch.get((b, sku), 0.0)
    for sku in SKUS_LADYCARE:
        d['lady_qty'] -= returns_by_branch.get((b, sku), 0.0)

print("=== SELLER TOTALS (Top performers) ===")
for s, d in sorted(sales_by_seller.items(), key=lambda x: x[1]['kat_qty'] + x[1]['ps_qty'], reverse=True)[:15]:
    print(f"{s:<25}: KAT={d['kat_qty']:>4.0f} (rev={d['kat_rev']:>10,.0f}) | PS={d['ps_qty']:>4.0f} | Lady={d['lady_qty']:>4.0f} (rev={d['lady_rev']:>10,.0f})")

print("\n=== BRANCH TOTALS ===")
for b, d in sorted(sales_by_branch.items()):
    print(f"{b:<30}: KAT={d['kat_qty']:>4.0f} | PS={d['ps_qty']:>4.0f} | Lady={d['lady_rev']:>12,.0f}")
