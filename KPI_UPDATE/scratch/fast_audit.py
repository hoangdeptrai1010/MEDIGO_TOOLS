import openpyxl
import os
import sys
import unicodedata
import re

sys.stdout.reconfigure(encoding='utf-8')

def remove_accents(input_str):
    if not input_str:
        return ""
    nfkd_form = unicodedata.normalize('NFKD', str(input_str))
    return "".join([c for c in nfkd_form if not unicodedata.combining(c)])

def normalize_name(s):
    if not s:
        return ""
    s = unicodedata.normalize('NFC', str(s))
    return " ".join(s.lower().strip().split())

def clean_stk(stk):
    if stk is None:
        return ""
    stk_str = str(stk).strip()
    if stk_str in ["#N/A", "#REF!", "None", "nan", ""]:
        return ""
    stk_clean = re.sub(r'[\s\.\-_]', '', stk_str)
    if stk_clean.endswith(".0"):
        stk_clean = stk_clean[:-2]
    return stk_clean

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_bl = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final.xlsx")
f_ngoc = os.path.join(latvat_dir, "Gởi Ms. Ngoc - 07.09. Bang Luong T8.2026.xlsx")

# 1. Read all banking info from Gởi Ms. Ngoc
wb_ngoc = openpyxl.load_workbook(f_ngoc, data_only=True, read_only=True)
master_records = {} # norm_name -> dict(stk, bank, sources)

for sname in wb_ngoc.sheetnames:
    ws = wb_ngoc[sname]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        continue
    # find headers
    header_idx = -1
    c_name, c_stk, c_bank = -1, -1, -1
    for r_i, row in enumerate(rows[:15]):
        for c_i, val in enumerate(row):
            v_str = str(val or '').strip().lower()
            if any(k == v_str or v_str.startswith(k) for k in ['họ tên', 'họ và tên', 'tên chủ tk', 'tên chủ tài khoản', 'tên ctv', 'tên nhân viên', 'họ và tên nhân viên']):
                if c_name == -1: c_name = c_i
            if any(k in v_str for k in ['số tk', 'số tài khoản']):
                if c_stk == -1: c_stk = c_i
            if any(k in v_str for k in ['ngân hàng', 'cn ngân hàng']):
                if c_bank == -1: c_bank = c_i
        if c_name != -1 and (c_stk != -1 or c_bank != -1):
            header_idx = r_i
            break
            
    if header_idx != -1:
        for row in rows[header_idx + 1:]:
            if c_name < len(row):
                name = row[c_name]
                stk = row[c_stk] if c_stk != -1 and c_stk < len(row) else None
                bank = row[c_bank] if c_bank != -1 and c_bank < len(row) else None
                if name and str(name).strip() and not str(name).strip().startswith("Tổng"):
                    c_s = clean_stk(stk)
                    c_b = str(bank).strip() if bank and str(bank).strip() not in ["#N/A", "#REF!"] else ""
                    if c_s or c_b:
                        norm = normalize_name(name)
                        norm_noacc = normalize_name(remove_accents(name))
                        for k in [norm, norm_noacc]:
                            if k not in master_records:
                                master_records[k] = {'stk': c_s, 'bank': c_b, 'sources': []}
                            if c_s and not master_records[k]['stk']:
                                master_records[k]['stk'] = c_s
                            if c_b and not master_records[k]['bank']:
                                master_records[k]['bank'] = c_b
                            master_records[k]['sources'].append(f"{sname}: {name} | {c_s} | {c_b}")

# Also scan BANGLUONGT8FIXV31 final.xlsx sheets TT CK LƯƠNG, TT CK CNT
wb_bl_ro = openpyxl.load_workbook(f_bl, data_only=True, read_only=True)
for sname in wb_bl_ro.sheetnames:
    if sname in ["Sheet1", "Med"]:
        continue
    ws = wb_bl_ro[sname]
    rows = list(ws.iter_rows(values_only=True))
    header_idx = -1
    c_name, c_stk, c_bank = -1, -1, -1
    for r_i, row in enumerate(rows[:15]):
        for c_i, val in enumerate(row):
            v_str = str(val or '').strip().lower()
            if any(k == v_str or v_str.startswith(k) for k in ['họ tên', 'họ và tên', 'tên chủ tk', 'tên chủ tài khoản', 'tên ctv', 'tên nhân viên']):
                if c_name == -1: c_name = c_i
            if any(k in v_str for k in ['số tk', 'số tài khoản']):
                if c_stk == -1: c_stk = c_i
            if any(k in v_str for k in ['ngân hàng', 'cn ngân hàng']):
                if c_bank == -1: c_bank = c_i
        if c_name != -1 and (c_stk != -1 or c_bank != -1):
            header_idx = r_i
            break
    if header_idx != -1:
        for row in rows[header_idx + 1:]:
            if c_name < len(row):
                name = row[c_name]
                stk = row[c_stk] if c_stk != -1 and c_stk < len(row) else None
                bank = row[c_bank] if c_bank != -1 and c_bank < len(row) else None
                if name and str(name).strip() and not str(name).strip().startswith("Tổng"):
                    c_s = clean_stk(stk)
                    c_b = str(bank).strip() if bank and str(bank).strip() not in ["#N/A", "#REF!"] else ""
                    if c_s or c_b:
                        norm = normalize_name(name)
                        norm_noacc = normalize_name(remove_accents(name))
                        for k in [norm, norm_noacc]:
                            if k not in master_records:
                                master_records[k] = {'stk': c_s, 'bank': c_b, 'sources': []}
                            if c_s and not master_records[k]['stk']:
                                master_records[k]['stk'] = c_s
                            if c_b and not master_records[k]['bank']:
                                master_records[k]['bank'] = c_b
                            master_records[k]['sources'].append(f"BL_{sname}: {name} | {c_s} | {c_b}")

print(f"Total entries in master_records: {len(master_records)}")

# Now test comparison against Sheet1
print("\n=== COMPARISON FOR Sheet1 ===")
ws_s1 = wb_bl_ro["Sheet1"]
s1_rows = list(ws_s1.iter_rows(values_only=True))
for r_i, row in enumerate(s1_rows[1:], 2):
    if len(row) >= 6:
        stt, store, name, acc_holder, stk, bank = row[0], row[1], row[2], row[3], row[4], row[5]
        if not name and not stk and not bank:
            continue
        c_s = clean_stk(stk)
        c_b = str(bank).strip() if bank else ""
        norm_n = normalize_name(name)
        norm_acc = normalize_name(acc_holder)
        
        ref = master_records.get(norm_n) or master_records.get(norm_acc)
        is_stk_ok = False
        is_bank_ok = False
        if ref:
            ref_stk = ref['stk']
            ref_bank = ref['bank']
            is_stk_ok = (c_s == ref_stk and c_s != "")
            # check bank fuzzy
            is_bank_ok = (remove_accents(c_b).lower() in remove_accents(ref_bank).lower() or remove_accents(ref_bank).lower() in remove_accents(c_b).lower()) and c_b not in ["#N/A", "#REF!", ""]
            print(f"Row {r_i:2d} | {str(name):<25} | STK: {str(stk):<16} (Ref: {ref_stk:<16}) -> {'OK' if is_stk_ok else 'WRONG/MISSING'} | Bank: {str(bank):<15} (Ref: {ref_bank:<15}) -> {'OK' if is_bank_ok else 'WRONG/MISSING'}")
        else:
            print(f"Row {r_i:2d} | {str(name):<25} | STK: {str(stk):<16} (NO REF) -> WRONG/MISSING | Bank: {str(bank):<15} (NO REF) -> WRONG/MISSING")

# Now test comparison against Med
print("\n=== COMPARISON FOR Med ===")
ws_med = wb_bl_ro["Med"]
med_rows = list(ws_med.iter_rows(values_only=True))
for r_i, row in enumerate(med_rows[9:], 10):
    if len(row) >= 10:
        code = row[2]
        name = row[3]
        store = row[5]
        stk = row[8]
        bank = row[9]
        if name and str(name).strip() and not str(name).strip().startswith("CHUỖI") and not str(name).strip().startswith("Tổng"):
            c_s = clean_stk(stk)
            c_b = str(bank).strip() if bank else ""
            norm_n = normalize_name(name)
            norm_c = normalize_name(code)
            
            ref = master_records.get(norm_n) or master_records.get(norm_c)
            is_stk_ok = False
            is_bank_ok = False
            if ref:
                ref_stk = ref['stk']
                ref_bank = ref['bank']
                is_stk_ok = (c_s == ref_stk and c_s != "")
                is_bank_ok = (remove_accents(c_b).lower() in remove_accents(ref_bank).lower() or remove_accents(ref_bank).lower() in remove_accents(c_b).lower()) and c_b not in ["#N/A", "#REF!", ""]
                if not is_stk_ok or not is_bank_ok:
                    print(f"Row {r_i:2d} | {str(name):<25} | STK: {str(stk):<16} (Ref: {ref_stk:<16}) -> {'OK' if is_stk_ok else 'WRONG/MISSING'} | Bank: {str(bank):<15} (Ref: {ref_bank:<15}) -> {'OK' if is_bank_ok else 'WRONG/MISSING'}")
                else:
                    pass
            else:
                print(f"Row {r_i:2d} | {str(name):<25} | STK: {str(stk):<16} (NO REF) -> WRONG/MISSING | Bank: {str(bank):<15} (NO REF) -> WRONG/MISSING")
