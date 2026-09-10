# -*- coding: utf-8 -*-
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

for p in ['thang8/BANGLUONGTHANG8.xlsx', 'thang8/BANGLUONGTHANG8_HOANG_.xlsx']:
    try:
        wb = openpyxl.load_workbook(p, data_only=True)
        ws_ck = wb['Thưởng CK']
        ws_bl = wb['BẢNG LƯƠNG']
        
        print(f"\n==================== FILE: {p} ====================")
        print(f"{'STT':<3} | {'Chi nhánh':<15} | {'Họ và tên':<24} | {'CK gốc (E)':<14} | {'Hệ số (D)':<10} | {'Total CK (C=E*D)':<18} | {'Cột AF (Bảng Lương)':<20} | {'Khớp?'}")
        print("-" * 125)
        
        for r in range(2, ws_ck.max_row + 1):
            cn = ws_ck.cell(r, 1).value
            name = ws_ck.cell(r, 2).value
            if not name: continue
            
            tot_c = ws_ck.cell(r, 3).value or 0
            hs_d = ws_ck.cell(r, 4).value or 0
            goc_e = ws_ck.cell(r, 5).value or 0
            
            af_val = 0
            for r_bl in range(3, 59):
                if str(ws_bl.cell(r_bl, 3).value or '').strip() == str(name).strip():
                    af_val = ws_bl.cell(r_bl, 32).value or 0
                    break
                    
            goc_str = f"{goc_e:,.0f} đ" if isinstance(goc_e, (int, float)) and goc_e > 0 else "-"
            tot_str = f"{tot_c:,.0f} đ" if isinstance(tot_c, (int, float)) and tot_c > 0 else "-"
            af_str = f"{af_val:,.0f} đ" if isinstance(af_val, (int, float)) and af_val > 0 else "-"
            
            match_str = "MATCH" if abs((tot_c or 0) - (af_val or 0)) < 1 else "DIFF!"
            
            if goc_e > 0 or af_val > 0:
                print(f"{r-1:<3} | {str(cn):<15} | {str(name):<24} | {goc_str:<14} | {hs_d:<10} | {tot_str:<18} | {af_str:<20} | {match_str}")
    except Exception as e:
        print(f"Error {p}: {e}")
