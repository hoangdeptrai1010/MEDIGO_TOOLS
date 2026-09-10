import openpyxl

wb = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx', data_only=True)
for sn in ['kpi dược sĩ', 'kpi nhà thuốc', 'Dự án T8', 'Hot Bill HN']:
    if sn in wb.sheetnames:
        ws = wb[sn]
        print(f"\n=== Sheet {sn} (rows={ws.max_row}, cols={ws.max_column}) ===")
        for r in range(1, min(6, ws.max_row + 1)):
            row_v = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
            print(f"  r{r}: {row_v}")
