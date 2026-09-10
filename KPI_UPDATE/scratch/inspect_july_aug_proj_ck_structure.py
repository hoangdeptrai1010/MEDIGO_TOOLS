# -*- coding: utf-8 -*-
import openpyxl, glob, os, sys
sys.stdout.reconfigure(encoding='utf-8')

print("=== 1. JULY BẢNG LƯƠNG (thang7/target/BẢNG LƯƠNG THÁNG 7 2026.xlsx) ===")
wb_j = openpyxl.load_workbook('thang7/target/BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=True)
ws_j_bl = wb_j['BẢNG LƯƠNG']

print(f"{'STT':<3} | {'Chi nhánh':<15} | {'Họ và tên':<24} | {'Col AD (Dự án)':<16} | {'Col AF (Thưởng CK)':<18} | {'Col AE (Thưởng KPI)':<18}")
print("-" * 105)
for r in range(3, ws_j_bl.max_row + 1):
    name = ws_j_bl.cell(r, 3).value
    if not name: continue
    cn = ws_j_bl.cell(r, 2).value or ''
    ad = ws_j_bl.cell(r, 30).value or 0
    af = ws_j_bl.cell(r, 32).value or 0
    ae = ws_j_bl.cell(r, 31).value or 0
    if ad or af or ae:
        ad_str = f"{ad:,.0f}" if isinstance(ad, (int, float)) else str(ad)
        af_str = f"{af:,.0f}" if isinstance(af, (int, float)) else str(af)
        ae_str = f"{ae:,.0f}" if isinstance(ae, (int, float)) else str(ae)
        print(f"{ws_j_bl.cell(r,1).value:<3} | {cn:<15} | {name:<24} | {ad_str:<16} | {af_str:<18} | {ae_str:<18}")

print("\n=== 2. JULY SHEETS: 'Dự án' & 'Thưởng CK' in July Payroll ===")
for s in ['Dự án', 'Thưởng CK', 'KPI']:
    if s in wb_j.sheetnames:
        ws = wb_j[s]
        print(f"\n--- Sheet {s} in July (sample rows) ---")
        for r in range(1, 10):
            row_vals = [f"C{c}:{ws.cell(r, c).value}" for c in range(1, min(15, ws.max_column + 1)) if ws.cell(r, c).value is not None]
            print(f"  Row {r}: {row_vals}")

print("\n=== 3. JULY KPI REPORT (baocaokpi_thang7_hoanthien.xlsx) ===")
wb_kpi7 = openpyxl.load_workbook('baocaokpi_thang7_hoanthien.xlsx', data_only=True)
print("KPI July sheetnames:", wb_kpi7.sheetnames)
for s in ['Báo cáo KPI Tháng 7', 'Dự án T7', 'CK', 'Dự án']:
    if s in wb_kpi7.sheetnames:
        ws = wb_kpi7[s]
        print(f"\n--- Sheet {s} in baocaokpi_thang7 (sample rows) ---")
        for r in range(1, 10):
            row_vals = [f"C{c}:{ws.cell(r, c).value}" for c in range(1, min(15, ws.max_column + 1)) if ws.cell(r, c).value is not None]
            print(f"  Row {r}: {row_vals}")
