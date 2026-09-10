import openpyxl

target_files = ['thang8/baocaokpi_thang8_hoanthien.xlsx', 'baocaokpi_thang8_hoanthien.xlsx']

for target in target_files:
    print(f'=== FIXING {target} ===')
    wb = openpyxl.load_workbook(target, data_only=False)
    
    # 1. Fix Hot Bill in 'Dự án T8'
    ws_da = wb['Dự án T8']
    for r in range(3, ws_da.max_row + 1):
        nv = ws_da.cell(r, 2).value
        if nv:
            # Fix formula in Col L to match branch using wildcard:
            # =IFERROR(SUMIFS('Hot Bill HN'!$G:$G, 'Hot Bill HN'!$C:$C, "*"&$A{r}&"*", 'Hot Bill HN'!$D:$D, $B{r}), 0)
            ws_da.cell(r, 12).value = f'=IFERROR(SUMIFS(\'Hot Bill HN\'!$G:$G, \'Hot Bill HN\'!$C:$C, "*"&$A{r}&"*", \'Hot Bill HN\'!$D:$D, $B{r}), 0)'
            # Ensure Total dự án in Col M is =IFERROR($J{r}+$K{r}+$L{r},0)
            ws_da.cell(r, 13).value = f'=IFERROR($J{r}+$K{r}+$L{r},0)'
            
    # 2. Fix Double Penalty in 'kpi dược sĩ' Col N
    ws_kpi = wb['kpi dược sĩ']
    for r in range(3, ws_kpi.max_row + 1):
        nv = ws_kpi.cell(r, 2).value
        if nv:
            ws_kpi.cell(r, 14).value = f'=K{r}+M{r}'
            
    wb.save(target)
    print(f'Saved fixes to {target}')
