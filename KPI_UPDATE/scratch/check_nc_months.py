import openpyxl

wb = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=True)
ws_nc = wb['Ngày công']

dates = []
for r in range(2, ws_nc.max_row + 1):
    d = ws_nc.cell(r, 3).value
    if d is not None:
        dates.append(d)

print(f"Total dates in Ngày công: {len(dates)}")
months = {}
for d in dates:
    m = getattr(d, 'month', str(d)[:7])
    months[m] = months.get(m, 0) + 1
print(f"Months breakdown in Ngày công: {months}")
