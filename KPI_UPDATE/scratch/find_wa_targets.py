import openpyxl, sys, io, glob

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

targets = [67046300, 92439700, 58824400, 19489000, 8501100, 9377800]

for f in glob.glob('d:/MEDIGO/KPI_UPDATE/thang8/**/*.xlsx', recursive=True):
    if '~$' in f: continue
    try:
        wb = openpyxl.load_workbook(f, data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            for r in range(1, ws.max_row+1):
                for c in range(1, ws.max_column+1):
                    v = ws.cell(r, c).value
                    if isinstance(v, (int, float)) and any(abs(v - t) < 1 for t in targets):
                        row_vals = [ws.cell(r, col_i).value for col_i in range(1, min(15, ws.max_column+1))]
                        print(f"File: {f} | Sheet: {sname} | Cell: {ws.cell(r,c).coordinate} | Val={v} | Row={row_vals}")
        wb.close()
    except Exception:
        pass
