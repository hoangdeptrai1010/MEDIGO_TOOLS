# -*- coding: utf-8 -*-
import openpyxl
import sys
sys.stdout.reconfigure(encoding='utf-8')

wb_kpi = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx', data_only=False)
if 'Dự án T8' in wb_kpi.sheetnames:
    ws = wb_kpi['Dự án T8']
    print('=== Dự án T8 Headers (Row 1-2) ===')
    for c in range(1, 15):
        print(f"Col {c} ({openpyxl.utils.get_column_letter(c)}): {ws.cell(1, c).value} | Row 2: {ws.cell(2, c).value}")

wb_bl = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=False)
ws_bl = wb_bl['BẢNG LƯƠNG']
print('\n=== BẢNG LƯƠNG Headers (Row 2-3) ===')
for c in range(5, 25):
    print(f"Col {c} ({openpyxl.utils.get_column_letter(c)}): {ws_bl.cell(2, c).value} | Row 3: {ws_bl.cell(3, c).value}")

if 'Ngày công' in wb_bl.sheetnames:
    ws_nc = wb_bl['Ngày công']
    print('\n=== Ngày công Headers (Row 1) & sample formulas (Row 3) ===')
    for c in range(1, 16):
        print(f"Col {c} ({openpyxl.utils.get_column_letter(c)}): {ws_nc.cell(1, c).value} | Row 3: {ws_nc.cell(3, c).value}")
