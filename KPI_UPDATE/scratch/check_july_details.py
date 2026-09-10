import openpyxl

wb_july = openpyxl.load_workbook('thang7/BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=False)

for sname in ['Giờ công', 'Ngày công']:
    ws = wb_july[sname]
    print(f"=== JULY {sname} ===")
    print(f"max_row = {ws.max_row}, max_column = {ws.max_column}")
    for r in range(1, 10):
        vals = [ws.cell(r, c).value for c in range(1, 17)]
        print(f"R{r}: {vals}")
