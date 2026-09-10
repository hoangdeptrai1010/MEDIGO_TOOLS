import openpyxl

wb = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=True)
ws_nc = wb['Ngày công']

print("=== Checking all unique employees and their row ranges in Ngày công ===")
curr_emp = None
start_r = None
emp_ranges = []

for r in range(2, ws_nc.max_row + 1):
    emp = ws_nc.cell(r, 2).value
    cn = ws_nc.cell(r, 1).value
    if emp != curr_emp:
        if curr_emp is not None:
            emp_ranges.append((curr_cn, curr_emp, start_r, r - 1))
        curr_emp = emp
        curr_cn = cn
        start_r = r

if curr_emp is not None:
    emp_ranges.append((curr_cn, curr_emp, start_r, ws_nc.max_row))

for cn, emp, s, e in emp_ranges:
    if emp:
        print(f"Rows {s:>4} - {e:>4}: {cn:<18} | {emp}")
