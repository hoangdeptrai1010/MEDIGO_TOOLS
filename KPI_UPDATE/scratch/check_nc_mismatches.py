import openpyxl
import sys
from collections import defaultdict

sys.stdout.reconfigure(encoding='utf-8')

for fname in ['thang8/BANGLUONGTHANG8.xlsx', 'thang8/bangluong_thang8_hoanthien.xlsx']:
    print(f"\n==================== {fname} ====================")
    wb = openpyxl.load_workbook(fname, data_only=True)
    ws_nc = wb['Ngày công']
    
    # Calculate sum of Col F (Giờ công chuẩn) by (Chi nhánh, Tên NV) in raw table (Cols A-F)
    raw_hours = defaultdict(float)
    raw_days = defaultdict(int)
    for r in range(3, ws_nc.max_row + 1):
        cn = ws_nc.cell(r, 1).value
        name = ws_nc.cell(r, 2).value
        h = ws_nc.cell(r, 6).value or 0
        cong = ws_nc.cell(r, 8).value or 0
        if name and str(name).strip():
            name = str(name).strip()
            cn = str(cn).strip() if cn else ''
            try:
                raw_hours[(cn, name)] += float(h)
                if float(cong) > 0:
                    raw_days[(cn, name)] += 1
            except:
                pass
                
    # Summary table (Cols K - O)
    summary_hours = defaultdict(float)
    summary_days = defaultdict(float)
    for r in range(2, 75):
        cn = ws_nc.cell(r, 11).value
        name = ws_nc.cell(r, 12).value
        d = ws_nc.cell(r, 13).value or 0
        h = ws_nc.cell(r, 14).value or 0
        if name and str(name).strip():
            name = str(name).strip()
            cn = str(cn).strip() if cn else ''
            try:
                summary_days[(cn, name)] = float(d)
                summary_hours[(cn, name)] = float(h)
            except:
                pass
                
    # Find mismatches
    print("--- Mismatches between Summary (Cols K-O) and Raw (Cols A-H) in 'Ngày công' ---")
    all_pairs = set(list(raw_hours.keys()) + list(summary_hours.keys()))
    mismatch_count = 0
    for (cn, name) in sorted(all_pairs):
        rh = round(raw_hours.get((cn, name), 0.0), 2)
        sh = round(summary_hours.get((cn, name), 0.0), 2)
        rd = round(raw_days.get((cn, name), 0.0), 1)
        sd = round(summary_days.get((cn, name), 0.0), 1)
        
        diff_h = round(rh - sh, 2)
        diff_d = round(rd - sd, 1)
        if abs(diff_h) > 0.05 or abs(diff_d) > 0.05:
            mismatch_count += 1
            print(f"[{cn}] {name:<25}: Raw H={rh:>7.2f} vs Sum H={sh:>7.2f} (diff={diff_h:>6.2f}) | Raw D={rd:>4.1f} vs Sum D={sd:>4.1f} (diff={diff_d:>4.1f})")
    print(f"Total mismatched pairs: {mismatch_count}")
