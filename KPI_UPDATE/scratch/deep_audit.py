import openpyxl
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

print("==================================================================")
print("1. SOI TOÀN BỘ FILE ĐỀ XUẤT HN: e_xuat_KPI_Quy_3.26.xlsx")
print("==================================================================")
hn_path = 'thang9/data/kpi_thang9/e_xuat_KPI_Quy_3.26.xlsx'
wb_hn = openpyxl.load_workbook(hn_path, data_only=True)
for sname in wb_hn.sheetnames:
    ws = wb_hn[sname]
    print(f"\n--- Sheet: {sname} (Rows: {ws.max_row}, Cols: {ws.max_column}) ---")
    for r in range(1, ws.max_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(v is not None for v in row_vals):
            print(f"  R{r:02d}: {row_vals}")

print("\n==================================================================")
print("2. SOI TOÀN BỘ FILE ĐỀ XUẤT HCM: KPI CNT HCM Tháng 09.xlsx (Sheet KPI tháng 09)")
print("==================================================================")
hcm_path = 'thang9/data/kpi_thang9/KPI CNT HCM Tháng 09.xlsx'
wb_hcm = openpyxl.load_workbook(hcm_path, data_only=True)
ws_hcm = wb_hcm['KPI tháng 09']
for r in range(1, ws_hcm.max_row + 1):
    row_vals = [ws_hcm.cell(r, c).value for c in range(1, min(ws_hcm.max_column + 1, 15))]
    if any(v is not None for v in row_vals):
        print(f"  R{r:02d}: {row_vals}")

print("\n==================================================================")
print("3. SOI TOÀN BỘ CÔNG THỨC SHEET 'kpi dược sĩ' VÀ 'kpi nhà thuốc' CỦA BENCHMARK THÁNG 7")
print("==================================================================")
b_path = 'goc/NHÀ THUỐC THÁNG 7 2026.xlsx'
wb_b = openpyxl.load_workbook(b_path, data_only=False)

ws_ds = wb_b['kpi dược sĩ']
print("\n--- Benchmark T7: kpi dược sĩ Header (Row 2) ---")
for c in range(1, 33):
    print(f"  Cột {c:02d} ({openpyxl.utils.get_column_letter(c)}): {ws_ds.cell(2, c).value}")

print("\n--- Benchmark T7: kpi dược sĩ Dòng 3 (Hàng Bông - Lê Ngọc Anh) ---")
for c in range(1, 33):
    print(f"  Cột {c:02d} ({openpyxl.utils.get_column_letter(c)} - {ws_ds.cell(2, c).value}): {ws_ds.cell(3, c).value}")

print("\n--- Benchmark T7: kpi dược sĩ Dòng 7 (Đường Láng - Hứa Thị Kim Thoa) ---")
for c in range(1, 33):
    print(f"  Cột {c:02d} ({openpyxl.utils.get_column_letter(c)} - {ws_ds.cell(2, c).value}): {ws_ds.cell(7, c).value}")

print("\n--- Benchmark T7: kpi dược sĩ Dòng 12 (Trường Sa - Nguyễn Trần Ngọc Phương) ---")
for c in range(1, 33):
    print(f"  Cột {c:02d} ({openpyxl.utils.get_column_letter(c)} - {ws_ds.cell(2, c).value}): {ws_ds.cell(12, c).value}")

ws_nt = wb_b['kpi nhà thuốc']
print("\n--- Benchmark T7: kpi nhà thuốc Header (Row 2) ---")
for c in range(1, 33):
    print(f"  Cột {c:02d} ({openpyxl.utils.get_column_letter(c)}): {ws_nt.cell(2, c).value}")

print("\n--- Benchmark T7: kpi nhà thuốc Dòng 3 (Hàng Bông) ---")
for c in range(1, 33):
    print(f"  Cột {c:02d} ({openpyxl.utils.get_column_letter(c)} - {ws_nt.cell(2, c).value}): {ws_nt.cell(3, c).value}")

print("\n--- Benchmark T7: kpi nhà thuốc Dòng 4 (Đường Láng) ---")
for c in range(1, 33):
    print(f"  Cột {c:02d} ({openpyxl.utils.get_column_letter(c)} - {ws_nt.cell(2, c).value}): {ws_nt.cell(4, c).value}")
