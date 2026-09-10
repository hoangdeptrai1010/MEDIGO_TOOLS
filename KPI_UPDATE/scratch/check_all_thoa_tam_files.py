import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

files = [
    'thang8/tinhcongnhungthuongchia.xlsx',
    'thang8/BANGLUONGTHANG8_HOANG_.xlsx',
    'thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx',
    'thang8/bangluong_thang8_hoanthien.xlsx'
]

for fp in files:
    try:
        wb = openpyxl.load_workbook(fp, data_only=True)
        ws = wb['BẢNG LƯƠNG']
        print(f"\n==================== File: {fp} ====================")
        for r in range(3, ws.max_row+1):
            name = str(ws.cell(r, 3).value or '')
            if 'thoa' in name.lower() or ('tâm' in name.lower() and 'đoàn' not in name.lower()):
                row_vals = [f"{openpyxl.utils.get_column_letter(c)}: {ws.cell(r, c).value}" for c in range(1, 26) if ws.cell(r, c).value is not None]
                print(f"Row {r:2d} ({name}):")
                print("  " + " | ".join(row_vals))
    except Exception as e:
        print(f"Error {fp}: {e}")
