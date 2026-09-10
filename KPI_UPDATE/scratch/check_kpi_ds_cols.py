import openpyxl

wb = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx', data_only=True)
ws = wb['kpi dược sĩ']
print("kpi dược sĩ row 2 headers:")
for c in range(1, 35):
    v = ws.cell(2, c).value
    if v is not None:
        print(f"Col {c} ({openpyxl.utils.get_column_letter(c)}): {repr(v)}")

print("\nkpi dược sĩ row 3 values:")
for c in range(1, 35):
    v = ws.cell(3, c).value
    if v is not None:
        print(f"Col {c} ({openpyxl.utils.get_column_letter(c)}): {repr(v)}")
