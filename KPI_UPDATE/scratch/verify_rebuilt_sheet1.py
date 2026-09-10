import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_bl = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final.xlsx")

wb = openpyxl.load_workbook(f_bl, data_only=False)
ws1 = wb["Sheet1"]

print("=== ALL ROWS IN REBUILT Sheet1 ===")
for r in range(1, ws1.max_row + 1):
    vals = [ws1.cell(r, c).value for c in range(1, 10)]
    c_stk = ws1.cell(r, 5)
    c_bank = ws1.cell(r, 6)
    c_amt = ws1.cell(r, 7)
    
    stk_color = c_stk.fill.start_color.rgb if c_stk.fill and c_stk.fill.start_color else ""
    bank_color = c_bank.fill.start_color.rgb if c_bank.fill and c_bank.fill.start_color else ""
    amt_color = c_amt.fill.start_color.rgb if c_amt.fill and c_amt.fill.start_color else ""
    
    flags = []
    if "FFFF9999" in str(stk_color): flags.append("RED_STK")
    if "FFFF9999" in str(bank_color): flags.append("RED_BANK")
    if "FFFF9999" in str(amt_color): flags.append("RED_AMT")
    
    flag_str = ", ".join(flags) if flags else "OK"
    
    if any(v is not None for v in vals):
        print(f"Row {r:2d} | {str(vals[1]):<18} | {str(vals[2]):<26} | STK: {str(vals[4]):<16} | Bank: {str(vals[5]):<20} | Amt: {str(vals[6]):<14} | Flag: {flag_str}")
