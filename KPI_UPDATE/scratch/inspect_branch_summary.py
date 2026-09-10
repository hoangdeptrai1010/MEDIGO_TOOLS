import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

# Let's inspect multiple source files around rows 74 to 90 of BẢNG LƯƠNG
for fname in ['thang8/tinhcongnhungthuongchia.xlsx', 'thang8/chiacongnhunggomthuong.xlsx', 'thang8/BANGLUONGTHANG8_HOANG_.xlsx', 'thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx']:
    try:
        wb = openpyxl.load_workbook(fname, data_only=True)
        ws = wb['BẢNG LƯƠNG']
        print(f"\n{'='*20} {fname} {'='*20}")
        print("Row 74 (Tổng):", [ws.cell(74, c).value for c in range(5, 12)])
        for r in range(76, 89):
            print(f"Row {r}:", [ws.cell(r, c).value for c in range(4, 11)])
        wb.close()
    except Exception as e:
        print(f"Could not read {fname}: {e}")
