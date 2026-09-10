import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb7 = openpyxl.load_workbook(r'thang7\target\BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=False)
wb7_v = openpyxl.load_workbook(r'thang7\target\BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=True)

ws = wb7['Thưởng CK']
wsv = wb7_v['Thưởng CK']

print("--- [TARGET BẢNG LƯƠNG T7] Sheet 'Thưởng CK' (Tất cả các cột) ---")
for r in range(1, 15):
    f_row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    v_row = [wsv.cell(r, c).value for c in range(1, ws.max_column + 1)]
    print(f"R{r} FORMULA: {f_row}")
    print(f"R{r} VALUE:   {v_row}")
