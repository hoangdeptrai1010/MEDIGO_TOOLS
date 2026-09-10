import openpyxl

wb_july = openpyxl.load_workbook('thang7/BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=False)

for sname in ['Giờ công', 'Ngày công']:
    ws = wb_july[sname]
    print(f"=== JULY SHEET: {sname} ===")
    print("Row 1 headers:")
    print([ws.cell(1, c).value for c in range(1, 20)])
    print("Row 2 data:")
    print([ws.cell(2, c).value for c in range(1, 20)])
    print("Row 3 data:")
    print([ws.cell(3, c).value for c in range(1, 20)])
