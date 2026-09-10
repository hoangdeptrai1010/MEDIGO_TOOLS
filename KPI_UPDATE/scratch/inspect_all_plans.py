import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

print("================================================================================")
print("1. KẾ HOẠCH KPI THÁNG 7 VÀ THÁNG 8 (plans/)")
print("================================================================================")

for p_name in ['plans/KeHoachKPI_2026-07.xlsx', 'plans/KeHoachKPI_2026-08.xlsx']:
    if os.path.exists(p_name):
        print(f"\n--- FILE: {p_name} ---")
        wb = openpyxl.load_workbook(p_name, data_only=True)
        print("Sheets:", wb.sheetnames)
        for s in wb.sheetnames:
            ws = wb[s]
            print(f"Sheet '{s}' ({ws.max_row} rows, {ws.max_column} cols):")
            for r in range(1, min(10, ws.max_row + 1)):
                row_vals = [ws.cell(r, c).value for c in range(1, min(10, ws.max_column + 1))]
                if any(v is not None for v in row_vals):
                    print(f"  R{r}: {row_vals}")

print("\n================================================================================")
print("2. CÁC FILE MINIKAT / DỰ ÁN TRONG THANG7 VÀ THANG8")
print("================================================================================")

for root, dirs, files in os.walk('thang7'):
    for f in files:
        if 'kat' in f.lower() or 'lady' in f.lower() or 'party' in f.lower() or 'avc' in f.lower() or 'ck' in f.lower():
            fp = os.path.join(root, f)
            print(f"\n--- [T7] {fp} ---")
            try:
                wb = openpyxl.load_workbook(fp, data_only=True)
                print("Sheets:", wb.sheetnames)
                for s in wb.sheetnames:
                    ws = wb[s]
                    print(f"Sheet '{s}': R1={list(ws.iter_rows(values_only=True))[0] if ws.max_row>=1 else 'empty'}")
                    if ws.max_row >= 2:
                        print(f"Sheet '{s}': R2={list(ws.iter_rows(values_only=True))[1]}")
            except Exception as e:
                print(f"Error: {e}")

for root, dirs, files in os.walk('thang8'):
    for f in files:
        if 'kat' in f.lower() or 'lady' in f.lower() or 'party' in f.lower() or 'avc' in f.lower() or 'ck' in f.lower():
            fp = os.path.join(root, f)
            print(f"\n--- [T8] {fp} ---")
            try:
                wb = openpyxl.load_workbook(fp, data_only=True)
                print("Sheets:", wb.sheetnames)
                for s in wb.sheetnames:
                    ws = wb[s]
                    print(f"Sheet '{s}': R1={list(ws.iter_rows(values_only=True))[0] if ws.max_row>=1 else 'empty'}")
                    if ws.max_row >= 2:
                        print(f"Sheet '{s}': R2={list(ws.iter_rows(values_only=True))[1]}")
            except Exception as e:
                print(f"Error: {e}")
