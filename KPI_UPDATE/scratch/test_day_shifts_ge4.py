import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'thang8\bangluong_thang8_hoanthien.xlsx', data_only=True)
ws_nc = wb['Ngày công']
ws_bl = wb['BẢNG LƯƠNG']

# Count day shifts with F >= 4h
day_shifts_ge_4h = {}
for r in range(3, ws_nc.max_row + 1):
    name = ws_nc.cell(r, 2).value
    ca = str(ws_nc.cell(r, 4).value or '')
    hours = ws_nc.cell(r, 6).value
    
    if name and ca != '-' and isinstance(hours, (int, float)) and hours >= 4.0 and 'đêm' not in ca.lower():
        n_str = str(name).strip()
        day_shifts_ge_4h[n_str] = day_shifts_ge_4h.get(n_str, 0) + 1

print(f"{'STT':<4} | {'Họ và tên':<26} | {'Chức danh':<10} | {'Giờ ngày (G)':<14} | {'G/8':<10} | {'Ca ngày >=4h':<14} | {'Hiện tại':<10} | {'Ca đêm':<8}")
print("-" * 115)

for r in range(3, ws_bl.max_row + 1):
    stt = ws_bl.cell(r, 1).value
    name = ws_bl.cell(r, 3).value
    chucdanh = ws_bl.cell(r, 4).value
    g = ws_bl.cell(r, 7).value or 0
    curr_j = ws_bl.cell(r, 10).value or 0
    k = ws_bl.cell(r, 11).value or 0
    
    if name and str(name).strip() and "Tổng" not in str(name):
        n_str = str(name).strip()
        g_val = float(g) if isinstance(g, (int, float)) else 0.0
        g_div_8 = g_val / 8.0
        ca_ge_4 = day_shifts_ge_4h.get(n_str, 0)
        
        print(f"{str(stt):<4} | {n_str:<26} | {str(chucdanh):<10} | {g_val:>12.2f} h | {g_div_8:>8.1f} ng | {ca_ge_4:>12} ca | {str(curr_j):>8} ng | {str(k):>6}")
