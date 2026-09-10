import openpyxl

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws_bl = wb['BẢNG LƯƠNG']

total_staff = 0
total_kpi = 0
total_da = 0
total_wa = 0
total_ck = 0
total_cd = 0
total_tru = 0

for r in range(3, ws_bl.max_row + 1):
    stt = ws_bl.cell(r, 1).value
    name = ws_bl.cell(r, 3).value
    if isinstance(stt, (int, float)) and name:
        total_staff += 1
        total_kpi += float(ws_bl.cell(r, 31).value or 0)
        total_da += float(ws_bl.cell(r, 30).value or 0)
        total_wa += float(ws_bl.cell(r, 29).value or 0)
        total_ck += float(ws_bl.cell(r, 32).value or 0)
        total_cd += float(ws_bl.cell(r, 33).value or 0)
        total_tru += float(ws_bl.cell(r, 36).value or 0)

print("=" * 60)
print(f"TỔNG HỢP TOÀN BỘ BẢNG LƯƠNG THÁNG 8 2026")
print("=" * 60)
print(f"- Tổng số lượt nhân sự: {total_staff}")
print(f"- Tổng thưởng KPI:      {total_kpi:>15,.0f} đ")
print(f"- Tổng thưởng Dự án:    {total_da:>15,.0f} đ")
print(f"- Tổng thưởng WhatsApp: {total_wa:>15,.0f} đ")
print(f"- Tổng thưởng CK:       {total_ck:>15,.0f} đ")
print(f"- Tổng thưởng Cận date: {total_cd:>15,.0f} đ")
print(f"- Tổng giảm trừ KPI:    {total_tru:>15,.0f} đ")
print("=" * 60)
