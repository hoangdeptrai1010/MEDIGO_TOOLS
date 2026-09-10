import openpyxl

wb = openpyxl.load_workbook('thang7/BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=False)
ws = wb['BẢNG LƯƠNG']
print(f"Max col: {ws.max_column}, Max row: {ws.max_row}")
for r in [1, 2]:
    non_empty = [(c, ws.cell(r, c).value) for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None]
    print(f"Row {r} has {len(non_empty)} cells:")
    for c, v in non_empty:
        print(f"  Col {c} ({openpyxl.utils.get_column_letter(c)}): {repr(v)[:50]}")
