import openpyxl
import os
import sys
import unicodedata
import re

sys.stdout.reconfigure(encoding='utf-8')

def remove_accents(input_str):
    if not input_str:
        return ""
    nfkd_form = unicodedata.normalize('NFKD', str(input_str))
    return "".join([c for c in nfkd_form if not unicodedata.combining(c)])

def normalize_name(s):
    if not s:
        return ""
    s = unicodedata.normalize('NFC', str(s))
    return " ".join(s.lower().strip().split())

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_bl = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final_backup.xlsx")
wb = openpyxl.load_workbook(f_bl, data_only=True)

ws_med = wb["Med"]
ws_s1 = wb["Sheet1"]

print("=== OLD Sheet1 Rows ===")
old_s1_records = []
for r in range(2, ws_s1.max_row + 1):
    stt = ws_s1.cell(r, 1).value
    store = ws_s1.cell(r, 2).value
    name = ws_s1.cell(r, 3).value
    acc_holder = ws_s1.cell(r, 4).value
    stk = ws_s1.cell(r, 5).value
    bank = ws_s1.cell(r, 6).value
    amount = ws_s1.cell(r, 7).value
    content = ws_s1.cell(r, 8).value
    
    if name is not None:
        rec = {
            'row': r,
            'stt': stt,
            'store': str(store).strip() if store else "",
            'name': str(name).strip(),
            'acc_holder': str(acc_holder).strip() if acc_holder else "",
            'stk': str(stk).strip() if stk is not None else "",
            'bank': str(bank).strip() if bank is not None else "",
            'amount': amount,
            'content': str(content).strip() if content else ""
        }
        old_s1_records.append(rec)
        print(f"Old Row {r:2d} | Store: {rec['store']:<18} | Name: {rec['name']:<25} | STK: {rec['stk']:<15} | Bank: {rec['bank']:<18} | Amt: {rec['amount']}")

print(f"\nTotal old Sheet1 employee rows: {len(old_s1_records)}")

print("\n=== NEW Med Employee Rows ===")
new_med_records = []
current_store = ""
for r in range(10, ws_med.max_row + 1):
    stt = ws_med.cell(r, 1).value
    code = ws_med.cell(r, 3).value
    name = ws_med.cell(r, 4).value
    store = ws_med.cell(r, 6).value
    role = ws_med.cell(r, 7).value
    stk = ws_med.cell(r, 9).value
    bank = ws_med.cell(r, 10).value
    amount = ws_med.cell(r, 85).value # Col CG: CK TK (or Col 80 / 83)
    
    if store and str(store).strip():
        current_store = str(store).strip()
        
    if name and str(name).strip() and not str(name).strip().startswith("CHUỖI") and not str(name).strip().startswith("Tổng"):
        rec = {
            'med_row': r,
            'stt': stt,
            'code': str(code).strip() if code else "",
            'name': str(name).strip(),
            'store': current_store,
            'role': str(role).strip() if role else "",
            'stk': str(stk).strip() if stk is not None else "",
            'bank': str(bank).strip() if bank is not None else "",
            'amount': amount
        }
        new_med_records.append(rec)
        print(f"Med Row {r:2d} | Store: {rec['store']:<18} | Name: {rec['name']:<25} | STK: {rec['stk']:<15} | Bank: {rec['bank']:<18} | Amt: {rec['amount']}")

print(f"\nTotal new Med employee rows: {len(new_med_records)}")
