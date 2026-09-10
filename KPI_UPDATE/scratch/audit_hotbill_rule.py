import openpyxl, datetime
from collections import defaultdict

# 1. Load catalog CK / CKHN
wb_plan = openpyxl.load_workbook('d:/MEDIGO/KPI_UPDATE/plans/KeHoachKPI_2026-08.xlsx', data_only=True)
ws_dm = wb_plan['Danh mục dự án']
ck_skus = {}
for r in range(2, ws_dm.max_row + 1):
    sku = str(ws_dm.cell(r, 1).value or '').strip()
    name = str(ws_dm.cell(r, 2).value or '').strip()
    grp = str(ws_dm.cell(r, 3).value or '').strip()
    if grp in ['CK', 'CKHN']:
        ck_skus[sku] = (grp, name)

print(f'Catalog has {len(ck_skus)} CK/CKHN SKUs.')

# 2. Inspect all items in Hanoi invoices
wb_inv = openpyxl.load_workbook('d:/MEDIGO/KPI_UPDATE/thang8/DATA/DanhSachChiTietHoaDon_3182026.xlsx', read_only=True)
ws_inv = wb_inv.active
it = ws_inv.iter_rows(values_only=True)
hdr = next(it)

# Group rows by invoice
inv_items = defaultdict(list)
for r in it:
    if not r or len(r) < 27: continue
    raw_cn = str(r[0] or '').strip()
    if 'hàng bông' in raw_cn.lower():
        cn = 'Hàng Bông'
    elif 'đường láng' in raw_cn.lower() or 'duong lang' in raw_cn.lower():
        cn = 'Đường Láng'
    else:
        continue
    
    ma_hd = str(r[1] or '').strip()
    dt = r[2]
    seller = str(r[7] or '').strip()
    sku = str(r[14] or '').strip()
    ten = str(r[16] or '').strip()
    qty = float(r[21] or 0)
    tt = float(r[26] or 0)
    inv_items[ma_hd].append({
        'cn': cn, 'seller': seller, 'time': dt, 'sku': sku, 'ten': ten, 'qty': qty, 'tt': tt
    })

print(f'Total Hanoi invoices: {len(inv_items)}')

# Check Hot Bills
hot_bills_found = []
staff_hb_count = defaultdict(int)
staff_hb_bonus = defaultdict(float)

for ma_hd, items in inv_items.items():
    cn = items[0]['cn']
    seller = items[0]['seller']
    dt = items[0]['time']
    
    # Check date range: 15/08/2026 to 31/08/2026
    if isinstance(dt, datetime.datetime):
        if dt < datetime.datetime(2026, 8, 15, 0, 0) or dt > datetime.datetime(2026, 8, 31, 23, 59, 59):
            continue
    
    # Sum CK / CKHN
    ck_rev = 0.0
    ck_qty = 0
    ck_details = []
    for itm in items:
        sku = itm['sku']
        ten = itm['ten']
        p = ten.split()[0].upper() if ten else ''
        is_ck = (sku in ck_skus) or p.startswith('CK') or p.startswith('CKHN')
        if is_ck:
            ck_rev += itm['tt']
            ck_qty += itm['qty']
            ck_details.append(f"{ten} (SL: {itm['qty']}, TT: {itm['tt']:,.0f})")
    
    if ck_rev >= 1000000:
        hot_bills_found.append({
            'ma_hd': ma_hd, 'cn': cn, 'seller': seller, 'time': dt,
            'ck_rev': ck_rev, 'ck_qty': ck_qty, 'bonus': 50000,
            'details': ' | '.join(ck_details)
        })
        staff_hb_count[(cn, seller)] += 1
        staff_hb_bonus[(cn, seller)] += 50000

print(f'Total Hot Bills qualified (15-31/8, CK+CKHN >= 1M): {len(hot_bills_found)}')
print('\nBreakdown by Pharmacist (Hà Nội):')
for (cn, seller), cnt in sorted(staff_hb_count.items(), key=lambda x: x[1], reverse=True):
    print(f'{cn:12} | {seller:22} | {cnt:2d} bills | Thưởng: {staff_hb_bonus[(cn, seller)]:>10,.0f} VNĐ')

# Compare with existing sheet 'Hot Bill HN' in baocaokpi_thang8_hoanthien.xlsx
wb_kpi = openpyxl.load_workbook('d:/MEDIGO/KPI_UPDATE/baocaokpi_thang8_hoanthien.xlsx', data_only=True)
ws_hb = wb_kpi['Hot Bill HN']
existing_bills = set()
for r in range(2, ws_hb.max_row + 1):
    c1 = ws_hb.cell(r, 1).value
    if c1:
        existing_bills.add(str(c1).strip())

print(f'\nExisting Hot Bills in baocaokpi sheet: {len(existing_bills)}')
found_set = {b['ma_hd'] for b in hot_bills_found}
missing = found_set - existing_bills
extra = existing_bills - found_set
print(f'Missing bills (found in raw data but not in report): {len(missing)}')
if missing:
    print('  Missing:', missing)
print(f'Extra bills (in report but not in raw data match): {len(extra)}')
if extra:
    print('  Extra:', extra)
