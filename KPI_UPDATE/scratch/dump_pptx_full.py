import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

for p in ['thang8/Pharmacy_retail_Store_KPIs_August 2026/Hà Nội/Dự án/Chương trình/Hanoi_Projects_August_2026_v1.pptx',
          'thang8/Pharmacy_retail_Store_KPIs_August 2026/Hồ Chí Minh/Dự án/Chương trình/HCM_Projects_August_2026_v1.pptx']:
    print(f"\n==================== PPTX: {p} ====================")
    prs = Presentation(p)
    for idx, slide in enumerate(prs.slides, 1):
        print(f"--- Slide {idx} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    print(para.text)
            if shape.has_table:
                table = shape.table
                for r in table.rows:
                    print([c.text.strip() for c in r.cells])
