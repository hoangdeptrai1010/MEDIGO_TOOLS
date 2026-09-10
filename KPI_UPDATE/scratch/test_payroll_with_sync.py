import openpyxl, sys, os
sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, 'd:/MEDIGO/KPI_UPDATE/TOOL_BANGLUONG')

from build_payroll_report import generate_payroll_report_perfect

kpi_file = 'd:/MEDIGO/KPI_UPDATE/thang8/baocaokpi_thang8_hoanthien.xlsx'
tmpl_file = 'd:/MEDIGO/KPI_UPDATE/thang8/tinhcongnhungthuongchia.xlsx'
out_file = 'd:/MEDIGO/KPI_UPDATE/thang8/BANGLUONGTHANG8_test_sync.xlsx'
inv_file = 'd:/MEDIGO/KPI_UPDATE/thang8/DATA/DanhSachChiTietHoaDon_3182026.xlsx'
ret_file = 'd:/MEDIGO/KPI_UPDATE/thang8/DATA/DanhSachChiTietTraHang_3182026.xlsx'
timecard_file = 'd:/MEDIGO/KPI_UPDATE/thang8/DATA/BangChiTietChamCong_thang8.xlsx'

res = generate_payroll_report_perfect(
    kpi_file=kpi_file,
    template_file=tmpl_file,
    output_file=out_file,
    month=8,
    inv_file=inv_file,
    ret_file=ret_file,
    timecard_file=timecard_file
)

print(f"Generated output: {res}")
wb = openpyxl.load_workbook(res, data_only=False)
ws_bl = wb['BẢNG LƯƠNG']
print(f"Sheet BẢNG LƯƠNG rows: {ws_bl.max_row}")

# Check all Hoàng Lâm Gia Bảo rows in BẢNG LƯƠNG
print("\n=== ROWS FOR HOÀNG LÂM GIA BẢO IN BẢNG LƯƠNG ===")
for r in range(3, ws_bl.max_row + 1):
    n = ws_bl.cell(r, 3).value
    if 'Gia Bảo' in str(n) or 'Gia Bao' in str(n):
        b = ws_bl.cell(r, 2).value
        role = ws_bl.cell(r, 4).value
        fml_g = ws_bl.cell(r, 7).value
        fml_h = ws_bl.cell(r, 8).value
        fml_r = ws_bl.cell(r, 18).value
        print(f"Row {r}: STT={ws_bl.cell(r, 1).value}, Branch={b}, Name={n}, Role={role}, Col G={fml_g}, Col H={fml_h}, Col R={fml_r}")
