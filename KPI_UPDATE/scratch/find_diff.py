import openpyxl
import subprocess
import os

wb_step5 = openpyxl.load_workbook('test_step5.xlsx')
wb_kpi = openpyxl.load_workbook('baocaokpi_thang7_hoanthien.xlsx')

print('Step 5 sheets:', wb_step5.sheetnames)
print('KPI sheets:', wb_kpi.sheetnames)

# Let's see what happens if we remove sheets one by one from wb_kpi and test recalc!
for sname in list(wb_kpi.sheetnames):
    wb_temp = openpyxl.load_workbook('baocaokpi_thang7_hoanthien.xlsx')
    if len(wb_temp.sheetnames) > 1:
        wb_temp.remove(wb_temp[sname])
        tname = f'test_without_{sname.replace(" ", "_")}.xlsx'
        wb_temp.save(tname)
        wb_temp.close()
        
        p = os.path.abspath(tname)
        res = subprocess.run(['powershell', '-ExecutionPolicy', 'Bypass', '-File', 'recalc_workbook.ps1', '-FilePath', p], capture_output=True, text=True)
        is_ok = 'Recalculation complete' in res.stdout
        print(f'Without sheet [{sname}]: {"SUCCESS" if is_ok else "FAILED"}')
