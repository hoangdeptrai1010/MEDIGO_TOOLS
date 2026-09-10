import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
ws = wb.active

for r in range(205, 212):
    shift = ws.cell(r, 7).value
    punches = {}
    for day in range(1, 32):
        col_in = 8 + (day-1)*2
        col_out = col_in + 1
        v_in = ws.cell(r, col_in).value
        v_out = ws.cell(r, col_out).value
        if v_in or v_out:
            punches[day] = f"{v_in} -> {v_out}"
    print(f"\nRow {r} (Hứa Thị Kim Thoa - {shift}): {len(punches)} ngày quẹt thẻ")
    for d, p in punches.items():
        print(f"  Day {d:02d}/08: {p}")
