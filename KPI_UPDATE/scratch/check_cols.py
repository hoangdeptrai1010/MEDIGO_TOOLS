import sys
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=False)
ws = wb['BẢNG LƯƠNG']

for c in range(1, 26):
    col_let = openpyxl.utils.get_column_letter(c)
    h = str(ws.cell(2, c).value or '').replace('\n', ' ')
    f = str(ws.cell(4, c).value or '')
    print(f'Col {c:2d} ({col_let:>2s}): Header = {h:<35s} | Formula = {f}')
