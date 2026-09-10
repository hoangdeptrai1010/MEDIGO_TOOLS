# -*- coding: utf-8 -*-
import sys, os
sys.stdout.reconfigure(encoding='utf-8')

# Try extracting from docx
try:
    import docx
    for p in ['thang8/Pharmacy_retail_Store_KPIs_August 2026/Hồ Chí Minh/KPIs/Chương trình/Chương trình KPI.docx',
              'thang8/Pharmacy_retail_Store_KPIs_August 2026/Hà Nội/KPIs/Chương trình/Chương trình KPI.docx']:
        if os.path.exists(p):
            print(f"\n=== DOCX: {p} ===")
            doc = docx.Document(p)
            for para in doc.paragraphs:
                if para.text.strip():
                    print(para.text)
            for table in doc.tables:
                for row in table.rows:
                    print([c.text.strip() for c in row.cells])
except Exception as e:
    print(f"Docx error: {e}")

# Try extracting from pptx
try:
    from pptx import Presentation
    for p in ['thang8/Pharmacy_retail_Store_KPIs_August 2026/Hồ Chí Minh/Dự án/Chương trình/HCM_Projects_August_2026_v1.pptx',
              'thang8/Pharmacy_retail_Store_KPIs_August 2026/Hà Nội/Dự án/Chương trình/Hanoi_Projects_August_2026_v1.pptx']:
        if os.path.exists(p):
            print(f"\n=== PPTX: {p} ===")
            prs = Presentation(p)
            for idx, slide in enumerate(prs.slides):
                print(f"--- Slide {idx+1} ---")
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            if paragraph.text.strip():
                                print(paragraph.text.strip())
                    elif shape.has_table:
                        for row in shape.table.rows:
                            print([cell.text.strip() for cell in row.cells])
except Exception as e:
    print(f"Pptx error: {e}")
