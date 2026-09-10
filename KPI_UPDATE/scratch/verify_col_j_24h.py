import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws = wb['BẢNG LƯƠNG']

print("--- KIỂM TRA CỘT J (SỐ NGÀY LÀM NGÀY: 24H = 1 NGÀY, 8H = 0.3 NGÀY) TRÊN SHEET BẢNG LƯƠNG ---")

for r in range(3, 20):
    stt = ws.cell(r, 1).value
    cn = ws.cell(r, 2).value
    name = ws.cell(r, 3).value
    gio_ngay = ws.cell(r, 7).value
    ngay_lam = ws.cell(r, 10).value
    print(f"Row {r:2d} | STT={stt} | {cn:<16} | {name:<26} | Giờ ca ngày (Col G): {gio_ngay:6.2f}h -> Số ngày (Col J): {ngay_lam:4.1f} ngày" if isinstance(gio_ngay, (int, float)) and isinstance(ngay_lam, (int, float)) else f"Row {r:2d} | {name} | {gio_ngay} -> {ngay_lam}")

print("\n--- Kiểm tra dòng tổng cộng Row 74 ---")
print(f"Tổng giờ ca ngày (Col G): {ws.cell(74, 7).value:,.2f} h")
print(f"Tổng số ngày làm ngày (Col J): {ws.cell(74, 10).value:,.1f} ngày")
