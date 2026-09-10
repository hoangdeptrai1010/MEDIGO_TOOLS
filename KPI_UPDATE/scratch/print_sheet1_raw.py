import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_bl = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final_backup.xlsx")
wb = openpyxl.load_workbook(f_bl, data_only=False)
ws1 = wb["Sheet1"]

print("=== ALL CELLS WITH FORMULAS OR VALUES IN Sheet1 ===")
for r in range(1, ws1.max_row + 1):
    row_vals = [ws1.cell(r, c).value for c in range(1, ws1.max_column + 1)]
    print(f"Row {r:2d}: {row_vals}")
