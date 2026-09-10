import openpyxl

wb = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', read_only=True, data_only=True)
print("Sheets in August:", wb.sheetnames)

for sn in ['BẢNG LƯƠNG', 'MiniKat - HN', 'MiniKat - HCM', 'whatsapp', 'KPI', 'Thưởng CK', 'Dự án', 'Giờ công', 'Ngày công', 'Cận date']:
    if sn in wb.sheetnames:
        ws = wb[sn]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if any(x is not None for x in row):
                rows.append((i+1, [x for x in row[:8]]))
            if len(rows) >= 4:
                break
        print(f"Sheet {sn}: sample:")
        for idx, r in rows:
            print(f"  r{idx}: {r}")
