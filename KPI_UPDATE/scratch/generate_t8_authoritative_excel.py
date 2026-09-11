import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from python_calamine import CalamineWorkbook
import os
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

out_dir = 'd:/MEDIGO/KPI_UPDATE/thang8/output'
os.makedirs(out_dir, exist_ok=True)
out_file = os.path.join(out_dir, 'TONG_HOP_HOA_DON_VA_LUONG_WHATSAPP_THANG_8.xlsx')

wb = openpyxl.Workbook()

# Styles
font_title = Font(name='Calibri', size=14, bold=True, color='1F4E79')
font_subtitle = Font(name='Calibri', size=10, italic=True, color='595959')
font_header_white = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
font_header_dark = Font(name='Calibri', size=11, bold=True, color='1E293B')
font_bold = Font(name='Calibri', size=11, bold=True)
font_regular = Font(name='Calibri', size=11)

fill_navy = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
fill_soft_blue = PatternFill(start_color='DDEBF7', end_color='DDEBF7', fill_type='solid')
fill_light_blue = PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid')
fill_emerald = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
fill_highlight = PatternFill(start_color='EBF9F1', end_color='EBF9F1', fill_type='solid')

border_thin = Border(
    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
)
border_header = Border(
    left=Side(style='thin', color='B0C4DE'), right=Side(style='thin', color='B0C4DE'),
    top=Side(style='medium', color='1F4E79'), bottom=Side(style='medium', color='1F4E79')
)
border_total = Border(
    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='1F4E79'), bottom=Side(style='double', color='1F4E79')
)

# ==============================================================================
# SHEET 1: TỔNG HỢP DOANH THU & THƯỞNG WHATSAPP THÁNG 8 (TỪ BẢNG LƯƠNG CHÍNH THỨC)
# ==============================================================================
ws1 = wb.active
ws1.title = 'Tổng hợp Lương T8 (WhatsApp)'
ws1.views.sheetView[0].showGridLines = True

ws1.merge_cells('A1:G1')
ws1['A1'] = 'BẢNG TỔNG HỢP DOANH THU & THƯỞNG ĐƠN WHATSAPP THÁNG 08/2026'
ws1['A1'].font = font_title
ws1['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws1.row_dimensions[1].height = 28

ws1.merge_cells('A2:G2')
ws1['A2'] = 'Số liệu chuẩn hóa từ Sheet Whatsapp trong Bảng lương Tháng 08/2026 chính thức'
ws1['A2'].font = font_subtitle
ws1['A2'].alignment = Alignment(horizontal='center', vertical='center')
ws1.row_dimensions[2].height = 18

headers1 = ['STT', 'Chi nhánh', 'Họ và tên Dược sĩ', 'Chức danh', 'Doanh thu đơn Whatsapp (VNĐ)', 'Tỷ lệ thưởng (%)', 'Tiền thưởng Whatsapp (VNĐ)']
ws1.append([])
ws1.append(headers1)
ws1.row_dimensions[4].height = 26

for c_idx in range(1, 8):
    cell = ws1.cell(4, c_idx)
    cell.font = font_header_white
    cell.fill = fill_navy
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = border_header

staff_wa_t8 = [
    {'stt': 1, 'branch': 'Đỗ Quang Đẩu', 'name': 'Ngô Thị Thanh Thắm', 'role': 'CHT', 'rev': 67046300, 'rate': 0.04},
    {'stt': 2, 'branch': 'Đỗ Quang Đẩu', 'name': 'Hoàng Thanh Thủy', 'role': 'DSXC', 'rev': 92439700, 'rate': 0.04},
    {'stt': 3, 'branch': 'Đỗ Quang Đẩu', 'name': 'Lê Thị Huyền Trân', 'role': 'DSXC', 'rev': 58824400, 'rate': 0.04},
    {'stt': 4, 'branch': 'Đỗ Quang Đẩu', 'name': 'Trần Thiên Phát', 'role': 'DSCD', 'rev': 9377800, 'rate': 0.04},
    {'stt': 5, 'branch': 'Đỗ Quang Đẩu', 'name': 'Phạm Thị Nghĩa Hương', 'role': 'DSCD', 'rev': 8501100, 'rate': 0.04},
    {'stt': 6, 'branch': 'Hàng Bông', 'name': 'Đinh Thị Khánh Ly', 'role': 'DSXC', 'rev': 19489000, 'rate': 0.015},
]

for idx, s in enumerate(staff_wa_t8, start=5):
    ws1.append([s['stt'], s['branch'], s['name'], s['role'], s['rev'], s['rate'], f'=E{idx}*F{idx}'])
    for c_idx in range(1, 8):
        cell = ws1.cell(idx, c_idx)
        cell.font = font_regular
        cell.border = border_thin
        if c_idx in (1, 4):
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif c_idx == 5:
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.number_format = '#,##0 "đ"'
        elif c_idx == 6:
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.number_format = '0.0%'
        elif c_idx == 7:
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.number_format = '#,##0 "đ"'
            cell.font = font_bold
            cell.fill = fill_emerald
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center')

last_row1 = 5 + len(staff_wa_t8)
ws1.cell(last_row1, 1, 'TỔNG CỘNG').font = font_bold
ws1.cell(last_row1, 1).alignment = Alignment(horizontal='center', vertical='center')
for c in range(1, 8):
    ws1.cell(last_row1, c).fill = fill_soft_blue
    ws1.cell(last_row1, c).border = border_total

c_tot_rev = ws1.cell(last_row1, 5, '=SUM(E5:E10)')
c_tot_rev.font = font_bold
c_tot_rev.number_format = '#,##0 "đ"'
c_tot_rev.alignment = Alignment(horizontal='right', vertical='center')

c_tot_bonus = ws1.cell(last_row1, 7, '=SUM(G5:G10)')
c_tot_bonus.font = font_bold
c_tot_bonus.number_format = '#,##0 "đ"'
c_tot_bonus.alignment = Alignment(horizontal='right', vertical='center')

for col in ws1.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws1.column_dimensions[col_letter].width = max(max_len + 5, 14)

# ==============================================================================
# SHEET 2: DANH SÁCH HÓA ĐƠN THÁNG 8 (01/08/2026 - 31/08/2026)
# ==============================================================================
ws2 = wb.create_sheet(title='Chi tiết HĐ Tháng 8 (KiotViet)')
ws2.views.sheetView[0].showGridLines = True

ws2.merge_cells('A1:I1')
ws2['A1'] = 'DANH SÁCH HÓA ĐƠN BÁN HÀNG THÁNG 08/2026 (THEO DƯỢC SĨ PHỤ TRÁCH)'
ws2['A1'].font = font_title
ws2['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws2.row_dimensions[1].height = 28

ws2.merge_cells('A2:I2')
ws2['A2'] = 'Trích xuất từ file gốc KiotViet Tháng 8: DanhSachChiTietHoaDon_3182026.xlsx (01/08/2026 - 31/08/2026)'
ws2['A2'].font = font_subtitle
ws2['A2'].alignment = Alignment(horizontal='center', vertical='center')
ws2.row_dimensions[2].height = 18

headers2 = ['STT', 'Chi nhánh', 'Mã hóa đơn', 'Thời gian bán', 'Người bán (Dược sĩ)', 'Kênh bán', 'Bảng giá', 'Mã khách hàng', 'Doanh thu (VNĐ)']
ws2.append([])
ws2.append(headers2)
ws2.row_dimensions[4].height = 26

for c_idx in range(1, len(headers2) + 1):
    cell = ws2.cell(4, c_idx)
    cell.font = font_header_white
    cell.fill = fill_navy
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = border_header

# Read August invoices from DanhSachChiTietHoaDon_3182026.xlsx
f_aug = 'd:/MEDIGO/KPI_UPDATE/thang8/DATA/DanhSachChiTietHoaDon_3182026.xlsx'
wb_cal = CalamineWorkbook.from_path(f_aug)
sheet = wb_cal.get_sheet_by_name(wb_cal.sheet_names[0])
rows = sheet.to_python()

sellers_set = {'Ngô Thị Thanh Thắm', 'Hoàng Thanh Thủy', 'Lê Thị Huyền Trân', 'Trần Thiên Phát', 'Phạm Thị Nghĩa Hương', 'Đinh Thị Khánh Ly'}
aug_invoices = {}

for r in rows[1:]:
    seller = str(r[7] or '').strip()
    if seller in sellers_set:
        code = str(r[1] or '').strip()
        if code not in aug_invoices:
            aug_invoices[code] = {
                'branch': str(r[0] or '').strip(),
                'code': code,
                'time': str(r[2])[:19] if r[2] else '',
                'seller': seller,
                'channel': str(r[8] or '').strip(),
                'price_list': str(r[6] or '').strip(),
                'cust_id': str(r[5] or '').strip(),
                'need_pay': float(r[12]) if r[12] is not None else 0.0,
            }

aug_invoices_list = sorted(aug_invoices.values(), key=lambda x: (x['branch'], x['time']), reverse=True)

start_row2 = 5
for idx, inv in enumerate(aug_invoices_list, start=1):
    r_idx = start_row2 + idx - 1
    ws2.append([
        idx,
        inv['branch'],
        inv['code'],
        inv['time'],
        inv['seller'],
        inv['channel'],
        inv['price_list'],
        inv['cust_id'],
        inv['need_pay']
    ])
    for c_idx in range(1, len(headers2) + 1):
        cell = ws2.cell(r_idx, c_idx)
        cell.font = font_regular
        cell.border = border_thin
        if c_idx in (1, 3):
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif c_idx in (4, 8):
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif c_idx == 9:
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.number_format = '#,##0 "đ"'
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center')

last_row2 = start_row2 + len(aug_invoices_list)
ws2.cell(last_row2, 1, 'TỔNG CỘNG').font = font_bold
ws2.cell(last_row2, 1).alignment = Alignment(horizontal='center', vertical='center')
for c in range(1, 10):
    ws2.cell(last_row2, c).fill = fill_soft_blue
    ws2.cell(last_row2, c).border = border_total

tot_cell2 = ws2.cell(last_row2, 9, f'=SUM(I{start_row2}:I{last_row2-1})')
tot_cell2.font = font_bold
tot_cell2.fill = fill_soft_blue
tot_cell2.border = border_total
tot_cell2.alignment = Alignment(horizontal='right', vertical='center')
tot_cell2.number_format = '#,##0 "đ"'

for col in ws2.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws2.column_dimensions[col_letter].width = max(max_len + 4, 14)
ws2.column_dimensions['A'].width = 8
ws2.column_dimensions['C'].width = 16
ws2.column_dimensions['D'].width = 20
ws2.column_dimensions['E'].width = 22
ws2.column_dimensions['I'].width = 18

# ==============================================================================
# SHEET 3: CHÍNH SÁCH THƯỞNG WHATSAPP THÁNG 8
# ==============================================================================
ws3 = wb.create_sheet(title='Chính sách Thưởng WhatsApp')
ws3.views.sheetView[0].showGridLines = True

ws3.merge_cells('A1:D1')
ws3['A1'] = 'QUY ĐỊNH CHƯƠNG TRÌNH THƯỞNG ĐƠN HÀNG QUA ỨNG DỤNG WHATSAPP'
ws3['A1'].font = font_title
ws3['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws3.row_dimensions[1].height = 28

headers3 = ['Mức KPI Doanh thu Cửa hàng', 'Ngưỡng Doanh thu (VNĐ)', 'Tỷ lệ thưởng Dược sĩ (%)', 'Ghi chú điều kiện']
ws3.append([])
ws3.append(headers3)
ws3.row_dimensions[3].height = 26

for c_idx in range(1, 5):
    cell = ws3.cell(3, c_idx)
    cell.font = font_header_white
    cell.fill = fill_navy
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = border_header

policy = [
    ('Mức 1', '150,000,000 đ', '3.0%', 'Cửa hàng đạt mốc 150 triệu -> Thưởng 3% doanh thu đơn cá nhân'),
    ('Mức 2', '200,000,000 đ', '4.0%', 'Cửa hàng đạt mốc 200 triệu -> Thưởng 4% doanh thu đơn cá nhân'),
    ('Mức 3', '280,000,000 đ', '6.0%', 'Cửa hàng đạt mốc 280 triệu -> Thưởng 6% doanh thu đơn cá nhân'),
    ('Cơ bản (Không đạt)', '< 150,000,000 đ', '1.5%', 'Cửa hàng không đạt KPI -> Thưởng mức cơ bản 1.5% doanh thu đơn cá nhân'),
]

for idx, p in enumerate(policy, start=4):
    ws3.append(list(p))
    for c_idx in range(1, 5):
        cell = ws3.cell(idx, c_idx)
        cell.font = font_regular
        cell.border = border_thin
        if c_idx in (1, 2, 3):
            cell.alignment = Alignment(horizontal='center', vertical='center')
            if c_idx == 3: cell.font = font_bold
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center')

# Rules notes
ws3.append([])
ws3.append(['ĐIỀU KIỆN GHI NHẬN ĐƠN HÀNG HỢP LỆ (THEO QUY CHẾ THÁNG 8):'])
ws3.cell(9, 1).font = font_bold

rules = [
    '1. Dược sĩ lên đơn trên phần mềm KiotViet, bắt buộc phải ghi chú "Whatsapp" trong đơn hàng.',
    '2. Khách hàng có tương tác trực tiếp qua WhatsApp (hỏi giá, tư vấn dùng thuốc, xác nhận đơn hàng...).',
    '3. Đơn hàng được tạo sau cuộc trao đổi và truy vết được (SĐT, tên khách, thời gian tạo, ghi chú trên KiotViet).',
    '4. Chỉ áp dụng cho tệp khách hàng quốc tế kết nối qua QR WhatsApp của nhà thuốc.'
]
for r in rules:
    ws3.append([r])
    cur_r = ws3.max_row
    ws3.cell(cur_r, 1).font = font_regular
    ws3.merge_cells(f'A{cur_r}:D{cur_r}')

for col in ws3.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws3.column_dimensions[col_letter].width = max(max_len + 5, 20)
ws3.column_dimensions['A'].width = 25
ws3.column_dimensions['B'].width = 25
ws3.column_dimensions['C'].width = 25
ws3.column_dimensions['D'].width = 65

wb.save(out_file)
print(f'--> Đã tạo thành công file chuẩn Tháng 8: {out_file}')
