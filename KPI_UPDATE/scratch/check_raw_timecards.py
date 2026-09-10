import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
ws = wb.active
print(f"Sheet: {ws.title}, max_row={ws.max_row}")

headers = [ws.cell(1, c).value for c in range(1, 20)]
print("Headers:", headers)

print("\n--- Chi tiết ca của Hứa Thị Kim Thoa trong DATA Chấm công ---")
thoa_shifts = []
for r in range(2, ws.max_row+1):
    name = str(ws.cell(r, 2).value or '')
    if 'thoa' in name.lower():
        cn = ws.cell(r, 1).value
        ca = ws.cell(r, 4).value
        gio = ws.cell(r, 5).value
        thoa_shifts.append((cn, ca, gio))
        print(f"Row {r:3d} | CN: {cn} | Ca: {ca} | Giờ: {gio}")

print(f"Tổng số bản ghi của Thoa: {len(thoa_shifts)}")

print("\n--- Chi tiết ca của Nguyễn Thị Tâm trong DATA Chấm công ---")
tam_shifts = []
for r in range(2, ws.max_row+1):
    name = str(ws.cell(r, 2).value or '')
    if 'tâm' in name.lower() and 'đoàn' not in name.lower():
        cn = ws.cell(r, 1).value
        ca = ws.cell(r, 4).value
        gio = ws.cell(r, 5).value
        tam_shifts.append((cn, ca, gio))
        print(f"Row {r:3d} | CN: {cn} | Ca: {ca} | Giờ: {gio}")

print(f"Tổng số bản ghi của Tâm: {len(tam_shifts)}")
