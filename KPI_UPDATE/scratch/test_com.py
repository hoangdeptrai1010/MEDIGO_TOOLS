import openpyxl
import subprocess
import os

def test_step(step_name, func):
    wb = openpyxl.load_workbook('goc/NHÀ THUỐC THÁNG 8 2026.xlsx')
    func(wb)
    wb.save('test_step.xlsx')
    wb.close()
    res = subprocess.run(['powershell', '-ExecutionPolicy', 'Bypass', '-File', 'recalc_workbook.ps1', '-FilePath', 'test_step.xlsx'], capture_output=True, text=True)
    success = 'successfully' in res.stdout
    if success:
        print(f'{step_name}: SUCCESS')
    else:
        print(f'{step_name}: FAILED')
        for line in res.stdout.splitlines():
            if 'Error' in line or 'Exception' in line:
                print('   ', line)

# Test 3: Add Hot Bill HN
def step_hot_bill(wb):
    ws_hb = wb.create_sheet('Hot Bill HN')
    ws_hb.append(['Chi nhánh', 'Thời gian', 'Mã hóa đơn', 'Dược sĩ', 'Doanh số', 'Mức', 'Thưởng', 'Ghi chú'])
    ws_hb.append(['Hàng Bông', '2026-08-15', 'HD123', 'Đinh Thị Lan Anh', 1200000, 1, 50000, 'CK'])

test_step('Step 3 (Hot Bill HN sheet)', step_hot_bill)

# Test 4: Modify project sheet formulas
def step_proj(wb):
    step_hot_bill(wb)
    ws = wb['Dự án T8']
    ws.cell(3, 10, '=IF($A3="Hàng Bông", 500000, 0)')
    ws.cell(3, 12, "=IFERROR(SUMIFS('Hot Bill HN'!$G:$G, 'Hot Bill HN'!$A:$A, $A3, 'Hot Bill HN'!$D:$D, $B3), 0)")

test_step('Step 4 (Project sheet formula with Hot Bill reference)', step_proj)
