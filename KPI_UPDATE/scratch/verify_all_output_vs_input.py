import openpyxl
import sys, io

if sys.stdout.encoding != 'utf-8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

sys.path.insert(0, 'TOOL_KPISHEET')
from builder_engine import parse_hcm_proposal, parse_hn_proposal

def verify_kpisheet(month_num):
    print(f"\n=================================================================")
    print(f"VERIFYING OUTPUT vs INPUT FOR MONTH {month_num}")
    print(f"=================================================================")
    
    hcm_p = 'TOOL_KPISHEET/KPI CNT HCM Tháng 09.xlsx'
    hn_p = 'TOOL_KPISHEET/e_xuat_KPI_Quy_3.26.xlsx'
    out_p = f'TOOL_KPISHEET/output/NHÀ THUỐC THÁNG {month_num} 2026.xlsx'
    
    hcm_staff, hcm_stores = parse_hcm_proposal(hcm_p, month_num)
    hn_staff, hn_stores = parse_hn_proposal(hn_p, month_num)
    
    expected_staff = {s['name']: s for s in hcm_staff + hn_staff}
    expected_stores = {**hcm_stores, **hn_stores}
    
    wb_out = openpyxl.load_workbook(out_p, data_only=False)
    
    # 1. Verify Sheet 'kpi dược sĩ'
    ws_ds = wb_out['kpi dược sĩ']
    mismatches = []
    
    for r in range(3, ws_ds.max_row + 1):
        st = ws_ds.cell(r, 1).value
        name = ws_ds.cell(r, 2).value
        role = ws_ds.cell(r, 3).value
        tb_bill = ws_ds.cell(r, 4).value
        target_thang = ws_ds.cell(r, 22).value # Col V
        
        if not name:
            continue
            
        if name not in expected_staff:
            mismatches.append(f"Row {r}: Staff '{name}' not found in input proposals!")
            continue
            
        exp = expected_staff[name]
        
        if exp['store'] != st:
            mismatches.append(f"Row {r} '{name}': Store mismatch: Out='{st}' vs Exp='{exp['store']}'")
            
        if exp['tb_bill'] > 0 and float(tb_bill or 0) != float(exp['tb_bill']):
            mismatches.append(f"Row {r} '{name}': TB Bill mismatch: Out={tb_bill} vs Exp={exp['tb_bill']}")
            
        if exp['kpi_thang'] > 0 and float(target_thang or 0) != float(exp['kpi_thang']):
            mismatches.append(f"Row {r} '{name}': KPI Tháng mismatch: Out={target_thang} vs Exp={exp['kpi_thang']}")

    # 2. Verify Sheet 'kpi nhà thuốc'
    ws_nt = wb_out['kpi nhà thuốc']
    for r in range(3, ws_nt.max_row + 1):
        st = ws_nt.cell(r, 2).value
        target_nt = ws_nt.cell(r, 3).value
        if not st:
            continue
        if st in expected_stores:
            exp_t = expected_stores[st]
            if float(target_nt or 0) != float(exp_t):
                mismatches.append(f"Store '{st}': Target NT mismatch: Out={target_nt} vs Exp={exp_t}")

    if mismatches:
        print(f"❌ Found {len(mismatches)} mismatches:")
        for m in mismatches:
            print("  -", m)
    else:
        print(f"✅ ALL {len(expected_staff)} STAFF & {len(expected_stores)} STORES MATCH 100% INPUT DATA PERFECTLY!")

verify_kpisheet(9)
verify_kpisheet(8)
