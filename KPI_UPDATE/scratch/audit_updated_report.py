import openpyxl

wb = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx', data_only=True)

print("=== 1. KIỂM TRA LỖI TOÀN BỘ CÁC SHEET ===")
for sname in wb.sheetnames:
    ws = wb[sname]
    errs = []
    for r in range(1, ws.max_row+1):
        for c in range(1, ws.max_column+1):
            val = str(ws.cell(r, c).value or '')
            if any(e in val for e in ['#NAME?', '#VALUE!', '#REF!', '#DIV/0!']):
                errs.append((r, c, val))
    status = f"{len(errs)} loi" if errs else "SACH 100%"
    print(f"Sheet {sname:18s} : {status}")

print("\n=== 2. SHEET DỰ ÁN T8 CỦA 8 DƯỢC SĨ HÀ NỘI ===")
ws_da = wb['Dự án T8']
for r in range(3, 11):
    name = str(ws_da.cell(r, 2).value or '')
    d = float(ws_da.cell(r, 4).value or 0)
    j = float(ws_da.cell(r, 10).value or 0)
    k = float(ws_da.cell(r, 11).value or 0)
    l = float(ws_da.cell(r, 12).value or 0)
    m = float(ws_da.cell(r, 13).value or 0)
    print(f"Row {r:2d} ({name:25s}): D={d:>12,.0f} | Bac={j:>10,.0f} | Them={k:>10,.0f} | HotBill={l:>8,.0f} | Total DA={m:>10,.0f}")

print("\n=== 3. SHEET KPI DƯỢC SĨ (ĐỒNG BỘ TOÀN BỘ NHÂN SỰ) ===")
ws_ds = wb['kpi dược sĩ']
for r in [3, 4, 5, 6, 7, 8, 9, 10, 24, 26]:
    name = str(ws_ds.cell(r, 2).value or '')
    cn = str(ws_ds.cell(r, 1).value or '')
    kpi = float(ws_ds.cell(r, 11).value or 0)
    hs = ws_ds.cell(r, 12).value
    da = float(ws_ds.cell(r, 13).value or 0)
    tot = float(ws_ds.cell(r, 14).value or 0)
    print(f"Row {r:2d} ({name:25s} @ {cn:12s}): KPI={kpi:>12,.0f} | He so={hs} | Thuong DA={da:>10,.0f} | TONG LUONG KPI={tot:>12,.0f}")
