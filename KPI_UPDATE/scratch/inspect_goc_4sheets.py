import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

f_nt7 = r"d:\MEDIGO\KPI_UPDATE\goc\NHÀ THUỐC THÁNG 7 2026.xlsx"
f_nt8 = r"d:\MEDIGO\KPI_UPDATE\goc\NHÀ THUỐC THÁNG 8 2026.xlsx"

for fname, fpath in [("NHÀ THUỐC THÁNG 7", f_nt7), ("NHÀ THUỐC THÁNG 8", f_nt8)]:
    print(f"\n=======================================================")
    print(f"FILE: {fname}")
    print(f"=======================================================")
    wb = openpyxl.load_workbook(fpath, data_only=True)
    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f"\n--- SHEET: '{sname}' (max_row={ws.max_row}, max_col={ws.max_column}) ---")
        for r in range(1, min(10, ws.max_row + 1)):
            row_vals = [ws.cell(r, c).value for c in range(1, min(25, ws.max_column + 1))]
            if any(v is not None for v in row_vals):
                print(f"Row {r:2d}: {row_vals}")
