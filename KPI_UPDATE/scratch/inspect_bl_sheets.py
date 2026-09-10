import openpyxl
import os
import sys
sys.stdout.reconfigure(encoding='utf-8')

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_bangluong = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final.xlsx")
wb = openpyxl.load_workbook(f_bangluong, data_only=True)
print("BANGLUONG sheetnames:", wb.sheetnames)

for s in ["Sheet1", "Med"]:
    if s in wb.sheetnames:
        ws = wb[s]
        print(f"\n--- SHEET: {s} (rows: {ws.max_row}, cols: {ws.max_column}) ---")
        for r in range(1, min(20, ws.max_row + 1)):
            vals = [ws.cell(r, c).value for c in range(1, min(20, ws.max_column + 1))]
            if any(v is not None for v in vals):
                print(f"Row {r}: {vals[:12]}")
