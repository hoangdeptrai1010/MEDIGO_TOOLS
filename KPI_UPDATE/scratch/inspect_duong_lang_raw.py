import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
ws = wb.active

print("=== TẤT CẢ NHÂN VIÊN VÀ CA LÀM VIỆC TẠI ĐƯỜNG LÁNG TRONG FILE CHẤM CÔNG GỐC ===")

curr_staff = ""
curr_code = ""
curr_role = ""

for r in range(2, ws.max_row+1):
    c_stt = ws.cell(r, 1).value
    c_code = ws.cell(r, 2).value
    c_name = ws.cell(r, 3).value
    c_dept = ws.cell(r, 4).value
    c_role = ws.cell(r, 5).value
    c_branch = ws.cell(r, 6).value
    c_shift = ws.cell(r, 7).value
    
    if c_name and str(c_name).strip():
        curr_staff = str(c_name).strip()
        curr_code = str(c_code or '').strip()
        curr_role = str(c_role or '').strip()
    
    if c_branch and 'đường láng' in str(c_branch).lower():
        # Count non-empty punch cells across days
        punches = []
        for day in range(1, 32):
            col_in = 8 + (day-1)*2
            col_out = col_in + 1
            v_in = ws.cell(r, col_in).value
            v_out = ws.cell(r, col_out).value
            if v_in or v_out:
                punches.append((day, v_in, v_out))
        print(f"Row {r:3d} | Mã: {curr_code:<8} | Tên: {curr_staff:<22} | Vị trí: {curr_role:<15} | Ca: {str(c_shift):<30} | Số ngày có quẹt thẻ: {len(punches)}")
