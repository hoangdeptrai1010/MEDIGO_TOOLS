# -*- coding: utf-8 -*-
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb_kpi = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx', data_only=True)
ws_da_kpi = wb_kpi['Dự án T8']

kpi_duan_map = {}
for r in range(3, ws_da_kpi.max_row + 1):
    name = ws_da_kpi.cell(r, 2).value
    if name:
        cn = ws_da_kpi.cell(r, 1).value
        role = ws_da_kpi.cell(r, 3).value
        ck_day = ws_da_kpi.cell(r, 8).value or 0
        cb_day = ws_da_kpi.cell(r, 9).value or 0
        thuong_tier = ws_da_kpi.cell(r, 10).value or 0
        thuong_them = ws_da_kpi.cell(r, 11).value or 0
        thuong_hot = ws_da_kpi.cell(r, 12).value or 0
        tot_da = ws_da_kpi.cell(r, 13).value or 0
        
        kpi_duan_map[str(name).strip()] = {
            'cn': cn,
            'role': role,
            'ck_day': ck_day,
            'cb_day': cb_day,
            'tier': thuong_tier,
            'them': thuong_them,
            'hot': thuong_hot,
            'total': tot_da
        }

print(f"=== BẢNG TỔNG HỢP THƯỞNG DỰ ÁN T8 TỪ BÁO CÁO KPI ({len(kpi_duan_map)} Dược sĩ) ===")
print(f"{'Nhân viên':<26} | {'Chi nhánh':<15} | {'Thưởng Tier':<14} | {'Thưởng thêm':<14} | {'Hot Bill HN':<14} | {'TOTAL DỰ ÁN':<14}")
print("-" * 105)

for name, d in kpi_duan_map.items():
    if d['total'] > 0:
        print(f"{name:<26} | {str(d['cn']):<15} | {d['tier']:>12,.0f} đ | {d['them']:>12,.0f} đ | {d['hot']:>12,.0f} đ | {d['total']:>12,.0f} đ")
