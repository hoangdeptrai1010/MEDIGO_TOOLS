import openpyxl, sys, io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('d:/MEDIGO/KPI_UPDATE/thang7/target/BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=True)
ws = wb['whatsapp']
print("=== THANG 7 SHEET whatsapp ===")
for r in range(1, min(50, ws.max_row+1)):
    row_vals = {c: ws.cell(r, c).value for c in range(1, ws.max_column+1) if ws.cell(r, c).value is not None}
    if row_vals:
        print(f"Row {r:2d}: {row_vals}")

wb.close()
