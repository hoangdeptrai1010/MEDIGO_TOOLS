import openpyxl

wb_aug = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=False)

for sname in ['KPI', 'Dự án', 'MiniKat - HN', 'MiniKat - HCM', 'whatsapp', 'Thưởng CK', 'Cận date', 'KPI trừ']:
    ws = wb_aug[sname]
    print(f"=== {sname} in thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx ===")
    print(f"max_row={ws.max_row}, max_col={ws.max_column}")
    for r in range(1, 5):
        vals = [ws.cell(r, c).value for c in range(1, 10)]
        if any(v is not None for v in vals):
            print(f"  R{r}: {vals}")
