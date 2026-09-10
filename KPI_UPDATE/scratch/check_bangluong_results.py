import openpyxl
import sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'd:\MEDIGO\KPI_UPDATE\thang8\bangluong_thang8_hoanthien.xlsx', data_only=True)
ws = wb['BẢNG LƯƠNG']

print(f"{'STT':<4} | {'Mã NV':<8} | {'Họ và tên':<26} | {'Chức danh':<10} | {'Ca ngày(J)':<10} | {'Ca đêm(K)':<10} | {'Tổng công(R)':<12} | {'PC Đêm(X)':<14} | {'Thực lĩnh(AR)':<14}")
print("-" * 115)

for r in range(1, 10):
    print(r, [ws.cell(r, c).value for c in range(1, 6)])
    stt = ws.cell(r, 1).value
    manv = ws.cell(r, 2).value
    name = ws.cell(r, 3).value
    chucdanh = ws.cell(r, 4).value
    cangay = ws.cell(r, 10).value
    cadem = ws.cell(r, 11).value
    tongcong = ws.cell(r, 18).value
    pc_dem = ws.cell(r, 24).value
    thuclinh = ws.cell(r, 44).value
    
    if name and str(name).strip() != "" and "Tổng" not in str(name):
        cangay_str = f"{cangay:.2f}" if isinstance(cangay, (int, float)) else str(cangay)
        cadem_str = f"{cadem:.2f}" if isinstance(cadem, (int, float)) else str(cadem)
        tongcong_str = f"{tongcong:.2f}" if isinstance(tongcong, (int, float)) else str(tongcong)
        pc_dem_str = f"{pc_dem:,.0f}" if isinstance(pc_dem, (int, float)) else str(pc_dem)
        thuclinh_str = f"{thuclinh:,.0f}" if isinstance(thuclinh, (int, float)) else str(thuclinh)
        
        print(f"{str(stt):<4} | {str(manv):<16} | {str(name):<26} | {str(chucdanh):<10} | {cangay_str:<10} | {cadem_str:<10} | {tongcong_str:<12} | {pc_dem_str:<14} | {thuclinh_str:<14}")
