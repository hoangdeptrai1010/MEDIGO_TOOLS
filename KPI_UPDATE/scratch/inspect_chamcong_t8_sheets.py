import openpyxl

wb = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"=== {sname} ===")
    for r in range(1, 10):
        row_vals = [ws.cell(r, c).value for c in range(1, 15)]
        if any(v is not None for v in row_vals):
            print(f"R{r}: {row_vals}")
