import os
import sys
import io
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

sys.path.insert(0, 'd:/MEDIGO/KPI_UPDATE/TOOL_KPISHEET')
from builder_engine import generate_kpisheet_package

sys.path.insert(0, 'd:/MEDIGO/KPI_UPDATE/TOOL_KPI')
from kpi_engine import execute_kpi_engine

print("================================================================")
print("TEST 1: CHẠY TOOL_KPISHEET GENERATE TEMPLATE THÁNG 9")
print("================================================================")

hcm_file = 'd:/MEDIGO/KPI_UPDATE/thang9/data/kpi_thang9/KPI CNT HCM Tháng 09.xlsx'
hn_file = 'd:/MEDIGO/KPI_UPDATE/thang9/data/kpi_thang9/e_xuat_KPI_Quy_3.26.xlsx'
tmpl_out = 'd:/MEDIGO/KPI_UPDATE/thang9/output/NHÀ THUỐC THÁNG 9 2026.xlsx'

res_kpisheet = generate_kpisheet_package(
    hcm_file_path=hcm_file,
    hn_file_path=hn_file,
    month_num=9,
    output_filepath=tmpl_out
)
print("--> Result:", res_kpisheet['filename'], "Staff count:", res_kpisheet['staff_count'])

# Inspect template
wb_tmpl = openpyxl.load_workbook(tmpl_out, data_only=False)

print("\n--- KIỂM TRA SHEET 'kpi nhà thuốc' (Template T9) ---")
ws_nt = wb_tmpl['kpi nhà thuốc']
for r in range(3, ws_nt.max_row + 1):
    cht = ws_nt.cell(r, 1).value
    st = ws_nt.cell(r, 2).value
    tgt = ws_nt.cell(r, 3).value
    if st:
        print(f"  Row {r}: {st} (CHT: {cht}) -> Target: {tgt:,.0f} đ")

print("\n--- KIỂM TRA SHEET 'kpi dược sĩ' (Template T9 - 10 NV Hà Nội) ---")
ws_ds = wb_tmpl['kpi dược sĩ']
for r in range(3, ws_ds.max_row + 1):
    st = ws_ds.cell(r, 1).value
    name = ws_ds.cell(r, 2).value
    role = ws_ds.cell(r, 3).value
    tb = ws_ds.cell(r, 4).value
    col_j = ws_ds.cell(r, 10).value
    col_k = ws_ds.cell(r, 11).value
    col_v = ws_ds.cell(r, 22).value
    if st in ['Hàng Bông', 'Đường Láng']:
        print(f"  Row {r:2d} | {st:10s} | {name:25s} | {role:10s} | Monthly: {col_v:>11,d} đ | TB: {tb}")

# Check formula sample
print("\nSample Formula Col J (Row 3):", ws_ds.cell(3, 10).value)
print("Sample Formula Col K (Row 3):", str(ws_ds.cell(3, 11).value)[:100] + "...")
print("Sample Formula Col F (Row 3):", str(ws_ds.cell(3, 6).value)[:100] + "...")

wb_tmpl.close()

print("\n================================================================")
print("TEST 2: CHẠY TOOL_KPI TÍNH TOÁN BÁO CÁO THÁNG 9")
print("================================================================")

# Find Month 9 invoices and return files
import glob
hd_files = glob.glob('d:/MEDIGO/KPI_UPDATE/thang9/data/*hoa_don*.xlsx') + glob.glob('d:/MEDIGO/KPI_UPDATE/thang9/data/*hoadon*.xlsx') + glob.glob('d:/MEDIGO/KPI_UPDATE/thang9/*hoa_don*.xlsx')
th_files = glob.glob('d:/MEDIGO/KPI_UPDATE/thang9/data/*tra_hang*.xlsx') + glob.glob('d:/MEDIGO/KPI_UPDATE/thang9/data/*trahang*.xlsx') + glob.glob('d:/MEDIGO/KPI_UPDATE/thang9/*tra_hang*.xlsx')

hd_file = hd_files[0] if hd_files else None
th_file = th_files[0] if th_files else None

print("Hóa đơn file:", hd_file)
print("Trả hàng file:", th_file)

if hd_file:
    plan_file = 'd:/MEDIGO/KPI_UPDATE/plans/KeHoachKPI_2026-09.xlsx'
    if not os.path.exists(plan_file):
        from build_monthly_plan_packages import build_plan_from_manager_proposals
        build_plan_from_manager_proposals(hcm_file, hn_file, month_num=9, output_filepath=plan_file)

    report_out = 'd:/MEDIGO/KPI_UPDATE/thang9/output/baocaokpi_thang9_hoanthien.xlsx'
    
    execute_kpi_engine(
        hoadon_file=hd_file,
        trahang_file=th_file,
        template_file=tmpl_out,
        plan_file=plan_file,
        output_file=report_out,
        target_month=9,
        target_year=2026
    )

    print("\n--- KIỂM TRA BÁO CÁO HOÀN THIỆN THÁNG 9 ---")
    wb_rep = openpyxl.load_workbook(report_out, data_only=False)
    ws_rep_ds = wb_rep['kpi dược sĩ']
    print("kpi dược sĩ Rows count:", ws_rep_ds.max_row)
    for r in range(3, min(15, ws_rep_ds.max_row + 1)):
        st = ws_rep_ds.cell(r, 1).value
        name = ws_rep_ds.cell(r, 2).value
        role = ws_rep_ds.cell(r, 3).value
        col_j = ws_rep_ds.cell(r, 10).value
        col_v = ws_rep_ds.cell(r, 22).value
        print(f"  Row {r:2d} | {st:10s} | {name:25s} | Role: {role:8s} | Monthly: {col_v:,.0f} | % KPI Formula: {col_j}")
        
    wb_rep.close()

print("\n================================================================")
print("TẤT CẢ CÁC BƯỚC ĐÃ HOÀN TẤT THÀNH CÔNG!")
print("================================================================")
