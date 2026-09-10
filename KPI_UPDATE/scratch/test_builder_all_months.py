import os
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

sys.path.insert(0, r"d:\MEDIGO\KPI_UPDATE\TOOL_KPISHEET")
from builder_engine import generate_kpisheet_package

hcm_p = r"d:\MEDIGO\KPI_UPDATE\TOOL_KPISHEET\KPI CNT HCM Tháng 09.xlsx"
hn_p = r"d:\MEDIGO\KPI_UPDATE\TOOL_KPISHEET\e_xuat_KPI_Quy_3.26.xlsx"
proj_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\Pharmacy_retail_Store_KPIs_August 2026"

print("--- Testing Generation for Month 8 ---")
res8 = generate_kpisheet_package(hcm_p, hn_p, month_num=8, project_folder=proj_dir)
print("Month 8 Result:", res8)

print("\n--- Testing Generation for Month 9 ---")
res9 = generate_kpisheet_package(hcm_p, hn_p, month_num=9, project_folder=proj_dir)
print("Month 9 Result:", res9)

# Verify sheet structure of both outputs
for fn in [res8['filepath'], res9['filepath']]:
    wb = openpyxl.load_workbook(fn, data_only=True)
    print(f"\nWorkbook: {os.path.basename(fn)}")
    print("Sheets:", wb.sheetnames)
    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f"  - Sheet '{sname}': {ws.max_row} rows, {ws.max_column} cols")
