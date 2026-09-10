import openpyxl

wb = openpyxl.load_workbook('thang8/baocaokpi_thang8_hoanthien.xlsx', data_only=True)

print("==========================================================================================")
print("AUDIT TOÀN BỘ BÁO CÁO KPI THÁNG 8 HOÀN THIỆN (thang8/baocaokpi_thang8_hoanthien.xlsx)")
print("==========================================================================================")

# 1. Check for formula errors across all sheets
for s in wb.sheetnames:
    ws = wb[s]
    errs = []
    for r in range(1, ws.max_row+1):
        for c in range(1, ws.max_column+1):
            val = str(ws.cell(r, c).value or '')
            if val in ['#NAME?', '#VALUE!', '#REF!', '#DIV/0!', '#N/A', '#NULL!']:
                errs.append((r, c, val))
    print(f"Sheet '{s:15s}': {len(errs)} errors found" + (f" -> {errs[:3]}" if errs else " -> ✅ SẠCH HOÀN TOÀN"))

# 2. Print CHT results
print("\n--- 1. KẾT QUẢ CỬA HÀNG TRƯỞNG (SHEET KPI NHÀ THUỐC) ---")
ws_nt = wb['kpi nhà thuốc']
for r in range(3, ws_nt.max_row+1):
    cht = ws_nt.cell(r, 1).value
    nt = ws_nt.cell(r, 2).value
    w = ws_nt.cell(r, 23).value
    g = ws_nt.cell(r, 7).value
    if cht:
        g_str = f"{g:,.0f} VNĐ" if isinstance(g, (int, float)) else str(g)
        print(f"CHT: {cht:25s} | NT: {nt:15s} | Xét Đạt: {str(w):10s} | Thưởng CHT: {g_str:>15s}")

# 3. Print Dược sĩ & Hot Bill results
print("\n--- 2. KẾT QUẢ DƯỢC SĨ TIÊU BIỂU & HOT BILL (SHEET KPI DƯỢC SĨ) ---")
ws_ds = wb['kpi dược sĩ']
key_staff = ['Vũ Thanh Hằng', 'Đinh Thị Lan Anh', 'Nguyễn Mạnh Tuấn', 'Đinh Thị Khánh Ly', 'Hứa Thị Kim Thoa', 
             'Lê Thị Soạn', 'Nguyễn Thị Tâm', 'Nguyễn Thị Mai Duyên', 'Trần Thị Kim Khánh', 'Hoàng Thanh Thủy',
             'Phan Công Vũ Tài', 'Trịnh Thị Phượng', 'Cao Trọng Nhân', 'Đỗ Thị Phương Thảo', 'Hoàng Lâm Gia Bảo',
             'Võ Ngọc Giàu Sang', 'Cù Thị Tường Vy']

print(f"{'Nhân viên':25s} | {'Chi nhánh':12s} | {'Thưởng KPI':>14s} | {'Thưởng Dự Án':>14s} | {'TỔNG THƯỞNG':>14s}")
print("-" * 88)
for r in range(3, ws_ds.max_row+1):
    nv = ws_ds.cell(r, 2).value
    if nv in key_staff:
        b = ws_ds.cell(r, 1).value
        k = ws_ds.cell(r, 11).value or 0
        m = ws_ds.cell(r, 13).value or 0
        tot = ws_ds.cell(r, 14).value or 0
        k_s = f"{k:,.0f}đ" if isinstance(k, (int, float)) and k > 0 else "0đ"
        m_s = f"{m:,.0f}đ" if isinstance(m, (int, float)) and m > 0 else "0đ"
        tot_s = f"{tot:,.0f}đ" if isinstance(tot, (int, float)) and tot > 0 else "0đ"
        print(f"{nv:25s} | {b:12s} | {k_s:>14s} | {m_s:>14s} | {tot_s:>14s}")
