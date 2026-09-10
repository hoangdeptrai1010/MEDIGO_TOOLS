import openpyxl

wb = openpyxl.load_workbook('thang7/BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=False)
ws = wb['Dự án']
print("Dự án headers:")
for c in range(1, 15):
    print(f"Col {c}: {ws.cell(1, c).value}")
print("\nDự án sample row 2:")
for c in range(1, 15):
    print(f"Col {c}: {ws.cell(2, c).value}")
