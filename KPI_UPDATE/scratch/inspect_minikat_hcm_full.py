# -*- coding: utf-8 -*-
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/BANGLUONGTHANG8.xlsx', data_only=True)
ws_mk = wb['MiniKat - HCM']

print(f"{'Row':<4} | {'Người bán (A)':<26} | {'SL Party (B)':<12} | {'SL KAT (C)':<10} | {'SL Lady (D)':<11} | {'DT Lady (E)':<12} | {'DT KAT (F)':<12} | {'Thưởng (J)':<12} | {'Chi nhánh (M)':<16} | {'Thưởng CHT (U)':<14}")
print("-" * 145)

for r in range(2, ws_mk.max_row+1):
    a = ws_mk.cell(r, 1).value
    b = ws_mk.cell(r, 2).value or 0
    c = ws_mk.cell(r, 3).value or 0
    d = ws_mk.cell(r, 4).value or 0
    e = ws_mk.cell(r, 5).value or 0
    f = ws_mk.cell(r, 6).value or 0
    j = ws_mk.cell(r, 10).value or 0
    m = ws_mk.cell(r, 13).value or ''
    u = ws_mk.cell(r, 21).value or 0
    
    if a or m:
        print(f"{r:<4} | {str(a):<26} | {b:<12} | {c:<10} | {d:<11} | {e:>10,.0f} đ | {f:>10,.0f} đ | {j:>10,.0f} đ | {str(m):<16} | {u:>10,.0f} đ")
