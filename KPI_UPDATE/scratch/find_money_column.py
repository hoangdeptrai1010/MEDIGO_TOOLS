import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_bl = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final_backup.xlsx")
wb = openpyxl.load_workbook(f_bl, data_only=True)
ws_med = wb["Med"]

print("=== All headers in Med (Row 9) ===")
for c in range(1, ws_med.max_column + 1):
    h = ws_med.cell(9, c).value
    if h:
        print(f"Col {c:3d} ({openpyxl.utils.get_column_letter(c):>3s}): {str(h).strip()}")

print("\n=== Check sample rows in Med for money columns ===")
# Row 13 in Med is Nguyễn Ngọc Anh Thư. In Sheet1 Row 2, Số tiền is 15736910.
# Row 12 in Med is Nguyễn Trần Ngọc Phương. In Sheet1 Row 3, Số tiền is 6373917.
# Row 15 in Med is Ngô Trần Tuyết Vy. In Sheet1 Row 4, Số tiền is 3401714.
# Row 16 in Med is Phạm Nguyễn Ngọc Quý. In Sheet1 Row 5, Số tiền is 5221977.

targets = [
    (13, 15736910, "Nguyễn Ngọc Anh Thư"),
    (12, 6373917, "Nguyễn Trần Ngọc Phương"),
    (15, 3401714, "Ngô Trần Tuyết Vy"),
    (16, 5221977, "Phạm Nguyễn Ngọc Quý")
]

for r_med, expected_val, name in targets:
    print(f"\nSearching for {name} (expected {expected_val}) in Med row {r_med}:")
    for c in range(1, ws_med.max_column + 1):
        v = ws_med.cell(r_med, c).value
        try:
            if v is not None and abs(float(v) - expected_val) < 1.0:
                h = ws_med.cell(9, c).value
                print(f"  FOUND MATCH at Col {c} ({openpyxl.utils.get_column_letter(c)}): header={h!r}, value={v}")
        except:
            pass
