import openpyxl

wb_cc = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
ws_th = wb_cc['Bảng tổng hợp chấm công']
wb_aug = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=True)
ws_gc = wb_aug['Giờ công']

print("--- Comparing Bảng tổng hợp chấm công with Giờ công in thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx ---")
print("Bảng tổng hợp chấm công rows 4-15:")
for r in range(4, 16):
    stt = ws_th.cell(r, 1).value
    manv = ws_th.cell(r, 2).value
    ten = ws_th.cell(r, 3).value
    cn = ws_th.cell(r, 6).value
    ca = ws_th.cell(r, 7).value
    gio_ca = ws_th.cell(r, 42).value # Col AP
    gio_tong = ws_th.cell(r, 43).value # Col AQ
    print(f"R{r}: STT={stt}, Ma={manv}, Ten={ten}, CN={cn}, Ca={ca}, GioCa={gio_ca}, GioTong={gio_tong}")

print("\nGiờ công in thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx rows 2-15:")
for r in range(2, 16):
    a = ws_gc.cell(r, 1).value
    b = ws_gc.cell(r, 2).value
    c = ws_gc.cell(r, 3).value
    d = ws_gc.cell(r, 4).value
    e = ws_gc.cell(r, 5).value
    print(f"R{r}: A={a}, B={b}, C={c}, D={d}, E={e}")
