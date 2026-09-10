import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws = wb['BẢNG LƯƠNG']

print("=== CHECK DISCREPANCIES IN BANGLUONGTHANG8 ===")

# Row 74 totals
tot_e = ws.cell(74, 5).value
tot_f = ws.cell(74, 6).value
tot_g = ws.cell(74, 7).value
tot_h = ws.cell(74, 8).value
print(f"Row 74: E={tot_e}, F={tot_f}, G={tot_g}, H={tot_h}")

# Row 88 totals (bảng chi nhánh)
b_tot_g = ws.cell(88, 7).value
b_tot_h = ws.cell(88, 8).value
b_tot_i = ws.cell(88, 9).value
print(f"Row 88 (Tổng chi nhánh): G={b_tot_g}, H={b_tot_h}, I={b_tot_i}")

print(f"Lệch G (Row 74 - Row 88): {tot_g - b_tot_g}")
print(f"Lệch H (Row 74 - Row 88): {tot_h - b_tot_h}")

# Compare each branch sum in rows 3..72 vs rows 77..87
for r_b in range(77, 88):
    b_name = ws.cell(r_b, 5).value
    g_val = ws.cell(r_b, 7).value
    h_val = ws.cell(r_b, 8).value
    
    # Calculate manual sum from rows 3..72
    staff_g = sum(ws.cell(r, 7).value for r in range(3, 73) if ws.cell(r, 2).value == b_name)
    staff_h = sum(ws.cell(r, 8).value for r in range(3, 73) if ws.cell(r, 2).value == b_name)
    
    diff_g = g_val - staff_g
    diff_h = h_val - staff_h
    print(f"Branch '{b_name}': G={g_val:.2f} (staff={staff_g:.2f}, diff={diff_g:.2f}) | H={h_val:.2f} (staff={staff_h:.2f}, diff={diff_h:.2f})")

wb.close()
