import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

# Staff main branch mapping
STAFF_MAIN_BRANCH_MAP = {
    'Phan Công Vũ Tài': 'Minh Châu',
    'Trịnh Thị Phượng': 'Minh Châu',
    'Hồ Thị Minh Hòa': 'Lê Bình',
    'Nguyễn Trần Ngọc Phương': 'Trường Sa',
    'Phạm Nguyễn Ngọc Quý': 'Trường Sa',
    'Dương Thị Huỳnh Như': 'Nam Hòa',
    'Hoàng Lâm Gia Bảo': 'Lê Bình',
    'Nguyễn Thị Huyền Trang': 'Nam Hòa',
    'Ngô Thị Ngọc Thủy': 'Rạch Bùng Binh',
    'Cù Thị Tường Vy': 'Rạch Bùng Binh',
    'Nguyễn Trí Nghĩa': 'Nguyễn Thị Thập',
    'Võ Ngọc Giàu Sang': 'Nguyễn Văn Quá',
}

wb = openpyxl.load_workbook(r'baocaokpi_thang8_hoanthien.xlsx', data_only=True)
ws = wb['Dự án T8']

staff_totals = {}

for r in range(3, ws.max_row + 1):
    cn = ws.cell(r, 1).value
    name = ws.cell(r, 2).value
    ny3 = ws.cell(r, 5).value or 0
    ck = ws.cell(r, 6).value or 0
    combo = ws.cell(r, 7).value or 0
    
    if name and str(name).strip():
        n_str = str(name).strip()
        main_cn = STAFF_MAIN_BRANCH_MAP.get(n_str, str(cn).strip())
        if n_str not in staff_totals:
            staff_totals[n_str] = {'branch': main_cn, 'ny3': 0, 'ck': 0, 'combo': 0}
        staff_totals[n_str]['ny3'] += ny3
        staff_totals[n_str]['ck'] += ck
        staff_totals[n_str]['combo'] += combo

print("--- TỔNG DOANH SỐ DỰ ÁN SAU KHI GỘP 100% XOAY CA CHO VŨ TÀI & CÁC BẠN ---")
for name in ['Phan Công Vũ Tài', 'Trịnh Thị Phượng', 'Trần Thị Kim Khánh', 'Hoàng Lâm Gia Bảo']:
    d = staff_totals[name]
    ck_day = d['ck'] / 31
    cb_day = d['combo'] / 31
    tot_day = (d['ck'] + d['combo'] + d['ny3']) / 31
    
    # HCM Tier check:
    # 4M: ck >= 1.5M & cb >= 700k
    # 2.8M: ck >= 1.2M & cb >= 650k
    # 2.2M: ck >= 1.0M & cb >= 580k
    # 1.6M: ck >= 850k & cb >= 480k
    # 1.0M: ck >= 630k & cb >= 350k
    tier = 0
    if ck_day >= 1500000 and cb_day >= 700000: tier = 4000000
    elif ck_day >= 1200000 and cb_day >= 650000: tier = 2800000
    elif ck_day >= 1000000 and cb_day >= 580000: tier = 2200000
    elif ck_day >= 850000 and cb_day >= 480000: tier = 1600000
    elif ck_day >= 630000 and cb_day >= 350000: tier = 1000000
    
    print(f"Nhân viên: {name:<20} | Chi nhánh: {d['branch']:<12}")
    print(f"  Tổng CK: {d['ck']:>12,.0f} | CK/ngày: {ck_day:>10,.0f}")
    print(f"  Tổng CB: {d['combo']:>12,.0f} | CB/ngày: {cb_day:>10,.0f}")
    print(f"  => MỨC THƯỞNG DỰ ÁN ĐẠT ĐƯỢC: {tier:>10,.0f} VNĐ\n")
