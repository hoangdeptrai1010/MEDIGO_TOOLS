import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

print("================================================================================")
print("1. THÁNG 7 TARGET - SHEET 'BẢNG LƯƠNG', 'Giờ công', 'Ngày công'")
print("================================================================================")

wb7 = openpyxl.load_workbook(r'thang7\target\BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=False)
wb7_v = openpyxl.load_workbook(r'thang7\target\BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=True)

ws7_bl = wb7['BẢNG LƯƠNG']
ws7_bl_v = wb7_v['BẢNG LƯƠNG']

print("BẢNG LƯƠNG T7 (Row 3 đến 8):")
for r in range(3, 9):
    name = ws7_bl.cell(r, 3).value
    g_f = ws7_bl.cell(r, 7).value # Col G (Số giờ ngày)
    g_v = ws7_bl_v.cell(r, 7).value
    h_f = ws7_bl.cell(r, 8).value # Col H (Số giờ đêm)
    h_v = ws7_bl_v.cell(r, 8).value
    j_f = ws7_bl.cell(r, 10).value # Col J
    j_v = ws7_bl_v.cell(r, 10).value
    k_f = ws7_bl.cell(r, 11).value # Col K
    k_v = ws7_bl_v.cell(r, 11).value
    r_f = ws7_bl.cell(r, 18).value # Col R
    r_v = ws7_bl_v.cell(r, 18).value
    print(f"Row {r} ({name}):")
    print(f"  Col G (Giờ ngày): F={g_f} | V={g_v}")
    print(f"  Col H (Giờ đêm):  F={h_f} | V={h_v}")
    print(f"  Col J (Ngày làm): F={j_f} | V={j_v}")
    print(f"  Col K (Làm đêm):  F={k_f} | V={k_v}")
    print(f"  Col R (Tổng công):F={r_f} | V={r_v}")

if 'Ngày công' in wb7.sheetnames:
    ws7_nc = wb7['Ngày công']
    ws7_nc_v = wb7_v['Ngày công']
    print("\nSheet 'Ngày công' T7 Header row 1 & formulas row 2..5:")
    print("  Header R1:", [ws7_nc.cell(1, c).value for c in range(1, ws7_nc.max_column + 1)])
    for r in range(2, 6):
        print(f"  R{r} F:", [ws7_nc.cell(r, c).value for c in range(1, 16)])
        print(f"  R{r} V:", [ws7_nc_v.cell(r, c).value for c in range(1, 16)])

print("\n================================================================================")
print("2. THÁNG 8 GỐC TEMPLATE - 'BẢNG LƯƠNG THÁNG 8 2026.xlsx'")
print("================================================================================")

wb8_g = openpyxl.load_workbook(r'thang8\BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=False)
wb8_gv = openpyxl.load_workbook(r'thang8\BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=True)

ws8_bl = wb8_g['BẢNG LƯƠNG']
ws8_bl_v = wb8_gv['BẢNG LƯƠNG']

for r in range(3, 9):
    name = ws8_bl.cell(r, 3).value
    g_f = ws8_bl.cell(r, 7).value
    g_v = ws8_bl_v.cell(r, 7).value
    h_f = ws8_bl.cell(r, 8).value
    h_v = ws8_bl_v.cell(r, 8).value
    j_f = ws8_bl.cell(r, 10).value
    j_v = ws8_bl_v.cell(r, 10).value
    k_f = ws8_bl.cell(r, 11).value
    k_v = ws8_bl_v.cell(r, 11).value
    r_f = ws8_bl.cell(r, 18).value
    r_v = ws8_bl_v.cell(r, 18).value
    print(f"Row {r} ({name}):")
    print(f"  Col G (Giờ ngày): F={g_f} | V={g_v}")
    print(f"  Col H (Giờ đêm):  F={h_f} | V={h_v}")
    print(f"  Col J (Ngày làm): F={j_f} | V={j_v}")
    print(f"  Col K (Làm đêm):  F={k_f} | V={k_v}")
    print(f"  Col R (Tổng công):F={r_f} | V={r_v}")

if 'Ngày công' in wb8_g.sheetnames:
    ws8_nc = wb8_g['Ngày công']
    ws8_nc_v = wb8_gv['Ngày công']
    print("\nSheet 'Ngày công' T8 Header row 1 & formulas row 2..5:")
    print("  Header R1:", [ws8_nc.cell(1, c).value for c in range(1, ws8_nc.max_column + 1)])
    for r in range(2, 6):
        print(f"  R{r} F:", [ws8_nc.cell(r, c).value for c in range(1, 16)])
        print(f"  R{r} V:", [ws8_nc_v.cell(r, c).value for c in range(1, 16)])
