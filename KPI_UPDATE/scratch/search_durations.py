import openpyxl

wb = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
for sname in wb.sheetnames:
    ws = wb[sname]
    for r in range(1, min(ws.max_row + 1, 500)):
        for c in range(1, min(ws.max_column + 1, 50)):
            v = str(ws.cell(r, c).value or '')
            if '245h26' in v or '3h1p' in v or '233h17' in v:
                print(f"Found in {sname} row {r}, col {c} ({openpyxl.utils.get_column_letter(c)}): {v}")
