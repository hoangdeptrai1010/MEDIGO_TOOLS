import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'baocaokpi_thang8_hoanthien.xlsx', data_only=True)
ws = wb['Dự án T8']

print("--- Vũ Tài trong sheet Dự án T8 của baocaokpi_thang8_hoanthien.xlsx ---")
for r in range(1, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    if any('Tài' in str(v) for v in vals):
        print(f"Row {r}: {vals[:14]}")
