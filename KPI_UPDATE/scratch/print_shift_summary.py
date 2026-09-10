# -*- coding: utf-8 -*-
import openpyxl
import sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws = wb['BẢNG LƯƠNG']

print(f"{'STT':<3} | {'Chi nhánh':<15} | {'Họ và tên':<24} | {'Chức danh':<8} | {'Giờ ngày(G)':<12} | {'Giờ đêm(H)':<12} | {'Ngày ngày(J)':<12} | {'Ca đêm(K)':<10} | {'Công TT(R)':<10} | {'PC Đêm(X)':<10}")
print("-" * 125)

for r in range(3, 59):
    stt = ws.cell(r, 1).value
    cn = ws.cell(r, 2).value or ''
    name = ws.cell(r, 3).value or ''
    role = ws.cell(r, 4).value or ''
    g = ws.cell(r, 7).value or 0
    h = ws.cell(r, 8).value or 0
    j = ws.cell(r, 10).value or 0
    k = ws.cell(r, 11).value or 0
    r_val = ws.cell(r, 18).value or 0
    x_val = ws.cell(r, 24).value or 0
    
    g_str = f"{g:,.2f}h" if isinstance(g, (int, float)) else str(g)
    h_str = f"{h:,.2f}h" if isinstance(h, (int, float)) else str(h)
    j_str = f"{j}"
    k_str = f"{k}"
    r_str = f"{r_val}"
    x_str = f"{x_val:,.0f}" if isinstance(x_val, (int, float)) else str(x_val)
    
    print(f"{stt:<3} | {cn:<15} | {name:<24} | {role:<8} | {g_str:<12} | {h_str:<12} | {j_str:<12} | {k_str:<10} | {r_str:<10} | {x_str:<10}")
