import openpyxl, subprocess

target_files = ['baocaokpi_thang8_hoanthien.xlsx', 'thang8/baocaokpi_thang8_hoanthien.xlsx']

for target in target_files:
    print(f'=== FIXING ALL UDF AND LINKAGES IN {target} ===')
    wb = openpyxl.load_workbook(target, data_only=False)
    
    # 1. Sheet 'Hot Bill HN' - ensure clean branch name in Col C
    if 'Hot Bill HN' in wb.sheetnames:
        ws_hb = wb['Hot Bill HN']
        for r in range(2, ws_hb.max_row + 1):
            b_val = ws_hb.cell(r, 3).value
            if b_val:
                b_str = str(b_val).strip()
                if 'Hàng Bông' in b_str: ws_hb.cell(r, 3).value = 'Hàng Bông'
                elif 'Đường Láng' in b_str: ws_hb.cell(r, 3).value = 'Đường Láng'
                
    # 2. Sheet 'Dự án T8' - ensure Hot Bill formula matches cleanly and Col M = J + K + L
    if 'Dự án T8' in wb.sheetnames:
        ws_da = wb['Dự án T8']
        for r in range(3, ws_da.max_row + 1):
            nv = ws_da.cell(r, 2).value
            if nv:
                ws_da.cell(r, 12).value = f'=IFERROR(SUMIFS(\'Hot Bill HN\'!$G:$G, \'Hot Bill HN\'!$C:$C, $A{r}, \'Hot Bill HN\'!$D:$D, $B{r}), 0)'
                ws_da.cell(r, 13).value = f'=IFERROR($J{r}+$K{r}+$L{r},0)'
                
    # 3. Sheet 'kpi dược sĩ' - fix Col N = K + M, replace _xludf.IFS and _xludf.SWITCH
    if 'kpi dược sĩ' in wb.sheetnames:
        ws_ds = wb['kpi dược sĩ']
        for r in range(1, ws_ds.max_row + 1):
            for c in range(1, ws_ds.max_column + 1):
                cell = ws_ds.cell(r, c)
                val = cell.value
                if hasattr(val, 'text') and isinstance(val.text, str):
                    val.text = val.text.replace('_xludf.IFS', '_xlfn.IFS').replace('_xludf.SWITCH', '_xlfn.SWITCH')
                elif isinstance(val, str) and val.startswith('='):
                    cell.value = val.replace('_xludf.IFS', '_xlfn.IFS').replace('_xludf.SWITCH', '_xlfn.SWITCH')
                    
        for r in range(3, ws_ds.max_row + 1):
            nv = ws_ds.cell(r, 2).value
            if nv:
                ws_ds.cell(r, 14).value = f'=K{r}+M{r}'

    # 4. Sheet 'kpi nhà thuốc' - replace _xludf.IFS
    if 'kpi nhà thuốc' in wb.sheetnames:
        ws_nt = wb['kpi nhà thuốc']
        for r in range(1, ws_nt.max_row + 1):
            for c in range(1, ws_nt.max_column + 1):
                cell = ws_nt.cell(r, c)
                val = cell.value
                if hasattr(val, 'text') and isinstance(val.text, str):
                    val.text = val.text.replace('_xludf.IFS', '_xlfn.IFS').replace('_xludf.SWITCH', '_xlfn.SWITCH')
                elif isinstance(val, str) and val.startswith('='):
                    cell.value = val.replace('_xludf.IFS', '_xlfn.IFS').replace('_xludf.SWITCH', '_xlfn.SWITCH')
                    
    wb.save(target)
    wb.close()
    print(f'Saved fixed workbook: {target}')

    # 5. Run recalc_workbook.ps1 to calculate and cache 100% values
    res = subprocess.run(['powershell', '-ExecutionPolicy', 'Bypass', '-File', 'recalc_workbook.ps1', '-FilePath', target], capture_output=True, text=True)
    success = 'successfully' in res.stdout
    print(f'Recalculation on {target}: {"SUCCESS" if success else "FAILED"}')
    if not success:
        for line in res.stdout.splitlines():
            if 'Error' in line:
                print('   ', line)
