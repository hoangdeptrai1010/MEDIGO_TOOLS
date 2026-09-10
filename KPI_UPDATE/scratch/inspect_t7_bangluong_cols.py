import openpyxl

wb = openpyxl.load_workbook('thang7/BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=False)

for name in ['BẢNG LƯƠNG', 'Bảng đánh giá']:
    ws = wb[name]
    print(f"\n=== SHEET: {name} ===")
    for r in range(1, 4):
        cols = [f"Col {c} ({ws.cell(r, c).coordinate}): {ws.cell(r, c).value}" for c in range(1, min(45, ws.max_column + 1)) if ws.cell(r, c).value is not None]
        print(f"Row {r}:")
        for x in cols:
            print("  ", x)
