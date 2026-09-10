import openpyxl

wb = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
ws = wb['Bảng chi tiết chấm công']

print("Row 2 headers:")
for c in range(1, ws.max_column + 1):
    v2 = ws.cell(2, c).value
    v3 = ws.cell(3, c).value
    if v2 is not None or v3 is not None:
        print(f"Col {c} ({openpyxl.utils.get_column_letter(c)}): r2={repr(v2)}, r3={repr(v3)}")

print("\nRow 4-10 samples:")
for r in range(4, 11):
    row_vals = {openpyxl.utils.get_column_letter(c): ws.cell(r, c).value for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None}
    print(f"Row {r}: {row_vals}")
