import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

def inspect_wb(path, label):
    print("=" * 80)
    print(f"INSPECTING: {label} ({path})")
    print("=" * 80)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        print("Sheet names:", wb.sheetnames)
        for sname in wb.sheetnames:
            if any(k in sname.lower() for k in ['ck', 'dự án', 'du an', 'thưởng', 'lương', 'kpi', 'tổng hợp', 'tong hop']):
                ws = wb[sname]
                print(f"\n--- Sheet: {sname} (max_row={ws.max_row}, max_col={ws.max_column}) ---")
                for r in range(1, min(6, ws.max_row + 1)):
                    row_vals = [str(ws.cell(r, c).value) if ws.cell(r, c).value is not None else "" for c in range(1, min(20, ws.max_column + 1))]
                    print(f"Row {r}: {row_vals}")
    except Exception as e:
        print(f"Error reading {path}: {e}")

inspect_wb(r'goc\NHÀ THUỐC THÁNG 7 2026.xlsx', 'GỐC NHÀ THUỐC T7')
inspect_wb(r'thang7\target\BẢNG LƯƠNG THÁNG 7 2026.xlsx', 'TARGET BẢNG LƯƠNG T7')
inspect_wb(r'baocaokpi_thang7_hoanthien.xlsx', 'BAOCAOKPI T7 HOÀN THIỆN')

inspect_wb(r'goc\NHÀ THUỐC THÁNG 8 2026.xlsx', 'GỐC NHÀ THUỐC T8')
inspect_wb(r'thang8\BẢNG LƯƠNG THÁNG 8 2026.xlsx', 'BẢNG LƯƠNG T8 MẪU')
inspect_wb(r'baocaokpi_thang8_hoanthien.xlsx', 'BAOCAOKPI T8 HOÀN THIỆN')
inspect_wb(r'thang8\bangluong_thang8_hoanthien.xlsx', 'BẢNG LƯƠNG T8 HOÀN THIỆN')
