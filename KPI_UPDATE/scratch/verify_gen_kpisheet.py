import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

f_gen = r"d:\MEDIGO\KPI_UPDATE\TOOL_KPISHEET\output\NHÀ THUỐC THÁNG 9 2026.xlsx"
wb = openpyxl.load_workbook(f_gen, data_only=False)

print("Generated sheets:", wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n=======================================================")
    print(f"SHEET: '{sname}' ({ws.max_row} rows, {ws.max_column} cols)")
    print(f"=======================================================")
    for r in range(1, min(6, ws.max_row + 1)):
        row_vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
        print(f"Row {r:2d}: {row_vals}")
