import openpyxl
import os, sys
sys.stdout.reconfigure(encoding='utf-8')

target = 'Hoàng Lâm Gia Bảo'

def search_wb(file_path):
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return
    print(f"\n=======================================================")
    print(f"FILE: {os.path.basename(file_path)}")
    print(f"=======================================================")
    wb = openpyxl.load_workbook(file_path, data_only=True)
    for s_name in wb.sheetnames:
        ws = wb[s_name]
        matches = []
        for r in range(1, ws.max_row + 1):
            row_vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 40))]
            row_str = ' '.join([str(v or '') for v in row_vals])
            if 'Gia Bảo' in row_str or 'Gia Bao' in row_str or 'Hoàng Lâm' in row_str or 'Hoang Lam' in row_str:
                header_vals = [ws.cell(1, c).value or ws.cell(2, c).value for c in range(1, min(ws.max_column + 1, 40))]
                col_data = [(header_vals[i] if i < len(header_vals) else f'Col{i+1}', val) for i, val in enumerate(row_vals) if val is not None]
                matches.append((r, col_data))
        if matches:
            print(f"\n--- Sheet: [{s_name}] ({len(matches)} rows matched) ---")
            for r, c_data in matches:
                print(f"  Row {r}:")
                for h, v in c_data:
                    print(f"     • {h}: {v}")

search_wb('d:/MEDIGO/KPI_UPDATE/thang8/FIXFORMAT/Bang luong tong T8-2026 - CNT.xlsx')
search_wb('d:/MEDIGO/KPI_UPDATE/thang8/bangluong_thang8_hoanthien.xlsx')
search_wb('d:/MEDIGO/KPI_UPDATE/thang8/DATA/BangChiTietChamCong_thang8.xlsx')
search_wb('d:/MEDIGO/KPI_UPDATE/thang8/DATA/TongSoPhatViPhamVaThuongNgay_thang8.xlsx')
