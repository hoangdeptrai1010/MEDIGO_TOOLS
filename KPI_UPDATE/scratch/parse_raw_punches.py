import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb_raw = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
ws_raw = wb_raw.active

print("=== SO SÁNH GIỜ LÀM VIỆC TỪ FILE CHẤM CÔNG GỐC CHO TÂM VÀ THOA ===")

# Compute exact hours from BangChiTietChamCong_thang8.xlsx
for r in [202, 203, 204, 205, 206, 207, 208, 209, 210, 211]:
    staff = ws_raw.cell(r, 3).value
    shift = ws_raw.cell(r, 7).value
    
    # Calculate hours for each day
    tot_hrs = 0.0
    valid_shifts_cnt = 0
    days_detail = []
    for day in range(1, 32):
        col_in = 8 + (day-1)*2
        col_out = col_in + 1
        v_in = str(ws_raw.cell(r, col_in).value or '').strip()
        v_out = str(ws_raw.cell(r, col_out).value or '').strip()
        if v_in and v_out:
            try:
                # parse HH:MM
                h_in, m_in = map(int, v_in.split(':'))
                h_out, m_out = map(int, v_out.split(':'))
                t_in = h_in * 60 + m_in
                t_out = h_out * 60 + m_out
                if t_out < t_in: # overnight
                    t_out += 24 * 60
                dur_m = t_out - t_in
                dur_h = dur_m / 60.0
                tot_hrs += dur_h
                if dur_h > 4:
                    valid_shifts_cnt += 1
                days_detail.append((day, v_in, v_out, dur_h))
            except Exception as e:
                pass
    print(f"\nRow {r:3d} | Nhân viên: {staff} | Ca: {shift}:")
    print(f"  -> Tổng số ngày quẹt: {len(days_detail)}, Số ca > 4h: {valid_shifts_cnt}, Tổng giờ: {tot_hrs:.2f}h")
