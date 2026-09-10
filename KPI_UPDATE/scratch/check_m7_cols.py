import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

print("================================================================================")
print("THÁNG 7: SO SÁNH THƯỞNG CK & DỰ ÁN GIỮA TARGET VÀ HOÀN THIỆN")
print("================================================================================")

wb7_target = openpyxl.load_workbook(r'thang7\target\BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=True)
ws7_bl = wb7_target['BẢNG LƯƠNG']

print("--- [TARGET BẢNG LƯƠNG THÁNG 7] Các cột Thưởng CK và Dự án ---")
# Let's find columns in header row 2
for c in range(1, ws7_bl.max_column + 1):
    header = ws7_bl.cell(2, c).value
    if header:
        print(f"Col {openpyxl.utils.get_column_letter(c)} ({c}): {str(header).strip()}")

print("\n--- In 15 dòng đầu của BẢNG LƯƠNG T7 Target ---")
for r in range(3, 18):
    name = ws7_bl.cell(r, 3).value
    # Let's inspect relevant columns: KPI (AD), Thưởng CK (AF), Thưởng Dự án (AH), Thực lĩnh (AR)
    # Check headers
    row_vals = {openpyxl.utils.get_column_letter(c): ws7_bl.cell(r, c).value for c in range(1, ws7_bl.max_column + 1) if ws7_bl.cell(r, c).value is not None}
    print(f"Row {r} ({name}): {row_vals}")

