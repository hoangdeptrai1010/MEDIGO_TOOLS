import openpyxl
import subprocess
import os

wb_orig = openpyxl.load_workbook('goc/NHÀ THUỐC THÁNG 7 2026.xlsx')
wb_kpi = openpyxl.load_workbook('baocaokpi_thang7_hoanthien.xlsx')

ws_kpi_ds = wb_kpi['kpi dược sĩ']

# Test replacing ws_ds with original template ws_ds
wb_test = openpyxl.load_workbook('baocaokpi_thang7_hoanthien.xlsx')
wb_test.remove(wb_test['kpi dược sĩ'])
# Copy original kpi dược sĩ from goc
ws_orig_ds = wb_orig['kpi dược sĩ']
# Create new sheet and copy cells
ws_new = wb_test.create_sheet('kpi dược sĩ', 1)
for (r, c), cell in ws_orig_ds._cells.items():
    ws_new.cell(r, c, cell.value)

wb_test.save('test_with_orig_ds.xlsx')
wb_test.close()

p = os.path.abspath('test_with_orig_ds.xlsx')
res = subprocess.run(['powershell', '-ExecutionPolicy', 'Bypass', '-File', 'recalc_workbook.ps1', '-FilePath', p], capture_output=True, text=True)
print('Test with orig ds return:', res.returncode)
print('Output:', res.stdout.strip()[:140])
