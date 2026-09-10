import openpyxl
import sys
from collections import defaultdict

sys.stdout.reconfigure(encoding='utf-8')

wb_formula = openpyxl.load_workbook('thang8/BANGLUONGTHANG8.xlsx', data_only=False)
wb_data = openpyxl.load_workbook('thang8/BANGLUONGTHANG8.xlsx', data_only=True)

ws_nc_f = wb_formula['Ngày công']
ws_nc_d = wb_data['Ngày công']
ws_gc_f = wb_formula['Giờ công']
ws_gc_d = wb_data['Giờ công']
ws_bl_f = wb_formula['BẢNG LƯƠNG']
ws_bl_d = wb_data['BẢNG LƯƠNG']

print("--- Formulas in Ngày công summary (Row 2 & 3) ---")
for c in range(11, 16):
    print(f"Col {c} ({ws_nc_f.cell(1, c).value}): formula={ws_nc_f.cell(2, c).value} | val={ws_nc_d.cell(2, c).value}")
for c in range(11, 16):
    print(f"Col {c} ({ws_nc_f.cell(1, c).value}): formula={ws_nc_f.cell(3, c).value} | val={ws_nc_d.cell(3, c).value}")

print("\n--- Formulas in Giờ công summary (Row 2 & 3) ---")
for c in range(7, 12):
    print(f"Col {c} ({ws_gc_f.cell(1, c).value}): formula={ws_gc_f.cell(2, c).value} | val={ws_gc_d.cell(2, c).value}")
for c in range(7, 12):
    print(f"Col {c} ({ws_gc_f.cell(1, c).value}): formula={ws_gc_f.cell(3, c).value} | val={ws_gc_d.cell(3, c).value}")

# Check raw hours by person and branch in Ngày công (cols A, B, F)
raw_nc_hours_by_staff = defaultdict(float)
raw_nc_hours_by_staff_branch = defaultdict(float)
raw_nc_days_by_staff = defaultdict(set)
raw_nc_days_by_staff_branch = defaultdict(set)

for r in range(3, ws_nc_d.max_row + 1):
    branch = ws_nc_d.cell(r, 1).value
    staff = ws_nc_d.cell(r, 2).value
    date_val = ws_nc_d.cell(r, 3).value
    shift = ws_nc_d.cell(r, 4).value
    hours = ws_nc_d.cell(r, 6).value or 0
    if staff and str(staff).strip():
        staff = str(staff).strip()
        branch = str(branch).strip() if branch else ''
        try:
            h = float(hours)
        except:
            h = 0.0
        raw_nc_hours_by_staff[staff] += h
        raw_nc_hours_by_staff_branch[(branch, staff)] += h
        if h > 0 and date_val:
            raw_nc_days_by_staff[staff].add(date_val)
            raw_nc_days_by_staff_branch[(branch, staff)].add(date_val)

# Summary table in Ngày công (cols K, L, M, N)
summary_nc_hours_by_staff = defaultdict(float)
summary_nc_days_by_staff = defaultdict(float)
summary_nc_pairs = []

for r in range(2, ws_nc_d.max_row + 1):
    branch = ws_nc_d.cell(r, 11).value
    staff = ws_nc_d.cell(r, 12).value
    days = ws_nc_d.cell(r, 13).value or 0
    hours = ws_nc_d.cell(r, 14).value or 0
    if staff and str(staff).strip():
        staff = str(staff).strip()
        branch = str(branch).strip() if branch else ''
        summary_nc_pairs.append((r, branch, staff))
        try:
            summary_nc_days_by_staff[staff] += float(days)
        except:
            pass
        try:
            summary_nc_hours_by_staff[staff] += float(hours)
        except:
            pass

# Summary table in Giờ công (cols G, H, I, J, K)
summary_gc_hours_by_staff = defaultdict(float)
for r in range(2, ws_gc_d.max_row + 1):
    branch = ws_gc_d.cell(r, 7).value
    staff = ws_gc_d.cell(r, 8).value
    day_h = ws_gc_d.cell(r, 9).value or 0
    night_h = ws_gc_d.cell(r, 10).value or 0
    tot_h = ws_gc_d.cell(r, 11).value or 0
    if staff and str(staff).strip():
        staff = str(staff).strip()
        try:
            summary_gc_hours_by_staff[staff] += float(tot_h)
        except:
            pass

print("\n=== SO SÁNH GIỜ CÔNG & NGÀY CÔNG (RAW vs SUMMARY NC vs SUMMARY GC) ===")
all_staff = sorted(list(set(list(raw_nc_hours_by_staff.keys()) + list(summary_nc_hours_by_staff.keys()))))
diff_count = 0
for s in all_staff:
    raw_h = round(raw_nc_hours_by_staff[s], 2)
    snc_h = round(summary_nc_hours_by_staff[s], 2)
    sgc_h = round(summary_gc_hours_by_staff[s], 2)
    raw_d = len(raw_nc_days_by_staff[s])
    snc_d = round(summary_nc_days_by_staff[s], 2)
    
    diff_h = round(raw_h - snc_h, 2)
    diff_d = round(raw_d - snc_d, 2)
    if abs(diff_h) > 0.05 or abs(diff_d) > 0.05:
        diff_count += 1
        print(f"NV: {s:<30} | Raw H: {raw_h:>7.2f} vs NC_Sum H: {snc_h:>7.2f} (diff={diff_h:>6.2f}) | GC_Sum H: {sgc_h:>7.2f} | Raw Days: {raw_d:>2} vs NC Days: {snc_d:>5.1f} (diff={diff_d:>4.1f})")

print(f"\nTổng số nhân viên có chênh lệch: {diff_count}")
