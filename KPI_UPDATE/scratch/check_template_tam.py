import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/tinhcongnhungthuongchia.xlsx', data_only=True)
ws_nc = wb['Ngày công']

print("=== TẤT CẢ DÒNG CỦA NGUYỄN THỊ TÂM TRONG SHEET NGÀY CÔNG (tinhcongnhungthuongchia.xlsx) ===")
for r in range(2, ws_nc.max_row+1):
    name = str(ws_nc.cell(r, 2).value or '')
    if 'tâm' in name.lower() and 'đoàn' not in name.lower():
        cn = ws_nc.cell(r, 1).value
        dt = ws_nc.cell(r, 3).value
        ca = ws_nc.cell(r, 4).value
        txt = ws_nc.cell(r, 5).value
        hrs = ws_nc.cell(r, 6).value
        print(f"Row {r:4d} | Ngày: {str(dt)[:10]} | Ca: {str(ca):<40} | Giờ: {hrs}")
