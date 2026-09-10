import sys
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

wb7 = openpyxl.load_workbook('thang7/BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=False)
ws7 = wb7['BẢNG LƯƠNG']

print('=== THÁNG 7 TEMPLATE ===')
for c in range(5, 12):
    col_let = openpyxl.utils.get_column_letter(c)
    h = str(ws7.cell(2, c).value or '').replace('\n', ' ')
    f = str(ws7.cell(3, c).value or '')
    print(f'Col {c:2d} ({col_let:>2s}): Header = {h:<35s} | Formula = {f}')

wb8 = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=False)
ws8 = wb8['BẢNG LƯƠNG']
print('\n=== THÁNG 8 TEMPLATE ===')
for c in range(5, 12):
    col_let = openpyxl.utils.get_column_letter(c)
    h = str(ws8.cell(2, c).value or '').replace('\n', ' ')
    f = str(ws8.cell(3, c).value or '')
    print(f'Col {c:2d} ({col_let:>2s}): Header = {h:<35s} | Formula = {f}')
