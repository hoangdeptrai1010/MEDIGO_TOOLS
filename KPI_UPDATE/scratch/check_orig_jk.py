# -*- coding: utf-8 -*-
import openpyxl, os, sys
sys.stdout.reconfigure(encoding='utf-8')

orig_path = 'thang8/Pharmacy_retail_Store_KPIs_August 2026/BẢNG LƯƠNG/BẢNG TỔNG HỢP LƯƠNG T08.2026.xlsx'
if os.path.exists(orig_path):
    wb = openpyxl.load_workbook(orig_path, data_only=False)
    for s in wb.sheetnames:
        if 'LƯƠNG' in s.upper():
            ws = wb[s]
            print(f"=== {orig_path} [{s}] ===")
            for r in range(2, 10):
                print(f"Row {r:2d}: Name={ws.cell(r,3).value} | J{r}={ws.cell(r,10).value} | K{r}={ws.cell(r,11).value} | R{r}={ws.cell(r,18).value}")

july_files = [f for f in os.listdir('thang7') if 'bangluong' in f.lower() or 'luong' in f.lower()]
for jf in july_files:
    p = os.path.join('thang7', jf)
    try:
        wb = openpyxl.load_workbook(p, data_only=False)
        for s in wb.sheetnames:
            if 'LƯƠNG' in s.upper():
                ws = wb[s]
                print(f"\n=== July file {p} [{s}] ===")
                for r in range(2, 10):
                    print(f"Row {r:2d}: Name={ws.cell(r,3).value} | J{r}={ws.cell(r,10).value} | K{r}={ws.cell(r,11).value} | R{r}={ws.cell(r,18).value}")
    except Exception as e:
        print(f"Error {p}: {e}")
