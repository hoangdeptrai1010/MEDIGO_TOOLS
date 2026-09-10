import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_bangluong = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final.xlsx")
wb = openpyxl.load_workbook(f_bangluong, data_only=True)
ws_med = wb["Med"]

print("=== Med rows > 88 ===")
for r in range(89, ws_med.max_row + 1):
    vals = [ws_med.cell(r, c).value for c in range(1, 15)]
    if any(v is not None for v in vals):
        name = ws_med.cell(r, 4).value
        stk = ws_med.cell(r, 9).value
        bank = ws_med.cell(r, 10).value
        dept = ws_med.cell(r, 6).value
        print(f"Row {r:3d} | Name: {str(name):<30} | Dept: {str(dept):<20} | STK: {str(stk):<20} | Bank: {str(bank):<25}")
