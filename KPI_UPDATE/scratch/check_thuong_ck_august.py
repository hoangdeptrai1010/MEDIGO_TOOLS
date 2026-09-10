import openpyxl

for p in ['thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', 'thang8/bangluong_thang8_hoanthien.xlsx']:
    wb = openpyxl.load_workbook(p, data_only=True)
    if 'Thưởng CK' in wb.sheetnames:
        ws = wb['Thưởng CK']
        print(f"=== {p} -> Thưởng CK ===")
        for r in range(1, 10):
            print([ws.cell(r, c).value for c in range(1, 12) if ws.cell(r, c).value is not None])
