import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('goc/NHÀ THUỐC THÁNG 8 2026.xlsx', data_only=False)

print("=== 1. KIỂM TRA SHEET 'DỰ ÁN T8' ===")
ws_da = wb['Dự án T8']
print(f"Cell A1: {ws_da.cell(1, 1).value}")
for c in range(1, 14):
    print(f"Col {openpyxl.utils.get_column_letter(c)}: {ws_da.cell(2, c).value} | Row 3 formula: {ws_da.cell(3, c).value}")

print("\n=== 2. KIỂM TRA SHEET 'KPI DƯỢC SĨ' ===")
ws_ds = wb['kpi dược sĩ']
print(f"Cell C1: {ws_ds.cell(1, 3).value}")
for c in range(1, 15):
    print(f"Col {openpyxl.utils.get_column_letter(c)}: {ws_ds.cell(2, c).value} | Row 3 formula: {ws_ds.cell(3, c).value}")

wb.close()
