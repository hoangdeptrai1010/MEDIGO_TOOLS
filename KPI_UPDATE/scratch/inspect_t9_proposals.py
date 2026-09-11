import openpyxl
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

hn_path = 'thang9/data/kpi_thang9/e_xuat_KPI_Quy_3.26.xlsx'
wb_hn = openpyxl.load_workbook(hn_path, data_only=True)
ws_t9 = wb_hn['T9.26']

print("=== CHI TIẾT SHEET T9.26 FILE ĐỀ XUẤT HÀ NỘI ===")
for r in range(1, ws_t9.max_row + 1):
    vals = [ws_t9.cell(r, c).value for c in range(1, ws_t9.max_column + 1)]
    if any(v is not None for v in vals):
        print(f"Row {r:02d}: {vals}")

hcm_path = 'thang9/data/kpi_thang9/KPI CNT HCM Tháng 09.xlsx'
wb_hcm = openpyxl.load_workbook(hcm_path, data_only=True)
ws_hcm = wb_hcm['KPI tháng 09']

print("\n=== CHI TIẾT SHEET KPI THÁNG 09 FILE ĐỀ XUẤT HCM ===")
for r in range(1, ws_hcm.max_row + 1):
    vals = [ws_hcm.cell(r, c).value for c in range(1, 10)]
    if any(v is not None for v in vals):
        print(f"Row {r:02d}: {vals}")
