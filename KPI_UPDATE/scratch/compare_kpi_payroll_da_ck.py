# -*- coding: utf-8 -*-
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb_kpi = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx', data_only=True)
wb_bl = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)

print("=== SHEETS IN baocaokpi_thang8_hoanthien.xlsx ===")
print(wb_kpi.sheetnames)

print("\n=== SHEETS IN thang8/bangluong_thang8_hoanthien.xlsx ===")
print(wb_bl.sheetnames)

ws_da_kpi = wb_kpi['Dự án T8'] if 'Dự án T8' in wb_kpi.sheetnames else None
ws_bl = wb_bl['BẢNG LƯƠNG']
ws_da_bl = wb_bl['Dự án'] if 'Dự án' in wb_bl.sheetnames else None
ws_ck_bl = wb_bl['Thưởng CK'] if 'Thưởng CK' in wb_bl.sheetnames else None

# Check KPI sheet Dự án T8
da_kpi_dict = {}
if ws_da_kpi:
    for r in range(3, ws_da_kpi.max_row + 1):
        name = ws_da_kpi.cell(r, 2).value
        tot_da = ws_da_kpi.cell(r, 13).value # Total dự án
        da_tier = ws_da_kpi.cell(r, 10).value # Thưởng dự án
        da_them = ws_da_kpi.cell(r, 11).value # Thưởng thêm
        da_hot = ws_da_kpi.cell(r, 12).value # Hot bill
        if name:
            da_kpi_dict[str(name).strip()] = {
                'total': tot_da,
                'tier': da_tier,
                'them': da_them,
                'hot': da_hot
            }

print("\n=== COMPARISON: BÁO CÁO KPI (Dự án T8) vs BẢNG LƯƠNG (Sheet Dự án & Cột AD/AF) ===")
print(f"{'STT':<3} | {'Chi nhánh':<15} | {'Họ và tên':<24} | {'Dự án (KPI)':<14} | {'Dự án (BL sheet)':<18} | {'Cột AD (BL)':<14} | {'Cột AF (Thưởng CK)':<18} | {'Khớp DA?'}")
print("-" * 135)

for r in range(3, 59):
    stt = ws_bl.cell(r, 1).value
    cn = ws_bl.cell(r, 2).value or ''
    name = str(ws_bl.cell(r, 3).value or '').strip()
    ad_val = ws_bl.cell(r, 30).value or 0 # Cột AD (Dự án)
    af_val = ws_bl.cell(r, 32).value or 0 # Cột AF (Thưởng CK)
    
    kpi_da = da_kpi_dict.get(name, {}).get('total', 0)
    
    # In sheet Dự án of bangluong
    bl_sheet_da = 0
    if ws_da_bl:
        for r_da in range(2, ws_da_bl.max_row + 1):
            if str(ws_da_bl.cell(r_da, 2).value or '').strip() == name:
                bl_sheet_da = ws_da_bl.cell(r_da, 3).value or 0
                break
    
    match_str = "MATCH" if (abs((kpi_da or 0) - (ad_val or 0)) < 1 and abs((bl_sheet_da or 0) - (ad_val or 0)) < 1) else "DIFF!"
    
    kpi_str = f"{kpi_da:,.0f}" if isinstance(kpi_da, (int, float)) else str(kpi_da)
    bl_sheet_str = f"{bl_sheet_da:,.0f}" if isinstance(bl_sheet_da, (int, float)) else str(bl_sheet_da)
    ad_str = f"{ad_val:,.0f}" if isinstance(ad_val, (int, float)) else str(ad_val)
    af_str = f"{af_val:,.0f}" if isinstance(af_val, (int, float)) else str(af_val)
    
    print(f"{stt:<3} | {cn:<15} | {name:<24} | {kpi_str:<14} | {bl_sheet_str:<18} | {ad_str:<14} | {af_str:<18} | {match_str}")
