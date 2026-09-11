import openpyxl
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

b_path = 'goc/NHÀ THUỐC THÁNG 7 2026.xlsx'
wb_b = openpyxl.load_workbook(b_path, data_only=False)
ws_ds = wb_b['kpi dược sĩ']

print("=== CÔNG THỨC CHÍNH XÁC TỪNG CỘT TRONG SHEET 'kpi dược sĩ' (THÁNG 7) ===")

for r in range(3, 10):
    b = ws_ds.cell(r, 1).value
    s = ws_ds.cell(r, 2).value
    print(f"\n--- DÒNG {r}: {b} - {s} ---")
    for c in range(1, 33):
        col_letter = openpyxl.utils.get_column_letter(c)
        val = ws_ds.cell(r, c).value
        header = ws_ds.cell(2, c).value
        f_text = getattr(val, 'text', str(val))
        print(f"  Cột {c:02d} ({col_letter} - {header}): {f_text}")
