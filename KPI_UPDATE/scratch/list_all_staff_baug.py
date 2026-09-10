import openpyxl

wb_aug = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=True)
ws_bl = wb_aug['BẢNG LƯƠNG']

print("=== thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx: All staff rows ===")
for r in range(3, ws_bl.max_row + 1):
    stt = ws_bl.cell(r, 1).value
    cn = ws_bl.cell(r, 2).value
    name = ws_bl.cell(r, 3).value
    role = ws_bl.cell(r, 4).value
    if stt is not None and isinstance(stt, (int, float)):
        print(f"R{r}: STT={stt}, CN={cn}, Name={name}, Role={role}")
