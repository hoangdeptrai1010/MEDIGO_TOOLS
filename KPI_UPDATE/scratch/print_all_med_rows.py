import openpyxl
import os
import sys
import unicodedata
import re

sys.stdout.reconfigure(encoding='utf-8')

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_bl = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final.xlsx")
wb_bl = openpyxl.load_workbook(f_bl, data_only=True)
ws_med = wb_bl["Med"]

print("=== ALL ROWS IN Med WITH STK AND BANK ===")
for r in range(10, ws_med.max_row + 1):
    name = ws_med.cell(r, 4).value
    code = ws_med.cell(r, 3).value
    stk = ws_med.cell(r, 9).value
    bank = ws_med.cell(r, 10).value
    store = ws_med.cell(r, 6).value
    
    if name is not None and str(name).strip() != "" and not str(name).strip().startswith("CHUỖI") and not str(name).strip().startswith("Tổng"):
        print(f"Row {r:3d} | Name: {str(name):<28} | Store: {str(store):<18} | STK: {str(stk):<22} | Bank: {str(bank):<25}")
