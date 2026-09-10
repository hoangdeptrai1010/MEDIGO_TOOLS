import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_bl = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final_backup.xlsx")
wb = openpyxl.load_workbook(f_bl, data_only=True)
ws_med = wb["Med"]

for r in range(10, 95):
    stt = ws_med.cell(r, 1).value
    code = ws_med.cell(r, 3).value
    name = ws_med.cell(r, 4).value
    store = ws_med.cell(r, 6).value
    print(f"Row {r:2d} | STT: {str(stt):<10} | Name: {str(name):<28} | Store: {str(store):<25}")
