import openpyxl, os, sys
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

def audit_and_rebuild_sheets(timecard_path, template_path, output_path):
    print("=== ĐỌC DỮ LIỆU TỪ MÁY CHẤM CÔNG ===")
    wb_tc = openpyxl.load_workbook(timecard_path, data_only=True)
    ws_tc = wb_tc.active

    BRANCH_ALIAS = {
        'NT 24H Trường Sa': 'Trường Sa',
        'NT 24h Đỗ Quang Đẩu': 'Đỗ Quang Đẩu',
        'NT 24H Nam Hòa': 'Nam Hòa',
        'NT 24H Minh Châu 1': 'Minh Châu',
        'NT 24H Lê Bình': 'Lê Bình',
        'NT 24H Nguyễn Chí Thanh': 'Nguyễn Chí Thanh',
        'NT 24H Nguyễn Thị Thập': 'Nguyễn Thị Thập',
        'NT 24H Nguyễn Văn Quá': 'Nguyễn Văn Quá',
        'NT 24H Rạch Bùng Binh': 'Rạch Bùng Binh',
        'NT Đường Láng 247': 'Đường Láng',
        'NT Hàng Bông 247': 'Hàng Bông',
    }

    # List of all punch shifts across all staff
    # Each shift: (branch, name, date, shift_name, dur_str, dur_hours, is_night, gt4)
    all_shifts = []
    # Dict of summary per (branch, name, shift_name) -> total_hours, formatted_str
    shift_summary_dict = {}

    curr_staff = ""
    curr_branch = ""
    curr_role = ""
    curr_code = ""

    for r in range(5, ws_tc.max_row + 1):
        c_name = ws_tc.cell(r, 3).value
        c_branch = ws_tc.cell(r, 6).value
        c_shift = ws_tc.cell(r, 7).value
        c_role = ws_tc.cell(r, 5).value
        c_code = ws_tc.cell(r, 2).value

        if c_name and str(c_name).strip():
            curr_staff = str(c_name).strip()
        if c_branch and str(c_branch).strip():
            curr_branch = BRANCH_ALIAS.get(str(c_branch).strip(), str(c_branch).strip())
        if c_role and str(c_role).strip():
            curr_role = str(c_role).strip()
        if c_code and str(c_code).strip():
            curr_code = str(c_code).strip()

        if not curr_staff or not c_shift:
            continue

        shift_name = str(c_shift).strip()

        for day in range(1, 32):
            col_in = 8 + (day - 1) * 2
            col_out = col_in + 1
            v_in = str(ws_tc.cell(r, col_in).value or '').strip()
            v_out = str(ws_tc.cell(r, col_out).value or '').strip()

            if v_in and v_out:
                try:
                    p_in = v_in.split(' ')[0].split(':')
                    p_out = v_out.split(' ')[0].split(':')
                    h_in, m_in = int(p_in[0]), int(p_in[1])
                    h_out, m_out = int(p_out[0]), int(p_out[1])

                    t_in = h_in * 60 + m_in
                    t_out = h_out * 60 + m_out

                    is_overnight = False
                    if 'đêm' in shift_name.lower() or 'qua đêm' in shift_name.lower() or t_out < t_in:
                        is_overnight = True
                        if t_out < t_in:
                            t_out += 24 * 60
                    elif t_out < t_in:
                        t_out += 24 * 60

                    dur_m = t_out - t_in
                    dur_h = round(dur_m / 60.0, 2)
                    dur_str = f"{dur_m // 60}h{dur_m % 60}p"

                    is_night = ('đêm' in shift_name.lower() or 'qua đêm' in shift_name.lower() or h_in >= 22 or h_in < 6)
                    gt4 = (dur_h > 4.0)

                    dt_val = datetime(2026, 8, day)

                    all_shifts.append({
                        'branch': curr_branch,
                        'name': curr_staff,
                        'date': dt_val,
                        'day_num': day,
                        'shift_name': shift_name,
                        'dur_str': dur_str,
                        'dur_h': dur_h,
                        'is_night': is_night,
                        'gt4': gt4
                    })

                    sum_key = (curr_branch, curr_staff, shift_name)
                    if sum_key not in shift_summary_dict:
                        shift_summary_dict[sum_key] = 0.0
                    shift_summary_dict[sum_key] += dur_h

                except Exception as e:
                    pass
    wb_tc.close()

    print(f"--> Tổng số ca chấm công chi tiết tìm thấy: {len(all_shifts)}")
    print(f"--> Tổng số nhóm ca tổng hợp trong Giờ công: {len(shift_summary_dict)}")

    # Load template
    wb = openpyxl.load_workbook(template_path, data_only=False)

    # 1. BẢNG LƯƠNG staff rows
    ws_bl = wb['BẢNG LƯƠNG']
    staff_rows = []
    for r in range(3, 73):
        cn = ws_bl.cell(r, 2).value
        name = ws_bl.cell(r, 3).value
        role = ws_bl.cell(r, 4).value
        if name and str(name).strip() and str(role).strip() != 'Tổng':
            staff_rows.append((r, str(cn).strip(), str(name).strip(), str(role).strip()))

    print(f"--> Số dòng nhân sự trong BẢNG LƯƠNG: {len(staff_rows)}")

    # 2. Xây dựng lại sheet 'Giờ công'
    ws_gc = wb['Giờ công']
    # Clear old detail rows from row 2 down in A:E
    for r in range(2, ws_gc.max_row + 1):
        for c in range(1, 6):
            ws_gc.cell(r, c).value = None

    # Write new detail rows into A:E
    row_idx = 2
    for (br, name, sh_name), tot_h in shift_summary_dict.items():
        tot_m = int(round(tot_h * 60))
        tot_str = f"{tot_m // 60}h{tot_m % 60}p"
        ws_gc.cell(row_idx, 1, br)
        ws_gc.cell(row_idx, 2, name)
        ws_gc.cell(row_idx, 3, sh_name)
        ws_gc.cell(row_idx, 4, tot_str)
        ws_gc.cell(row_idx, 5, tot_h)
        row_idx += 1

    # Clear and rebuild summary table in G:K
    for r in range(2, max(ws_gc.max_row + 1, 100)):
        for c in range(7, 12):
            ws_gc.cell(r, c).value = None

    ws_gc.cell(1, 7, "Chi nhánh")
    ws_gc.cell(1, 8, "Tên Nhân viên")
    ws_gc.cell(1, 9, "Giờ ca ngày")
    ws_gc.cell(1, 10, "Giờ ca đêm")
    ws_gc.cell(1, 11, "Tổng giờ")

    for idx, (bl_r, br, name, role) in enumerate(staff_rows, start=2):
        ws_gc.cell(idx, 7, br)
        ws_gc.cell(idx, 8, name)
        ws_gc.cell(idx, 9, f"=SUMIFS(E:E, B:B, H{idx}, A:A, G{idx}) - J{idx}")
        ws_gc.cell(idx, 10, f"=SUMIFS(E:E, B:B, H{idx}, C:C, \"*đêm*\", A:A, G{idx})")
        ws_gc.cell(idx, 11, f"=I{idx}+J{idx}")

    # 3. Xây dựng lại sheet 'Ngày công'
    ws_nc = wb['Ngày công']
    # Clear old detail rows A:I
    for r in range(2, ws_nc.max_row + 1):
        for c in range(1, 10):
            ws_nc.cell(r, c).value = None

    ws_nc.cell(1, 1, "Chi nhánh")
    ws_nc.cell(1, 2, "Tên Nhân viên")
    ws_nc.cell(1, 3, "Ngày")
    ws_nc.cell(1, 4, "Tên ca")
    ws_nc.cell(1, 5, "Giờ làm thực tế")
    ws_nc.cell(1, 6, "Giờ công chuẩn")
    ws_nc.cell(1, 7, "Ca <= 4h")
    ws_nc.cell(1, 8, "Công thực tế")
    ws_nc.cell(1, 9, "Ca đêm chuẩn")

    for idx, sh in enumerate(all_shifts, start=2):
        ws_nc.cell(idx, 1, sh['branch'])
        ws_nc.cell(idx, 2, sh['name'])
        ws_nc.cell(idx, 3, sh['date'])
        ws_nc.cell(idx, 4, sh['shift_name'])
        ws_nc.cell(idx, 5, sh['dur_str'])
        ws_nc.cell(idx, 6, sh['dur_h'])
        # Cột G: Ca <= 4h
        ws_nc.cell(idx, 7, f'=IF(AND(ISNUMBER(F{idx}), F{idx}<=4), 1, 0)')
        # Cột H: Công thực tế (> 4h và là ca >4h đầu tiên trong ngày của người đó tại chi nhánh)
        ws_nc.cell(idx, 8, f'=IF(AND(ISNUMBER(F{idx}), F{idx}>4, COUNTIFS($A$2:A{idx}, A{idx}, $B$2:B{idx}, B{idx}, $C$2:C{idx}, C{idx}, $F$2:F{idx}, ">4")=1), 1, 0)')
        # Cột I: Ca đêm chuẩn (> 4 tiếng và là ca đêm)
        ws_nc.cell(idx, 9, f'=IF(AND(ISNUMBER(F{idx}), F{idx}>4, ISNUMBER(SEARCH("đêm", D{idx}))), 1, 0)')

    # Clear and rebuild summary table in K:O
    for r in range(2, max(ws_nc.max_row + 1, 100)):
        for c in range(11, 16):
            ws_nc.cell(r, c).value = None

    ws_nc.cell(1, 11, "Chi nhánh")
    ws_nc.cell(1, 12, "Tên nhân viên")
    ws_nc.cell(1, 13, "Ngày công chuẩn")
    ws_nc.cell(1, 14, "Giờ công")
    ws_nc.cell(1, 15, "Trung bình giờ/công")

    for idx, (bl_r, br, name, role) in enumerate(staff_rows, start=2):
        ws_nc.cell(idx, 11, br)
        ws_nc.cell(idx, 12, name)
        ws_nc.cell(idx, 13, f"=SUMIFS(H:H, A:A, K{idx}, B:B, L{idx})")
        ws_nc.cell(idx, 14, f"='Giờ công'!K{idx}")
        ws_nc.cell(idx, 15, f'=IF(M{idx}>0, N{idx}/M{idx}, "")')

    wb.save(output_path)
    wb.close()
    print(f"--> Đã lưu file kiểm tra: {output_path}")

audit_and_rebuild_sheets(
    'thang8/DATA/BangChiTietChamCong_thang8.xlsx',
    'thang8/tinhcongnhungthuongchia.xlsx',
    'thang8/test_rebuilt_timecards.xlsx'
)
