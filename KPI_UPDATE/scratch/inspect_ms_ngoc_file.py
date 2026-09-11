import openpyxl, sys, io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

fpath = 'd:/MEDIGO/KPI_UPDATE/thang8/latvat/Gởi Ms. Ngoc - 07.09. Bang Luong T8.2026.xlsx'
wb = openpyxl.load_workbook(fpath, data_only=True)
print("=== SHEETS IN Gởi Ms. Ngoc ===")
print(wb.sheetnames)

for sname in ['Nghia', 'Bac Thu', 'DSDB', 'PL nha thuoc', 'CNT']:
    if sname in wb.sheetnames:
        ws = wb[sname]
        print(f"\n--- Sheet: {sname} (rows={ws.max_row}, cols={ws.max_column}) ---")
        for r in range(1, min(20, ws.max_row+1)):
            vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column+1))]
            if any(vals):
                print(f"  Row {r:2d}:", [v for v in vals if v is not None])

wb.close()
