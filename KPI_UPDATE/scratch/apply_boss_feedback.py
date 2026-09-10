# -*- coding: utf-8 -*-
import openpyxl, os, sys
sys.stdout.reconfigure(encoding='utf-8')

for p in ['thang8/BANGLUONGTHANG8.xlsx']:
    if not os.path.exists(p): continue
    print(f"--> Updating {p}...")
    wb = openpyxl.load_workbook(p)
    
    # Update Ngày công sheet
    if 'Ngày công' in wb.sheetnames:
        ws_nc = wb['Ngày công']
        for r in range(3, ws_nc.max_row + 1):
            if ws_nc.cell(r, 1).value or ws_nc.cell(r, 2).value:
                # Ca đêm chuẩn: ca qua đêm (22:00 - 06:00) >= 4h
                ws_nc.cell(r, 9, f'=IF(AND(NOT(ISBLANK(C{r})), D{r}<>"-", ISNUMBER(F{r}), F{r}>=4, ISNUMBER(SEARCH("qua đêm", D{r}))), 1, 0)')
                # Ca ngày chuẩn: ca 6h - 22h >= 4h
                ws_nc.cell(r, 10, f'=IF(AND(NOT(ISBLANK(C{r})), D{r}<>"-", ISNUMBER(F{r}), F{r}>=4, NOT(ISNUMBER(SEARCH("qua đêm", D{r})))), 1, 0)')

    # Update Giờ công sheet
    if 'Giờ công' in wb.sheetnames:
        ws_gc = wb['Giờ công']
        for r in range(2, 58):
            ws_gc.cell(r, 9, f'=SUMIFS(E:E, B:B, H{r}, C:C, "<>-") - J{r}')
            ws_gc.cell(r, 10, f'=SUMIFS(E:E, B:B, H{r}, C:C, "*qua đêm*")')
            ws_gc.cell(r, 11, f"=I{r}+J{r}")

    # Update BẢNG LƯƠNG sheet
    ws_bl = wb['BẢNG LƯƠNG']
    for r in range(3, 59):
        # Cột J: Quy đổi 24h = 1 ngày, chặn J + K <= 31
        ws_bl.cell(r, 10, f"=MIN(ROUND(G{r}/24, 0), MAX(0, 31-K{r}))")
        # Cột K: Ca đêm >= 4h
        ws_bl.cell(r, 11, f"=SUMIFS('Ngày công'!I:I, 'Ngày công'!B:B, C{r})")
        # Cột X: Phụ cấp sức khỏe ca đêm MAX 1,500,000 đ
        ws_bl.cell(r, 24, f'=IF(ISNUMBER(SEARCH("DSCD", D{r})), MIN(1500000, ROUND(1500000/28*K{r}, 0)), IF(K{r}>20, MIN(1500000, ROUND(1500000/28*K{r}, 0)), 0))')
        
    wb.save(p)
    print(f"✅ Đã lưu file: {p}")
