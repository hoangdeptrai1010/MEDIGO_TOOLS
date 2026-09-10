import openpyxl

wb = openpyxl.load_workbook('thang7/BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=False)
ws = wb['BẢNG LƯƠNG']

for c in range(43, 54):
    r1 = ws.cell(1, c).value
    r2 = ws.cell(2, c).value
    r3 = ws.cell(3, c).value
    print(f"Col {c} ({openpyxl.utils.get_column_letter(c)}): r1='{r1}' | r2='{r2}' | r3='{r3}'")
