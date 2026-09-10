import openpyxl

wb = openpyxl.load_workbook('thang7/BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=False)
ws = wb['BẢNG LƯƠNG']

for c in range(1, 54):
    h = ws.cell(2, c).value
    f = ws.cell(3, c).value
    print(f"Col {c} ({openpyxl.utils.get_column_letter(c)}): Header='{h}' | Row3 Formula='{f}'")
