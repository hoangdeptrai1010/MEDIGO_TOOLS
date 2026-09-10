import openpyxl

wb = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx', data_only=True)
ws_ds = wb['kpi dược sĩ']

print(f"{'Row':<4} | {'Họ tên':<25} | {'Chi nhánh':<12} | {'F (Chỉ tiêu)':>12} | {'G (Thực tế)':>12} | {'G >= F?':>8} | {'Thưởng KPI':>12}")
print("-" * 95)

for r in range(3, ws_ds.max_row+1):
    name = ws_ds.cell(r, 2).value
    if not name: continue
    cn = ws_ds.cell(r, 1).value
    f = ws_ds.cell(r, 6).value or 0
    g = ws_ds.cell(r, 7).value or 0
    kpi = ws_ds.cell(r, 11).value or 0
    
    pass_cond = "ĐẠT" if g >= f else "TRƯỢT"
    print(f"{r:<4} | {name:<25} | {str(cn):<12} | {f:>12,.0f} | {g:>12,.0f} | {pass_cond:>8} | {kpi:>12,.0f}")
