import openpyxl

wb_july = openpyxl.load_workbook('thang7/BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=True)
ws_gc = wb_july['Giờ công']
print("Total rows in July Giờ công:", ws_gc.max_row)

for r in range(2, 25):
    row_vals = [ws_gc.cell(r, c).value for c in range(1, 12)]
    print(f"R{r}: {row_vals}")
