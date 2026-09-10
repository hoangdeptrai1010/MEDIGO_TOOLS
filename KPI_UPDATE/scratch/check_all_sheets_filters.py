import openpyxl

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=False)

print("=== CHECKING HIDDEN ROWS & FILTERS IN ALL SHEETS ===")
for sname in wb.sheetnames:
    ws = wb[sname]
    hidden_count = sum(1 for dim in ws.row_dimensions.values() if dim.hidden)
    has_filter = ws.auto_filter.ref is not None
    filter_cols = len(ws.auto_filter.filterColumn) if ws.auto_filter else 0
    print(f"Sheet '{sname:<20}': Total rows={ws.max_row:<5}, Hidden rows={hidden_count:<5}, Has AutoFilter={has_filter}, Filter columns={filter_cols}")
