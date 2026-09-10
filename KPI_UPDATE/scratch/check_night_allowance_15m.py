# -*- coding: utf-8 -*-
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/BANGLUONGTHANG8.xlsx', data_only=True)
ws_bl = wb['BẢNG LƯƠNG']

print("=== KIỂM TRA PHỤ CẤP CA ĐÊM (MAX 1,500,000 Đ) ===")
print(f"{'STT':<3} | {'Chi nhánh':<15} | {'Họ và tên':<24} | {'Chức danh':<10} | {'Ca đêm (K)':<10} | {'PC Đêm cũ':<14} | {'PC Đêm mới (Max 1.5M)':<22}")
print("-" * 115)

for r in range(3, 59):
    stt = ws_bl.cell(r, 1).value
    cn = ws_bl.cell(r, 2).value or ''
    name = ws_bl.cell(r, 3).value or ''
    role = str(ws_bl.cell(r, 4).value or '').strip()
    k = ws_bl.cell(r, 11).value or 0
    k_num = int(k) if isinstance(k, (int, float)) else 0
    
    # Old calculation
    old_pc = 0
    if "DSCD" in role:
        old_pc = round(1500000 / 28 * k_num)
    elif k_num > 20:
        old_pc = round(1500000 / 28 * k_num)
        
    # New calculation with MAX 1,500,000
    new_pc = 0
    if "DSCD" in role:
        new_pc = min(1500000, round(1500000 / 28 * k_num))
    elif k_num > 20:
        new_pc = min(1500000, round(1500000 / 28 * k_num))
        
    if k_num > 0 or new_pc > 0:
        print(f"{stt:<3} | {cn:<15} | {name:<24} | {role:<10} | {k_num:<10} | {old_pc:>10,.0f} đ | {new_pc:>14,.0f} đ")
