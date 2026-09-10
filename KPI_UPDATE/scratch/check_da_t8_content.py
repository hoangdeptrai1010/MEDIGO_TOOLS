import openpyxl

wb = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx', data_only=True)
for sname in ['Dự án T8', 'Hot Bill HN']:
    ws = wb[sname]
    print(f"=== {sname} ===")
    for r in range(1, 6):
        print([ws.cell(r, c).value for c in range(1, 16) if ws.cell(r, c).value is not None])
