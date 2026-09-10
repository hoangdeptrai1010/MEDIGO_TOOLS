import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

print("================================================================================")
print("KIỂM TRA CÔNG THỨC & NGUỒN DỮ LIỆU TỪ BAOCAOKPI SANG BẢNG LƯƠNG")
print("================================================================================")

wb_kpi = openpyxl.load_workbook(r'baocaokpi_thang8_hoanthien.xlsx', data_only=False)
ws_ds = wb_kpi['kpi dược sĩ']

print("--- [baocaokpi_thang8_hoanthien.xlsx] Cột trong sheet 'kpi dược sĩ' ---")
for c in range(1, 16):
    print(f"Col {openpyxl.utils.get_column_letter(c)} ({c}) Header: {ws_ds.cell(2, c).value}")

print("\nCông thức Row 3 (Lê Thị Soạn) trong sheet 'kpi dược sĩ':")
for c in [8, 9, 10, 11, 12, 13, 14]:
    print(f"  Col {openpyxl.utils.get_column_letter(c)} ({ws_ds.cell(2, c).value}): Formula={ws_ds.cell(3, c).value}")

wb_bl = openpyxl.load_workbook(r'thang8\bangluong_thang8_hoanthien.xlsx', data_only=False)
wb_bl_val = openpyxl.load_workbook(r'thang8\bangluong_thang8_hoanthien.xlsx', data_only=True)

ws_bl = wb_bl['BẢNG LƯƠNG']
ws_bl_val = wb_bl_val['BẢNG LƯƠNG']

# Find Lê Thị Soạn in BẢNG LƯƠNG
for r in range(3, ws_bl.max_row + 1):
    name = ws_bl.cell(r, 3).value
    if name == 'Lê Thị Soạn':
        print(f"\n--- [BẢNG LƯƠNG] Dòng Lê Thị Soạn (Row {r}) ---")
        for col_let in ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI']:
            c_idx = openpyxl.utils.column_index_from_string(col_let)
            h_text = ws_bl.cell(2, c_idx).value
            f_val = ws_bl.cell(r, c_idx).value
            v_val = ws_bl_val.cell(r, c_idx).value
            print(f"  Cột {col_let} ({str(h_text).strip()}): Formula={f_val} | Value={v_val}")
