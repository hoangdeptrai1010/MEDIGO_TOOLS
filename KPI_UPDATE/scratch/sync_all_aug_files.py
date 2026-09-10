# -*- coding: utf-8 -*-
import openpyxl, os, sys, time
sys.stdout.reconfigure(encoding='utf-8')

target_files = [
    'thang8/BANGLUONGTHANG8.xlsx',
    'thang8/BANGLUONGTHANG8_HOANG_.xlsx',
    'thang8/bangluong_thang8_hoanthien.xlsx'
]

for p in target_files:
    if not os.path.exists(p):
        continue
    print(f"\n--> Updating formulas in {p}...")
    try:
        wb = openpyxl.load_workbook(p)
        
        # 1. Update Ngày công
        if 'Ngày công' in wb.sheetnames:
            ws_nc = wb['Ngày công']
            for r in range(3, ws_nc.max_row + 1):
                if ws_nc.cell(r, 1).value or ws_nc.cell(r, 2).value:
                    # Ca đêm chuẩn (22:00 - 06:00 >= 4h)
                    ws_nc.cell(r, 9, f'=IF(AND(NOT(ISBLANK(C{r})), D{r}<>"-", ISNUMBER(F{r}), F{r}>=4, ISNUMBER(SEARCH("qua đêm", D{r}))), 1, 0)')
                    # Ca ngày chuẩn (06:00 - 22:00 >= 4h)
                    ws_nc.cell(r, 10, f'=IF(AND(NOT(ISBLANK(C{r})), D{r}<>"-", ISNUMBER(F{r}), F{r}>=4, NOT(ISNUMBER(SEARCH("qua đêm", D{r})))), 1, 0)')

        # 2. Update Giờ công
        if 'Giờ công' in wb.sheetnames:
            ws_gc = wb['Giờ công']
            for r in range(2, 58):
                # Sáng: Toàn bộ giờ 6h-22h (trừ ca qua đêm)
                ws_gc.cell(r, 9, f'=SUMIFS(E:E, B:B, H{r}, C:C, "<>-") - J{r}')
                # Đêm: Toàn bộ ca qua đêm 22h-6h
                ws_gc.cell(r, 10, f'=SUMIFS(E:E, B:B, H{r}, C:C, "*qua đêm*")')
                ws_gc.cell(r, 11, f"=I{r}+J{r}")

        # 3. Update BẢNG LƯƠNG
        if 'BẢNG LƯƠNG' in wb.sheetnames:
            ws_bl = wb['BẢNG LƯƠNG']
            ws_bl.cell(2, 10, "Số ngày làm ngày")
            ws_bl.cell(2, 11, "Số ca làm đêm")
            for r in range(3, 59):
                # Cột G: Số giờ ca ngày (6h - 22h)
                ws_bl.cell(r, 7, f"=SUMIFS('Giờ công'!I:I, 'Giờ công'!H:H, C{r})")
                # Cột H: Số giờ ca đêm (22h - 6h)
                ws_bl.cell(r, 8, f"=SUMIFS('Giờ công'!J:J, 'Giờ công'!H:H, C{r})")
                # Cột J: Số ngày làm ngày (24h = 1 ngày)
                ws_bl.cell(r, 10, f"=MIN(ROUND(G{r}/24, 0), MAX(0, 31-K{r}))")
                # Cột K: Số ca làm đêm >= 4h
                ws_bl.cell(r, 11, f"=SUMIFS('Ngày công'!I:I, 'Ngày công'!B:B, C{r})")
                # Cột X: Phụ cấp ca đêm MAX 1,500,000 đ
                ws_bl.cell(r, 24, f'=IF(ISNUMBER(SEARCH("DSCD", D{r})), MIN(1500000, ROUND(1500000/28*K{r}, 0)), IF(K{r}>20, MIN(1500000, ROUND(1500000/28*K{r}, 0)), 0))')

        wb.save(p)
        print(f"✅ Đã lưu thành công: {p}")
    except PermissionError:
        print(f"❌ PermissionError: File {p} đang bị khóa do đang mở trong Excel!")
    except Exception as e:
        print(f"Lỗi {p}: {e}")
