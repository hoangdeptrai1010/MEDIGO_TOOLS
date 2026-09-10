import openpyxl

for p in ['thang7/BẢNG LƯƠNG THÁNG 7 2026.xlsx', 'thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx']:
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb['BẢNG LƯƠNG']
    print(f"=== {p} ===")
    for r in range(3, 8):
        name = ws.cell(r, 3).value
        cn = ws.cell(r, 2).value
        gn = ws.cell(r, 5).value # Col E
        gd = ws.cell(r, 6).value # Col F
        g_sang = ws.cell(r, 7).value # Col G
        g_dem = ws.cell(r, 8).value # Col H
        so_ngay = ws.cell(r, 10).value # Col J
        so_dem = ws.cell(r, 11).value # Col K
        ngay_cong = ws.cell(r, 18).value # Col R
        print(f"{cn} - {name}: G_sang={g_sang}, G_dem={g_dem}, So_ngay={so_ngay}, So_dem={so_dem}, Ngay_cong={ngay_cong}")
