import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/test_rebuilt_timecards.xlsx', data_only=False)

ws_bl = wb['BẢNG LƯƠNG']
ws_nc = wb['Ngày công']
ws_gc = wb['Giờ công']

print("=== KIỂM TRA MỘT SỐ NHÂN SỰ TIÊU BIỂU ===")

# Check staff in BL
for r in range(3, 73):
    cn = ws_bl.cell(r, 2).value
    name = ws_bl.cell(r, 3).value
    role = ws_bl.cell(r, 4).value
    g_f = ws_bl.cell(r, 7).value
    h_f = ws_bl.cell(r, 8).value
    j_f = ws_bl.cell(r, 10).value
    k_f = ws_bl.cell(r, 11).value
    r_f = ws_bl.cell(r, 18).value
    
    if name in ['Nguyễn Thị Tâm', 'Hứa Thị Kim Thoa', 'Hồ Thị Minh Hòa', 'Hoàng Lâm Gia Bảo', 'Trần Thiên Phát', 'Lê Thị Xuyến']:
        print(f"Row {r}: [{cn}] {name} ({role})")
        print(f"  Col G (Giờ ngày): {g_f}")
        print(f"  Col H (Giờ đêm): {h_f}")
        print(f"  Col J (Ngày làm ngày = G/24): {j_f}")
        print(f"  Col K (Ca làm đêm): {k_f}")
        print(f"  Col R (Ngày công thực tế): {r_f}")

wb.close()
