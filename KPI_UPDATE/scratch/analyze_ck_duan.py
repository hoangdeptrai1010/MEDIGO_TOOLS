import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

def analyze_m7():
    print("======================== THÁNG 7 ========================")
    wb_goc = openpyxl.load_workbook(r'goc\NHÀ THUỐC THÁNG 7 2026.xlsx', data_only=True)
    print("Sheets goc T7:", wb_goc.sheetnames)
    
    # In goc T7, let's look at Dự án T7, CK, etc.
    for sname in wb_goc.sheetnames:
        if 'dự án' in sname.lower() or 'ck' in sname.lower() or 'kpi' in sname.lower() or 'lương' in sname.lower():
            ws = wb_goc[sname]
            print(f"\n[GỐC T7] Sheet: {sname}")
            for r in range(1, min(10, ws.max_row + 1)):
                row_vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
                if any(v is not None for v in row_vals):
                    print(f"  R{r}: {row_vals}")

    wb_bl_t7 = openpyxl.load_workbook(r'thang7\target\BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=True)
    print("\nSheets Target Bảng Lương T7:", wb_bl_t7.sheetnames)
    for sname in ['BẢNG LƯƠNG', 'Thưởng CK', 'Dự án', 'KPI']:
        if sname in wb_bl_t7.sheetnames:
            ws = wb_bl_t7[sname]
            print(f"\n[BL T7 Target] Sheet: {sname}")
            for r in range(1, min(10, ws.max_row + 1)):
                row_vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
                if any(v is not None for v in row_vals):
                    print(f"  R{r}: {row_vals}")

def analyze_m8():
    print("\n======================== THÁNG 8 ========================")
    wb_goc8 = openpyxl.load_workbook(r'goc\NHÀ THUỐC THÁNG 8 2026.xlsx', data_only=True)
    print("Sheets goc T8:", wb_goc8.sheetnames)
    for sname in wb_goc8.sheetnames:
        if 'dự án' in sname.lower() or 'ck' in sname.lower() or 'kpi' in sname.lower() or 'lương' in sname.lower():
            ws = wb_goc8[sname]
            print(f"\n[GỐC T8] Sheet: {sname}")
            for r in range(1, min(10, ws.max_row + 1)):
                row_vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
                if any(v is not None for v in row_vals):
                    print(f"  R{r}: {row_vals}")

    wb_bl_t8 = openpyxl.load_workbook(r'thang8\bangluong_thang8_hoanthien.xlsx', data_only=True)
    print("\nSheets BL T8 Hoàn thiện:", wb_bl_t8.sheetnames)
    for sname in ['BẢNG LƯƠNG', 'Thưởng CK', 'Dự án', 'KPI']:
        if sname in wb_bl_t8.sheetnames:
            ws = wb_bl_t8[sname]
            print(f"\n[BL T8 Hoàn thiện] Sheet: {sname}")
            for r in range(1, min(10, ws.max_row + 1)):
                row_vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
                if any(v is not None for v in row_vals):
                    print(f"  R{r}: {row_vals}")

analyze_m7()
analyze_m8()
