import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

plan_file = 'plans/KeHoachKPI_2026-08.xlsx'
print(f"=== ĐỌC CHI TIẾT BẢN KẾ HOẠCH THÁNG 8: {plan_file} ===")

wb = openpyxl.load_workbook(plan_file, data_only=True)
print("Các sheets trong file Kế hoạch:", wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n--- Sheet: {sname} (Rows: {ws.max_row}, Cols: {ws.max_column}) ---")
    for r in range(1, min(ws.max_row + 1, 20)):
        row_vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 15))]
        if any(v is not None for v in row_vals):
            print(f"Row {r:2d}: {row_vals}")

wb.close()
