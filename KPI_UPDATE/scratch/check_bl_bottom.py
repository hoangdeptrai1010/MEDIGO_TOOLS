import openpyxl

wb = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=False)
ws_bl = wb['BẢNG LƯƠNG']

print("=== thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx bottom rows (72-88) ===")
for r in range(72, 89):
    vals = [ws_bl.cell(r, c).value for c in range(1, 10) if ws_bl.cell(r, c).value is not None]
    print(f"R{r}: {vals}")
