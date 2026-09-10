import openpyxl

wb = openpyxl.load_workbook('thang7/DATAKIOT/hoadon.xlsx', data_only=True, read_only=True)
ws = wb.active
headers = next(ws.iter_rows(max_row=1, values_only=True))
col_map = {h: idx for idx, h in enumerate(headers) if h}

c_seller = col_map.get('Người bán', 6)
c_date = col_map.get('Thời gian', 2)
c_sku = col_map.get('Mã hàng', 13)
c_name = col_map.get('Tên hàng', 15)
c_exp = col_map.get('Hạn sử dụng', 18)
c_val = col_map.get('Thành tiền', 25)

hang_items = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r: continue
    s = str(r[c_seller] or '').strip()
    if s == 'Vũ Thanh Hằng':
        exp = r[c_exp]
        val = float(r[c_val] or 0)
        if exp and val > 0:
            exp_str = str(exp)[:10]
            if exp_str <= '2027-02-28':
                hang_items.append({
                    'date': str(r[c_date])[:10],
                    'exp': exp_str,
                    'sku': r[c_sku],
                    'name': r[c_name],
                    'val': val
                })

print(f"Vu Thanh Hang items count: {len(hang_items)}, sum: {sum(it['val'] for it in hang_items):,.0f}đ")
print("Target is: 2,510,500đ")

# Group by SKU
by_sku = {}
for it in hang_items:
    sku = it['sku']
    if sku not in by_sku:
        by_sku[sku] = {'val': 0, 'name': it['name'], 'exp': it['exp'], 'count': 0}
    by_sku[sku]['val'] += it['val']
    by_sku[sku]['count'] += 1

for sku, d in sorted(by_sku.items(), key=lambda x: x[1]['val'], reverse=True):
    print(f"SKU {sku:12s} | HSD: {d['exp']} | {d['val']:>10,.0f}đ ({d['count']:2d} lần) | {d['name'][:35]}")
