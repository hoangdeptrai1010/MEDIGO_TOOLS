import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

def compare_m7():
    print("==================================================================================")
    print("THÁNG 7: BÁO CÁO KPI GỐC vs TARGET BẢNG LƯƠNG vs HOÀN THIỆN")
    print("==================================================================================")
    wb_kpi = openpyxl.load_workbook(r'baocaokpi_thang7_hoanthien.xlsx', data_only=True)
    wb_goc = openpyxl.load_workbook(r'goc\NHÀ THUỐC THÁNG 7 2026.xlsx', data_only=True)
    wb_bl = openpyxl.load_workbook(r'thang7\target\BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=True)

    # In goc T7, let's see Dự án T7 sheet
    ws_goc_da = wb_goc['Dự án T7'] if 'Dự án T7' in wb_goc.sheetnames else None
    goc_da_dict = {}
    if ws_goc_da:
        for r in range(3, ws_goc_da.max_row + 1):
            name = ws_goc_da.cell(r, 2).value
            # Total du an col 13 (M)
            tot_da = ws_goc_da.cell(r, 13).value
            ck = ws_goc_da.cell(r, 6).value
            if name:
                goc_da_dict[str(name).strip()] = (ck, tot_da)

    # In BL T7, let's see Thưởng CK (Cột AF / sheet Thưởng CK) and Dự án (Cột AD / sheet Dự án)
    ws_bl_s = wb_bl['BẢNG LƯƠNG']
    ws_bl_ck = wb_bl['Thưởng CK']
    ws_bl_da = wb_bl['Dự án']

    print(f"{'Tên':<25} | {'Gốc Dự án (Col M)':<18} | {'BL Dự án (Col AD)':<18} | {'BL Thưởng CK (Col AF)':<20}")
    print("-" * 90)
    for r in range(3, ws_bl_s.max_row + 1):
        name = ws_bl_s.cell(r, 3).value
        if name and str(name).strip():
            name_str = str(name).strip()
            da_goc = goc_da_dict.get(name_str, (None, None))[1]
            da_bl = ws_bl_s.cell(r, 30).value # Col AD (30)
            ck_bl = ws_bl_s.cell(r, 32).value # Col AF (32)
            print(f"{name_str:<25} | {str(da_goc):<18} | {str(da_bl):<18} | {str(ck_bl):<20}")

def compare_m8():
    print("\n==================================================================================")
    print("THÁNG 8: BÁO CÁO KPI GỐC vs BẢNG LƯƠNG HOÀN THIỆN")
    print("==================================================================================")
    wb_kpi = openpyxl.load_workbook(r'baocaokpi_thang8_hoanthien.xlsx', data_only=True)
    wb_goc = openpyxl.load_workbook(r'goc\NHÀ THUỐC THÁNG 8 2026.xlsx', data_only=True)
    wb_bl = openpyxl.load_workbook(r'thang8\bangluong_thang8_hoanthien.xlsx', data_only=True)

    ws_goc_da = wb_goc['Dự án T8'] if 'Dự án T8' in wb_goc.sheetnames else None
    goc_da_dict = {}
    if ws_goc_da:
        for r in range(3, ws_goc_da.max_row + 1):
            name = ws_goc_da.cell(r, 2).value
            tot_da = ws_goc_da.cell(r, 13).value
            ck = ws_goc_da.cell(r, 6).value
            if name:
                goc_da_dict[str(name).strip()] = (ck, tot_da)

    ws_kpi_da = wb_kpi['Dự án T8'] if 'Dự án T8' in wb_kpi.sheetnames else None
    kpi_da_dict = {}
    if ws_kpi_da:
        for r in range(3, ws_kpi_da.max_row + 1):
            name = ws_kpi_da.cell(r, 2).value
            tot_da = ws_kpi_da.cell(r, 13).value
            if name:
                kpi_da_dict[str(name).strip()] = tot_da

    ws_bl_s = wb_bl['BẢNG LƯƠNG']
    print(f"{'Tên':<26} | {'Gốc DA T8':<12} | {'KPI DA T8':<12} | {'BL DA (AD)':<12} | {'BL CK (AF)':<15} | {'MiniKat AA':<12}")
    print("-" * 105)
    for r in range(3, ws_bl_s.max_row + 1):
        name = ws_bl_s.cell(r, 3).value
        if name and str(name).strip() and "Tổng" not in str(name):
            name_str = str(name).strip()
            da_goc = goc_da_dict.get(name_str, (None, None))[1]
            da_kpi = kpi_da_dict.get(name_str, None)
            da_bl = ws_bl_s.cell(r, 30).value # AD (30)
            ck_bl = ws_bl_s.cell(r, 32).value # AF (32)
            minikat_aa = ws_bl_s.cell(r, 27).value # AA (27)
            
            da_goc_str = f"{da_goc:,.0f}" if isinstance(da_goc, (int, float)) else str(da_goc)
            da_kpi_str = f"{da_kpi:,.0f}" if isinstance(da_kpi, (int, float)) else str(da_kpi)
            da_bl_str = f"{da_bl:,.0f}" if isinstance(da_bl, (int, float)) else str(da_bl)
            ck_bl_str = f"{ck_bl:,.0f}" if isinstance(ck_bl, (int, float)) else str(ck_bl)
            mk_str = f"{minikat_aa:,.0f}" if isinstance(minikat_aa, (int, float)) else str(minikat_aa)
            
            print(f"{name_str:<26} | {da_goc_str:<12} | {da_kpi_str:<12} | {da_bl_str:<12} | {ck_bl_str:<15} | {mk_str:<12}")

compare_m7()
compare_m8()
