import openpyxl

for p in ['thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', 'thang8/bangluong_thang8_hoanthien.xlsx']:
    try:
        wb = openpyxl.load_workbook(p, data_only=True)
        print(f"=== {p} ===")
        if 'whatsapp' in wb.sheetnames:
            ws = wb['whatsapp']
            for r in range(1, min(10, ws.max_row + 1)):
                print(f"  whatsapp r{r}: {[ws.cell(r, c).value for c in range(1, 6)]}")
        if 'MiniKat - HN' in wb.sheetnames:
            ws = wb['MiniKat - HN']
            for r in range(1, min(6, ws.max_row + 1)):
                print(f"  MiniKat-HN r{r}: {[ws.cell(r, c).value for c in range(1, 11)]}")
        if 'Hot Bill HN' in wb.sheetnames:
            print("  Hot Bill HN sheet exists!")
    except Exception as e:
        print(f"Error {p}: {e}")
