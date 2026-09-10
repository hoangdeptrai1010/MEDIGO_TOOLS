import openpyxl

wb_cc = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
ws_ct = wb_cc['Bảng chi tiết chấm công']

count = 0
for r in range(4, ws_ct.max_row + 1):
    d = ws_ct.cell(r, 6).value
    if d is None or str(d).strip() == '':
        vals = {openpyxl.utils.get_column_letter(c): ws_ct.cell(r, c).value for c in range(1, 24) if ws_ct.cell(r, c).value is not None}
        print(f"Row {r}: {vals}")
        count += 1
        if count >= 10:
            break
