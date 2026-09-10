# -*- coding: utf-8 -*-
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws_bl = wb['BẢNG LƯƠNG']
ws_da = wb['Dự án']
ws_kpi = wb['KPI']
ws_ck = wb['Thưởng CK']

print(f"{'STT':<3} | {'Chi nhánh':<15} | {'Họ và tên':<24} | {'Dự án (AD)':<14} | {'Thưởng KPI (AE)':<16} | {'Thưởng CK (AF)':<16} | {'Tổng 3 khoản':<14}")
print("-" * 115)

for r in range(3, 59):
    stt = ws_bl.cell(r, 1).value
    cn = ws_bl.cell(r, 2).value or ''
    name = ws_bl.cell(r, 3).value or ''
    ad = ws_bl.cell(r, 30).value or 0
    ae = ws_bl.cell(r, 31).value or 0
    af = ws_bl.cell(r, 32).value or 0
    tot = (ad or 0) + (ae or 0) + (af or 0)
    
    ad_str = f"{ad:,.0f} đ" if isinstance(ad, (int, float)) and ad > 0 else "-"
    ae_str = f"{ae:,.0f} đ" if isinstance(ae, (int, float)) and ae > 0 else "-"
    af_str = f"{af:,.0f} đ" if isinstance(af, (int, float)) and af > 0 else "-"
    tot_str = f"{tot:,.0f} đ" if isinstance(tot, (int, float)) and tot > 0 else "-"
    
    if tot > 0:
        print(f"{stt:<3} | {cn:<15} | {name:<24} | {ad_str:<14} | {ae_str:<16} | {af_str:<16} | {tot_str:<14}")
