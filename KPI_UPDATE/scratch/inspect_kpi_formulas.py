import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

f_nt8 = r"d:\MEDIGO\KPI_UPDATE\goc\NHÀ THUỐC THÁNG 8 2026.xlsx"
wb = openpyxl.load_workbook(f_nt8, data_only=False)

for sname in ['kpi dược sĩ', 'kpi nhà thuốc', 'Dự án T8']:
    ws = wb[sname]
    print(f"\n==========================================================================")
    print(f"FORMULAS IN SHEET: '{sname}'")
    print(f"==========================================================================")
    for r in range(2, 6):
        row_vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 35))]
        print(f"Row {r:2d}:")
        for c, v in enumerate(row_vals, 1):
            if v is not None:
                print(f"  [{c:2d} | {openpyxl.utils.get_column_letter(c):>3s}]: {repr(v)}")
