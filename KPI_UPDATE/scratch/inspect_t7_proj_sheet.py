import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb7 = openpyxl.load_workbook(r'baocaokpi_thang7_hoanthien.xlsx', data_only=True)
ws7 = wb7['Dự án T7']

print("--- baocaokpi_thang7_hoanthien.xlsx Sheet 'Dự án T7' Headers ---")
for c in range(1, ws7.max_column + 1):
    h = ws7.cell(2, c).value
    if h:
        print(f"Col {openpyxl.utils.get_column_letter(c)} ({c}): {h}")

print("\n--- 5 rows data in Dự án T7 ---")
for r in range(3, 8):
    print([ws7.cell(r, c).value for c in range(1, 15)])
