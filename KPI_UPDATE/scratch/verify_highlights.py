import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_bl = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final.xlsx")

wb = openpyxl.load_workbook(f_bl)
ws1 = wb["Sheet1"]

print("=== VERIFYING HIGHLIGHTED CELLS IN Sheet1 ===")
for r in range(2, ws1.max_row + 1):
    c_name = ws1.cell(r, 3).value
    c_stk = ws1.cell(r, 5)
    c_bank = ws1.cell(r, 6)
    
    stk_color = c_stk.fill.start_color.rgb if c_stk.fill and c_stk.fill.start_color else None
    bank_color = c_bank.fill.start_color.rgb if c_bank.fill and c_bank.fill.start_color else None
    
    if c_name:
        is_highlighted = (stk_color is not None or bank_color is not None)
        print(f"Row {r:2d} | Name: {c_name:<25} | STK: {str(c_stk.value):<10} ({stk_color}) | Bank: {str(c_bank.value):<10} ({bank_color}) | Highlighted: {is_highlighted}")
