import openpyxl
import pandas as pd
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

latvat_dir = r"d:\MEDIGO\KPI_UPDATE\thang8\latvat"
f_bangluong = os.path.join(latvat_dir, "BANGLUONGT8FIXV31 final.xlsx")
f_ngoc = os.path.join(latvat_dir, "Gởi Ms. Ngoc - 07.09. Bang Luong T8.2026.xlsx")
f_dsnv = r"d:\MEDIGO\KPI_UPDATE\TOOL_BANGLUONG\danhsachnhanvien.xlsx"

print("--- Inspecting BANGLUONGT8FIXV31 final.xlsx ---")
wb_bl = openpyxl.load_workbook(f_bangluong, data_only=True)
print("Sheets in BANGLUONG:", wb_bl.sheetnames)

for sname in ["Sheet1", "Med"]:
    if sname in wb_bl.sheetnames:
        ws = wb_bl[sname]
        print(f"\nSample data from sheet {sname} (max_row={ws.max_row}, max_col={ws.max_column}):")
        for r in range(1, min(15, ws.max_row + 1)):
            row_vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
            if any(v is not None for v in row_vals):
                print(f"Row {r}: {row_vals[:12]}")

print("\n--- Inspecting Gởi Ms. Ngoc - 07.09. Bang Luong T8.2026.xlsx ---")
wb_ngoc = openpyxl.load_workbook(f_ngoc, data_only=True)
print("Sheets in Ms. Ngoc:", wb_ngoc.sheetnames)

for sname in wb_ngoc.sheetnames:
    ws = wb_ngoc[sname]
    print(f"\nSheet '{sname}' sample headers (max_row={ws.max_row}):")
    for r in range(1, min(12, ws.max_row + 1)):
        row_vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
        if any(v is not None for v in row_vals):
            print(f"  Row {r}: {row_vals[:10]}")

if os.path.exists(f_dsnv):
    print("\n--- Inspecting danhsachnhanvien.xlsx ---")
    wb_dsnv = openpyxl.load_workbook(f_dsnv, data_only=True)
    print("Sheets in danhsachnhanvien:", wb_dsnv.sheetnames)
    for sname in wb_dsnv.sheetnames:
        ws = wb_dsnv[sname]
        print(f"\nSheet '{sname}' sample headers (max_row={ws.max_row}):")
        for r in range(1, min(10, ws.max_row + 1)):
            row_vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
            if any(v is not None for v in row_vals):
                print(f"  Row {r}: {row_vals[:10]}")
