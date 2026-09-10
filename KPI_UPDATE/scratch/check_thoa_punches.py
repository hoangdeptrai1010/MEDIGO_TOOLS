import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
ws = wb.active

for r in range(205, 212):
    sh = ws.cell(r, 7).value
    punches = []
    for day in range(1, 32):
        col_in = 8 + (day - 1) * 2
        col_out = col_in + 1
        v_in = ws.cell(r, col_in).value
        v_out = ws.cell(r, col_out).value
        if v_in and v_out:
            punches.append((day, v_in, v_out))
    print(f"Row {r}: Shift='{sh}' -> {len(punches)} punches: {punches}")

wb.close()
