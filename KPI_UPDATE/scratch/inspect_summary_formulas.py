import openpyxl

wb = openpyxl.load_workbook('thang8/BẢNG LƯƠNG THÁNG 8 2026.xlsx', data_only=False)

# Let's check sheet Ngày công formulas in summary table:
ws_nc = wb['Ngày công']
print("Ngày công summary table formulas (rows 2-8):")
for r in range(2, 9):
    print(f"R{r}: K={repr(ws_nc.cell(r,11).value)}, L={repr(ws_nc.cell(r,12).value)}, M={repr(ws_nc.cell(r,13).value)}, N={repr(ws_nc.cell(r,14).value)}, O={repr(ws_nc.cell(r,15).value)}")

# Let's check sheet Giờ công formulas in summary table:
ws_gc = wb['Giờ công']
print("\nGiờ công summary table formulas (rows 2-8):")
for r in range(2, 9):
    print(f"R{r}: G={repr(ws_gc.cell(r,7).value)}, H={repr(ws_gc.cell(r,8).value)}, I={repr(ws_gc.cell(r,9).value)}, J={repr(ws_gc.cell(r,10).value)}, K={repr(ws_gc.cell(r,11).value)}")
