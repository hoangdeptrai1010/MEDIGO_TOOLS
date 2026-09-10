import openpyxl, itertools

wb_tgt = openpyxl.load_workbook('thang7/target/BẢNG LƯƠNG THÁNG 7 2026.xlsx', data_only=True)
ws_tgt = wb_tgt['Cận date']
targets = {}
for r in range(2, ws_tgt.max_row+1):
    n = ws_tgt.cell(r, 8).value
    v = ws_tgt.cell(r, 9).value
    if n and v is not None:
        targets[str(n).strip()] = float(v)
wb_tgt.close()

wb = openpyxl.load_workbook('thang7/DATAKIOT/hoadon.xlsx', data_only=True, read_only=True)
ws = wb.active
headers = next(ws.iter_rows(max_row=1, values_only=True))
col_map = {h: idx for idx, h in enumerate(headers) if h}

c_seller = col_map.get('Người bán', 6)
c_date = col_map.get('Thời gian', 2)
c_sku = col_map.get('Mã hàng', 13)
c_name = col_map.get('Tên hàng', 15)
c_exp = col_map.get('Hạn sử dụng', 18)
c_val = col_map.get('Thành tiền', 25)

seller_items = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r: continue
    s = str(r[c_seller] or '').strip()
    if s in targets:
        exp = r[c_exp]
        val = float(r[c_val] or 0)
        if exp and val > 0 and str(exp)[:10] <= '2027-02-28':
            if s not in seller_items:
                seller_items[s] = []
            seller_items[s].append({
                'date': str(r[c_date])[:10],
                'exp': str(exp)[:10],
                'sku': r[c_sku],
                'name': r[c_name],
                'val': val
            })
wb.close()

all_matched_skus = set()

for s, target in targets.items():
    items = seller_items.get(s, [])
    # Group by SKU
    by_sku = {}
    for it in items:
        sku = it['sku']
        if sku not in by_sku:
            by_sku[sku] = {'val': 0, 'name': it['name'], 'exp': it['exp']}
        by_sku[sku]['val'] += it['val']
    
    sku_list = [(v['val'], k, v['name'], v['exp']) for k, v in by_sku.items() if v['val'] <= target + 1]
    
    found = False
    for r in range(1, min(len(sku_list) + 1, 10)):
        for comb in itertools.combinations(sku_list, r):
            if abs(sum(c[0] for c in comb) - target) < 1.0:
                print(f"\n✅ {s} (Target: {target:,.0f}đ): MATCHED with {len(comb)} SKUs:")
                for c in comb:
                    print(f"   - {c[1]}: {c[0]:>10,.0f}đ | {c[2]} (HSD: {c[3]})")
                    all_matched_skus.add(c[1])
                found = True
                break
        if found:
            break
    if not found:
        print(f"\n❌ {s} (Target: {target:,.0f}đ): NOT FOUND exact combination (total items: {len(items)}, total sum: {sum(it['val'] for it in items):,.0f}đ)")

print(f"\n=======================================================")
print(f"TỔNG SỐ SKU CẬN DATE ĐƯỢC XÁC ĐỊNH: {len(all_matched_skus)}")
print("=======================================================")
