import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('goc/NHÀ THUỐC THÁNG 8 2026.xlsx', data_only=False)
ws_ds = wb['kpi dược sĩ']

for c_idx in [6, 7, 10, 11, 13]:
    header = ws_ds.cell(2, c_idx).value
    col_let = openpyxl.utils.get_column_letter(c_idx)
    f_val = ws_ds.cell(3, c_idx).value
    txt = f_val.text if hasattr(f_val, 'text') else f_val
    print(f"\n--- Cột {col_let} ({header}) ---")
    print(txt)

wb.close()
