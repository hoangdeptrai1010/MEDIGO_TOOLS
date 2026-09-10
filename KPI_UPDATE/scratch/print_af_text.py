import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx', data_only=False)
ws = wb['kpi dược sĩ']

for r in range(3, 7):
    name = ws.cell(r, 2).value
    role = ws.cell(r, 3).value
    f_val = ws.cell(r, 6).value
    j_val = ws.cell(r, 10).value
    k_val = ws.cell(r, 11).value
    print(f"Row {r} ({name} - {role}):")
    print(f"  Col F text: {f_val.text if hasattr(f_val, 'text') else f_val}")
    print(f"  Col J text: {j_val.text if hasattr(j_val, 'text') else j_val}")
    print(f"  Col K text: {k_val.text if hasattr(k_val, 'text') else k_val}")

wb.close()
