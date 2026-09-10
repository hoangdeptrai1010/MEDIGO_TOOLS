import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws_bl = wb['BẢNG LƯƠNG']

print("--- Staff rows in BẢNG LƯƠNG for Hoàng Lâm Gia Bảo, Phan Công Vũ Tài, Nguyễn Trí Nghĩa, Võ Ngọc Giàu Sang ---")
for r in range(3, ws_bl.max_row + 1):
    stt = ws_bl.cell(r, 1).value
    cn = ws_bl.cell(r, 2).value
    name = ws_bl.cell(r, 3).value
    role = ws_bl.cell(r, 4).value
    h_day = ws_bl.cell(r, 7).value
    h_night = ws_bl.cell(r, 8).value
    days = ws_bl.cell(r, 10).value
    if name in ['Hoàng Lâm Gia Bảo', 'Phan Công Vũ Tài', 'Nguyễn Trí Nghĩa', 'Võ Ngọc Giàu Sang']:
        print(f"Row {r:2d} (STT {stt}): CN={str(cn):<18} | NV={name:<20} | Role={role:<10} | Ca ngày={h_day} | Ca đêm={h_night} | Số ngày={days}")
