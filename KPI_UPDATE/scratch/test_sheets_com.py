import openpyxl, subprocess

wb_src = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx')

for s in wb_src.sheetnames:
    wb_test = openpyxl.Workbook()
    wb_test.remove(wb_test.active)
    # copy sheet
    # simpler: just create workbook with only that sheet
    wb_single = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx')
    for other in wb_single.sheetnames:
        if other != s:
            wb_single.remove(wb_single[other])
    wb_single.save(f'test_sheet_{s[:5]}.xlsx')
    wb_single.close()
    res = subprocess.run(['powershell', '-ExecutionPolicy', 'Bypass', '-File', 'recalc_workbook.ps1', '-FilePath', f'test_sheet_{s[:5]}.xlsx'], capture_output=True, text=True)
    success = 'successfully' in res.stdout
    print(f'Sheet {s}: {"SUCCESS" if success else "FAILED"}')
    if not success:
        for line in res.stdout.splitlines():
            if 'Error' in line:
                print('   ', line)
