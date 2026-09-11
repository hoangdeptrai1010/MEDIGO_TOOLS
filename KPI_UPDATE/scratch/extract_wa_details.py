import openpyxl, sys, io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('d:/MEDIGO/KPI_UPDATE/thang8/output/bangluong_thang8_hoanthien.xlsx', data_only=True)

# 1. Sheet whatsapp
ws_wa = wb['whatsapp']
wa_data = []
for r in range(2, ws_wa.max_row + 1):
    name = ws_wa.cell(r, 1).value
    rev = ws_wa.cell(r, 2).value
    rate = ws_wa.cell(r, 3).value
    bonus = ws_wa.cell(r, 4).value
    if name and rev and float(rev) > 0:
        wa_data.append({
            'name': str(name).strip(),
            'revenue': float(rev),
            'rate': float(rate),
            'bonus': float(bonus)
        })

print("=== NHÂN SỰ CÓ DOANH SỐ & THƯỞNG WHATSAPP THÁNG 8 ===")
for idx, d in enumerate(wa_data, 1):
    print(f"{idx}. {d['name']:25s} | Doanh thu: {d['revenue']:>12,f} đ | Hệ số: {d['rate']*100:.1f}% | Thưởng: {d['bonus']:>10,f} đ")

# 2. Map branch and role from BẢNG LƯƠNG
ws_bl = wb['BẢNG LƯƠNG']
headers = [ws_bl.cell(1, c).value for c in range(1, ws_bl.max_column+1)]
print("\nHeaders in BẢNG LƯƠNG:", [h for h in headers if h][:10])

c_stt = 1
c_cn = next((c for c, h in enumerate(headers, 1) if h and 'chi nhánh' in str(h).lower()), 2)
c_name = next((c for c, h in enumerate(headers, 1) if h and 'tên nhân viên' in str(h).lower()), 3)
c_role = next((c for c, h in enumerate(headers, 1) if h and 'chức danh' in str(h).lower()), 4)

name_info = {}
for r in range(2, ws_bl.max_row + 1):
    cn_v = ws_bl.cell(r, c_cn).value
    nm_v = ws_bl.cell(r, c_name).value
    role_v = ws_bl.cell(r, c_role).value
    if nm_v:
        name_info[str(nm_v).strip()] = {
            'branch': str(cn_v or '').strip(),
            'role': str(role_v or '').strip()
        }

wb.close()

print("\n=== TỔNG HỢP CHI TIẾT WHATSAPP CÓ CHI NHÁNH & CHỨC DANH ===")
for d in wa_data:
    info = name_info.get(d['name'], {'branch': '', 'role': ''})
    d['branch'] = info['branch']
    d['role'] = info['role']
    print(f"CN: {d['branch']:15s} | DS: {d['name']:25s} | Chức danh: {d['role']:10s} | DT Whatsapp: {d['revenue']:>12,f} đ | % Thưởng: {d['rate']*100:.1f}% | Tiền thưởng: {d['bonus']:>10,f} đ")
