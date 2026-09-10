import openpyxl

wb = openpyxl.load_workbook('baocaokpi_thang8_hoanthien.xlsx', data_only=True)
ws = wb['Dự án T8']

for r in range(3, ws.max_row + 1):
    n = ws.cell(r, 2).value
    if 'Hoàng Thanh Thủy' in str(n):
        print(f"Dự án T8 row {r}: {[ws.cell(r, c).value for c in range(1, 15)]}")
