import openpyxl

wb = openpyxl.load_workbook('scratch/test_baug_recalc.xlsx', data_only=True)
ws_bl = wb['BẢNG LƯƠNG']

for r in [55, 56]:
    stt = ws_bl.cell(r, 1).value
    cn = ws_bl.cell(r, 2).value
    name = ws_bl.cell(r, 3).value
    cv = ws_bl.cell(r, 4).value
    gn = ws_bl.cell(r, 5).value
    gd = ws_bl.cell(r, 6).value
    sn = ws_bl.cell(r, 10).value
    sd = ws_bl.cell(r, 11).value
    nc = ws_bl.cell(r, 18).value
    print(f"R{r} ({cn} - {name} - {cv}): Giờ ngày={gn}, Giờ đêm={gd}, Ngày làm={sn}, Ngày đêm={sd}, Ngày công={nc}")
