import openpyxl

for p in [
    'baocaokpi_thang8_hoanthien.xlsx',
    'thang8/DATA/BangChiTietChamCong_thang8.xlsx',
    'thang8/DATA/DanhSachChiTietTraHang_3182026.xlsx',
    'thang8/DATA/TongSoPhatViPhamVaThuongNgay_thang8.xlsx'
]:
    wb = openpyxl.load_workbook(p, read_only=True)
    print(f"=== {p} ===")
    print("Sheets:", wb.sheetnames)
    ws = wb.active
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if any(x is not None for x in r):
            rows.append((i+1, r[:10]))
        if len(rows) >= 4:
            break
    for idx, r in rows:
        print(f"  r{idx}: {r}")
