import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws = wb['BẢNG LƯƠNG']

print("=== KIỂM TRA SỐ CA ĐÊM (CỘT K: CHỈ TÍNH KHI > 4 TIẾNG) TRÊN SHEET BẢNG LƯƠNG ===")
for r in range(3, ws.max_row+1):
    name = str(ws.cell(r, 3).value or '')
    if 'thoa' in name.lower() or ('tâm' in name.lower() and 'đoàn' not in name.lower()) or r in [3, 4, 5, 6, 7, 8, 9]:
        stt = ws.cell(r, 1).value
        cn = ws.cell(r, 2).value
        role = ws.cell(r, 4).value
        gio_ngay = ws.cell(r, 7).value
        gio_dem = ws.cell(r, 8).value
        col_j = ws.cell(r, 10).value
        col_k = ws.cell(r, 11).value
        col_r = ws.cell(r, 18).value
        col_x = ws.cell(r, 24).value
        print(f"Row {r:2d} | STT={stt} | CN: {cn:<14} | Tên: {name:<24} | Vị trí: {role:<8} | Giờ ngày: {gio_ngay:6.2f}h | Giờ đêm: {gio_dem:6.2f}h | Cột J: {col_j} | Cột K (Ca đêm >4h): {col_k} | Phụ cấp đêm (Col X): {col_x:,.0f} đ" if isinstance(col_x, (int, float)) else f"Row {r:2d} | {name} | K={col_k}")
