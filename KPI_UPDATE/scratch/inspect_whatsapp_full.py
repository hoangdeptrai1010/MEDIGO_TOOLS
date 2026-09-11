import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

# 1. Inspect Sheet 'whatsapp' in bangluong_thang8_hoanthien.xlsx
wb_bl = openpyxl.load_workbook('d:/MEDIGO/KPI_UPDATE/thang8/output/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws_wa = wb_bl['whatsapp']
print(f"=== BANGLUONG Sheet 'whatsapp' (max_row={ws_wa.max_row}, max_col={ws_wa.max_column}) ===")
for r in range(1, min(40, ws_wa.max_row + 1)):
    row_vals = [ws_wa.cell(r, c).value for c in range(1, min(20, ws_wa.max_column + 1))]
    if any(row_vals):
        print(f"Row {r:2d}:", [v for v in row_vals if v is not None])

# Also check other columns if there are raw order lists in sheet 'whatsapp'
for c in range(5, ws_wa.max_column + 1):
    col_header = ws_wa.cell(1, c).value or ws_wa.cell(2, c).value
    if col_header:
        print(f"Col {c} Header: {col_header}")

wb_bl.close()

# 2. Inspect Invoices file for Whatsapp notes
wb_hd = openpyxl.load_workbook('d:/MEDIGO/KPI_UPDATE/thang8/DATA/DanhSachChiTietHoaDon_3182026.xlsx', read_only=True)
ws_hd = wb_hd.active
rows_iter = ws_hd.iter_rows(values_only=True)
header = next(rows_iter)
print("\n=== HOA DON HEADERS ===")
for idx, h in enumerate(header):
    print(f"  Col {idx}: {h}")

wb_hd.close()
