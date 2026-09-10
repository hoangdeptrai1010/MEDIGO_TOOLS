import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

f_nt7 = r"d:\MEDIGO\KPI_UPDATE\goc\NHÀ THUỐC THÁNG 7 2026.xlsx"
f_nt8 = r"d:\MEDIGO\KPI_UPDATE\goc\NHÀ THUỐC THÁNG 8 2026.xlsx"
f_plan7 = r"d:\MEDIGO\KPI_UPDATE\plans\KeHoachKPI_2026-07.xlsx"
f_plan8 = r"d:\MEDIGO\KPI_UPDATE\plans\KeHoachKPI_2026-08.xlsx"

print("=== SHEETS IN NHÀ THUỐC THÁNG 7 2026.xlsx ===")
wb7 = openpyxl.load_workbook(f_nt7, data_only=True, read_only=True)
print(wb7.sheetnames)

print("\n=== SHEETS IN NHÀ THUỐC THÁNG 8 2026.xlsx ===")
wb8 = openpyxl.load_workbook(f_nt8, data_only=True, read_only=True)
print(wb8.sheetnames)

print("\n=== SHEETS IN KeHoachKPI_2026-07.xlsx ===")
wb_p7 = openpyxl.load_workbook(f_plan7, data_only=True, read_only=True)
print(wb_p7.sheetnames)

print("\n=== SHEETS IN KeHoachKPI_2026-08.xlsx ===")
wb_p8 = openpyxl.load_workbook(f_plan8, data_only=True, read_only=True)
print(wb_p8.sheetnames)
