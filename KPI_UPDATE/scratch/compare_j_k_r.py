# -*- coding: utf-8 -*-
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/bangluong_thang8_hoanthien.xlsx', data_only=True)
ws = wb['BẢNG LƯƠNG']

print(f"{'STT':<3} | {'Chi nhánh':<15} | {'Họ và tên':<24} | {'Chức danh':<8} | {'Giờ ngày':<9} | {'Giờ đêm':<9} | {'J (G/8)':<8} | {'K (đêm)':<8} | {'R (công)':<8} | {'R - K':<8} | {'J+K > 31?'}")
print("-" * 125)

for r in range(3, 59):
    stt = ws.cell(r, 1).value
    cn = ws.cell(r, 2).value or ''
    name = ws.cell(r, 3).value or ''
    role = ws.cell(r, 4).value or ''
    g = ws.cell(r, 7).value or 0
    h = ws.cell(r, 8).value or 0
    j = ws.cell(r, 10).value or 0
    k = ws.cell(r, 11).value or 0
    r_val = ws.cell(r, 18).value or 0
    
    g_num = float(g) if isinstance(g, (int, float)) else 0.0
    h_num = float(h) if isinstance(h, (int, float)) else 0.0
    j_num = int(j) if isinstance(j, (int, float)) else 0
    k_num = int(k) if isinstance(k, (int, float)) else 0
    r_num = int(r_val) if isinstance(r_val, (int, float)) else 0
    
    r_minus_k = max(0, r_num - k_num)
    is_over = "OVER 31! (" + str(j_num + k_num) + ")" if (j_num + k_num > 31) else "OK (" + str(j_num + k_num) + ")"
    
    print(f"{stt:<3} | {cn:<15} | {name:<24} | {role:<8} | {g_num:<9.1f} | {h_num:<9.1f} | {j_num:<8} | {k_num:<8} | {r_num:<8} | {r_minus_k:<8} | {is_over}")
