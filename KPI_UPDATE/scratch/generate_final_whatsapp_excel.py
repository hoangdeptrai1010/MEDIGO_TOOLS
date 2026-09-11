import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from python_calamine import CalamineWorkbook
import os
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

out_dir = 'd:/MEDIGO/KPI_UPDATE/thang8/output'
os.makedirs(out_dir, exist_ok=True)
out_file = os.path.join(out_dir, 'TONG_HOP_HOA_DON_WHATSAPP_CHUAN.xlsx')

wb = openpyxl.Workbook()

# Styles
font_title = Font(name='Calibri', size=14, bold=True, color='1F4E79')
font_subtitle = Font(name='Calibri', size=10, italic=True, color='595959')
font_header_white = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
font_header_dark = Font(name='Calibri', size=11, bold=True, color='1E293B')
font_bold = Font(name='Calibri', size=11, bold=True)
font_regular = Font(name='Calibri', size=11)

fill_navy = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
fill_soft_blue = PatternFill(start_color='DDEBF7', end_color='DDEBF7', fill_type='solid')
fill_light_blue = PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid')
fill_emerald = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
fill_wa_highlight = PatternFill(start_color='EBF9F1', end_color='EBF9F1', fill_type='solid')

border_thin = Border(
    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
)
border_header = Border(
    left=Side(style='thin', color='B0C4DE'), right=Side(style='thin', color='B0C4DE'),
    top=Side(style='medium', color='1F4E79'), bottom=Side(style='medium', color='1F4E79')
)
border_total = Border(
    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='1F4E79'), bottom=Side(style='double', color='1F4E79')
)

# 1. Quét chính xác các hóa đơn có chữ whatsapp trong Ghi chú
files = [
    'd:/MEDIGO/KPI_UPDATE/thang9/data/phieunhap_tra/hoadon1092026.xlsx',
    'C:/Users/10102/Downloads/DanhSachChiTietHoaDon_KV10092026-105951-828.xlsx',
    'd:/MEDIGO/KPI_UPDATE/thang7/DATAKIOT/DanhSachChiTietHoaDon_KV03092026-134834-071.xlsx',
]

wa_invoices_dict = {}

for f in files:
    if not os.path.exists(f): continue
    wb_cal = CalamineWorkbook.from_path(f)
    sheet = wb_cal.get_sheet_by_name(wb_cal.sheet_names[0])
    rows = sheet.to_python()
    hdr = rows[0]
    note_col = -1
    for idx, h in enumerate(hdr):
        if h and str(h).strip().lower() == 'ghi chú':
            note_col = idx
            break
    
    for r in rows[1:]:
        if note_col != -1 and note_col < len(r) and r[note_col]:
            note_val = str(r[note_col]).strip()
            n_low = note_val.lower()
            if 'whatsapp' in n_low or 'whats app' in n_low or 'watsapp' in n_low:
                code = str(r[1]).strip()
                if code not in wa_invoices_dict:
                    wa_invoices_dict[code] = {
                        'branch': str(r[0] or '').strip(),
                        'code': code,
                        'time': str(r[6])[:19] if len(r)>6 else '',
                        'seller': str(r[20] or '').strip() if len(r)>20 else '',
                        'cust': str(r[12] or '').strip() if len(r)>12 else '',
                        'channel': str(r[21] or '').strip() if len(r)>21 else '',
                        'note': note_val,
                        'need_pay': float(r[43]) if len(r)>43 and r[43] is not None else 0.0
                    }

wa_invoices_list = sorted(wa_invoices_dict.values(), key=lambda x: (x['branch'], x['time']), reverse=True)

# ==============================================================================
# SHEET 1: DANH SÁCH HÓA ĐƠN GHI CHÚ WHATSAPP
# ==============================================================================
ws1 = wb.active
ws1.title = 'Chi tiết Hóa đơn WhatsApp'
ws1.views.sheetView[0].showGridLines = True

ws1.merge_cells('A1:H1')
ws1['A1'] = 'DANH SÁCH CHI TIẾT HÓA ĐƠN CÓ GHI CHÚ "WHATSAPP"'
ws1['A1'].font = font_title
ws1['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws1.row_dimensions[1].height = 28

ws1.merge_cells('A2:H2')
ws1['A2'] = 'Chỉ trích xuất các đơn hàng có ghi chú chữ "WhatsApp" trên phần mềm KiotViet'
ws1['A2'].font = font_subtitle
ws1['A2'].alignment = Alignment(horizontal='center', vertical='center')
ws1.row_dimensions[2].height = 18

headers1 = ['STT', 'Chi nhánh', 'Mã hóa đơn', 'Nội dung Ghi chú WhatsApp', 'Thời gian bán', 'Dược sĩ bán', 'Khách hàng', 'Doanh thu (VNĐ)']
ws1.append([])
ws1.append(headers1)
ws1.row_dimensions[4].height = 26

for c_idx in range(1, len(headers1) + 1):
    cell = ws1.cell(4, c_idx)
    cell.font = font_header_white
    cell.fill = fill_navy
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = border_header

start_row1 = 5
for idx, inv in enumerate(wa_invoices_list, start=1):
    r_idx = start_row1 + idx - 1
    ws1.append([
        idx,
        inv['branch'],
        inv['code'],
        inv['note'],
        inv['time'],
        inv['seller'],
        inv['cust'],
        inv['need_pay']
    ])
    for c_idx in range(1, len(headers1) + 1):
        cell = ws1.cell(r_idx, c_idx)
        cell.font = font_regular
        cell.border = border_thin
        cell.fill = fill_wa_highlight
        if c_idx in (1, 3):
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif c_idx == 4:
            cell.alignment = Alignment(horizontal='left', vertical='center')
            cell.font = font_bold
        elif c_idx == 5:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif c_idx == 8:
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.number_format = '#,##0 "đ"'
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center')

last_row1 = start_row1 + len(wa_invoices_list)
ws1.cell(last_row1, 1, 'TỔNG CỘNG').font = font_bold
ws1.cell(last_row1, 1).alignment = Alignment(horizontal='center', vertical='center')
for c in range(1, 8):
    ws1.cell(last_row1, c).fill = fill_soft_blue
    ws1.cell(last_row1, c).border = border_total

tot_cell = ws1.cell(last_row1, 8, f'=SUM(H{start_row1}:H{last_row1-1})')
tot_cell.font = font_bold
tot_cell.fill = fill_soft_blue
tot_cell.border = border_total
tot_cell.alignment = Alignment(horizontal='right', vertical='center')
tot_cell.number_format = '#,##0 "đ"'

for col in ws1.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws1.column_dimensions[col_letter].width = max(max_len + 4, 14)
ws1.column_dimensions['A'].width = 8
ws1.column_dimensions['C'].width = 16
ws1.column_dimensions['D'].width = 40
ws1.column_dimensions['E'].width = 20
ws1.column_dimensions['H'].width = 18

# ==============================================================================
# SHEET 2: TỔNG HỢP THƯỞNG WHATSAPP LƯƠNG THÁNG 8
# ==============================================================================
ws2 = wb.create_sheet(title='Tổng hợp Lương T8 (WhatsApp)')
ws2.views.sheetView[0].showGridLines = True

ws2.merge_cells('A1:G1')
ws2['A1'] = 'BẢNG TỔNG HỢP DOANH THU & THƯỞNG ĐƠN WHATSAPP TRONG LƯƠNG THÁNG 8'
ws2['A1'].font = font_title
ws2['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws2.row_dimensions[1].height = 28

ws2.merge_cells('A2:G2')
ws2['A2'] = 'Số liệu tổng hợp từ Bảng lương Tháng 8 chính thức của hệ thống Medigo'
ws2['A2'].font = font_subtitle
ws2['A2'].alignment = Alignment(horizontal='center', vertical='center')
ws2.row_dimensions[2].height = 18

headers2 = ['STT', 'Chi nhánh', 'Họ và tên Dược sĩ', 'Chức danh', 'Doanh thu đơn Whatsapp (VNĐ)', 'Tỷ lệ thưởng (%)', 'Tiền thưởng Whatsapp (VNĐ)']
ws2.append([])
ws2.append(headers2)
ws2.row_dimensions[4].height = 26

for c_idx in range(1, 8):
    cell = ws2.cell(4, c_idx)
    cell.font = font_header_white
    cell.fill = fill_navy
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = border_header

staff_wa_t8 = [
    {'stt': 1, 'branch': 'Đỗ Quang Đẩu', 'name': 'Ngô Thị Thanh Thắm', 'role': 'CHT', 'rev': 67046300, 'rate': 0.04},
    {'stt': 2, 'branch': 'Đỗ Quang Đẩu', 'name': 'Hoàng Thanh Thủy', 'role': 'DSXC', 'rev': 92439700, 'rate': 0.04},
    {'stt': 3, 'branch': 'Đỗ Quang Đẩu', 'name': 'Lê Thị Huyền Trân', 'role': 'DSXC', 'rev': 58824400, 'rate': 0.04},
    {'stt': 4, 'branch': 'Đỗ Quang Đẩu', 'name': 'Trần Thiên Phát', 'role': 'DSCD', 'rev': 9377800, 'rate': 0.04},
    {'stt': 5, 'branch': 'Đỗ Quang Đẩu', 'name': 'Phạm Thị Nghĩa Hương', 'role': 'DSCD', 'rev': 8501100, 'rate': 0.04},
    {'stt': 6, 'branch': 'Hàng Bông', 'name': 'Đinh Thị Khánh Ly', 'role': 'DSXC', 'rev': 19489000, 'rate': 0.015},
]

for idx, s in enumerate(staff_wa_t8, start=5):
    ws2.append([s['stt'], s['branch'], s['name'], s['role'], s['rev'], s['rate'], f'=E{idx}*F{idx}'])
    for c_idx in range(1, 8):
        cell = ws2.cell(idx, c_idx)
        cell.font = font_regular
        cell.border = border_thin
        if c_idx in (1, 4):
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif c_idx == 5:
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.number_format = '#,##0 "đ"'
        elif c_idx == 6:
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.number_format = '0.0%'
        elif c_idx == 7:
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.number_format = '#,##0 "đ"'
            cell.font = font_bold
            cell.fill = fill_emerald
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center')

last_row2 = 5 + len(staff_wa_t8)
ws2.cell(last_row2, 1, 'TỔNG CỘNG').font = font_bold
ws2.cell(last_row2, 1).alignment = Alignment(horizontal='center', vertical='center')
for c in range(1, 8):
    ws2.cell(last_row2, c).fill = fill_soft_blue
    ws2.cell(last_row2, c).border = border_total

c_tot_rev = ws2.cell(last_row2, 5, '=SUM(E5:E10)')
c_tot_rev.font = font_bold
c_tot_rev.number_format = '#,##0 "đ"'
c_tot_rev.alignment = Alignment(horizontal='right', vertical='center')

c_tot_bonus = ws2.cell(last_row2, 7, '=SUM(G5:G10)')
c_tot_bonus.font = font_bold
c_tot_bonus.number_format = '#,##0 "đ"'
c_tot_bonus.alignment = Alignment(horizontal='right', vertical='center')

for col in ws2.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws2.column_dimensions[col_letter].width = max(max_len + 5, 14)

# ==============================================================================
# SHEET 3: THỐNG KÊ THEO DƯỢC SĨ BÁN ĐƠN WHATSAPP
# ==============================================================================
ws3 = wb.create_sheet(title='Thống kê DS bán WhatsApp')
ws3.views.sheetView[0].showGridLines = True

ws3.merge_cells('A1:F1')
ws3['A1'] = 'THỐNG KÊ ĐƠN HÀNG GHI CHÚ WHATSAPP THEO DƯỢC SĨ'
ws3['A1'].font = font_title
ws3['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws3.row_dimensions[1].height = 28

headers3 = ['STT', 'Chi nhánh', 'Dược sĩ bán hàng', 'Số lượng đơn WhatsApp', 'Tổng doanh thu (VNĐ)', 'Ghi chú']
ws3.append([])
ws3.append(headers3)
ws3.row_dimensions[3].height = 26

for c_idx in range(1, 7):
    cell = ws3.cell(3, c_idx)
    cell.font = font_header_white
    cell.fill = fill_navy
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = border_header

# Group by seller & branch from wa_invoices_list
seller_agg = {}
for inv in wa_invoices_list:
    key = (inv['branch'], inv['seller'])
    if key not in seller_agg:
        seller_agg[key] = {'count': 0, 'rev': 0.0}
    seller_agg[key]['count'] += 1
    seller_agg[key]['rev'] += inv['need_pay']

sorted_seller_agg = sorted(seller_agg.items(), key=lambda x: (x[0][0], -x[1]['rev']))

start_row3 = 4
for idx, ((branch, seller), data) in enumerate(sorted_seller_agg, start=1):
    r_idx = start_row3 + idx - 1
    ws3.append([idx, branch, seller, data['count'], data['rev'], 'Hóa đơn ghi chú WhatsApp'])
    for c_idx in range(1, 7):
        cell = ws3.cell(r_idx, c_idx)
        cell.font = font_regular
        cell.border = border_thin
        if c_idx in (1, 4):
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif c_idx == 5:
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.number_format = '#,##0 "đ"'
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center')

last_row3 = start_row3 + len(sorted_seller_agg)
ws3.cell(last_row3, 1, 'TỔNG CỘNG').font = font_bold
ws3.cell(last_row3, 1).alignment = Alignment(horizontal='center', vertical='center')
for c in range(1, 7):
    ws3.cell(last_row3, c).fill = fill_soft_blue
    ws3.cell(last_row3, c).border = border_total

c3_cnt = ws3.cell(last_row3, 4, f'=SUM(D{start_row3}:D{last_row3-1})')
c3_cnt.font = font_bold
c3_cnt.alignment = Alignment(horizontal='center', vertical='center')

c3_tot = ws3.cell(last_row3, 5, f'=SUM(E{start_row3}:E{last_row3-1})')
c3_tot.font = font_bold
c3_tot.alignment = Alignment(horizontal='right', vertical='center')
c3_tot.number_format = '#,##0 "đ"'

for col in ws3.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws3.column_dimensions[col_letter].width = max(max_len + 5, 14)

wb.save(out_file)
print(f'--> Đã tạo thành công file chuẩn WhatsApp: {out_file}')
