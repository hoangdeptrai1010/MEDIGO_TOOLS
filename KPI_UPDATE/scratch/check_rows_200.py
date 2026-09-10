import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('thang8/DATA/BangChiTietChamCong_thang8.xlsx', data_only=True)
ws = wb.active

for r in range(200, 230):
    c_name = ws.cell(r, 3).value
    sh = ws.cell(r, 7).value
    br = ws.cell(r, 6).value
    print(f"Row {r}: Name='{c_name}' | Branch='{br}' | Shift='{sh}'")

wb.close()
