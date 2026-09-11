import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import unicodedata
import re

# ==============================================================================
# EXACT ORIGINAL COMPANY STYLING TOKENS (TIMES NEW ROMAN & AUTHENTIC PASTEL PALETTE)
# 100% EXTRACTED FROM 'goc/NHÀ THUỐC THÁNG 7 2026.xlsx'
# ==============================================================================
FONT_NAME = 'Times New Roman'

# Fonts
font_header = Font(name=FONT_NAME, size=11, bold=True, color='000000')
font_header_12 = Font(name=FONT_NAME, size=12, bold=True, color='000000')
font_header_12_italic = Font(name=FONT_NAME, size=12, bold=True, italic=True, color='000000')
font_bold = Font(name=FONT_NAME, size=11, bold=True, color='000000')
font_regular = Font(name=FONT_NAME, size=11, bold=False, color='000000')
font_italic_bold = Font(name=FONT_NAME, size=11, bold=True, italic=True, color='000000')
font_date = Font(name=FONT_NAME, size=12, bold=True, color='000000')

# Specific fonts for Hot Bill
font_hot_bill_header = Font(name="Segoe UI", size=10, bold=True, color='000000')
font_hot_bill_data = Font(name='Arial', size=10, bold=False, color='000000')

# Authentic Pastel Color Palette extracted 100% from Gốc Tháng 7:
fill_pastel_yellow = PatternFill(start_color='FFFFF2CC', end_color='FFFFF2CC', fill_type='solid') # #FFF2CC
fill_pastel_pink   = PatternFill(start_color='FFEAD1DC', end_color='FFEAD1DC', fill_type='solid') # #EAD1DC
fill_pastel_green  = PatternFill(start_color='FFD9EAD3', end_color='FFD9EAD3', fill_type='solid') # #D9EAD3
fill_pastel_blue   = PatternFill(start_color='FFCFE2F3', end_color='FFCFE2F3', fill_type='solid') # #CFE2F3
fill_nt_green      = PatternFill(start_color='FFB6D7A8', end_color='FFB6D7A8', fill_type='solid') # #B6D7A8
fill_nt_yellow     = PatternFill(start_color='FFFFE599', end_color='FFFFE599', fill_type='solid') # #FFE599
fill_nt_blue       = PatternFill(start_color='FFC9DAF8', end_color='FFC9DAF8', fill_type='solid') # #C9DAF8
fill_nt_orange     = PatternFill(start_color='FFF9CB9C', end_color='FFF9CB9C', fill_type='solid') # #F9CB9C
fill_nt_salmon     = PatternFill(start_color='FFE6B8AF', end_color='FFE6B8AF', fill_type='solid') # #E6B8AF
fill_nt_orchid     = PatternFill(start_color='FFD5A6BD', end_color='FFD5A6BD', fill_type='solid') # #D5A6BD
fill_nt_coral      = PatternFill(start_color='FFDD7E6B', end_color='FFDD7E6B', fill_type='solid') # #DD7E6B
fill_transparent   = PatternFill(fill_type=None)

# Borders
border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
border_header = Border(
    left=Side(style='thin', color='BFBFBF'),
    right=Side(style='thin', color='BFBFBF'),
    top=Side(style='thin', color='BFBFBF'),
    bottom=Side(style='thin', color='BFBFBF')
)

def strip_accents_lower(s):
    if not s:
        return ""
    nfkd = unicodedata.normalize('NFKD', str(s))
    res = "".join([c for c in nfkd if not unicodedata.combining(c)])
    return res.lower().replace('đ', 'd').replace('Đ', 'd').strip()

def safe_merge_cells(ws, cell_range):
    for rng in list(ws.merged_cells.ranges):
        if str(rng) == cell_range:
            return
    try:
        ws.merge_cells(cell_range)
    except Exception:
        pass

def apply_full_kpi_styling(wb):
    """
    Applies 100% exact authentic Medigo styling (Times New Roman, Pastel Fills, Column Widths, Freeze Panes, Merged Headers)
    matching the 4 benchmark sheets from 'goc/NHÀ THUỐC THÁNG 7 2026.xlsx'.
    """
    for sname in wb.sheetnames:
        ws = wb[sname]
        try:
            ws.views.sheetView[0].showGridLines = True
        except Exception:
            pass

        norm_name = strip_accents_lower(sname)
        total_cols = max(ws.max_column, 38)

        # -------------------------------------------------------------
        # 1. SHEET: 'data'
        # -------------------------------------------------------------
        if norm_name == 'data':
            ws.freeze_panes = None
            ws.row_dimensions[1].height = 27.75
            
            data_col_widths = {
                'A': 16.33, 'B': 23.89, 'C': 18.11, 'D': 13.0,
                'E': 13.0, 'F': 13.0, 'G': 13.0, 'H': 13.0,
                'I': 11.78, 'J': 13.11
            }
            for c_i in range(1, total_cols + 1):
                col_l = get_column_letter(c_i)
                ws.column_dimensions[col_l].width = data_col_widths.get(col_l, 13.0)

            for c_i in range(1, total_cols + 1):
                cell = ws.cell(1, c_i)
                cell.font = font_header
                cell.fill = fill_transparent
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = border_header

            for r in range(2, ws.max_row + 1):
                if not ws.cell(r, 1).value and not ws.cell(r, 2).value:
                    continue
                ws.row_dimensions[r].height = 15.75
                for c_i in range(1, total_cols + 1):
                    cell = ws.cell(r, c_i)
                    cell.font = font_regular
                    cell.border = border_thin
                    cell.fill = fill_transparent
                    if c_i in (1, 2):
                        cell.alignment = Alignment(horizontal='left', vertical='center')
                    elif c_i in (3, 5, 7):
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                        cell.number_format = '#,##0'
                    elif c_i == 9: # Ngày
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.alignment = Alignment(horizontal='right', vertical='center')
                        cell.number_format = '#,##0'

        # -------------------------------------------------------------
        # 2. SHEET: 'kpi dược sĩ'
        # -------------------------------------------------------------
        elif 'kpi duoc si' in norm_name:
            ws.freeze_panes = 'D3'
            ws.row_dimensions[1].height = 19.5
            ws.row_dimensions[2].height = 30.75
            
            # Exact widths from goc T7
            ds_col_widths = {
                'A': 15.66, 'B': 25.0, 'C': 13.11, 'D': 14.44, 'E': 14.44,
                'F': 18.66, 'G': 15.0, 'H': 19.78, 'I': 14.44, 'J': 11.78,
                'K': 11.22, 'L': 10.78, 'M': 11.22, 'N': 15.11, 'O': 10.78,
                'P': 13.0, 'Q': 13.0, 'R': 10.78, 'S': 13.0, 'T': 13.0,
                'U': 15.66, 'V': 15.66, 'W': 13.0, 'X': 13.0, 'Y': 13.0,
                'Z': 13.0, 'AA': 13.0, 'AB': 13.0, 'AC': 12.66, 'AD': 13.0,
                'AE': 13.0
            }
            for c_i in range(1, total_cols + 1):
                col_l = get_column_letter(c_i)
                ws.column_dimensions[col_l].width = ds_col_widths.get(col_l, 13.0)

            # Row 1 Merges: A1:B1, O1:Q1, R1:T1
            safe_merge_cells(ws, 'A1:B1')
            safe_merge_cells(ws, 'O1:Q1')
            safe_merge_cells(ws, 'R1:T1')

            # Row 1 Styling
            c1_note = ws.cell(1, 1)
            c1_note.value = c1_note.value or 'Dữ liệu cập nhật đến ngày'
            c1_note.font = font_header_12_italic
            c1_note.alignment = Alignment(horizontal='center', vertical='center')
            
            c1_date = ws.cell(1, 3)
            c1_date.font = font_header_12
            c1_date.alignment = Alignment(horizontal='center', vertical='center')

            # Row 1 Subtotals (O1:Q1 & R1:T1)
            cell_o1 = ws.cell(1, 15)
            if not cell_o1.value:
                cell_o1.value = 'Mức hoành thành KPI '
            cell_o1.font = font_header
            cell_o1.fill = fill_pastel_yellow
            cell_o1.alignment = Alignment(horizontal='center', vertical='center')

            cell_r1 = ws.cell(1, 18)
            if not cell_r1.value:
                cell_r1.value = 'Mức hoành thành KPI '
            cell_r1.font = font_header
            cell_r1.fill = fill_pastel_yellow
            cell_r1.alignment = Alignment(horizontal='center', vertical='center')

            for c_i in (16, 17, 19, 20):
                ws.cell(1, c_i).fill = fill_pastel_yellow

            # Row 2 Headers
            for c_i in range(1, total_cols + 1):
                cell = ws.cell(2, c_i)
                cell.font = font_header
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = border_header
                
                if c_i in (1, 2, 3):
                    cell.fill = fill_pastel_yellow
                elif c_i in range(4, 10):
                    cell.fill = fill_pastel_pink
                elif c_i in range(10, 15):
                    cell.fill = fill_pastel_green
                elif c_i in (15, 16, 17, 18, 19, 20):
                    cell.fill = fill_nt_orange
                elif c_i in (21, 22):
                    cell.fill = fill_nt_salmon
                elif c_i in (23, 24, 25):
                    cell.fill = fill_pastel_yellow
                elif c_i in (26, 27, 28):
                    cell.fill = fill_pastel_green
                else:
                    cell.fill = fill_pastel_pink

            # Row 3+ Data Rows
            for r in range(3, ws.max_row + 1):
                if not ws.cell(r, 1).value and not ws.cell(r, 2).value:
                    continue
                ws.row_dimensions[r].height = 13.8
                for c_i in range(1, total_cols + 1):
                    cell = ws.cell(r, c_i)
                    cell.font = font_regular
                    cell.border = border_thin
                    cell.fill = fill_transparent
                    
                    # Alignments & Number formats
                    if c_i in (1, 2):
                        cell.alignment = Alignment(horizontal='left', vertical='center')
                    elif c_i in (3, 10, 12, 15, 16, 17, 18, 19, 20, 23, 26, 29):
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.alignment = Alignment(horizontal='right', vertical='center')

                    if c_i in (10, 15, 16, 17, 18, 19, 20):
                        cell.number_format = '0.0%' if c_i == 10 else '0%'
                    elif c_i in (4, 5, 6, 7, 8, 9, 11, 13, 14, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30, 31):
                        cell.number_format = '#,##0'

                    # Col 14 (Thưởng Dự án + KPI + HH): Fill Soft Green + Bold
                    if c_i == 14:
                        cell.fill = fill_pastel_green
                        cell.font = font_bold

        # -------------------------------------------------------------
        # 3. SHEET: 'kpi nhà thuốc'
        # -------------------------------------------------------------
        elif 'kpi nha thuoc' in norm_name:
            ws.freeze_panes = 'C1'
            ws.row_dimensions[1].height = 20.25
            ws.row_dimensions[2].height = 39.75
            
            # Exact widths from goc T7
            nt_col_widths = {
                'A': 22.89, 'B': 15.0, 'C': 16.11, 'D': 15.66, 'E': 16.11,
                'F': 13.11, 'G': 13.0, 'H': 12.78, 'I': 12.66, 'J': 13.0,
                'K': 12.66, 'L': 12.66, 'M': 13.0, 'N': 13.0, 'O': 13.0,
                'P': 13.0, 'Q': 13.0, 'R': 13.0, 'S': 13.0, 'T': 13.0,
                'U': 13.0, 'V': 15.89, 'W': 14.25, 'X': 12.63, 'Y': 13.0,
                'Z': 13.0, 'AA': 13.0, 'AB': 13.0, 'AC': 13.0, 'AD': 13.0,
                'AE': 13.0, 'AF': 13.0
            }
            for c_i in range(1, total_cols + 1):
                col_l = get_column_letter(c_i)
                ws.column_dimensions[col_l].width = nt_col_widths.get(col_l, 13.0)

            # Row 1 Merges: D1:E1
            safe_merge_cells(ws, 'D1:E1')

            c1_note = ws.cell(1, 4)
            c1_note.value = c1_note.value or 'Dữ liệu cập nhật đến ngày'
            c1_note.font = font_header_12
            c1_note.alignment = Alignment(horizontal='center', vertical='center')

            for c_dt in (6, 20):
                cell = ws.cell(1, c_dt)
                cell.font = font_header_12
                cell.alignment = Alignment(horizontal='center', vertical='center')

            # Row 2 Headers
            for c_i in range(1, total_cols + 1):
                cell = ws.cell(2, c_i)
                cell.font = font_header
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = border_header
                
                if c_i in (1, 2, 3):
                    cell.fill = fill_nt_green
                elif c_i in (4, 5):
                    cell.fill = fill_nt_yellow
                elif c_i in (6, 7):
                    cell.fill = fill_nt_green
                elif c_i in range(8, 13):
                    cell.fill = fill_nt_blue
                elif c_i in range(13, 17):
                    cell.fill = fill_pastel_pink
                elif c_i in range(17, 21):
                    cell.fill = fill_nt_green
                elif c_i == 21:
                    cell.fill = fill_nt_salmon
                elif c_i in range(22, 27):
                    cell.fill = fill_nt_yellow
                elif c_i in (27, 28, 29):
                    cell.fill = fill_nt_orchid
                elif c_i in (30, 31, 32):
                    cell.fill = fill_nt_coral
                else:
                    cell.fill = fill_transparent

            # Row 3+ Data Rows
            for r in range(3, ws.max_row + 1):
                if not ws.cell(r, 1).value and not ws.cell(r, 2).value:
                    continue
                ws.row_dimensions[r].height = 15.75
                for c_i in range(1, total_cols + 1):
                    cell = ws.cell(r, c_i)
                    cell.font = font_regular
                    cell.border = border_thin
                    cell.fill = fill_transparent
                    
                    if c_i == 1:
                        cell.alignment = Alignment(horizontal='left', vertical='center')
                    elif c_i in (2, 6, 8, 13, 17, 22, 23, 24, 27, 30):
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.alignment = Alignment(horizontal='right', vertical='center')

                    if c_i in (6, 22):
                        cell.number_format = '0.00%'
                    elif c_i in (8, 13, 17):
                        cell.number_format = '0%'
                    elif c_i in (3, 4, 5, 7, 9, 10, 11, 12, 14, 15, 16, 18, 19, 20, 21, 24, 25, 26, 27, 28, 29, 30, 31, 32):
                        cell.number_format = '#,##0'

                    if c_i == 7: # Thưởng CHT
                        cell.font = font_bold

        # -------------------------------------------------------------
        # 4. SHEET: 'Dự án T...'
        # -------------------------------------------------------------
        elif norm_name.startswith('du an') or norm_name.startswith('dự án'):
            ws.freeze_panes = 'C3'
            ws.row_dimensions[1].height = 19.5
            ws.row_dimensions[2].height = 33.75
            
            # Exact widths from goc T7
            proj_col_widths = {
                'A': 15.66, 'B': 20.78, 'C': 13.0, 'D': 14.22, 'E': 13.0,
                'F': 13.0, 'G': 13.0, 'H': 13.0, 'I': 15.44, 'J': 14.22,
                'K': 13.0, 'L': 13.0, 'M': 14.22
            }
            for c_i in range(1, total_cols + 1):
                col_l = get_column_letter(c_i)
                ws.column_dimensions[col_l].width = proj_col_widths.get(col_l, 13.0)

            c1_date = ws.cell(1, 1)
            c1_date.font = font_header_12
            c1_date.fill = fill_pastel_yellow

            for c_i in range(2, 4):
                ws.cell(1, c_i).fill = fill_pastel_yellow
            for c_i in range(4, 14):
                ws.cell(1, c_i).fill = fill_pastel_green

            # Row 2 Headers
            for c_i in range(1, total_cols + 1):
                cell = ws.cell(2, c_i)
                cell.font = font_header
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = border_header
                
                if c_i in (1, 2, 3):
                    cell.fill = fill_pastel_yellow
                elif c_i in range(4, 14):
                    cell.fill = fill_pastel_green
                else:
                    cell.fill = fill_pastel_pink

            # Row 3+ Data Rows
            for r in range(3, ws.max_row + 1):
                if not ws.cell(r, 1).value and not ws.cell(r, 2).value:
                    continue
                ws.row_dimensions[r].height = 13.8
                for c_i in range(1, total_cols + 1):
                    cell = ws.cell(r, c_i)
                    cell.font = font_regular
                    cell.border = border_thin
                    cell.fill = fill_transparent
                    
                    if c_i == 1:
                        cell.fill = fill_pastel_blue
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    elif c_i == 2:
                        cell.fill = fill_pastel_blue
                        cell.alignment = Alignment(horizontal='left', vertical='center')
                    elif c_i == 3:
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.alignment = Alignment(horizontal='right', vertical='center')
                        cell.number_format = '#,##0'

                    # Total dự án (Col 12 or 13)
                    if c_i in (12, 13) and ('total' in str(ws.cell(2, c_i).value or '').lower() or c_i == 12 and not ws.cell(2, 13).value):
                        cell.fill = fill_pastel_blue
                        cell.font = font_bold
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    elif c_i == 13 and 'total' in str(ws.cell(2, c_i).value or '').lower():
                        cell.fill = fill_pastel_blue
                        cell.font = font_bold
                        cell.alignment = Alignment(horizontal='center', vertical='center')

        # -------------------------------------------------------------
        # 5. SHEET: 'Hot Bill HN'
        # -------------------------------------------------------------
        elif 'hot bill' in norm_name:
            ws.freeze_panes = None
            ws.row_dimensions[1].height = 12.75
            hot_col_widths = {
                'A': 12.5, 'B': 18.0, 'C': 16.75, 'D': 18.38,
                'E': 16.25, 'F': 29.5, 'G': 19.5, 'H': 62.5
            }
            for col_l, w in hot_col_widths.items():
                ws.column_dimensions[col_l].width = w

            for c_i in range(1, 9):
                cell = ws.cell(1, c_i)
                cell.font = font_hot_bill_header
                cell.fill = fill_transparent
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = border_header

            for r in range(2, ws.max_row + 1):
                if not ws.cell(r, 1).value:
                    continue
                ws.row_dimensions[r].height = 39.0
                for c_i in range(1, 9):
                    cell = ws.cell(r, c_i)
                    cell.font = font_hot_bill_data
                    cell.border = border_thin
                    cell.fill = fill_transparent
                    if c_i in (1, 2, 3, 4, 5):
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    elif c_i in (6, 7):
                        cell.alignment = Alignment(horizontal='right', vertical='center')
                        cell.number_format = '#,##0'
                    else:
                        cell.alignment = Alignment(horizontal='left', vertical='center')
