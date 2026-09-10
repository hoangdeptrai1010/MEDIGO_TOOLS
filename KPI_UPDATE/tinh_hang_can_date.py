#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
================================================================================
TOOL TÍNH CẬN DATE & THƯỞNG 5% - PHIÊN BẢN CHÍNH THỨC
================================================================================
Logic nghiệp vụ kép:
  1. SKU phải nằm trong DANH MỤC CẬN DATE ĐƯỢC DUYỆT
  2. HSD phải nằm trong khoảng 6 tháng (tính theo lịch) từ ngày bán

Hai hướng tính:
  BACKWARD (--backward, mặc định):
    EDATE(purchase_date, -6) <= expiry_date <= purchase_date
    → Áp dụng cho hàng đã hết hạn hoặc sắp hết hạn tại thời điểm bán.

  FORWARD (--forward):
    purchase_date <= expiry_date <= EDATE(purchase_date, +6)
    → Áp dụng cho hàng có HSD trong 6 tháng tới tính từ ngày bán.
    → Phù hợp dữ liệu thực tế bảng lương tháng 7/2026.

  bonus = product_amount × 5%  (nếu is_near_expiry = TRUE)

Ưu tiên: ACCURACY > AUTOMATION
  - Không tự suy diễn logic
  - Không tự điều chỉnh kết quả
  - Mọi kết quả phải truy ngược được từng dòng
================================================================================
"""

import os
import sys
import calendar
import datetime
import argparse
from typing import Dict, List, Tuple, Any, Optional, Set
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ==============================================================================
# MODULE 1: DATE HELPERS
# ==============================================================================

def edate(dt: datetime.date, months: int) -> Optional[datetime.date]:
    """
    Tính mốc ngày theo lịch chính xác, tương đương EDATE() trong Excel.
    Xử lý tự động ngày cuối tháng (VD: 31/08 - 6 tháng = 28/02).
    TUYỆT ĐỐI KHÔNG dùng 180 ngày.
    """
    if dt is None:
        return None
    year = dt.year + (dt.month + months - 1) // 12
    month = (dt.month + months - 1) % 12 + 1
    max_days = calendar.monthrange(year, month)[1]
    day = min(dt.day, max_days)
    return datetime.date(year, month, day)


def parse_date(val: Any) -> Tuple[Optional[datetime.date], str]:
    """
    Parse ngày tháng từ nhiều định dạng:
      - datetime.datetime / datetime.date
      - 'DD/MM/YYYY', 'YYYY-MM-DD', 'DD-MM-YYYY'
      - Text kèm giờ: 'YYYY-MM-DD HH:MM:SS'
    Returns: (date_object, status)
      status: 'OK', 'MISSING', 'INVALID'
    """
    if val is None or (isinstance(val, str) and val.strip() == ''):
        return None, "MISSING"

    if isinstance(val, datetime.datetime):
        return val.date(), "OK"
    if isinstance(val, datetime.date):
        return val, "OK"

    s = str(val).strip()
    # Loại bỏ phần giờ nếu có
    if ' ' in s:
        s = s.split(' ')[0]

    formats = ['%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%Y/%m/%d', '%m/%d/%Y']
    for fmt in formats:
        try:
            return datetime.datetime.strptime(s, fmt).date(), "OK"
        except ValueError:
            continue

    return None, "INVALID"


# ==============================================================================
# MODULE 2: DATA LOADING
# ==============================================================================

# Mapping tên cột -> key nội bộ (semantic matching)
COLUMN_ALIASES = {
    'order_id':      ['Mã hóa đơn', 'Mã HD', 'OrderID', 'Order ID'],
    'branch':        ['Chi nhánh', 'Branch', 'Cửa hàng'],
    'purchase_date': ['Thời gian', 'Ngày mua', 'Ngày bán', 'PurchaseDate', 'Date'],
    'seller':        ['Người bán', 'Nhân viên bán', 'Seller', 'Người tạo'],
    'product_id':    ['Mã hàng', 'SKU', 'Mã SP', 'ProductID'],
    'product_name':  ['Tên hàng', 'Tên sản phẩm', 'ProductName'],
    'expiry_date':   ['Hạn sử dụng', 'HSD', 'ExpiryDate', 'Ngày hết hạn'],
    'quantity':      ['Số lượng', 'SL', 'Quantity'],
    'unit_price':    ['Đơn giá', 'Giá bán', 'UnitPrice'],
    'amount':        ['Thành tiền', 'Tổng tiền', 'Amount', 'Total'],
}


def _find_column_index(headers: tuple, key: str) -> Optional[int]:
    """Tìm index của cột dựa trên semantic matching."""
    aliases = COLUMN_ALIASES.get(key, [])
    for idx, h in enumerate(headers):
        if h is None:
            continue
        h_clean = str(h).strip()
        for alias in aliases:
            if h_clean.lower() == alias.lower():
                return idx
    return None


def load_sales_data(filepath: str) -> Tuple[List[dict], dict]:
    """
    Đọc file hóa đơn Excel.
    Returns: (list_of_row_dicts, col_map)
    """
    print(f"\n--> Đang đọc dữ liệu hóa đơn: {filepath}")

    wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
    ws = wb.active
    headers = next(ws.iter_rows(max_row=1, values_only=True))

    # Auto-detect columns
    col_map = {}
    for key in COLUMN_ALIASES:
        idx = _find_column_index(headers, key)
        if idx is not None:
            col_map[key] = idx
            alias_name = headers[idx]
            print(f"   + {key:18s} -> Cột {idx:2d} [{alias_name}]")

    # Validate required columns
    required = ['seller', 'purchase_date', 'product_id', 'expiry_date', 'amount']
    missing = [k for k in required if k not in col_map]
    if missing:
        wb.close()
        raise ValueError(f"Không tìm thấy các cột bắt buộc: {missing}")

    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r is None:
            continue
        row = {}
        for key, idx in col_map.items():
            if idx < len(r):
                row[key] = r[idx]
            else:
                row[key] = None
        rows.append(row)

    wb.close()
    print(f"   Tổng số dòng đọc được: {len(rows):,}")
    return rows, col_map


def load_approved_skus(filepath: str) -> Set[str]:
    """
    Đọc danh mục SKU cận date được duyệt từ file Excel.
    File cần có cột 'SKU' hoặc 'Mã hàng' và cột 'Trạng thái' (nếu có).
    Chỉ load SKU có trạng thái ACTIVE (hoặc tất cả nếu không có cột trạng thái).
    """
    print(f"\n--> Đang đọc danh mục SKU cận date: {filepath}")

    wb = openpyxl.load_workbook(filepath, data_only=True)

    # Tìm sheet phù hợp
    target_sheet = None
    for sn in wb.sheetnames:
        if 'cận' in sn.lower() or 'date' in sn.lower() or 'danh mục' in sn.lower() or 'sku' in sn.lower():
            target_sheet = sn
            break
    if target_sheet is None:
        target_sheet = wb.sheetnames[0]

    ws = wb[target_sheet]
    print(f"   Sheet: {target_sheet}")

    # Tìm cột SKU và Trạng thái
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    sku_col = None
    status_col = None
    for idx, h in enumerate(headers):
        if h is None:
            continue
        h_lower = str(h).strip().lower()
        if h_lower in ['sku', 'mã hàng', 'mã sp', 'productid']:
            sku_col = idx
        if h_lower in ['trạng thái', 'status', 'trang thai']:
            status_col = idx

    if sku_col is None:
        wb.close()
        raise ValueError("Không tìm thấy cột SKU trong file danh mục cận date!")

    approved = set()
    for r in range(2, ws.max_row + 1):
        sku_val = ws.cell(r, sku_col + 1).value
        if sku_val is None:
            continue

        sku = str(sku_val).strip()
        if not sku:
            continue

        # Kiểm tra trạng thái nếu có
        if status_col is not None:
            status = str(ws.cell(r, status_col + 1).value or '').strip().upper()
            if status and status != 'ACTIVE':
                continue

        approved.add(sku)

    wb.close()
    print(f"   Tổng số SKU được duyệt: {len(approved)}")
    return approved


def load_payroll(filepath: str, sheet_name: str = 'Cận date') -> Dict[str, float]:
    """
    Đọc dữ liệu cận date từ bảng lương để đối soát.
    Returns: dict {seller_name: near_date_sales_amount}
    """
    print(f"\n--> Đang đọc bảng lương để đối soát: {filepath}")

    wb = openpyxl.load_workbook(filepath, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        print(f"   ⚠️ Không tìm thấy sheet '{sheet_name}' trong bảng lương.")
        return {}

    ws = wb[sheet_name]
    payroll = {}

    # Tìm cột tên nhân viên và doanh thu cận date
    # Trong bảng target: cột H = Tên nhân viên, cột I = Doanh thu cận date
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 8).value  # Cột H
        amount = ws.cell(r, 9).value  # Cột I
        if name and amount is not None:
            payroll[str(name).strip()] = float(amount)

    wb.close()
    print(f"   Tổng số nhân viên trong bảng lương: {len(payroll)}")
    return payroll


# ==============================================================================
# MODULE 3: NORMALIZATION
# ==============================================================================

def normalize_sku(sku: Any) -> str:
    """Chuẩn hóa SKU: strip whitespace, giữ nguyên format (không mất số 0 đầu)."""
    if sku is None:
        return ''
    return str(sku).strip()


def normalize_data(raw_rows: List[dict]) -> List[dict]:
    """
    Chuẩn hóa toàn bộ dữ liệu:
      - SKU: strip whitespace
      - Seller: strip whitespace
      - Dates: parse thành datetime.date
      - Amount: parse thành float
    """
    normalized = []
    for row in raw_rows:
        item = {}

        item['order_id'] = str(row.get('order_id', '') or '').strip()
        item['branch'] = str(row.get('branch', '') or '').strip()
        item['seller'] = str(row.get('seller', '') or '').strip()
        item['sku'] = normalize_sku(row.get('product_id'))
        item['product_name'] = str(row.get('product_name', '') or '').strip()

        # Parse dates
        pur_dt, pur_status = parse_date(row.get('purchase_date'))
        exp_dt, exp_status = parse_date(row.get('expiry_date'))
        item['purchase_date'] = pur_dt
        item['purchase_date_status'] = pur_status
        item['expiry_date'] = exp_dt
        item['expiry_date_status'] = exp_status

        # Parse amounts
        try:
            item['quantity'] = float(row.get('quantity') or 0)
        except (ValueError, TypeError):
            item['quantity'] = 0
        try:
            item['unit_price'] = float(row.get('unit_price') or 0)
        except (ValueError, TypeError):
            item['unit_price'] = 0
        try:
            item['amount'] = float(row.get('amount') or 0)
        except (ValueError, TypeError):
            item['amount'] = 0

        normalized.append(item)

    return normalized


# ==============================================================================
# MODULE 4: BUSINESS LOGIC
# ==============================================================================

def check_near_expiry(
    sku: str,
    purchase_date: Optional[datetime.date],
    expiry_date: Optional[datetime.date],
    purchase_date_status: str,
    expiry_date_status: str,
    approved_skus: Set[str],
    amount: float,
    rate: float = 0.05,
    direction: str = 'backward'
) -> dict:
    """
    Kiểm tra xem một dòng sản phẩm có phải hàng cận date được tính thưởng hay không.

    Logic kép:
      1. SKU phải nằm trong danh mục được duyệt (approved_skus)
      2. Điều kiện thời gian (tùy direction):
         BACKWARD: EDATE(purchase_date, -6) <= expiry_date <= purchase_date
         FORWARD:  purchase_date <= expiry_date <= EDATE(purchase_date, +6)

    Returns dict với các key:
      is_near_expiry, approved_sku, within_6months, status, bonus, reason
    """
    result = {
        'is_near_expiry': False,
        'approved_sku': False,
        'within_6months': False,
        'start_date': None,
        'status': '',
        'bonus': 0.0,
        'reason': ''
    }

    # Kiểm tra dữ liệu thiếu
    if purchase_date_status == 'MISSING':
        result['status'] = 'MISSING_PURCHASE_DATE'
        result['reason'] = 'Thiếu ngày mua hàng'
        return result

    if purchase_date_status == 'INVALID':
        result['status'] = 'INVALID_PURCHASE_DATE'
        result['reason'] = 'Ngày mua không hợp lệ'
        return result

    if expiry_date_status == 'MISSING':
        result['status'] = 'MISSING_EXPIRY_DATE'
        result['reason'] = 'Thiếu hạn sử dụng (HSD)'
        return result

    if expiry_date_status == 'INVALID':
        result['status'] = 'INVALID_EXPIRY_DATE'
        result['reason'] = 'HSD không hợp lệ'
        return result

    # STEP 5: Kiểm tra SKU trong danh mục
    sku_approved = sku in approved_skus
    result['approved_sku'] = sku_approved

    # STEP 4: Tính mốc 6 tháng theo hướng
    if direction == 'forward':
        range_start = purchase_date
        range_end = edate(purchase_date, 6)
    else:  # backward
        range_start = edate(purchase_date, -6)
        range_end = purchase_date

    result['start_date'] = range_start
    result['range_end'] = range_end

    # Kiểm tra HSD nằm trong khoảng
    within = range_start <= expiry_date <= range_end
    result['within_6months'] = within

    pur_str = purchase_date.strftime('%d/%m/%Y')
    exp_str = expiry_date.strftime('%d/%m/%Y')
    start_str = range_start.strftime('%d/%m/%Y')
    end_str = range_end.strftime('%d/%m/%Y')

    # STEP 6: Xác định kết quả
    if sku_approved and within:
        result['is_near_expiry'] = True
        result['status'] = 'CẬN DATE'
        result['bonus'] = amount * rate
        result['reason'] = (
            f"CẬN DATE: SKU thuộc danh mục được duyệt "
            f"và HSD {exp_str} nằm trong khoảng {start_str} - {end_str}"
        )
    elif not sku_approved and within:
        result['status'] = 'KHÔNG TÍNH'
        result['reason'] = f"SKU không thuộc danh mục cận date được duyệt"
    elif sku_approved and not within:
        result['status'] = 'KHÔNG TÍNH'
        if expiry_date > range_end:
            result['reason'] = (
                f"HSD {exp_str} ngoài khoảng {start_str} - {end_str} "
                f"(quá xa)"
            )
        elif expiry_date < range_start:
            result['reason'] = (
                f"HSD {exp_str} trước mốc {start_str} "
                f"(quá hạn)"
            )
        else:
            result['reason'] = f"HSD ngoài khoảng {start_str} - {end_str}"
    else:
        result['status'] = 'KHÔNG TÍNH'
        result['reason'] = 'SKU không thuộc danh mục và/hoặc HSD ngoài khoảng'

    return result


# ==============================================================================
# MODULE 5: BONUS CALCULATION
# ==============================================================================

def calculate_bonus(items: List[dict], approved_skus: Set[str], rate: float = 0.05,
                    direction: str = 'backward') -> List[dict]:
    """
    Tính thưởng 5% cho từng dòng sản phẩm.
    QUAN TRỌNG: Tính từng dòng, KHÔNG tính tổng đơn hàng.
    """
    results = []
    for item in items:
        result = check_near_expiry(
            sku=item['sku'],
            purchase_date=item['purchase_date'],
            expiry_date=item['expiry_date'],
            purchase_date_status=item['purchase_date_status'],
            expiry_date_status=item['expiry_date_status'],
            approved_skus=approved_skus,
            amount=item['amount'],
            rate=rate,
            direction=direction
        )

        # Merge item data with result
        enriched = {**item, **result}
        results.append(enriched)

    return results


# ==============================================================================
# MODULE 6: AGGREGATION
# ==============================================================================

def aggregate_by_seller(results: List[dict]) -> List[dict]:
    """
    Tổng hợp theo người bán.
    Returns list of dicts sorted by seller name.
    """
    sellers = {}
    for item in results:
        seller = item['seller']
        if not seller:
            continue

        if seller not in sellers:
            sellers[seller] = {
                'seller': seller,
                'order_count': set(),
                'near_expiry_count': 0,
                'total_near_expiry_sales': 0.0,
                'total_bonus': 0.0,
            }

        sellers[seller]['order_count'].add(item.get('order_id', ''))

        if item['is_near_expiry']:
            sellers[seller]['near_expiry_count'] += 1
            sellers[seller]['total_near_expiry_sales'] += item['amount']
            sellers[seller]['total_bonus'] += item['bonus']

    # Convert order_count sets to counts
    summary = []
    for s in sorted(sellers.values(), key=lambda x: x['seller']):
        s['order_count'] = len(s['order_count'])
        if s['near_expiry_count'] > 0:
            summary.append(s)

    return summary


# ==============================================================================
# MODULE 7: RECONCILIATION
# ==============================================================================

def reconcile_with_payroll(
    seller_summary: List[dict],
    payroll: Dict[str, float],
    all_results: List[dict]
) -> List[dict]:
    """
    Đối soát kết quả tool với bảng lương.
    KHÔNG TỰ SỬA KẾT QUẢ. Chỉ hiển thị chênh lệch.
    """
    if not payroll:
        return []

    print("\n" + "=" * 70)
    print("📊 ĐỐI SOÁT VỚI BẢNG LƯƠNG:")
    print("=" * 70)

    reconciliation = []
    tool_by_seller = {s['seller']: s['total_near_expiry_sales'] for s in seller_summary}

    # Tất cả sellers từ cả 2 nguồn
    all_sellers = set(payroll.keys()) | set(tool_by_seller.keys())

    match_count = 0
    diff_count = 0

    for seller in sorted(all_sellers):
        tool_val = tool_by_seller.get(seller, 0)
        payroll_val = payroll.get(seller, 0)
        diff = tool_val - payroll_val

        status = '✅ KHỚP' if abs(diff) < 1.0 else '❌ LỆCH'
        if abs(diff) < 1.0:
            match_count += 1
        else:
            diff_count += 1

        rec = {
            'seller': seller,
            'tool_amount': tool_val,
            'payroll_amount': payroll_val,
            'difference': diff,
            'status': status,
            'detail_skus': []
        }

        # Nếu lệch, liệt kê các SKU gây lệch
        if abs(diff) >= 1.0:
            seller_items = [r for r in all_results if r['seller'] == seller and r['is_near_expiry']]
            for item in seller_items:
                rec['detail_skus'].append(f"{item['sku']}: {item['amount']:,.0f}đ")

        reconciliation.append(rec)
        print(f"  {seller:30s} | Tool: {tool_val:>12,.0f}đ | BLương: {payroll_val:>12,.0f}đ | {status} ({diff:+,.0f}đ)")

    print(f"\n  Kết quả: {match_count} khớp, {diff_count} lệch / {len(all_sellers)} tổng")
    return reconciliation


# ==============================================================================
# MODULE 8: EXPORT EXCEL
# ==============================================================================

def export_excel(
    results: List[dict],
    seller_summary: List[dict],
    reconciliation: List[dict],
    output_file: str,
    near_only: bool = True
):
    """
    Xuất báo cáo ra Excel với 2-3 sheets:
      1. TỔNG HỢP CẬN DATE
      2. CẬN DATE (chi tiết)
      3. ĐỐI SOÁT (nếu có)
    """
    print(f"\n--> Đang xuất báo cáo ra Excel: {output_file}")

    wb = openpyxl.Workbook()

    # Styles
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    sub_header_fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
    match_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    diff_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    near_fill = PatternFill(start_color='DAEEF3', end_color='DAEEF3', fill_type='solid')
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center')
    border_thin = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    # ----- SHEET 1: TỔNG HỢP CẬN DATE -----
    ws_sum = wb.active
    ws_sum.title = "TỔNG HỢP CẬN DATE"

    sum_headers = ['STT', 'Người Bán', 'Số Đơn', 'Số SP Cận Date',
                   'Doanh Thu Hàng Cận Date', 'Thưởng 5%']
    ws_sum.append(sum_headers)
    for c in range(1, len(sum_headers) + 1):
        cell = ws_sum.cell(1, c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = border_thin

    for idx, s in enumerate(seller_summary, 1):
        ws_sum.cell(idx + 1, 1, idx).alignment = align_center
        ws_sum.cell(idx + 1, 2, s['seller']).alignment = align_left
        ws_sum.cell(idx + 1, 3, s['order_count']).alignment = align_center
        ws_sum.cell(idx + 1, 4, s['near_expiry_count']).alignment = align_center

        c_sales = ws_sum.cell(idx + 1, 5, s['total_near_expiry_sales'])
        c_sales.alignment = align_right
        c_sales.number_format = '#,##0'

        c_bonus = ws_sum.cell(idx + 1, 6, s['total_bonus'])
        c_bonus.alignment = align_right
        c_bonus.number_format = '#,##0'

        for c in range(1, 7):
            ws_sum.cell(idx + 1, c).border = border_thin

    # Dòng tổng
    tot_row = len(seller_summary) + 2
    ws_sum.cell(tot_row, 1, '').border = border_thin
    c_tot_label = ws_sum.cell(tot_row, 2, 'TỔNG CỘNG')
    c_tot_label.font = Font(name='Calibri', size=11, bold=True, color='1F4E79')
    c_tot_label.border = border_thin

    ws_sum.cell(tot_row, 3, '').border = border_thin
    ws_sum.cell(tot_row, 4, sum(s['near_expiry_count'] for s in seller_summary)).border = border_thin

    c_tot_sales = ws_sum.cell(tot_row, 5, sum(s['total_near_expiry_sales'] for s in seller_summary))
    c_tot_sales.font = Font(name='Calibri', size=11, bold=True, color='C00000')
    c_tot_sales.number_format = '#,##0'
    c_tot_sales.border = border_thin

    c_tot_bonus = ws_sum.cell(tot_row, 6, sum(s['total_bonus'] for s in seller_summary))
    c_tot_bonus.font = Font(name='Calibri', size=11, bold=True, color='C00000')
    c_tot_bonus.number_format = '#,##0'
    c_tot_bonus.border = border_thin

    # ----- SHEET 2: CẬN DATE (chi tiết) -----
    ws_dt = wb.create_sheet(title="CẬN DATE")

    detail_headers = [
        'Mã Hóa Đơn', 'Ngày Mua', 'Người Bán', 'Mã Hàng (SKU)', 'Tên Sản Phẩm',
        'HSD', 'Ngày Bắt Đầu 6 Tháng', 'Thành Tiền',
        'SKU Được Duyệt', 'Trong 6 Tháng', 'Kết Quả', 'Tiền Thưởng 5%', 'Lý Do'
    ]
    ws_dt.append(detail_headers)
    for c in range(1, len(detail_headers) + 1):
        cell = ws_dt.cell(1, c)
        cell.font = header_font
        cell.fill = sub_header_fill
        cell.alignment = align_center
        cell.border = border_thin

    r_out = 2
    for item in results:
        if near_only and item.get('status') != 'CẬN DATE':
            continue

        ws_dt.cell(r_out, 1, item.get('order_id', '')).alignment = align_center
        ws_dt.cell(r_out, 2, item['purchase_date'].strftime('%d/%m/%Y') if item.get('purchase_date') else '').alignment = align_center
        ws_dt.cell(r_out, 3, item.get('seller', '')).alignment = align_left
        ws_dt.cell(r_out, 4, item.get('sku', '')).alignment = align_center
        ws_dt.cell(r_out, 5, item.get('product_name', '')).alignment = align_left
        ws_dt.cell(r_out, 6, item['expiry_date'].strftime('%d/%m/%Y') if item.get('expiry_date') else '').alignment = align_center
        ws_dt.cell(r_out, 7, item['start_date'].strftime('%d/%m/%Y') if item.get('start_date') else '').alignment = align_center

        c_amt = ws_dt.cell(r_out, 8, item.get('amount', 0))
        c_amt.alignment = align_right
        c_amt.number_format = '#,##0'

        ws_dt.cell(r_out, 9, 'YES' if item.get('approved_sku') else 'NO').alignment = align_center
        ws_dt.cell(r_out, 10, 'YES' if item.get('within_6months') else 'NO').alignment = align_center
        ws_dt.cell(r_out, 11, item.get('status', '')).alignment = align_center

        c_bonus = ws_dt.cell(r_out, 12, item.get('bonus', 0))
        c_bonus.alignment = align_right
        c_bonus.number_format = '#,##0'

        ws_dt.cell(r_out, 13, item.get('reason', '')).alignment = align_left

        # Tô màu dòng cận date
        if item.get('is_near_expiry'):
            for c in range(1, 14):
                ws_dt.cell(r_out, c).fill = near_fill

        for c in range(1, 14):
            ws_dt.cell(r_out, c).border = border_thin

        r_out += 1

    # ----- SHEET 3: ĐỐI SOÁT (nếu có) -----
    if reconciliation:
        ws_rec = wb.create_sheet(title="ĐỐI SOÁT")
        rec_headers = ['Người Bán', 'Tool Tính', 'Bảng Lương', 'Chênh Lệch', 'Trạng Thái']
        ws_rec.append(rec_headers)
        for c in range(1, len(rec_headers) + 1):
            cell = ws_rec.cell(1, c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = border_thin

        for idx, rec in enumerate(reconciliation, 2):
            ws_rec.cell(idx, 1, rec['seller']).alignment = align_left
            c_tool = ws_rec.cell(idx, 2, rec['tool_amount'])
            c_tool.number_format = '#,##0'
            c_tool.alignment = align_right

            c_pay = ws_rec.cell(idx, 3, rec['payroll_amount'])
            c_pay.number_format = '#,##0'
            c_pay.alignment = align_right

            c_diff = ws_rec.cell(idx, 4, rec['difference'])
            c_diff.number_format = '#,##0'
            c_diff.alignment = align_right

            ws_rec.cell(idx, 5, rec['status']).alignment = align_center

            # Tô màu
            fill = match_fill if abs(rec['difference']) < 1.0 else diff_fill
            for c in range(1, 6):
                ws_rec.cell(idx, c).fill = fill
                ws_rec.cell(idx, c).border = border_thin

    # Auto-width
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(min(max_len + 3, 40), 10)

    wb.save(output_file)
    wb.close()
    print(f"✅ Đã lưu file kết quả: {output_file}")


# ==============================================================================
# MODULE 9: SELF TESTS (10 TEST CASES BẮT BUỘC)
# ==============================================================================

def run_self_tests():
    """Chạy 10 test cases bắt buộc theo spec."""
    print("=" * 70)
    print("🧪 ĐANG CHẠY BỘ KIỂM THỬ TỰ ĐỘNG (10 TEST CASES)")
    print("=" * 70)

    approved = {'SP001', 'SP002', 'SP003'}

    # Test 1: HSD = PurchaseDate - 6 tháng (biên dưới) -> CẬN DATE nếu SKU được duyệt
    r1 = check_near_expiry('SP001', datetime.date(2026, 7, 15), datetime.date(2026, 1, 15),
                           'OK', 'OK', approved, 200000)
    assert r1['is_near_expiry'] == True, f"Test 1 FAILED: {r1}"
    assert r1['bonus'] == 10000.0, f"Test 1 FAILED bonus: {r1['bonus']}"
    print("✅ Test 1: HSD = NgàyMua - 6 tháng → CẬN DATE (biên dưới)")

    # Test 2: HSD = PurchaseDate (biên trên) -> CẬN DATE nếu SKU được duyệt
    r2 = check_near_expiry('SP001', datetime.date(2026, 7, 15), datetime.date(2026, 7, 15),
                           'OK', 'OK', approved, 200000)
    assert r2['is_near_expiry'] == True, f"Test 2 FAILED: {r2}"
    print("✅ Test 2: HSD = NgàyMua → CẬN DATE (biên trên)")

    # Test 3: HSD < PurchaseDate - 6 tháng -> KHÔNG CẬN DATE
    r3 = check_near_expiry('SP001', datetime.date(2026, 7, 15), datetime.date(2026, 1, 10),
                           'OK', 'OK', approved, 200000)
    assert r3['is_near_expiry'] == False, f"Test 3 FAILED: {r3}"
    print("✅ Test 3: HSD < NgàyMua - 6 tháng → KHÔNG CẬN DATE")

    # Test 4: HSD > PurchaseDate -> KHÔNG CẬN DATE
    r4 = check_near_expiry('SP001', datetime.date(2026, 7, 15), datetime.date(2026, 8, 20),
                           'OK', 'OK', approved, 200000)
    assert r4['is_near_expiry'] == False, f"Test 4 FAILED: {r4}"
    print("✅ Test 4: HSD > NgàyMua → KHÔNG CẬN DATE")

    # Test 5: SKU không nằm trong danh mục -> KHÔNG CẬN DATE
    r5 = check_near_expiry('SP999', datetime.date(2026, 7, 15), datetime.date(2026, 5, 20),
                           'OK', 'OK', approved, 200000)
    assert r5['is_near_expiry'] == False, f"Test 5 FAILED: {r5}"
    assert r5['approved_sku'] == False, f"Test 5 FAILED approved: {r5['approved_sku']}"
    print("✅ Test 5: SKU không trong danh mục → KHÔNG CẬN DATE")

    # Test 6: Đơn 10 SP, chỉ 2 SKU được tính -> chỉ 2 SP tính tiền
    items = []
    for i in range(10):
        items.append({
            'sku': f'SP00{i+1}' if i < 2 else f'SP99{i}',
            'purchase_date': datetime.date(2026, 7, 15),
            'expiry_date': datetime.date(2026, 5, 20),
            'purchase_date_status': 'OK',
            'expiry_date_status': 'OK',
            'amount': 100000,
            'seller': 'Test', 'order_id': 'HD001',
            'branch': '', 'product_name': f'Product {i+1}',
            'quantity': 1, 'unit_price': 100000
        })
    results = calculate_bonus(items, approved)
    near_count = sum(1 for r in results if r['is_near_expiry'])
    total_bonus = sum(r['bonus'] for r in results)
    assert near_count == 2, f"Test 6 FAILED: near_count={near_count}"
    assert total_bonus == 10000.0, f"Test 6 FAILED: bonus={total_bonus}"
    print("✅ Test 6: Đơn 10 SP, chỉ 2 SKU được duyệt → chỉ 2 SP tính thưởng")

    # Test 7: Cùng SKU xuất hiện nhiều lần -> kiểm tra từng dòng
    items7 = [
        {'sku': 'SP001', 'purchase_date': datetime.date(2026, 7, 15),
         'expiry_date': datetime.date(2026, 5, 20), 'purchase_date_status': 'OK',
         'expiry_date_status': 'OK', 'amount': 100000,
         'seller': 'A', 'order_id': 'HD01', 'branch': '', 'product_name': 'P1',
         'quantity': 1, 'unit_price': 100000},
        {'sku': 'SP001', 'purchase_date': datetime.date(2026, 7, 15),
         'expiry_date': datetime.date(2027, 3, 20), 'purchase_date_status': 'OK',
         'expiry_date_status': 'OK', 'amount': 200000,
         'seller': 'A', 'order_id': 'HD02', 'branch': '', 'product_name': 'P1',
         'quantity': 1, 'unit_price': 200000},
    ]
    results7 = calculate_bonus(items7, approved)
    assert results7[0]['is_near_expiry'] == True, "Test 7a FAILED"
    assert results7[1]['is_near_expiry'] == False, "Test 7b FAILED"
    print("✅ Test 7: Cùng SKU nhiều lần → kiểm tra từng dòng theo HSD thực tế")

    # Test 8: Thiếu HSD -> MISSING_EXPIRY_DATE
    r8 = check_near_expiry('SP001', datetime.date(2026, 7, 15), None,
                           'OK', 'MISSING', approved, 500000)
    assert r8['status'] == 'MISSING_EXPIRY_DATE', f"Test 8 FAILED: {r8['status']}"
    assert r8['bonus'] == 0.0, f"Test 8 FAILED bonus: {r8['bonus']}"
    print("✅ Test 8: Thiếu HSD → MISSING_EXPIRY_DATE")

    # Test 9: Thiếu ngày mua -> MISSING_PURCHASE_DATE
    r9 = check_near_expiry('SP001', None, datetime.date(2026, 5, 20),
                           'MISSING', 'OK', approved, 500000)
    assert r9['status'] == 'MISSING_PURCHASE_DATE', f"Test 9 FAILED: {r9['status']}"
    print("✅ Test 9: Thiếu ngày mua → MISSING_PURCHASE_DATE")

    # Test 10: SKU có khoảng trắng -> chuẩn hóa
    sku_raw = "  SP001  "
    sku_normalized = normalize_sku(sku_raw)
    assert sku_normalized == "SP001", f"Test 10 FAILED: '{sku_normalized}'"
    print("✅ Test 10: SKU ' SP001 ' → chuẩn hóa thành 'SP001'")

    # Bonus: Test EDATE chính xác ngày cuối tháng
    assert edate(datetime.date(2026, 8, 31), -6) == datetime.date(2026, 2, 28), "EDATE test FAILED"
    print("✅ Bonus: EDATE(31/08/2026, -6) = 28/02/2026")

    print("\n🎉 TOÀN BỘ 10/10 TEST CASES ĐẠT 100%!\n")


# ==============================================================================
# MODULE 10: CLI INTERFACE
# ==============================================================================

def main():
    parser = argparse.ArgumentParser(
        description="Tool Tính Cận Date & Thưởng 5% - Phiên bản chính thức",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Ví dụ sử dụng:
  python tinh_hang_can_date.py --test
  python tinh_hang_can_date.py -m 7 --forward
  python tinh_hang_can_date.py -i hoadon.xlsx -s danh_muc.xlsx -p bangluong.xlsx -o output.xlsx --forward
        """
    )
    parser.add_argument("--input", "-i", type=str, default=None,
                        help="Đường dẫn file hóa đơn")
    parser.add_argument("--skus", "-s", type=str, default=None,
                        help="Đường dẫn file danh mục SKU cận date được duyệt")
    parser.add_argument("--payroll", "-p", type=str, default=None,
                        help="Đường dẫn file bảng lương để đối soát (tùy chọn)")
    parser.add_argument("--output", "-o", type=str, default=None,
                        help="Đường dẫn file Excel xuất kết quả")
    parser.add_argument("--month", "-m", type=int, default=None,
                        help="Tháng cần xử lý (VD: 7, 8)")
    parser.add_argument("--rate", "-r", type=float, default=0.05,
                        help="Tỷ lệ hoa hồng (mặc định 0.05 = 5%%)")
    parser.add_argument("--all-items", action="store_true",
                        help="Xuất toàn bộ dòng ở sheet chi tiết (kể cả không cận date)")
    parser.add_argument("--forward", action="store_true",
                        help="Dùng hướng 6 tháng TỚI: purchase_date <= HSD <= EDATE(purchase_date,+6)")
    parser.add_argument("--test", action="store_true",
                        help="Chạy bộ kiểm thử tự động (10 Test Cases)")

    args = parser.parse_args()

    # Chạy tests
    if args.test:
        run_self_tests()
        return

    base_dir = os.path.dirname(os.path.abspath(__file__))
    m = args.month or 8

    # --- Tìm file hóa đơn ---
    input_path = args.input
    if not input_path:
        candidates = [
            os.path.join(base_dir, f'thang{m}', 'DATAKIOT', 'hoadon.xlsx'),
            os.path.join(base_dir, f'thang{m}', 'DanhSachChiTietHoaDon_3182026.xlsx'),
            os.path.join(base_dir, f'thang{m}', 'hoadon.xlsx'),
        ]
        for cp in candidates:
            if os.path.exists(cp):
                input_path = cp
                break

    if not input_path or not os.path.exists(input_path):
        print(f"❌ Lỗi: Không tìm thấy file hóa đơn cho tháng {m}!")
        print(f"   Sử dụng --input để chỉ định đường dẫn file.")
        sys.exit(1)

    # --- Tìm file danh mục SKU ---
    sku_path = args.skus
    if not sku_path:
        candidates_sku = [
            os.path.join(base_dir, f'thang{m}', f'DanhMucCanDate_Thang{m}.xlsx'),
            os.path.join(base_dir, f'thang{m}', 'DanhMucCanDate.xlsx'),
            os.path.join(base_dir, 'DanhMucCanDate.xlsx'),
        ]
        for cp in candidates_sku:
            if os.path.exists(cp):
                sku_path = cp
                break

    if not sku_path or not os.path.exists(sku_path):
        print(f"❌ Lỗi: Không tìm thấy file danh mục SKU cận date!")
        print(f"   Sử dụng --skus để chỉ định đường dẫn file danh mục.")
        print(f"   File cần có sheet với cột: SKU | Tên hàng | Trạng thái")
        sys.exit(1)

    # --- Output ---
    output_path = args.output
    if not output_path:
        output_path = os.path.join(base_dir, f'thang{m}', f'BaoCao_CanDate_Thang{m}_V2.xlsx')

    # --- Payroll (tùy chọn) ---
    payroll_path = args.payroll
    if not payroll_path:
        candidates_pay = [
            os.path.join(base_dir, f'thang{m}', 'target', f'BẢNG LƯƠNG THÁNG {m} 2026.xlsx'),
        ]
        for cp in candidates_pay:
            if os.path.exists(cp):
                payroll_path = cp
                break

    # ========== CHẠY TESTS TRƯỚC ==========
    run_self_tests()

    # ========== STEP 1: ĐỌC DỮ LIỆU ==========
    raw_rows, col_map = load_sales_data(input_path)
    approved_skus = load_approved_skus(sku_path)

    payroll_data = {}
    if payroll_path and os.path.exists(payroll_path):
        payroll_data = load_payroll(payroll_path)

    # ========== STEP 2-3: CHUẨN HÓA ==========
    normalized = normalize_data(raw_rows)

    # ========== STEP 4-6: TÍNH TOÁN ==========
    direction = 'forward' if args.forward else 'backward'
    print(f"\n   Hướng tính 6 tháng: {direction.upper()}")
    results = calculate_bonus(normalized, approved_skus, rate=args.rate, direction=direction)

    # ========== STEP 7: TỔNG HỢP ==========
    seller_summary = aggregate_by_seller(results)

    # ========== VALIDATION CUỐI CÙNG (Sec XVI) ==========
    total_rows = len(results)
    total_with_exp = sum(1 for r in results if r.get('expiry_date_status') == 'OK')
    total_approved = sum(1 for r in results if r.get('approved_sku'))
    total_within = sum(1 for r in results if r.get('within_6months'))
    total_near = sum(1 for r in results if r.get('is_near_expiry'))
    total_sales = sum(r['amount'] for r in results if r.get('is_near_expiry'))
    total_bonus = sum(r['bonus'] for r in results if r.get('is_near_expiry'))
    total_sellers = len(seller_summary)

    print("\n" + "=" * 70)
    print("📊 KẾT QUẢ TÍNH TOÁN HÀNG CẬN DATE:")
    print("=" * 70)
    print(f"  Tổng số dòng hóa đơn:                  {total_rows:>10,}")
    print(f"  Số dòng có HSD hợp lệ:                 {total_with_exp:>10,}")
    print(f"  Số dòng SKU thuộc danh mục được duyệt:  {total_approved:>10,}")
    print(f"  Số dòng HSD trong khoảng 6 tháng:       {total_within:>10,}")
    print(f"  Số dòng cuối cùng được tính thưởng:      {total_near:>10,}")
    print(f"  Tổng doanh thu cận date:        {total_sales:>15,.0f} VNĐ")
    print(f"  👉 TỔNG THƯỞNG 5%:              {total_bonus:>15,.0f} VNĐ")
    print(f"  Số nhân viên được thưởng:               {total_sellers:>10,} người")
    print("=" * 70)

    # Các SKU chưa có trong danh mục nhưng có HSD trong 6 tháng
    unknown_skus = set()
    for r in results:
        if not r.get('approved_sku') and r.get('within_6months') and r.get('amount', 0) > 0:
            unknown_skus.add(r.get('sku', ''))
    if unknown_skus:
        print(f"\n⚠️ {len(unknown_skus)} SKU có HSD trong 6 tháng nhưng KHÔNG trong danh mục:")
        for sku in sorted(list(unknown_skus))[:20]:
            print(f"   - {sku}")
        if len(unknown_skus) > 20:
            print(f"   ... và {len(unknown_skus) - 20} SKU khác")

    # ========== ĐỐI SOÁT ==========
    reconciliation = reconcile_with_payroll(seller_summary, payroll_data, results)

    # ========== XUẤT EXCEL ==========
    export_excel(results, seller_summary, reconciliation, output_path,
                 near_only=not args.all_items)


if __name__ == "__main__":
    main()
