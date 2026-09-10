import http.server
import socketserver
import urllib.parse
import json
import os
import sys
import unicodedata
import io
import webbrowser
import threading
import datetime
import re
import tempfile
import base64
import openpyxl
from collections import defaultdict
from kpi_engine import execute_kpi_engine, MonthlyPlan, to_float
from build_monthly_plan_packages import build_monthly_plan_file, scan_folder_for_project_data
from build_payroll_report import generate_payroll_report_perfect

if sys.stdout.encoding != 'utf-8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

PORT = 8765
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_PATH = os.path.join(BASE_DIR, 'monthly_config.json')

def load_config():
    if os.path.exists(CONFIG_PATH):
        try:
            with open(CONFIG_PATH, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            pass
    return {
        "active_month": 8,
        "branches": [],
        "staff_roles": {},
        "monthly_policies": {}
    }

def save_config(cfg):
    try:
        with open(CONFIG_PATH, 'w', encoding='utf-8') as f:
            json.dump(cfg, f, ensure_ascii=False, indent=2)
        return True
    except Exception as e:
        print(f"Error saving config: {e}")
        return False

def inspect_project_file_bytes(file_bytes):
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name
    
    try:
        wb = openpyxl.load_workbook(tmp_path, data_only=True)
        sheet_names = wb.sheetnames
        
        stores = []
        staff_by_branch = defaultdict(list)
        project_skus = []
        rules = []
        
        # 1. Parse 'KPI nhà thuốc' or store list
        if 'KPI nhà thuốc' in sheet_names:
            ws = wb['KPI nhà thuốc']
            for r in range(2, ws.max_row + 1):
                st_id = ws.cell(r, 1).value
                branch = ws.cell(r, 2).value
                region = ws.cell(r, 3).value
                target = to_float(ws.cell(r, 4).value)
                cht = ws.cell(r, 5).value
                m1 = to_float(ws.cell(r, 6).value)
                m2 = to_float(ws.cell(r, 7).value)
                m3 = to_float(ws.cell(r, 8).value)
                if branch:
                    stores.append({
                        'store_id': str(st_id or ''),
                        'branch': str(branch).strip(),
                        'region': str(region or 'HCM').strip(),
                        'target': target,
                        'cht': str(cht or '').strip(),
                        'm1': m1, 'm2': m2, 'm3': m3
                    })
                    
        # 2. Parse 'Nhân sự'
        if 'Nhân sự' in sheet_names:
            ws = wb['Nhân sự']
            for r in range(2, ws.max_row + 1):
                s_id = ws.cell(r, 1).value
                s_name = ws.cell(r, 2).value
                role = ws.cell(r, 3).value or 'Dược sĩ'
                st_id = ws.cell(r, 4).value
                branch = ws.cell(r, 5).value
                region = ws.cell(r, 6).value
                if s_name and branch:
                    b_str = str(branch).strip()
                    staff_by_branch[b_str].append({
                        'id': str(s_id or ''),
                        'name': str(s_name).strip(),
                        'role': str(role).strip(),
                        'region': str(region or '').strip()
                    })
                    
        # 3. Parse 'Danh mục dự án'
        if 'Danh mục dự án' in sheet_names:
            ws = wb['Danh mục dự án']
            for r in range(2, ws.max_row + 1):
                sku = ws.cell(r, 1).value
                name = ws.cell(r, 2).value
                grp = ws.cell(r, 3).value
                if sku or name:
                    project_skus.append({
                        'sku': str(sku or '').strip(),
                        'name': str(name or '').strip(),
                        'group': str(grp or 'CK').strip()
                    })
                    
        # Fallback: if template format (like 'kpi nhà thuốc' or branch names as sheets)
        if not stores:
            for sname in sheet_names:
                if sname not in ['data', 'kpi dược sĩ', 'kpi nhà thuốc', 'Dự án T8', 'Hot Bill HN', 'Cài đặt']:
                    stores.append({
                        'store_id': sname,
                        'branch': sname,
                        'region': 'HN' if sname in ['Hàng Bông', 'Đường Láng'] else 'HCM',
                        'target': 0,
                        'cht': ''
                    })
                    
        wb.close()
        
        # Attach staff to stores
        total_staff_count = 0
        total_target_all = 0
        for st in stores:
            b_name = st['branch']
            matched_staff = staff_by_branch.get(b_name, [])
            if not matched_staff:
                for kb, slist in staff_by_branch.items():
                    if kb in b_name or b_name in kb:
                        matched_staff = slist
                        break
            st['staff_list'] = matched_staff
            st['staff_count'] = len(matched_staff)
            total_staff_count += len(matched_staff)
            total_target_all += st.get('target', 0)
            
        # Expected report sheets
        expected_sheets = [
            {"name": "kpi dược sĩ", "desc": f"Bảng tính KPI chi tiết của {total_staff_count} Dược sĩ chia theo {len(stores)} Nhà thuốc"},
            {"name": "kpi nhà thuốc", "desc": f"Bảng tổng hợp doanh thu & KPI của {len(stores)} Nhà thuốc toàn hệ thống"},
            {"name": "Dự án", "desc": f"Chi tiết doanh số {len(project_skus)} sản phẩm thuộc Danh mục Dự án (CK, Combo...)"},
            {"name": "Hot Bill", "desc": "Danh sách hóa đơn đạt mốc thưởng nóng ngày"},
            {"name": "data", "desc": "Bảng dữ liệu giao dịch chi tiết sau khi nạp Hóa đơn & Trả hàng"}
        ]
        
        # Categories breakdown
        cat_counts = defaultdict(int)
        for it in project_skus:
            cat_counts[it['group']] += 1
            
        return {
            'success': True,
            'filename_type': 'Gói Kế Hoạch Dự Án KPI' if 'KPI nhà thuốc' in sheet_names else 'Template Mẫu Báo Cáo',
            'raw_sheets': sheet_names,
            'store_count': len(stores),
            'stores': stores,
            'total_staff': total_staff_count,
            'total_target': total_target_all,
            'project_skus_count': len(project_skus),
            'project_categories': dict(cat_counts),
            'project_skus_sample': project_skus[:20],
            'expected_sheets': expected_sheets
        }
    finally:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except Exception:
                pass

def inspect_timecard_bytes(file_bytes):
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name
    try:
        wb = openpyxl.load_workbook(tmp_path, data_only=True)
        ws = wb.active
        staff_set = set()
        branch_set = set()
        shift_rows = 0
        night_shift_rows = 0
        
        for r in range(2, ws.max_row + 1):
            name_v = ws.cell(r, 3).value or ws.cell(r, 2).value
            branch_v = ws.cell(r, 6).value or ws.cell(r, 5).value or ws.cell(r, 4).value
            shift_v = ws.cell(r, 7).value or ws.cell(r, 6).value
            
            if name_v and str(name_v).strip():
                staff_set.add(str(name_v).strip())
            if branch_v and str(branch_v).strip():
                b_clean = str(branch_v).strip()
                if 'NT' in b_clean or '24H' in b_clean or 'Nhà thuốc' in b_clean:
                    branch_set.add(b_clean)
            if shift_v and str(shift_v).strip():
                shift_rows += 1
                if 'đêm' in str(shift_v).lower() or 'tối' in str(shift_v).lower():
                    night_shift_rows += 1
                    
        wb.close()
        return {
            'success': True,
            'total_staff': len(staff_set),
            'staff_names': sorted(list(staff_set))[:30],
            'total_branches': len(branch_set),
            'branches': sorted(list(branch_set)),
            'total_shift_rows': shift_rows,
            'night_shift_rows': night_shift_rows
        }
    finally:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except Exception:
                pass

def process_kpi_from_upload(hoadon_bytes, trahang_bytes=None, plan_bytes=None, template_bytes=None, report_date_str=None, config_override=None, month_num=8):
    cfg = config_override if config_override else load_config()
    month_selected = int(month_num or cfg.get("active_month", 8))
    
    with tempfile.TemporaryDirectory() as tmpdir:
        # 1. Hóa đơn
        hd_path = os.path.join(tmpdir, "hoadon_input.xlsx")
        with open(hd_path, "wb") as f:
            f.write(hoadon_bytes)
            
        inv_ym = detect_data_month_from_invoices(hd_path)
        if inv_ym:
            month_selected = inv_ym[1]

        # 2. Trả hàng
        th_path = None
        if trahang_bytes:
            th_path = os.path.join(tmpdir, "trahang_input.xlsx")
            with open(th_path, "wb") as f:
                f.write(trahang_bytes)
                
        # 3. Kế hoạch (Plan)
        if plan_bytes:
            plan_path = os.path.join(tmpdir, "plan_input.xlsx")
            with open(plan_path, "wb") as f:
                f.write(plan_bytes)
        else:
            plan_cand_list = [
                os.path.join(BASE_DIR, 'plans', f'KeHoachKPI_2026-{month_selected:02d}.xlsx'),
                os.path.join(BASE_DIR, 'plans', f'NHÀ THUỐC THÁNG {month_selected} 2026.xlsx'),
                os.path.join(BASE_DIR, 'plans', 'KeHoachKPI_2026-08.xlsx')
            ]
            plan_path = next((c for c in plan_cand_list if os.path.exists(c)), plan_cand_list[0])
            
        # 4. Template mẫu
        if template_bytes:
            tmpl_path = os.path.join(tmpdir, "template_input.xlsx")
            with open(tmpl_path, "wb") as f:
                f.write(template_bytes)
        else:
            tmpl_cand_list = [
                os.path.join(BASE_DIR, 'goc', f'NHÀ THUỐC THÁNG {month_selected} 2026.xlsx'),
                os.path.join(BASE_DIR, 'plans', f'NHÀ THUỐC THÁNG {month_selected} 2026.xlsx'),
                os.path.join(BASE_DIR, 'goc', 'NHÀ THUỐC THÁNG 8 2026.xlsx')
            ]
            tmpl_path = next((c for c in tmpl_cand_list if os.path.exists(c)), tmpl_cand_list[0])
            
        out_path = os.path.join(tmpdir, f"baocaokpi_thang{month_selected}_hoanthien.xlsx")
        
        # Chạy Core KPI Engine trên file người dùng tải lên
        stats = execute_kpi_engine(hd_path, th_path, tmpl_path, plan_path, out_path, report_date_str, target_month=month_selected)
        
        with open(out_path, "rb") as f:
            generated_bytes = f.read()

    return stats, generated_bytes

def process_payroll_from_upload(timecard_bytes=None, kpi_bytes=None, hoadon_bytes=None, trahang_bytes=None, template_bytes=None, month_num=8):
    month_selected = int(month_num or 8)
    m_folder = os.path.join(BASE_DIR, f'thang{month_selected}')
    
    with tempfile.TemporaryDirectory() as tmpdir:
        # 1. Chấm công
        timecard_path = None
        if timecard_bytes:
            timecard_path = os.path.join(tmpdir, "chamcong_input.xlsx")
            with open(timecard_path, "wb") as f:
                f.write(timecard_bytes)
        else:
            cand = os.path.join(m_folder, 'DATA', f'BangChiTietChamCong_thang{month_selected}.xlsx')
            if os.path.exists(cand): timecard_path = cand
            else:
                cand2 = os.path.join(m_folder, f'BangChiTietChamCong_thang{month_selected}.xlsx')
                if os.path.exists(cand2): timecard_path = cand2

        # 2. KPI Report
        kpi_path = None
        if kpi_bytes:
            kpi_path = os.path.join(tmpdir, "kpi_input.xlsx")
            with open(kpi_path, "wb") as f:
                f.write(kpi_bytes)
        else:
            cand = os.path.join(BASE_DIR, f'baocaokpi_thang{month_selected}_hoanthien.xlsx')
            if os.path.exists(cand): kpi_path = cand
            else:
                cand2 = os.path.join(m_folder, f'baocaokpi_thang{month_selected}_hoanthien.xlsx')
                if os.path.exists(cand2): kpi_path = cand2

        # 3. Hóa đơn
        hd_path = None
        if hoadon_bytes:
            hd_path = os.path.join(tmpdir, "hoadon_input.xlsx")
            with open(hd_path, "wb") as f:
                f.write(hoadon_bytes)
        else:
            cand = os.path.join(m_folder, 'DATA', 'DanhSachChiTietHoaDon_3182026.xlsx')
            if os.path.exists(cand): hd_path = cand

        # 4. Trả hàng
        th_path = None
        if trahang_bytes:
            th_path = os.path.join(tmpdir, "trahang_input.xlsx")
            with open(th_path, "wb") as f:
                f.write(trahang_bytes)
        else:
            cand = os.path.join(m_folder, 'DATA', 'DanhSachChiTietTraHang_3182026.xlsx')
            if os.path.exists(cand): th_path = cand

        # 5. Template bảng lương
        tmpl_path = None
        if template_bytes:
            tmpl_path = os.path.join(tmpdir, "template_payroll_input.xlsx")
            with open(tmpl_path, "wb") as f:
                f.write(template_bytes)
        else:
            cand_list = [
                os.path.join(m_folder, 'tinhcongnhungthuongchia.xlsx'),
                os.path.join(m_folder, 'chiacongnhunggomthuong.xlsx'),
                os.path.join(m_folder, f'BANGLUONGTHANG{month_selected}.xlsx'),
                os.path.join(m_folder, f'BẢNG LƯƠNG THÁNG {month_selected} 2026.xlsx'),
                os.path.join(BASE_DIR, f'thang{month_selected}', f'BẢNG LƯƠNG THÁNG {month_selected} 2026.xlsx')
            ]
            for c in cand_list:
                if os.path.exists(c):
                    tmpl_path = c
                    break

        out_path = os.path.join(tmpdir, f"BANGLUONGTHANG{month_selected}_hoanthien.xlsx")
        
        # Chạy engine tạo bảng lương
        final_payroll_file = generate_payroll_report_perfect(
            kpi_file=kpi_path,
            template_file=tmpl_path,
            output_file=out_path,
            month=month_selected,
            inv_file=hd_path,
            ret_file=th_path,
            timecard_file=timecard_path
        )
        
        with open(final_payroll_file, "rb") as f:
            generated_bytes = f.read()

        # Parse basic summary stats for UI
        wb_check = openpyxl.load_workbook(final_payroll_file, data_only=True)
        ws_bl = wb_check['BẢNG LƯƠNG'] if 'BẢNG LƯƠNG' in wb_check.sheetnames else None
        
        payroll_stats = {
            'month': month_selected,
            'total_rows': 0,
            'total_night_shifts': 0,
            'total_day_hours': 0.0,
            'total_night_hours': 0.0,
            'total_ot_hours': 0.0,
            'total_night_allowance': 0.0,
            'total_kpi_reward': 0.0,
            'total_ck_reward': 0.0,
            'total_project_reward': 0.0,
            'staff_sample': []
        }
        
        if ws_bl:
            for r in range(4, ws_bl.max_row + 1):
                name = ws_bl.cell(r, 3).value
                branch = ws_bl.cell(r, 2).value
                role = ws_bl.cell(r, 4).value
                day_hrs = ws_bl.cell(r, 5).value or 0
                night_hrs = ws_bl.cell(r, 6).value or 0
                ot_hrs = ws_bl.cell(r, 9).value or 0
                night_shifts = ws_bl.cell(r, 11).value or 0
                night_allowance = ws_bl.cell(r, 24).value or 0
                kpi_rw = ws_bl.cell(r, 31).value or 0
                ck_rw = ws_bl.cell(r, 32).value or 0
                proj_rw = ws_bl.cell(r, 30).value or 0
                
                if name and str(name).strip() and str(name).strip() != 'Tổng':
                    payroll_stats['total_rows'] += 1
                    try:
                        payroll_stats['total_day_hours'] += float(day_hrs)
                        payroll_stats['total_night_hours'] += float(night_hrs)
                        payroll_stats['total_ot_hours'] += float(ot_hrs)
                        payroll_stats['total_night_shifts'] += int(night_shifts)
                        payroll_stats['total_night_allowance'] += float(night_allowance)
                        payroll_stats['total_kpi_reward'] += float(kpi_rw)
                        payroll_stats['total_ck_reward'] += float(ck_rw)
                        payroll_stats['total_project_reward'] += float(proj_rw)
                    except Exception:
                        pass
                        
                    if len(payroll_stats['staff_sample']) < 25:
                        payroll_stats['staff_sample'].append({
                            'name': str(name).strip(),
                            'branch': str(branch or '').strip(),
                            'role': str(role or '').strip(),
                            'day_hrs': float(day_hrs) if isinstance(day_hrs, (int, float)) else 0,
                            'night_hrs': float(night_hrs) if isinstance(night_hrs, (int, float)) else 0,
                            'night_shifts': int(night_shifts) if isinstance(night_shifts, (int, float)) else 0,
                            'night_allowance': float(night_allowance) if isinstance(night_allowance, (int, float)) else 0,
                            'kpi_reward': float(kpi_rw) if isinstance(kpi_rw, (int, float)) else 0,
                            'ck_reward': float(ck_rw) if isinstance(ck_rw, (int, float)) else 0
                        })
        wb_check.close()

    return payroll_stats, generated_bytes

LAST_GENERATED_EXCEL = None
LAST_GENERATED_PAYROLL = None
LAST_GENERATED_PLAN = None

class KPIRequestHandler(http.server.SimpleHTTPRequestHandler):
    def do_GET(self):
        global LAST_GENERATED_EXCEL, LAST_GENERATED_PLAN
        parsed_url = urllib.parse.urlparse(self.path)
        
        if parsed_url.path in ['/', '/index.html']:
            self.send_response(200)
            self.send_header('Content-Type', 'text/html; charset=utf-8')
            self.end_headers()
            html_path = os.path.join(BASE_DIR, 'index.html')
            with open(html_path, 'rb') as f:
                self.wfile.write(f.read())
                
        elif parsed_url.path == '/api/config':
            cfg = load_config()
            self.send_response(200)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.end_headers()
            self.wfile.write(json.dumps(cfg, ensure_ascii=False).encode('utf-8'))

        elif parsed_url.path == '/download':
            query_params = urllib.parse.parse_qs(parsed_url.query)
            custom_name = query_params.get('filename', ['baocaokpi_hoanthien.xlsx'])[0]
            if not custom_name.endswith('.xlsx'):
                custom_name += '.xlsx'
            
            safe_name = urllib.parse.quote(custom_name)
            ascii_name = unicodedata.normalize('NFKD', custom_name).encode('ascii', 'ignore').decode('ascii') or 'baocaokpi.xlsx'
            
            data_to_send = LAST_GENERATED_EXCEL
            if data_to_send:
                self.send_response(200)
                self.send_header('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                self.send_header('Content-Disposition', f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{safe_name}')
                self.send_header('Content-Length', str(len(data_to_send)))
                self.end_headers()
                self.wfile.write(data_to_send)
            else:
                self.send_error(404, "Chưa có file nào được tạo. Vui lòng tải file lên và bấm Xử lý trước.")

        elif parsed_url.path == '/download-plan':
            query_params = urllib.parse.parse_qs(parsed_url.query)
            custom_name = query_params.get('filename', ['KeHoachKPI.xlsx'])[0]
            if not custom_name.endswith('.xlsx'):
                custom_name += '.xlsx'
            
            safe_name = urllib.parse.quote(custom_name)
            ascii_name = unicodedata.normalize('NFKD', custom_name).encode('ascii', 'ignore').decode('ascii') or 'KeHoachKPI.xlsx'
            
            data_to_send = LAST_GENERATED_PLAN
            if data_to_send:
                self.send_response(200)
                self.send_header('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                self.send_header('Content-Disposition', f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{safe_name}')
                self.send_header('Content-Length', str(len(data_to_send)))
                self.end_headers()
                self.wfile.write(data_to_send)
            else:
                self.send_error(404, "Chưa có file kế hoạch nào được tạo.")

        elif parsed_url.path == '/download-payroll':
            query_params = urllib.parse.parse_qs(parsed_url.query)
            custom_name = query_params.get('filename', ['BANGLUONG_hoanthien.xlsx'])[0]
            if not custom_name.endswith('.xlsx'):
                custom_name += '.xlsx'
            
            safe_name = urllib.parse.quote(custom_name)
            ascii_name = unicodedata.normalize('NFKD', custom_name).encode('ascii', 'ignore').decode('ascii') or 'bangluong.xlsx'
            
            data_to_send = LAST_GENERATED_PAYROLL
            if data_to_send:
                self.send_response(200)
                self.send_header('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                self.send_header('Content-Disposition', f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{safe_name}')
                self.send_header('Content-Length', str(len(data_to_send)))
                self.end_headers()
                self.wfile.write(data_to_send)
            else:
                self.send_error(404, "Chưa có file Bảng Lương nào được tạo. Vui lòng nạp file chấm công và bấm Tính Toán Bảng Lương.")
        else:
            super().do_GET()

    def do_POST(self):
        global LAST_GENERATED_EXCEL, LAST_GENERATED_PAYROLL, LAST_GENERATED_PLAN
        if self.path == '/api/build-plan':
            content_length = int(self.headers.get('Content-Length', 0))
            post_data = self.rfile.read(content_length)
            try:
                payload = json.loads(post_data.decode('utf-8'))
                month_num = int(payload.get('month', 8))
                extra_programs = payload.get('extra_programs', [])
                folder_path = payload.get('folder_path', None)
                custom_branches = payload.get('custom_branches', None)
                files_uploaded = payload.get('files', [])  # list of {name, base64}
                
                with tempfile.TemporaryDirectory() as tmpdir:
                    # If user uploaded raw files, write them into tmpdir
                    scan_dir = tmpdir
                    if files_uploaded:
                        for f_obj in files_uploaded:
                            fname = f_obj.get('name', 'file.xlsx')
                            fb64 = f_obj.get('base64', '')
                            if fb64:
                                dst = os.path.join(tmpdir, fname)
                                os.makedirs(os.path.dirname(dst), exist_ok=True)
                                with open(dst, 'wb') as f_out:
                                    f_out.write(base64.b64decode(fb64))
                    elif folder_path and os.path.exists(folder_path):
                        scan_dir = folder_path
                    else:
                        # Fallback default month folder if exists
                        def_folder = os.path.join(BASE_DIR, f'thang{month_num}', f'Pharmacy_retail_Store_KPIs_August 2026')
                        if not os.path.exists(def_folder):
                            def_folder = os.path.join(BASE_DIR, f'thang{month_num}')
                        if os.path.exists(def_folder):
                            scan_dir = def_folder
                            
                    out_plan = os.path.join(tmpdir, f'KeHoachKPI_2026-{month_num:02d}.xlsx')
                    plan_info = build_monthly_plan_file(
                        month_num=month_num,
                        input_folder=scan_dir,
                        extra_programs=extra_programs,
                        custom_branches=custom_branches,
                        output_filepath=out_plan
                    )
                    
                    with open(out_plan, 'rb') as f:
                        LAST_GENERATED_PLAN = f.read()
                        
                response = {
                    'success': True,
                    'message': f'Đã dựng thành công Gói Kế Hoạch KPI Tháng {month_num}!',
                    'filename': f'KeHoachKPI_2026-{month_num:02d}.xlsx',
                    'plan_base64': base64.b64encode(LAST_GENERATED_PLAN).decode('utf-8'),
                    'info': plan_info
                }
                json_bytes = json.dumps(response, ensure_ascii=False, default=str).encode('utf-8')
                self.send_response(200)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(json_bytes)))
                self.end_headers()
                self.wfile.write(json_bytes)
            except Exception as e:
                import traceback
                err_msg = traceback.format_exc()
                self.send_response(500)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(err_msg)))
                self.end_headers()
                self.wfile.write(json.dumps({'success': False, 'error': str(e), 'trace': err_msg}, ensure_ascii=False).encode('utf-8'))

        elif self.path == '/api/inspect-timecard':
            content_length = int(self.headers.get('Content-Length', 0))
            post_data = self.rfile.read(content_length)
            try:
                payload = json.loads(post_data.decode('utf-8'))
                tc_b64 = payload.get('timecard_base64', '')
                if not tc_b64:
                    self.send_response(400)
                    self.send_header('Content-Type', 'application/json; charset=utf-8')
                    self.end_headers()
                    self.wfile.write(json.dumps({'success': False, 'error': 'Vui lòng chọn hoặc kéo thả file Chấm công!'}, ensure_ascii=False).encode('utf-8'))
                    return
                
                tc_bytes = base64.b64decode(tc_b64)
                inspection_result = inspect_timecard_bytes(tc_bytes)
                
                json_bytes = json.dumps(inspection_result, ensure_ascii=False, default=str).encode('utf-8')
                self.send_response(200)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(json_bytes)))
                self.end_headers()
                self.wfile.write(json_bytes)
            except Exception as e:
                import traceback
                err_msg = traceback.format_exc()
                self.send_response(500)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(err_msg)))
                self.end_headers()
                self.wfile.write(json.dumps({'success': False, 'error': str(e), 'trace': err_msg}, ensure_ascii=False).encode('utf-8'))

        elif self.path == '/api/generate-payroll':
            content_length = int(self.headers.get('Content-Length', 0))
            post_data = self.rfile.read(content_length)
            try:
                payload = json.loads(post_data.decode('utf-8'))
                timecard_b64 = payload.get('timecard_base64', '')
                kpi_b64 = payload.get('kpi_base64', '')
                hoadon_b64 = payload.get('hoadon_base64', '')
                trahang_b64 = payload.get('trahang_base64', '')
                template_b64 = payload.get('template_base64', '')
                month_num = int(payload.get('month', 8))
                
                timecard_bytes = base64.b64decode(timecard_b64) if timecard_b64 else None
                kpi_bytes = base64.b64decode(kpi_b64) if kpi_b64 else None
                hoadon_bytes = base64.b64decode(hoadon_b64) if hoadon_b64 else None
                trahang_bytes = base64.b64decode(trahang_b64) if trahang_b64 else None
                template_bytes = base64.b64decode(template_b64) if template_b64 else None
                
                # Check if we should fallback to LAST_GENERATED_EXCEL for kpi_bytes if not passed
                if not kpi_bytes and LAST_GENERATED_EXCEL:
                    kpi_bytes = LAST_GENERATED_EXCEL
                    
                stats, payroll_bytes = process_payroll_from_upload(
                    timecard_bytes=timecard_bytes,
                    kpi_bytes=kpi_bytes,
                    hoadon_bytes=hoadon_bytes,
                    trahang_bytes=trahang_bytes,
                    template_bytes=template_bytes,
                    month_num=month_num
                )
                
                LAST_GENERATED_PAYROLL = payroll_bytes
                
                response = {
                    'success': True,
                    'message': f'Đã tính toán & tạo thành công Bảng Lương Tháng {month_num} hoàn thiện!',
                    'stats': stats,
                    'filename': f'BANGLUONGTHANG{month_num}_hoanthien.xlsx'
                }
                
                json_bytes = json.dumps(response, ensure_ascii=False, default=str).encode('utf-8')
                self.send_response(200)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(json_bytes)))
                self.end_headers()
                self.wfile.write(json_bytes)
            except Exception as e:
                import traceback
                err_msg = traceback.format_exc()
                self.send_response(500)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(err_msg)))
                self.end_headers()
                self.wfile.write(json.dumps({'success': False, 'error': str(e), 'trace': err_msg}, ensure_ascii=False).encode('utf-8'))

        elif self.path == '/api/inspect-project':
            content_length = int(self.headers.get('Content-Length', 0))
            post_data = self.rfile.read(content_length)
            try:
                payload = json.loads(post_data.decode('utf-8'))
                proj_b64 = payload.get('project_base64', '')
                if not proj_b64:
                    self.send_response(400)
                    self.send_header('Content-Type', 'application/json; charset=utf-8')
                    self.end_headers()
                    self.wfile.write(json.dumps({'success': False, 'error': 'Vui lòng chọn hoặc kéo thả file Dự án/Kế hoạch!'}, ensure_ascii=False).encode('utf-8'))
                    return
                
                proj_bytes = base64.b64decode(proj_b64)
                inspection_result = inspect_project_file_bytes(proj_bytes)
                
                json_bytes = json.dumps(inspection_result, ensure_ascii=False, default=str).encode('utf-8')
                self.send_response(200)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(json_bytes)))
                self.end_headers()
                self.wfile.write(json_bytes)
            except Exception as e:
                import traceback
                err_msg = traceback.format_exc()
                self.send_response(500)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.end_headers()
                self.wfile.write(json.dumps({'success': False, 'error': str(e), 'trace': err_msg}, ensure_ascii=False).encode('utf-8'))

        elif self.path == '/api/config':
            content_length = int(self.headers.get('Content-Length', 0))
            post_data = self.rfile.read(content_length)
            try:
                new_cfg = json.loads(post_data.decode('utf-8'))
                save_config(new_cfg)
                self.send_response(200)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.end_headers()
                self.wfile.write(json.dumps({'success': True, 'message': 'Đã lưu cấu hình kế hoạch thành công!'}, ensure_ascii=False).encode('utf-8'))
            except Exception as e:
                self.send_response(500)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.end_headers()
                self.wfile.write(json.dumps({'success': False, 'error': str(e)}, ensure_ascii=False).encode('utf-8'))

        elif self.path == '/api/generate':
            content_length = int(self.headers.get('Content-Length', 0))
            post_data = self.rfile.read(content_length)
            
            try:
                payload = json.loads(post_data.decode('utf-8'))
                
                hoadon_b64 = payload.get('hoadon_base64', '')
                trahang_b64 = payload.get('trahang_base64', '')
                plan_b64 = payload.get('plan_base64', '')
                template_b64 = payload.get('template_base64', '')
                report_date = payload.get('report_date', None)
                month_num = payload.get('month', 8)
                config_override = payload.get('config', None)
                
                if not hoadon_b64:
                    response = {
                        'success': False,
                        'error': 'Vui lòng chọn hoặc kéo thả file Hóa đơn (.xlsx) trước khi bấm Xử lý!'
                    }
                    self.send_response(400)
                    self.send_header('Content-Type', 'application/json; charset=utf-8')
                    self.end_headers()
                    self.wfile.write(json.dumps(response, ensure_ascii=False).encode('utf-8'))
                    return

                hoadon_bytes = base64.b64decode(hoadon_b64)
                trahang_bytes = base64.b64decode(trahang_b64) if trahang_b64 else None
                plan_bytes = base64.b64decode(plan_b64) if plan_b64 else None
                template_bytes = base64.b64decode(template_b64) if template_b64 else None

                stats, excel_bytes = process_kpi_from_upload(
                    hoadon_bytes=hoadon_bytes,
                    trahang_bytes=trahang_bytes,
                    plan_bytes=plan_bytes,
                    template_bytes=template_bytes,
                    report_date_str=report_date,
                    config_override=config_override,
                    month_num=month_num
                )
                
                LAST_GENERATED_EXCEL = excel_bytes

                response = {
                    'success': True,
                    'message': f'Đã xử lý thành công file Hóa đơn! Tạo xong Báo cáo KPI Tháng {stats["month"]} (ngày {stats["report_date"]}).',
                    'stats': stats,
                    'filename': f'baocaokpi_thang{stats["month"]}_hoanthien.xlsx'
                }
                
                json_bytes = json.dumps(response, ensure_ascii=False, default=str).encode('utf-8')
                self.send_response(200)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(json_bytes)))
                self.end_headers()
                self.wfile.write(json_bytes)

            except Exception as e:
                import traceback
                err_msg = traceback.format_exc()
                print('Error processing upload:', err_msg)
                response = {
                    'success': False,
                    'error': str(e),
                    'trace': err_msg
                }
                err_bytes = json.dumps(response, ensure_ascii=False, default=str).encode('utf-8')
                self.send_response(500)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(err_bytes)))
                self.end_headers()
                self.wfile.write(err_bytes)

def start_server():
    server_address = ('', PORT)
    with socketserver.TCPServer(server_address, KPIRequestHandler) as httpd:
        print(f'=====================================================')
        print(f'🚀 MEDIGO KPI WEB APP (PURE FILE UPLOAD MODE)')
        print(f'👉 Mở trình duyệt tại: http://localhost:{PORT}')
        print(f'=====================================================')
        httpd.serve_forever()

if __name__ == '__main__':
    threading.Timer(1.0, lambda: webbrowser.open(f'http://localhost:{PORT}')).start()
    start_server()
