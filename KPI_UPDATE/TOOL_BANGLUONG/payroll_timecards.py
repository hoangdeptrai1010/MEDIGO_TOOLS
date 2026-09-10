"""
MODULE 1: BÓC TÁCH MÁY CHẤM CÔNG, GIỜ CÔNG & NGÀY CÔNG
- Bóc tách dữ liệu từ file BangChiTietChamCong_thangX.xlsx
- Quy tắc chuẩn:
  + Ca > 4h: Tính 1 ca công thực tế
  + Ca <= 4h: Không tính là 1 ca, nhưng vẫn cộng giờ vào tổng giờ làm
  + Phân biệt ca ngày vs ca đêm (tối qua đêm)
- Đổ dữ liệu và đồng bộ vào 2 sheet: 'Giờ công' và 'Ngày công'
"""

import openpyxl
import datetime
from collections import defaultdict

BRANCH_ALIAS = {
    'NT 24H Trường Sa': 'Trường Sa',
    'NT 24h Đỗ Quang Đẩu': 'Đỗ Quang Đẩu',
    'NT 24H Nam Hòa': 'Nam Hòa',
    'NT 24H Minh Châu 1': 'Minh Châu',
    'NT 24H Lê Bình': 'Lê Bình',
    'NT 24H Nguyễn Chí Thanh': 'Nguyễn Chí Thanh',
    'NT 24H Nguyễn Thị Thập': 'Nguyễn Thị Thập',
    'NT 24H Nguyễn Văn Quá': 'Nguyễn Văn Quá',
    'NT 24H Rạch Bùng Binh': 'Rạch Bùng Binh',
    'NT Đường Láng 247': 'Đường Láng',
    'NT Hàng Bông 247': 'Hàng Bông',
}

def parse_time_str(t_str):
    if not t_str or not isinstance(t_str, str):
        return None
    s = t_str.strip()
    try:
        parts = s.split(':')
        if len(parts) >= 2:
            return int(parts[0]) * 60 + int(parts[1])
    except:
        pass
    return None

def calc_duration_hours(t_in_str, t_out_str, is_night_shift=False):
    t_in = parse_time_str(t_in_str)
    t_out = parse_time_str(t_out_str)
    if t_in is None or t_out is None:
        return 0.0, "0h0p"
    
    if is_night_shift:
        if t_out < t_in:
            diff_min = (1440 - t_in) + t_out
        else:
            diff_min = t_out - t_in
    else:
        if t_out >= t_in:
            diff_min = t_out - t_in
        else:
            diff_min = (1440 - t_in) + t_out
            
    hours = diff_min / 60.0
    h_int = int(diff_min // 60)
    m_int = int(diff_min % 60)
    return hours, f"{h_int}h{m_int}p"

def process_timecard_sheets(wb_out, timecard_path, month=8):
    if not timecard_path or not openpyxl.load_workbook(timecard_path, data_only=True):
        print(f"--> [Chấm công] Không có file chấm công, bỏ qua cập nhật Giờ công & Ngày công.")
        return {}

    print(f"--> [Module Chấm Công] Đang bóc tách dữ liệu quẹt thẻ từ: {timecard_path}")
    wb_tc = openpyxl.load_workbook(timecard_path, data_only=True)
    ws_tc = wb_tc.active

    all_shifts = []
    shift_summary_dict = {}

    curr_staff = ""
    curr_branch = ""
    curr_role = ""
    curr_code = ""

    for r in range(5, ws_tc.max_row + 1):
        c_name = ws_tc.cell(r, 3).value
        c_branch = ws_tc.cell(r, 6).value
        c_shift = ws_tc.cell(r, 7).value
        c_role = ws_tc.cell(r, 5).value
        c_code = ws_tc.cell(r, 2).value

        if c_name and str(c_name).strip():
            curr_staff = str(c_name).strip()
        if c_branch and str(c_branch).strip():
            curr_branch = BRANCH_ALIAS.get(str(c_branch).strip(), str(c_branch).strip())
        if c_role and str(c_role).strip():
            curr_role = str(c_role).strip()
        if c_code and str(c_code).strip():
            curr_code = str(c_code).strip()

        if not curr_staff or not c_shift:
            continue

        shift_name = str(c_shift).strip()
        is_night = any(k in shift_name.lower() for k in ['đêm', 'tối'])

        shift_total_hours = 0.0
        shift_total_minutes = 0

        for day in range(1, 32):
            col_in = 8 + (day - 1) * 2
            col_out = col_in + 1

            v_in = ws_tc.cell(r, col_in).value
            v_out = ws_tc.cell(r, col_out).value

            if v_in and v_out:
                hrs, h_str = calc_duration_hours(str(v_in), str(v_out), is_night_shift=is_night)
                if hrs > 0:
                    shift_total_hours += hrs
                    is_sub_4h = 1 if (hrs <= 4.0) else 0
                    cong_thuc_te = 0 if is_sub_4h == 1 else 1
                    ca_dem_chuan = 1 if (is_night and cong_thuc_te == 1) else 0

                    dt_val = datetime.date(2026, month, day)
                    all_shifts.append({
                        'branch': curr_branch,
                        'name': curr_staff,
                        'date': dt_val,
                        'shift_name': shift_name,
                        'actual_hours_str': h_str,
                        'standard_hours': round(hrs, 2),
                        'is_sub_4h': is_sub_4h,
                        'cong_thuc_te': cong_thuc_te,
                        'ca_dem_chuan': ca_dem_chuan,
                        'is_night': is_night
                    })

        tot_min = int(round(shift_total_hours * 60))
        h_part = tot_min // 60
        m_part = tot_min % 60
        moi_ca_str = f"{h_part}h{m_part}p"
        quy_doi_val = round(shift_total_hours, 2)

        key = (curr_branch, curr_staff, shift_name)
        shift_summary_dict[key] = {
            'moi_ca': moi_ca_str,
            'quy_doi': quy_doi_val,
            'is_night': is_night
        }

    # Đổ dữ liệu vào sheet 'Ngày công'
    if 'Ngày công' in wb_out.sheetnames and all_shifts:
        ws_nc = wb_out['Ngày công']
        for r_idx, s in enumerate(all_shifts, start=2):
            ws_nc.cell(r_idx, 1, s['branch'])
            ws_nc.cell(r_idx, 2, s['name'])
            ws_nc.cell(r_idx, 3, s['date'])
            ws_nc.cell(r_idx, 4, s['shift_name'])
            ws_nc.cell(r_idx, 5, s['actual_hours_str'])
            ws_nc.cell(r_idx, 6, s['standard_hours'])
            ws_nc.cell(r_idx, 7, s['is_sub_4h'])
            ws_nc.cell(r_idx, 8, s['cong_thuc_te'])
            ws_nc.cell(r_idx, 9, s['ca_dem_chuan'])

    # Đổ dữ liệu vào sheet 'Giờ công'
    if 'Giờ công' in wb_out.sheetnames and shift_summary_dict:
        ws_gc = wb_out['Giờ công']
        for r in range(2, ws_gc.max_row + 1):
            br = ws_gc.cell(r, 1).value
            nm = ws_gc.cell(r, 2).value
            sh = ws_gc.cell(r, 3).value
            if br and nm and sh:
                key = (str(br).strip(), str(nm).strip(), str(sh).strip())
                if key in shift_summary_dict:
                    ws_gc.cell(r, 4, shift_summary_dict[key]['moi_ca'])
                    ws_gc.cell(r, 5, shift_summary_dict[key]['quy_doi'])

    print(f"✅ Đã xử lý {len(all_shifts)} ca quẹt thẻ vào sheet 'Ngày công' và 'Giờ công' thành công!")
    return shift_summary_dict
