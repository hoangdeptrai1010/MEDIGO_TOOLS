import http.server
import socketserver
import urllib.parse
import json
import os
import sys
import io
import webbrowser
import threading
import tempfile
import base64
import openpyxl
_CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
if _CURRENT_DIR not in sys.path:
    sys.path.insert(0, _CURRENT_DIR)

from build_payroll_report import generate_payroll_report_perfect

if sys.stdout.encoding != 'utf-8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

PORT = 8766
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PARENT_DIR = os.path.dirname(BASE_DIR)

def inspect_timecard_bytes(file_bytes):
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name
    try:
        wb = openpyxl.load_workbook(tmp_path, data_only=True)
        ws = wb.active
        staff_set = set()
        branch_set = set()
        shift_rows = 0
        night_shift_rows = 0
        
        for r in range(2, ws.max_row + 1):
            name_v = ws.cell(r, 3).value or ws.cell(r, 2).value
            branch_v = ws.cell(r, 6).value or ws.cell(r, 5).value or ws.cell(r, 4).value
            shift_v = ws.cell(r, 7).value or ws.cell(r, 6).value
            
            if name_v and str(name_v).strip():
                staff_set.add(str(name_v).strip())
            if branch_v and str(branch_v).strip():
                b_clean = str(branch_v).strip()
                if 'NT' in b_clean or '24H' in b_clean or 'Nhà thuốc' in b_clean:
                    branch_set.add(b_clean)
            if shift_v and str(shift_v).strip():
                shift_rows += 1
                if 'đêm' in str(shift_v).lower() or 'tối' in str(shift_v).lower():
                    night_shift_rows += 1
                    
        wb.close()
        return {
            'success': True,
            'total_staff': len(staff_set),
            'staff_names': sorted(list(staff_set))[:30],
            'total_branches': len(branch_set),
            'branches': sorted(list(branch_set)),
            'total_shift_rows': shift_rows,
            'night_shift_rows': night_shift_rows
        }
    finally:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except Exception:
                pass

def process_payroll_from_upload(timecard_bytes=None, kpi_bytes=None, hoadon_bytes=None, trahang_bytes=None, template_bytes=None, month_num=8):
    month_selected = int(month_num or 8)
    m_folder = os.path.join(PARENT_DIR, f'thang{month_selected}')
    
    with tempfile.TemporaryDirectory() as tmpdir:
        # 1. Chấm công
        timecard_path = None
        if timecard_bytes:
            timecard_path = os.path.join(tmpdir, "chamcong_input.xlsx")
            with open(timecard_path, "wb") as f:
                f.write(timecard_bytes)
        else:
            cand_list = [
                os.path.join(m_folder, 'DATA', f'BangChiTietChamCong_thang{month_selected}.xlsx'),
                os.path.join(m_folder, f'BangChiTietChamCong_thang{month_selected}.xlsx'),
                os.path.join(PARENT_DIR, f'thang{month_selected}', 'DATA', f'BangChiTietChamCong_thang{month_selected}.xlsx')
            ]
            timecard_path = next((c for c in cand_list if os.path.exists(c)), None)

        # 2. KPI Report
        kpi_path = None
        if kpi_bytes:
            kpi_path = os.path.join(tmpdir, "kpi_input.xlsx")
            with open(kpi_path, "wb") as f:
                f.write(kpi_bytes)
        else:
            cand_list = [
                os.path.join(PARENT_DIR, f'baocaokpi_thang{month_selected}_hoanthien.xlsx'),
                os.path.join(m_folder, f'baocaokpi_thang{month_selected}_hoanthien.xlsx')
            ]
            kpi_path = next((c for c in cand_list if os.path.exists(c)), None)

        # 3. Hóa đơn
        hd_path = None
        if hoadon_bytes:
            hd_path = os.path.join(tmpdir, "hoadon_input.xlsx")
            with open(hd_path, "wb") as f:
                f.write(hoadon_bytes)
        else:
            cand_list = [
                os.path.join(m_folder, 'DATA', 'DanhSachChiTietHoaDon_3182026.xlsx'),
                os.path.join(m_folder, 'DanhSachChiTietHoaDon_3182026.xlsx')
            ]
            hd_path = next((c for c in cand_list if os.path.exists(c)), None)

        # 4. Trả hàng
        th_path = None
        if trahang_bytes:
            th_path = os.path.join(tmpdir, "trahang_input.xlsx")
            with open(th_path, "wb") as f:
                f.write(trahang_bytes)
        else:
            cand_list = [
                os.path.join(m_folder, 'DATA', 'DanhSachChiTietTraHang_3182026.xlsx'),
                os.path.join(m_folder, 'DanhSachChiTietTraHang_3182026.xlsx')
            ]
            th_path = next((c for c in cand_list if os.path.exists(c)), None)

        # 5. Template bảng lương
        tmpl_path = None
        if template_bytes:
            tmpl_path = os.path.join(tmpdir, "template_payroll_input.xlsx")
            with open(tmpl_path, "wb") as f:
                f.write(template_bytes)
        else:
            cand_list = [
                os.path.join(m_folder, 'tinhcongnhungthuongchia.xlsx'),
                os.path.join(m_folder, 'chiacongnhunggomthuong.xlsx'),
                os.path.join(m_folder, f'BANGLUONGTHANG{month_selected}.xlsx'),
                os.path.join(m_folder, f'BẢNG LƯƠNG THÁNG {month_selected} 2026.xlsx'),
                os.path.join(PARENT_DIR, f'thang{month_selected}', f'BẢNG LƯƠNG THÁNG {month_selected} 2026.xlsx')
            ]
            tmpl_path = next((c for c in cand_list if os.path.exists(c)), None)

        out_path = os.path.join(tmpdir, f"BANGLUONGTHANG{month_selected}_hoanthien.xlsx")
        
        final_payroll_file = generate_payroll_report_perfect(
            kpi_file=kpi_path,
            template_file=tmpl_path,
            output_file=out_path,
            month=month_selected,
            inv_file=hd_path,
            ret_file=th_path,
            timecard_file=timecard_path
        )
        
        with open(final_payroll_file, "rb") as f:
            generated_bytes = f.read()

        wb_check = openpyxl.load_workbook(final_payroll_file, data_only=True)
        ws_bl = wb_check['BẢNG LƯƠNG'] if 'BẢNG LƯƠNG' in wb_check.sheetnames else None
        
        payroll_stats = {
            'month': month_selected,
            'total_rows': 0,
            'total_night_shifts': 0,
            'total_day_hours': 0.0,
            'total_night_hours': 0.0,
            'total_ot_hours': 0.0,
            'total_night_allowance': 0.0,
            'total_kpi_reward': 0.0,
            'total_ck_reward': 0.0,
            'total_project_reward': 0.0,
            'staff_sample': []
        }
        
        if ws_bl:
            for r in range(4, ws_bl.max_row + 1):
                name = ws_bl.cell(r, 3).value
                branch = ws_bl.cell(r, 2).value
                role = ws_bl.cell(r, 4).value
                day_hrs = ws_bl.cell(r, 5).value or 0
                night_hrs = ws_bl.cell(r, 6).value or 0
                ot_hrs = ws_bl.cell(r, 9).value or 0
                night_shifts = ws_bl.cell(r, 11).value or 0
                night_allowance = ws_bl.cell(r, 24).value or 0
                kpi_rw = ws_bl.cell(r, 31).value or 0
                ck_rw = ws_bl.cell(r, 32).value or 0
                proj_rw = ws_bl.cell(r, 30).value or 0
                
                if name and str(name).strip() and str(name).strip() != 'Tổng':
                    payroll_stats['total_rows'] += 1
                    try:
                        payroll_stats['total_day_hours'] += float(day_hrs)
                        payroll_stats['total_night_hours'] += float(night_hrs)
                        payroll_stats['total_ot_hours'] += float(ot_hrs)
                        payroll_stats['total_night_shifts'] += int(night_shifts)
                        payroll_stats['total_night_allowance'] += float(night_allowance)
                        payroll_stats['total_kpi_reward'] += float(kpi_rw)
                        payroll_stats['total_ck_reward'] += float(ck_rw)
                        payroll_stats['total_project_reward'] += float(proj_rw)
                    except Exception:
                        pass
                        
                    if len(payroll_stats['staff_sample']) < 25:
                        payroll_stats['staff_sample'].append({
                            'name': str(name).strip(),
                            'branch': str(branch or '').strip(),
                            'role': str(role or '').strip(),
                            'day_hrs': float(day_hrs) if isinstance(day_hrs, (int, float)) else 0,
                            'night_hrs': float(night_hrs) if isinstance(night_hrs, (int, float)) else 0,
                            'night_shifts': int(night_shifts) if isinstance(night_shifts, (int, float)) else 0,
                            'night_allowance': float(night_allowance) if isinstance(night_allowance, (int, float)) else 0,
                            'kpi_reward': float(kpi_rw) if isinstance(kpi_rw, (int, float)) else 0,
                            'ck_reward': float(ck_rw) if isinstance(ck_rw, (int, float)) else 0
                        })
        wb_check.close()

    return payroll_stats, generated_bytes

LAST_GENERATED_PAYROLL = None

class PayrollRequestHandler(http.server.SimpleHTTPRequestHandler):
    def do_GET(self):
        global LAST_GENERATED_PAYROLL
        parsed_url = urllib.parse.urlparse(self.path)
        
        if parsed_url.path in ['/', '/index.html']:
            self.send_response(200)
            self.send_header('Content-Type', 'text/html; charset=utf-8')
            self.end_headers()
            html_path = os.path.join(BASE_DIR, 'index.html')
            with open(html_path, 'rb') as f:
                self.wfile.write(f.read())

        elif parsed_url.path == '/download-payroll':
            query_params = urllib.parse.parse_qs(parsed_url.query)
            custom_name = query_params.get('filename', ['BANGLUONG_hoanthien.xlsx'])[0]
            if not custom_name.endswith('.xlsx'):
                custom_name += '.xlsx'
            
            safe_name = urllib.parse.quote(custom_name)
            data_to_send = LAST_GENERATED_PAYROLL
            if data_to_send:
                self.send_response(200)
                self.send_header('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                self.send_header('Content-Disposition', f'attachment; filename="{custom_name}"; filename*=UTF-8\'\'{safe_name}')
                self.send_header('Content-Length', str(len(data_to_send)))
                self.end_headers()
                self.wfile.write(data_to_send)
            else:
                self.send_error(404, "Chưa có file Bảng Lương nào được tạo. Vui lòng nạp file và bấm Tính Toán Bảng Lương.")
        else:
            super().do_GET()

    def do_POST(self):
        global LAST_GENERATED_PAYROLL
        if self.path == '/api/inspect-timecard':
            content_length = int(self.headers.get('Content-Length', 0))
            post_data = self.rfile.read(content_length)
            try:
                payload = json.loads(post_data.decode('utf-8'))
                tc_b64 = payload.get('timecard_base64', '')
                if not tc_b64:
                    self.send_response(400)
                    self.send_header('Content-Type', 'application/json; charset=utf-8')
                    self.end_headers()
                    self.wfile.write(json.dumps({'success': False, 'error': 'Vui lòng chọn file Chấm công!'}, ensure_ascii=False).encode('utf-8'))
                    return
                
                tc_bytes = base64.b64decode(tc_b64)
                inspection_result = inspect_timecard_bytes(tc_bytes)
                json_bytes = json.dumps(inspection_result, ensure_ascii=False, default=str).encode('utf-8')
                self.send_response(200)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(json_bytes)))
                self.end_headers()
                self.wfile.write(json_bytes)
            except Exception as e:
                import traceback
                err_msg = traceback.format_exc()
                self.send_response(500)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(err_msg)))
                self.end_headers()
                self.wfile.write(json.dumps({'success': False, 'error': str(e), 'trace': err_msg}, ensure_ascii=False).encode('utf-8'))

        elif self.path == '/api/generate-payroll':
            content_length = int(self.headers.get('Content-Length', 0))
            post_data = self.rfile.read(content_length)
            try:
                payload = json.loads(post_data.decode('utf-8'))
                timecard_b64 = payload.get('timecard_base64', '')
                kpi_b64 = payload.get('kpi_base64', '')
                hoadon_b64 = payload.get('hoadon_base64', '')
                trahang_b64 = payload.get('trahang_base64', '')
                template_b64 = payload.get('template_base64', '')
                month_num = int(payload.get('month', 8))
                
                timecard_bytes = base64.b64decode(timecard_b64) if timecard_b64 else None
                kpi_bytes = base64.b64decode(kpi_b64) if kpi_b64 else None
                hoadon_bytes = base64.b64decode(hoadon_b64) if hoadon_b64 else None
                trahang_bytes = base64.b64decode(trahang_b64) if trahang_b64 else None
                template_bytes = base64.b64decode(template_b64) if template_b64 else None
                
                stats, payroll_bytes = process_payroll_from_upload(
                    timecard_bytes=timecard_bytes,
                    kpi_bytes=kpi_bytes,
                    hoadon_bytes=hoadon_bytes,
                    trahang_bytes=trahang_bytes,
                    template_bytes=template_bytes,
                    month_num=month_num
                )
                
                LAST_GENERATED_PAYROLL = payroll_bytes
                
                response = {
                    'success': True,
                    'message': f'Đã tính toán & tạo thành công Bảng Lương Tháng {month_num} hoàn thiện!',
                    'stats': stats,
                    'filename': f'BANGLUONGTHANG{month_num}_hoanthien.xlsx'
                }
                
                json_bytes = json.dumps(response, ensure_ascii=False, default=str).encode('utf-8')
                self.send_response(200)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(json_bytes)))
                self.end_headers()
                self.wfile.write(json_bytes)
            except Exception as e:
                import traceback
                err_msg = traceback.format_exc()
                self.send_response(500)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(err_msg)))
                self.end_headers()
                self.wfile.write(json.dumps({'success': False, 'error': str(e), 'trace': err_msg}, ensure_ascii=False).encode('utf-8'))

def start_server():
    server_address = ('', PORT)
    with socketserver.TCPServer(server_address, PayrollRequestHandler) as httpd:
        print(f'=====================================================')
        print(f'💰 MEDIGO TOOL BẢNG LƯƠNG & CHẤM CÔNG SERVER RUNNING')
        print(f'👉 Mở trình duyệt tại: http://localhost:{PORT}')
        print(f'=====================================================')
        httpd.serve_forever()

if __name__ == '__main__':
    threading.Timer(1.0, lambda: webbrowser.open(f'http://localhost:{PORT}')).start()
    start_server()
