"""
MODULE 3: TÍNH THƯỞNG HÀNG CẬN DATE (5% DOANH SỐ HÀNG HSD <= 6 THÁNG)
- Quét các dòng sản phẩm có HSD <= 6 tháng so với ngày bán
- Trừ các hóa đơn trả hàng tương ứng
- Nhân tỷ lệ thưởng 5% cho từng dược sĩ theo chi nhánh
- Đổ dữ liệu và gắn công thức vào sheet 'Cận date'
"""

import openpyxl
import datetime
import os
from collections import defaultdict

BRANCH_MAP = {
    'NT Hàng Bông 247': 'Hàng Bông', 'NT Hàng Bông': 'Hàng Bông', 'Hàng Bông': 'Hàng Bông',
    'NT Đường Láng 247': 'Đường Láng', 'NT Đường Láng': 'Đường Láng', 'Đường Láng': 'Đường Láng',
    'NT 24H Minh Châu 1': 'Minh Châu', 'Minh Châu 1': 'Minh Châu', 'Minh Châu': 'Minh Châu',
    'NT 24H Nam Hòa': 'Nam Hòa', 'Nam Hòa': 'Nam Hòa',
    'NT 24H Nguyễn Chí Thanh': 'Nguyễn Chí Thanh', 'Nguyễn Chí Thanh': 'Nguyễn Chí Thanh',
    'NT 24H Nguyễn Thị Thập': 'Nguyễn Thị Thập', 'Nguyễn Thị Thập': 'Nguyễn Thị Thập',
    'NT 24H Nguyễn Văn Quá': 'Nguyễn Văn Quá', 'Nguyễn Văn Quá': 'Nguyễn Văn Quá',
    'NT 24H Rạch Bùng Binh': 'Rạch Bùng Binh', 'Rạch Bùng Binh': 'Rạch Bùng Binh',
    'NT 24H Trường Sa': 'Trường Sa', 'Trường Sa': 'Trường Sa',
    'NT 24h Đỗ Quang Đẩu': 'Đỗ Quang Đẩu', 'Đỗ Quang Đẩu': 'Đỗ Quang Đẩu',
    'NT 24H Lê Bình': 'Lê Bình', 'Lê Bình': 'Lê Bình',
}

def normalize_branch(b):
    if not b: return 'Trường Sa'
    b_clean = str(b).strip()
    return BRANCH_MAP.get(b_clean, b_clean)

def to_num(v, default=0.0):
    if v is None: return default
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace(',', '')
    try: return float(s)
    except: return default

def parse_date_val(v):
    if not v: return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.date() if isinstance(v, datetime.datetime) else v
    s = str(v).strip()
    if not s or s.lower() in ('none', 'nan', 'null', ''):
        return None
    if '-' in s:
        parts = s.split(' ')[0].split('-')
        if len(parts) == 3:
            try:
                if len(parts[0]) == 4:
                    return datetime.date(int(parts[0]), int(parts[1]), int(parts[2]))
                else:
                    return datetime.date(int(parts[2]), int(parts[1]), int(parts[0]))
            except: pass
    elif '/' in s:
        parts = s.split(' ')[0].split('/')
        if len(parts) == 3:
            try:
                if len(parts[0]) == 4:
                    return datetime.date(int(parts[0]), int(parts[1]), int(parts[2]))
                else:
                    return datetime.date(int(parts[2]), int(parts[1]), int(parts[0]))
            except: pass
    return None

def scan_candate_sales(inv_file, ret_file=None):
    candate_data = defaultdict(float)
    if not inv_file:
        return candate_data

    try:
        wb_inv = openpyxl.load_workbook(inv_file, data_only=True, read_only=True)
        ws_inv = wb_inv.active
        hdr = next(ws_inv.iter_rows(min_row=1, max_row=1, values_only=True))
        col_map = {str(h).strip().lower(): i for i, h in enumerate(hdr) if h}

        branch_idx = col_map.get('chi nhánh', 0)
        seller_idx = col_map.get('người bán', 20 if len(hdr) > 20 else 7)
        sale_date_idx = col_map.get('thời gian', 6 if len(hdr) > 6 else 2)
        exp_date_idx = col_map.get('hạn sử dụng', col_map.get('hsd', 17))
        rev_idx = col_map.get('thành tiền', 15 if len(hdr) > 15 else 5)

        for r in ws_inv.iter_rows(min_row=2, values_only=True):
            seller = str(r[seller_idx]).strip() if seller_idx < len(r) and r[seller_idx] else ''
            br = normalize_branch(r[branch_idx]) if branch_idx < len(r) else 'Trường Sa'
            sale_d = parse_date_val(r[sale_date_idx]) if sale_date_idx < len(r) else None
            exp_d = parse_date_val(r[exp_date_idx]) if exp_date_idx < len(r) else None
            rev = to_num(r[rev_idx]) if rev_idx < len(r) else 0.0

            if seller and sale_d and exp_d and rev > 0:
                diff_days = (exp_d - sale_d).days
                if diff_days <= 183: # <= 6 tháng
                    candate_data[(br, seller)] += rev

        wb_inv.close()

        # Quét trừ trả hàng cận date nếu có
        if ret_file and os.path.exists(ret_file):
            wb_ret = openpyxl.load_workbook(ret_file, data_only=True, read_only=True)
            ws_ret = wb_ret.active
            hdr_ret = next(ws_ret.iter_rows(min_row=1, max_row=1, values_only=True))
            col_map_ret = {str(h).strip().lower(): i for i, h in enumerate(hdr_ret) if h}

            r_br_idx = col_map_ret.get('chi nhánh', 0)
            r_seller_idx = col_map_ret.get('người bán', 6)
            r_date_idx = col_map_ret.get('thời gian', 2)
            r_exp_idx = col_map_ret.get('hạn sử dụng', col_map_ret.get('hsd', 33))
            r_qty_idx = col_map_ret.get('số lượng', 35)
            r_price_idx = col_map_ret.get('giá bán', 36)
            r_disc_idx = col_map_ret.get('giảm giá', 37)

            for r in ws_ret.iter_rows(min_row=2, values_only=True):
                seller = str(r[r_seller_idx]).strip() if r_seller_idx < len(r) and r[r_seller_idx] else ''
                br = normalize_branch(r[r_br_idx]) if r_br_idx < len(r) else 'Trường Sa'
                ret_d = parse_date_val(r[r_date_idx]) if r_date_idx < len(r) else None
                exp_d = parse_date_val(r[r_exp_idx]) if r_exp_idx < len(r) else None
                
                qty = to_num(r[r_qty_idx]) if r_qty_idx < len(r) else 0.0
                price = to_num(r[r_price_idx]) if r_price_idx < len(r) else 0.0
                disc = to_num(r[r_disc_idx]) if r_disc_idx < len(r) else 0.0
                ret_val = (qty * price) - disc

                if seller and ret_d and exp_d and ret_val > 0:
                    diff_days = (exp_d - ret_d).days
                    if diff_days <= 183:
                        candate_data[(br, seller)] -= ret_val

            wb_ret.close()
    except Exception as e:
        print(f"Warning scanning Near-Date sales: {e}")

    return candate_data

def process_candate_sheet(wb_out, inv_file=None, ret_file=None):
    if 'Cận date' not in wb_out.sheetnames:
        return

    print(f"--> [Module Hàng Cận Date] Đang tính toán doanh số cận date 6 tháng...")
    ws_cd = wb_out['Cận date']
    cd_data = scan_candate_sales(inv_file, ret_file)

    for r in range(2, ws_cd.max_row + 1):
        br = ws_cd.cell(r, 1).value
        nm = ws_cd.cell(r, 2).value
        if br and nm and str(nm).strip() and str(nm).strip() != 'Tổng':
            key = (normalize_branch(br), str(nm).strip())
            rev = cd_data.get(key, 0.0)
            ws_cd.cell(r, 4, rev).number_format = '#,##0'
            ws_cd.cell(r, 5, 0.05).number_format = '0.00%'
            ws_cd.cell(r, 3, f"=D{r}*E{r}").number_format = '#,##0'

    print("✅ Đã cập nhật xong sheet 'Cận date'!")
